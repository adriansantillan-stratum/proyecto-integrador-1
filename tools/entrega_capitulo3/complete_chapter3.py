"""Fill chapter 3 while preserving the academic DOCX package outside its slot."""
from __future__ import annotations
import argparse,copy,hashlib,io,json,re,zipfile
from pathlib import Path
from urllib.parse import quote
import os
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path(__file__).resolve().parents[2]
REFERENCE=ROOT/'docs/Integrador I_Grupo3_Correccion_Integral.docx'
OUTPUT=ROOT/'docs/Integrador I_Grupo3_Capitulo3_Final.docx'
DATA=ROOT/'tools/entrega_capitulo3/chapter3_content.json'
OUT=ROOT/'docs/entrega_capitulo3'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
TITLES=['Casos de uso','Arquitectura de software','Diagrama de componentes','Diagrama de clases','Modelamiento de base de datos','Entorno visual de los formularios','Código fuente','Manual de instalación']

def digest(b): return hashlib.sha256(b).hexdigest()
def pname(p): return ''.join(p.xpath('.//w:t/text()',namespaces=NS))
def safe_path(value):
 p=(ROOT/str(value).replace('\\','/')).resolve()
 if not p.is_relative_to(ROOT) or not p.is_file(): raise ValueError('Missing project source: '+str(value))
 return p
def flat(v):
 if isinstance(v,list):return '\n'.join(flat(i) for i in v)
 if isinstance(v,dict):return '\n'.join(str(k)+': '+flat(x) for k,x in v.items())
 return str(v)

class Chapter:
 def __init__(self,source):
  self.doc=Document(source); self.start=len(self.doc.element.body)-1
  self.width=self.doc.sections[0].page_width-self.doc.sections[0].left_margin-self.doc.sections[0].right_margin
  self.images=[];self.references=[];self.fig=0;self.section=0;self.counts={'paragraphs':0,'tables':0,'images':0,'codeFragments':0}
 def p(self,text,bold=False,style=None):
  p=self.doc.add_paragraph(style=style)
  p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.15
  r=p.add_run(flat(text));r.bold=bold
  if str(text).startswith(('Material complementario:','Referencia:')):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  self.counts['paragraphs']+=1
  return p
 def link(self,label,path):
  target=quote(os.path.relpath(path,OUTPUT.parent).replace('\\','/'),safe='/')
  rid=self.doc.part.relate_to(target,RT.HYPERLINK,is_external=True)
  p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(7)
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
  color=OxmlElement('w:color');color.set(qn('w:val'),'0057A3');pr.append(color)
  size=OxmlElement('w:sz');size.set(qn('w:val'),'20');pr.append(size);r.append(pr)
  t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
  return p
 def h(self,text,level=3):
  p=self.doc.add_heading(text,level=level);p.paragraph_format.keep_with_next=True
  p.paragraph_format.space_before=Pt(11);p.paragraph_format.space_after=Pt(6)
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
  return p
 def table(self,headers,rows,widths=None):
  if not rows:return
  t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  portions=widths or [1/len(headers)]*len(headers)
  portions=[float(x)/sum(portions) for x in portions]
  for i,values in enumerate([headers]+list(rows)):
   cells=t.rows[0].cells if i==0 else t.add_row().cells
   if headers[0]=='Función o símbolo real':t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   if len(values)!=len(headers):raise ValueError('Table column mismatch')
   if i==0 and headers[0]!='Función o símbolo real':t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
   for k,(cell,val) in enumerate(zip(cells,values)):
    cell.width=int(self.width*portions[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;cell.text=flat(val)
    pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     x=OxmlElement('w:'+side);x.set(qn('w:w'),'70');x.set(qn('w:type'),'dxa');m.append(x)
    pr.append(m)
    if i==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');pr.append(shade)
    for p in cell.paragraphs:
     p.alignment=WD_ALIGN_PARAGRAPH.LEFT
     p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.05
     for r in p.runs:r.font.size=Pt(10);r.bold=i==0
   # Permit long case descriptions to flow; no fixed row heights.
  for c,ratio in zip(t.columns,portions):c.width=int(self.width*ratio)
  borders=OxmlElement('w:tblBorders')
  for side in ['top','left','bottom','right','insideH','insideV']:
   b=OxmlElement('w:'+side)
   for k,v in [('val','single'),('sz','4'),('color','BFBFBF')]:b.set(qn('w:'+k),v)
   borders.append(b)
  t._tbl.tblPr.append(borders);self.counts['tables']+=1
  self.doc.add_paragraph().paragraph_format.space_after=Pt(0)
 def image(self,file,caption):
  path=safe_path(file)
  with Image.open(path) as im:w,h=im.size
  width=min(self.width,Inches(7.4)*w/h)
  p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  s=p.add_run().add_picture(str(path),width=int(width));s._inline.docPr.set('descr',caption)
  self.fig+=1
  c=self.p(f'Figura 3.{self.fig}. {caption}');c.paragraph_format.keep_with_next=True
  for run in c.runs:run.font.size=Pt(10);run.italic=True
  self.link('Fuente: desarrollo de EcoRoute. Abrir imagen original completa.',path)
  self.images.append({'file':path.relative_to(ROOT).as_posix(),'sha256':digest(path.read_bytes()),'caption':caption})
  self.counts['images']+=1
 def code(self,file,start,end,explanation):
  path=safe_path(file);lines=path.read_text(encoding='utf-8-sig').splitlines()
  start=int(start);end=int(end)
  if not 1<=start<=end<=len(lines) or end-start>35:raise ValueError('Code excerpt bounds')
  self.p(f'{file} · líneas {start} a {end}',bold=True).paragraph_format.keep_with_next=True
  for i in range(start-1,end):
   p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
   r=p.add_run(f'{i+1:>4}  '+lines[i]);r.font.name='Consolas';r.font.size=Pt(8.5)
  self.p(explanation)
  self.references.append({'file':file,'start':start,'end':end,'sha256':digest(path.read_bytes())});self.counts['codeFragments']+=1
 def case(self,case):
  self.h(case['id']+' '+case['title'],3)
  self.table(['Campo','Especificación'],[
   ['Actor y proceso',case['actor']+' · '+', '.join(case['processes'])],
   ['Precondiciones',case['preconditions']],
   ['Flujo principal',[str(i)+'. '+step for i,step in enumerate(case['basicFlow'],1)]],
   ['Alternativas y excepciones',[str(i)+'. '+step for i,step in enumerate(case['alternatives'],1)]],
   ['Postcondición',case['postcondition']],
   ['Requerimientos',', '.join(case.get('requirements',[]))+'\n'+case.get('requirementMappingNote','')],
   ['Evidencia',case.get('evidence','')],
  ],[.24,.76])
  for source in case.get('sources',[]):
   path=safe_path(source['file']);lines=path.read_text(encoding='utf-8-sig').splitlines();line=int(source['line'])
   if source['anchor'] not in lines[line-1]:
    found=[i for i,text in enumerate(lines,1) if source['anchor'] in text]
    if not found:raise ValueError('Source anchor missing: '+source['file'])
    line=found[0]
   self.p('Referencia: '+source['file']+' · línea '+str(line)+' · '+source['anchor'])
 def block(self,b):
  kind=b.get('kind',b.get('type','paragraph'))
  if kind in ('paragraph','prose','text'):self.p(b.get('text',b.get('prose','')))
  elif kind=='heading':self.h(b.get('text',b.get('title')),int(b.get('level',3)))
  elif kind=='table':self.table(b['headers'],b['rows'],b.get('widths'))
  elif kind in ('image','diagram','figure'):self.image(b.get('file',b.get('path')),b['caption'])
  elif kind in ('code','snippet'):self.code(b['file'],b['start'],b['end'],b.get('explanation',''))
  elif kind in ('list','steps'):
   for i,x in enumerate(b.get('items',b.get('steps',[])),1):self.p(f'{i}. '+flat(x))
  else:raise ValueError('Unsupported content kind: '+kind)
 def nodes(self):return list(self.doc.element.body)[self.start:-1]

def complete(data):
 source=REFERENCE.read_bytes();base=zipfile.ZipFile(io.BytesIO(source))
 original_xml=etree.fromstring(base.read('word/document.xml'));body=original_xml.find(qn('w:body'))
 paragraphs=list(body)
 # The design chapter lost its direct numbering in a prior revision. Restore
 # its existing list membership so development is chapter 3, not chapter 2.
 design=next(i for i,e in enumerate(paragraphs) if e.tag==qn('w:p') and re.search(r'CAP[IÍ]TULO\s*2',pname(e),re.I))
 development=next(i for i,e in enumerate(paragraphs) if e.tag==qn('w:p') and re.search(r'CAP[IÍ]TULO\s*3',pname(e),re.I))
 normalized=[]
 for e in paragraphs[design:development]:
  style=e.find('w:pPr/w:pStyle',NS)
  if e.tag!=qn('w:p') or style is None or style.get(qn('w:val')) not in ('Ttulo1','Ttulo2'):continue
  pr=e.find(qn('w:pPr'));num=pr.find(qn('w:numPr'))
  if num is None:num=etree.SubElement(pr,qn('w:numPr'))
  for child in list(num):num.remove(child)
  etree.SubElement(num,qn('w:ilvl')).set(qn('w:val'),'0' if style.get(qn('w:val'))=='Ttulo1' else '1')
  etree.SubElement(num,qn('w:numId')).set(qn('w:val'),'22')
  normalized.append(pname(e))
 chapter3=next(e for e in paragraphs if e.tag==qn('w:p') and re.search(r'CAP[IÍ]TULO\s*3',pname(e),re.I))
 chapter4=next(e for e in paragraphs if e.tag==qn('w:p') and re.search(r'CAP[IÍ]TULO\s*4',pname(e),re.I))
 a,b=paragraphs.index(chapter3),paragraphs.index(chapter4)
 slots=[e for e in paragraphs[a+1:b] if e.tag==qn('w:p') and e.find('w:pPr/w:pStyle',NS) is not None]
 if len(slots)!=8:raise ValueError('Expected eight chapter 3 slots')
 sections=data['sections']
 if len(sections)!=8:raise ValueError('Chapter 3 must contain eight sections')
 gen=Chapter(REFERENCE);new_by_slot=[]
 gen.p('Este capítulo desarrolla EcoRoute como soporte de los procesos de última milla de Grupo MICOTRANS S. A. C. Se describe la versión de fuentes entregada el 27 de octubre de 2026, su arquitectura, persistencia, interfaces y procedimiento de instalación. Las capturas y la prueba integrada del 26 de octubre corresponden a un entorno local con datos sintéticos; las comprobaciones técnicas del 27 de octubre se distinguen de esa evidencia y de la aceptación empresarial.')
 intro_nodes=gen.nodes().copy()
 for i,section in enumerate(sections):
  old_count=len(gen.doc.element.body)-1
  for p in section.get('intro',[]):gen.p(p)
  for block in section.get('blocks',[]):gen.block(block)
  for case_id in section.get('useCaseIds',[]):gen.case(next(c for c in data['useCases'] if c['id']==case_id))
  new_by_slot.append(list(gen.doc.element.body)[old_count:-1])
 # Reuse existing heading slots and numbering/bookmarks; insert content only.
 for e in reversed(intro_nodes):chapter3.addnext(copy.deepcopy(e))
 for slot,nodes in zip(slots,new_by_slot):
  for e in reversed(nodes):slot.addnext(copy.deepcopy(e))
 buffer=io.BytesIO();gen.doc.save(buffer);working=zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
 original_rels=etree.fromstring(base.read('word/_rels/document.xml.rels'))
 known={e.get('Id') for e in original_rels}
 for e in etree.fromstring(working.read('word/_rels/document.xml.rels')):
  if e.get('Id') not in known:original_rels.append(copy.deepcopy(e))
 payloads={name:base.read(name) for name in base.namelist()}
 payloads['word/document.xml']=etree.tostring(original_xml,xml_declaration=True,encoding='UTF-8',standalone=True)
 payloads['word/_rels/document.xml.rels']=etree.tostring(original_rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 types=etree.fromstring(base.read('[Content_Types].xml'));typekeys={(e.tag,e.get('Extension'),e.get('PartName')) for e in types}
 for e in etree.fromstring(working.read('[Content_Types].xml')):
  if (e.tag,e.get('Extension'),e.get('PartName')) not in typekeys:types.append(copy.deepcopy(e))
 payloads['[Content_Types].xml']=etree.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
 for name in working.namelist():
  if name not in payloads:payloads[name]=working.read(name)
 untouched=[name for name in base.namelist() if name not in ('word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml')]
 assert all(payloads[n]==base.read(n) for n in untouched)
 OUTPUT.parent.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(OUTPUT,'w',zipfile.ZIP_DEFLATED) as z:
  for name,value in payloads.items():z.writestr(name,value)
 report={'reference':REFERENCE.name,'referenceSha256':digest(source),'output':OUTPUT.name,'outputSha256':digest(OUTPUT.read_bytes()),'preservedOriginalParts':len(untouched),'allPreservedPartsByteIdenticalBeforeWordFieldRefresh':True,'chapter3Sections':TITLES,'counts':gen.counts,'images':gen.images,'codeReferences':gen.references,'otherChaptersRewritten':False,'existingDesignHeadingNumberingRestored':normalized}
 OUT.mkdir(parents=True,exist_ok=True);(OUT/'chapter3_build_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps({'output':OUTPUT.name,'counts':gen.counts,'preservedParts':len(untouched)},ensure_ascii=False))

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--validate-only',action='store_true');args=parser.parse_args()
 data=json.loads(DATA.read_text(encoding='utf-8-sig'))
 if args.validate_only:
  assert len(data['sections'])==8
  for section in data['sections']:
   for block in section.get('blocks',[]):
    if block.get('kind') in ('image','diagram','figure','code','snippet'):safe_path(block.get('file',block.get('path')))
  print('Eight sections and referenced files validated; no Word authored.')
 else:complete(data)
