"""Build integrated exposition, implementation manual and revised academic chapter.

All claims come from saved execution evidence or are explicitly identified as specifications.
The historical audit and previous Word/PPT versions remain untouched.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime
import json,re,html,hashlib,xml.etree.ElementTree as ET
from collections import defaultdict
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/correccion_integral'
E=OUT/'evidencias'
DIAG=OUT/'diagramas'
INV=json.loads((OUT/'inventario_trazabilidad.json').read_text(encoding='utf8'))
API=json.loads((E/'regresion_api.json').read_text(encoding='utf8'))
SCHEMA=json.loads((E/'esquema_fisico_verificado.json').read_text(encoding='utf8'))
ERD=json.loads((DIAG/'erd_verificado_manifest.json').read_text(encoding='utf8'))
MIGRATION_RANGE='V1 a V'+str(max(int(m['version']) for m in SCHEMA['migrations'] if m.get('version') and str(m['version']).isdigit()))
FONTSIZE=11
NARRATIVE={
'P01': ('Gestión de pedidos', 'Dispatcher y cliente solicitante',
    'La solicitud se revisa antes de reservar una ruta. El Dispatcher comprueba datos, dirección, cobertura y condiciones con el cliente. El registro técnico inicial queda PENDING y su confirmación comercial se guarda por separado en order_business_data. Un borrador sin aceptación, dirección o cobertura verificadas no pasa a asignación. El RUC es opcional y se valida su formato de once dígitos, no su existencia ante SUNAT.',
    'Cuando los datos no son utilizables, el Dispatcher solicita corrección o propone una alternativa. El rechazo, desistimiento o inviabilidad comercial termina esa solicitud sin despacho. Cuando existe conformidad, P01 entrega a P02 el mismo pedido registrado, sin duplicarlo. Cambiar la ventana de una entrega invalida la aceptación previa y obliga a confirmar el nuevo acuerdo.'),
'P02': ('Gestión de operaciones de reparto', 'Dispatcher, Jefatura logística y Conductor',
    'El Dispatcher consolida pedidos confirmados y selecciona conductor, vehículo y fecha. Crea una ruta PLANNED y asocia los IDs existentes mediante PATCH /orders/{id}/route. El servidor comprueba recursos activos y evita dos reservas activas del mismo recurso para la misma fecha. El recorrido consulta OSRM sobre la red vial; su resultado no incluye tráfico en tiempo real ni garantiza una solución global óptima con capacidades y ventanas.',
    'La Jefatura revisa carga y controles de salida mediante una cuenta ADMIN autorizada. El Conductor, con DRIVER y vínculo de identidad real, acepta su propia hoja. Solo después se admite IN_PROGRESS. No se despacha una ruta vacía. Si cambia la asignación o el acuerdo comercial, la revisión anterior deja de ser válida y debe repetirse. La inspección física y la aceptación real del cliente siguen siendo responsabilidades humanas.'),
'P03': ('Gestión de operaciones de reparto', 'Conductor y Dispatcher',
    'El Conductor consulta su recorrido, atiende cada parada y registra el resultado. DELIVERED no es una opción libre del cambio de estado: requiere una prueba de entrega con fotografía, firma, receptor, DNI de ocho dígitos, GPS válido y fecha de captura. El servidor registra además su propia hora de recepción y los hashes SHA256. Los archivos son privados y se recuperan a través de endpoints autenticados.',
    'Sin conexión, la aplicación conserva la operación en una cola privada por cuenta antes de intentar enviarla. Pendiente de sincronización no significa entrega confirmada por el servidor. Al recuperar conexión con el proceso activo, reabrir la aplicación o reanudarla, se intenta sincronizar. El servidor resuelve reenvíos por pedido devolviendo el POD existente. No se garantiza ejecución de un servicio nativo con la aplicación cerrada por el sistema operativo.'),
'P04': ('Gestión de operaciones de reparto', 'Conductor, Dispatcher y Jefatura logística',
    'Un resultado FAILED o RETURNED exige motivo y abre una incidencia del pedido. El Dispatcher revisa su sustento y coordina la solución con las personas afectadas. RETRY requiere que la jornada siga activa y crea otro intento IN_TRANSIT. REPLAN requiere una ruta PLANNED y crea otro intento ASSIGNED vinculado al original mediante order_attempts. El resultado del pedido original permanece intacto.',
    'La reprogramación exige nueva aceptación del cliente y permite acordar una nueva ventana antes de aprobar la jornada. Si el reintento en la jornada ya no es viable, debe reprogramarse; no se reabre una ruta cerrada. La Jefatura autoriza CLOSE_FAILED para un resultado fallido o devuelto. CLOSE_CASE cierra una incidencia general o una disputa sin convertir una entrega acreditada en fallida. La resolución se audita e impide duplicar intentos al reenviarse.'),
'P05': ('Gestión del desempeño logístico', 'Dispatcher y Jefatura logística',
    'El período y el conductor seleccionados delimitan todos los resultados de analítica. La fecha operativa se obtiene de la ruta; en pedidos sin ruta se toma la creación convertida a America/Lima. Los indicadores exponen numerador, denominador, porcentaje, meta y brecha. Si no hay denominador, el dato es nulo y se muestra sin información, no cero ni una mejora ficticia.',
    'La Jefatura interpreta las desviaciones y registra una acción con causa, responsable y fecha. La acción pasa de OPEN a IN_PROGRESS y solo se verifica con un valor de resultado y una nota por un usuario autorizado. Ese registro acredita la decisión y la medición declarada; no demuestra por sí mismo causalidad. Las metas tienen historial. Power BI dispone de vistas de lectura previstas, pero su conexión real y la validación empresarial siguen siendo actividades separadas.'),
'P06': ('Gestión de recursos logísticos', 'Jefatura logística, Administrador y Dispatcher',
    'La empresa determina si el personal y el vehículo están aptos antes de asignarlos. Las pantallas de Conductores y Vehículos registran datos, capacidades y condición activa. El servidor rechaza la planificación con recursos inactivos y los conflictos de reserva por fecha. El vínculo con una cuenta DRIVER es una operación de identidad separada del alta de datos del conductor.',
    'Las revisiones documentales, licencias vigentes, mantenimiento físico y disponibilidad real requieren documentos y decisiones humanas. isActive no equivale a un control automático de vencimientos. Si el recurso no cumple, no se habilita para la jornada y la gestión debe corregirlo o seleccionar un sustituto. No se presenta un CRUD como certificación de aptitud.'),
'P07': ('Gestión de operaciones de reparto', 'Conductor, Dispatcher y Jefatura logística',
    'El Conductor comparte muestras GPS con los identificadores reales de la ruta. Cada ping requiere identidad vinculada, vehículo propio y una jornada en ejecución. Se conserva el historial y se publica el punto por Redis. El mapa web abre un WebSocket y autentica su primer mensaje con JWT y vehicleId; una conexión anónima no recibe ubicaciones y se cierra con política de acceso.',
    'Si el canal en vivo falla, la interfaz dispone de consulta periódica y debe indicar la situación. El Dispatcher analiza pérdida de señal o desvíos y contacta al Conductor antes de escalar a P04. La falta de GPS no acredita automáticamente robo, retraso ni entrega. Los umbrales y protocolos de respuesta empresarial deben acordarse y validarse en el piloto.'),
'P08': ('Gestión de seguridad de acceso', 'Jefatura logística y Administrador',
    'La Jefatura justifica la necesidad del acceso y el Administrador asigna los permisos autorizados. Keycloak es la autoridad de identidad y roles. El backend valida firma, emisor, vigencia y revocación del JWT. Para un conductor compara el sub validado con drivers.external_id; no confía en un driverId elegido libremente por la aplicación.',
    'El panel modifica roles efectivos y habilitación en Keycloak mediante un cliente de servicio local autorizado. Las cuentas técnicas no son administrables como personas desde ese panel y se impide que un administrador se retire su propio acceso. Los cambios revocan tokens anteriores y sesiones y dejan auditoría. Las tablas auth_* heredadas no gobiernan estos permisos. Keycloak mantiene una base distinta de la base logística.'),
'P09': ('Gestión de compras', 'Dispatcher, Administrador y Jefatura logística',
    'El negocio identifica la necesidad, comprueba presupuesto, selecciona proveedor y autoriza el gasto antes de adquirir. La aplicación aporta el registro de compra, proveedor, importe, categoría, observaciones y estado, con consulta y modificación. Los comprobantes físicos, aprobación presupuestaria y recepción conforme deben acompañar la decisión empresarial.',
    'Si no hay aprobación, se ajusta la solicitud o se cancela con motivo. Si la recepción no es conforme, se gestiona devolución o corrección con el proveedor. El CRUD existente solo apoya una parte de P09; no constituye un ERP, un almacén automatizado ni una aprobación bancaria. Esta frontera se mantiene explícita en el mapeo y las pruebas de aceptación.')}

def clean_title(text):
    return re.sub(r'[^\w\sáéíóúÁÉÍÓÚñÑ]',' ',text).replace('_',' ').strip()

def setup(doc):
    section=doc.sections[0]
    section.page_width,section.page_height=Inches(8.5),Inches(11)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Caption']:
        style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(FONTSIZE if name=='Normal' else 25 if name=='Title' else 17 if name=='Heading 1' else 13 if name.startswith('Heading') else 10)
        style.paragraph_format.space_after=Pt(7)
        if name.startswith('Heading'):style.paragraph_format.keep_with_next=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('EcoRoute  Página ')
    page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
    for run in footer.runs:run.font.name='Arial';run.font.size=Pt(9)
    title=doc.styles['Title']
    for node in title.element.xpath('.//w:pBdr'):node.getparent().remove(node)

def para(doc,text='',style=None):
    requested_style=style
    if style and style not in doc.styles:style=None
    p=doc.add_paragraph(style=style)
    # Inline emphasis is intentionally restrained. Paths remain readable as plain source references.
    run=p.add_run(re.sub(r'\*\*|`','',str(text)))
    if requested_style=='Caption':run.font.size=Pt(10)
    return p

def heading(doc,text,level=1):return doc.add_heading(clean_title(text),level)

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    section=doc.sections[-1]
    available=(section.page_width-section.left_margin-section.right_margin)/914400
    if widths and sum(widths)>available:widths=[width*available/sum(widths) for width in widths]
    if widths:
        for col,width in zip(t.columns,widths):col.width=Inches(width)
    for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
    repeated=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeated)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text='' if txt is None else str(txt)
    for index,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');margins.append(node)
            tcPr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'204D70' if index==0 else 'F2F5F8' if index%2==0 else 'FFFFFF');tcPr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.keep_with_next=index==0
                for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(255,255,255) if index==0 else RGBColor(0,0,0);r.bold=index==0
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    para(doc)
    return t

def figure(doc,path,caption,width=6.65,max_height=8):
    if not path.exists():return False
    with Image.open(path) as image: ratio=image.height/image.width
    section=doc.sections[-1]
    available=(section.page_width-section.left_margin-section.right_margin)/914400
    width=min(width,max_height/ratio,available)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    para(doc,caption,'Caption')
    return True

def markdown(doc,path):
    lines=path.read_text(encoding='utf8').splitlines();i=0;code=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if line.startswith('```'):code=not code;continue
        if code:
            p=para(doc,line);p.paragraph_format.space_after=Pt(1)
            for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8)
            continue
        if not line:continue
        if line.startswith('#'):
            level=min(3,len(line)-len(line.lstrip('#')));heading(doc,line.lstrip('# '),level);continue
        if line.startswith('|') and i<len(lines) and re.match(r'^\|?[\s:|-]+$',lines[i].strip()):
            headers=[x.strip() for x in line.strip('|').split('|')];i+=1;rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')];rows.append((row+['']*len(headers))[:len(headers)]);i+=1
            table(doc,headers,rows);continue
        line=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',line)
        para(doc,line)

def cover(doc,title,purpose):
    doc.add_paragraph(title,'Title')
    para(doc,'EcoRoute para Grupo MICOTRANS SAC','Subtitle')
    para(doc,purpose)
    para(doc,'Revisión técnica del 20 de octubre de 2026 en el entorno local. Las pruebas técnicas y los registros de ensayo no sustituyen la aceptación por personas de MICOTRANS.')

def related_traces(process):return [t for t in INV['traces'] if t['processId']==process]

def current_captures():
    candidates=[]
    for result in (OUT/'capturas/web').glob('*/results.json'):
        data=json.loads(result.read_text(encoding='utf8'))
        if data.get('summary',{}).get('failed')==0 and not data.get('fatal'):
            candidates.append((data.get('finishedAt',''),result.parent,data))
    files=[]
    if candidates:
        _,folder,data=max(candidates,key=lambda entry:entry[0])
        files=[folder/name for name in data.get('screenshots',[]) if (folder/name).exists()]
    files+=sorted((OUT/'capturas/mobile').glob('*.png'))
    supplement=OUT/'capturas/web_suplementarias/results.json'
    if supplement.exists() and json.loads(supplement.read_text(encoding='utf8')).get('result')=='PASS':
        files+=sorted(supplement.parent.glob('*.png'))
    # Keep original execution evidence intact. A later, independently checked
    # capture may replace its presentation copy without rewriting old results.
    for manifest in sorted((OUT/'capturas').glob('cierre_visual_*/results.json')):
        data=json.loads(manifest.read_text(encoding='utf-8-sig'))
        if data.get('result')!='PASS':continue
        for replacement in data.get('replacements',[]):
            target=(manifest.parent/replacement['replacement']).resolve()
            if not target.is_relative_to(manifest.parent.resolve()) or not target.exists():
                raise ValueError(f'Invalid replacement capture in {manifest}')
            matches=[i for i,file in enumerate(files) if file.name==replacement['original']]
            if len(matches)!=1:
                raise ValueError(f'Expected one original capture: {replacement["original"]}')
            files[matches[0]]=target
        for capture in data.get('captures',[]):
            if capture.get('original') is not None:continue
            target=(manifest.parent/capture['replacement']).resolve()
            if not target.is_relative_to(manifest.parent.resolve()) or not target.exists():
                raise ValueError(f'Invalid additional capture in {manifest}')
            if target not in files:files.append(target)
    return files

def capture_process(file):
    name=file.stem.lower()
    if '/mobile/' in file.as_posix():
        if any(x in name for x in ['login','perfil']):return 'P08'
        if any(x in name for x in ['fallo','incidencia']):return 'P04'
        if any(x in name for x in ['routes','route_map','aceptacion','cierre_ruta']):return 'P02'
        return 'P03'
    if 'login' in name or 'users' in name:return 'P08'
    if 'driver' in name or 'vehicle' in name:return 'P06'
    if 'purchase' in name:return 'P09'
    if 'notification' in name:return 'P07'
    if 'incident' in name:return 'P04'
    # "mobile" contains "bi"; dispatch acceptance is P02, not business BI.
    if 'route' in name or name.startswith('p02'):return 'P02'
    if any(x in name for x in ['bi','dashboard','report','action','home','causas']):return 'P05'
    return 'P01'

def capture_gallery(doc,level=2):
    files=current_captures()
    if not files:
        para(doc,'Las capturas y sus manifiestos se consultan en el paquete correccion_integral.')
        return
    notes={
      'P01':'Alta sin ruta, revisión comercial, bloqueo de asignación no confirmada, CSV y nuevo acuerdo de ventana. El mismo pedido continúa en P02.',
      'P02':'Planificación y revisión de salida. Mis rutas y el mapa móvil apoyan también la supervisión P07. El despacho queda bloqueado hasta la aceptación del conductor.',
      'P03':'Atención de parada, fotografía, DNI de ensayo y firma, persistencia sin red, reapertura de la aplicación, sincronización e historial. La fecha futura de la ruta corresponde únicamente al escenario sintético aislado.',
      'P04':'Registro, evaluación y cierre de incidencia. RETRY y REPLAN crean intentos vinculados y conservan el resultado original.',
      'P05':'Panel operativo, listados, cinco vistas analíticas, metas y gestión de acciones. Las continuaciones muestran los gráficos inferiores de cada vista.',
      'P06':'Datos y formularios de conductor y vehículo. La inspección física y documental es una responsabilidad humana.',
      'P07':'Notificaciones y seguimiento. El mapa compartido aparece en P02; una alerta requiere evaluación humana antes de escalar a incidencia.',
      'P08':'Inicio de sesión, perfil y consulta de roles efectivos. Las capturas de usuarios son de lectura y no representan nuevas aprobaciones de acceso de la empresa.',
      'P09':'Consulta y registro administrativo de compra. El presupuesto y la recepción conforme se documentan fuera del CRUD.'}
    for process in INV['processes']:
        pid=process['id'];group=[file for file in files if capture_process(file)==pid]
        heading(doc,pid+' '+process['name']+': recorrido de pantallas',level)
        para(doc,notes[pid])
        for file in group:
            caption=file.stem.replace('_',' ')
            caption=re.sub(r'scroll(\d+)',r'continuación vertical \1 px',caption)
            caption=caption.replace('screen ','Vista ').replace('synthetic','escenario de ensayo')
            figure(doc,file,caption,width=3 if '/mobile/' in file.as_posix() else 6.65,max_height=6.6)
        if not group:para(doc,'Consultar la ficha de soporte y los estados del proceso en el catálogo navegable. Las tareas humanas no tienen una pantalla automatizada propia.')

def process_section(doc,process,with_nodes=False,level=2):
    pid=process['id'];macro,actors,flow,exceptions=NARRATIVE[pid]
    heading(doc,f'{pid} {process["name"]}',level)
    para(doc,f'Macroproceso: {macro}. Responsables funcionales: {actors}. Los nombres identifican responsabilidades del negocio que deben confirmarse con la empresa, no nuevos puestos contratados.')
    para(doc,flow);para(doc,exceptions)
    source=ROOT/'docs/entregables_procesos_ecoroute/negocio_revisado/imagenes'/f'{pid}_negocio.png'
    figure(doc,source,f'BPMN de {pid}. Vista general. El archivo BPMN editable y el visualizador permiten ampliar el detalle de todos los nodos. Los controles y finales también se describen en la tabla siguiente.',max_height=3)
    rows=[]
    for node in process['nodes']:
        if node['type'] in ('startEvent','endEvent','exclusiveGateway') and not with_nodes:continue
        outputs=[]
        for condition in node['conditions']:
            target=next((f.get('to') for f in node['salidas'] if f['flowId']==condition['flowId']),None)
            if condition.get('label'):outputs.append(f'{condition["label"]}: {target}')
        trace=next((t for t in INV['traces'] if t['nodeId']==node['id']),{})
        links=trace.get('links',[])
        endpoints=sorted({ep.get('id',str(ep)) if isinstance(ep,dict) else str(ep) for link in links for ep in link.get('endpoints',[])})
        mapping=trace.get('mapping',node.get('mapping','manual'))
        support='Decisión humana o intercambio externo' if mapping in ('manual','external') else 'Control del flujo' if node['type'] in ('startEvent','endEvent','exclusiveGateway') else ', '.join(endpoints) or 'Consulte vínculo de código en el inventario'
        rows.append((node['id'],node['actor'] or 'Control',node['name'], '; '.join(outputs) or support))
    table(doc,['Nodo','Responsable','Actividad o decisión','Condición o soporte'],rows,[.8,1.05,2.25,2.55])
    tests=[f'{t["id"]} {t["name"]}' for t in API['tests'] if pid in t['process']]
    para(doc,'Evidencia técnica relacionada: '+('; '.join(tests) if tests else 'Este proceso conserva controles humanos y requiere ejecución de sus casos UAT.'))

def dictionary(doc):
    heading(doc,'Diccionario físico actualizado',1)
    allowed=set(ERD['objects'])
    for item in SCHEMA['tables']:
        if item['name'] not in allowed:continue
        heading(doc,item['name'],2)
        constraints=item.get('constraints') or []
        para(doc,('Vista analítica de lectura. ' if item['kind']=='v' else 'Tabla del modelo de negocio. ')+(item.get('comment') or ''))
        table(doc,['Campo','Tipo PostgreSQL','Nulo','Valor por defecto'],[(c['name'],c['type'],'Sí' if c['nullable'] else 'No',c.get('default') or '') for c in item['columns']],[1.95,1.75,.45,2.5])
        if constraints:para(doc,'Restricciones: '+'; '.join(c['definition'] for c in constraints))
    para(doc,'CORP, auth_* heredadas, Flyway y el respaldo local de bytes no se incorporan como entidades centrales del negocio. Su existencia física se conserva en el script completo. Las claves lógicas de Keycloak no se representan como una FK PostgreSQL inexistente.')

def evidence_section(doc):
    heading(doc,'Pruebas ejecutadas y aceptación pendiente',1)
    summary=API.get('summary',{})
    para(doc,f'La regresión de API registrada contiene {summary.get("total",len(API["tests"]))} casos y {summary.get("passed",0)} aprobaciones técnicas. El archivo de resultados conserva fecha, escenario y detalle. Son ensayos con datos sintéticos, no resultados empresariales de mejora logística.')
    table(doc,['Caso','Proceso','Escenario','Resultado'],[(t['id'],t['process'],t['name'],t['result']) for t in API['tests']],[.5,.7,4.7,.75])
    para(doc,'La comprobación física contrasta historial de estados, outbox, notificaciones y vínculos de intento con los IDs de ensayo. Dos objetos de la entrega de regresión se recuperan y comparan mediante SHA256 con los valores de PostgreSQL. La prueba de reinicio de LocalStack valida recuperación de los mismos bytes por el respaldo local, no una supuesta persistencia nativa de S3.')
    if (E/'prueba_emulador_offline.json').exists():
        mobile=json.loads((E/'prueba_emulador_offline.json').read_text(encoding='utf8'))
        para(doc,f'El recorrido Android aislado registra {mobile["summary"]["passed"]} comprobaciones aprobadas y {mobile["summary"]["failed"]} fallidas en su versión final. El pedido sintético 27 conserva foto de cámara, firma trazada, receptor de ensayo y GPS. Sin red queda pendiente; tras force-stop y reapertura mantiene la cola. Al reconectar confirma DELIVERED, conserva un solo POD y ambos hashes coinciden con los archivos recuperados. No se borraron los datos de los emuladores originales. El informe anterior del verificador se conserva por separado: comparaba campos de hash ausentes del DTO; la comprobación final los lee de PostgreSQL.')
    if (E/'comprobaciones_finales_api.json').exists():
        para(doc,'La comprobación adicional de la API registra 18/18 escenarios aprobados para la vista previa CSV y validación de compras. El preview no persiste pedidos; las comparaciones de conteo y hash antes/después delimitan esa observación. Los importes de compra cero o negativos se rechazan con HTTP 400.')
    if (E/'calidad_mobile.json').exists():
        para(doc,'El informe calidad_mobile.json conserva la salida y código de finalización de analyze, la suite Flutter y la integración opt-in contra el backend local. Estas pruebas de cliente y contratos no sustituyen la aprobación del conductor de MICOTRANS ni certifican precisión GPS en campo.')
    if (E/'prueba_emulador_no_entrega.json').exists():
        para(doc,'El segundo escenario Android registra 4/4 comprobaciones: FAILED enviado desde la interfaz con motivo de ensayo, actor e historial persistidos, incidencia vinculada y cierre COMPLETED de la ruta sintética cuando ambas paradas terminan. La consulta de contraste es de solo lectura y conserva los IDs utilizados.')
    for file,title in [('web_fix_notes.md','Pruebas de la consola web'),('mobile_fix_notes.md','Pruebas de la aplicación móvil')]:
        p=ROOT/'tools/audit_e2e'/file
        if file.startswith('mobile') and not p.exists():p=ROOT/'mobile-app/tools/audit_e2e/mobile_fix_notes.md'
        if p.exists():
            heading(doc,title,2)
            para(doc,f'Registro de ejecución y límites: {p.relative_to(ROOT)}. Las capturas con datos simulados y los tests con fakes verifican el cliente, pero no sustituyen su integración con el backend real.')
    para(doc,'El plan UAT contiene escenarios por los nueve procesos y casos transversales. Sus resultados empresariales están pendientes. La plantilla de acta no incluye firmas ni aprobaciones ficticias. Las actividades manuales deben probarse con las personas que las realizan.')

def build_manual():
    doc=Document();setup(doc)
    cover(doc,'Manual integral de EcoRoute','Este manual explica el flujo de negocio, su implementación web y móvil, las reglas del servidor y la evidencia técnica disponible. Permite localizar cada tarea BPMN en las pantallas y en las funciones que la soportan, conservando la diferencia entre automatización y decisión humana.')
    heading(doc,'Arquitectura y límites del desarrollo',1)
    para(doc,'EcoRoute utiliza una consola React con TypeScript y una aplicación Flutter. Ambas consumen un backend modular Spring Boot WebFlux. El backend tiene un único despliegue y separa dominio, casos de uso, puertos y adaptadores; no se describe como un conjunto de microservicios. Algunos controladores heredados consultan directamente repositorios y los servicios AWS mantienen dependencias de infraestructura explícitas, por lo que la separación hexagonal no se presenta como perfecta en toda la base de código.')
    figure(doc,DIAG/'arquitectura_corregida.png','Arquitectura actual de componentes y comunicaciones. HTTP se utiliza en el entorno local; una puesta en producción exige TLS, identidad y servicios gestionados correctamente configurados.')
    para(doc,'React organiza las páginas, los hooks de consulta, el contexto de autenticación y los servicios HTTP. Flutter usa BLoC para eventos y estados, casos de uso, repositorios y datasources Dio. PostgreSQL conserva operación y auditoría. Redis distribuye GPS y mantiene controles de revocación y tasa. OSRM devuelve geometría y orden sobre una red vial local. Los eventos de entrega usan outbox transaccional, SQS y notificaciones persistentes.')
    para(doc,'Keycloak está fuera de la base logística. El contenedor local usa start-dev y su base H2 persistida en el volumen de Keycloak. Ese modo no es la configuración de producción recomendada para este proyecto. El cliente ecoroute-operations-service está limitado al realm local autorizado; la credencial se conserva en .env ignorado por Git y excluido de las imágenes de construcción.')
    for name,title in [('secuencia_identidad','Autenticación y permisos efectivos'),('secuencia_pedido_despacho','Del pedido a la salida de la jornada'),('secuencia_pod','Prueba de entrega y sincronización'),('secuencia_excepciones','Gestión de los intentos de entrega'),('secuencia_gps','Ubicaciones y supervisión'),('secuencia_bi','Indicadores y acciones de mejora')]:
        heading(doc,title,2);figure(doc,DIAG/f'{name}.png',title)
    heading(doc,'Procesos del negocio y trazabilidad BPMN',1)
    para(doc,'Los seis macroprocesos agrupan nueve procesos de objetivo concreto. Gestión es una denominación válida de agrupación; las actividades y procesos hoja se nombran mediante un verbo y un resultado. EcoRoute apoya a las personas, pero no ocupa un carril como actor del negocio. Las decisiones Sí y No, las convergencias y los finales corresponden al XML BPMN conservado.')
    for process in INV['processes']:process_section(doc,process,True)
    heading(doc,'Pantallas y explicación del código',1)
    narrative=OUT/'Manual_Pantallas_y_Codigo.md'
    if narrative.exists():markdown(doc,narrative)
    else:
        for screen in INV['screens']:
            heading(doc,screen['name'],2);para(doc,screen['id'])
            para(doc,'Rutas: '+', '.join(r['path'] for r in screen['routes'])+'; framework: '+screen['framework'])
    heading(doc,'Capturas del recorrido de usuario',1)
    para(doc,'Las capturas siguientes pertenecen al entorno local. La consola oculta datos personales en la evidencia y conserva el diseño de sus controles. Las imágenes de móvil utilizan un dispositivo temporal de pruebas. Sus resultados deben leerse junto con el manifiesto; ninguna captura equivale a aceptación empresarial.')
    capture_gallery(doc)
    heading(doc,'Modelo físico del negocio',1)
    para(doc,'El ERD se obtiene del catálogo instalado en PostgreSQL y no de tablas inferidas por nombres. Se incluyen columnas, tipos, nulabilidad y restricciones reales. Las líneas físicas son únicamente claves foráneas declaradas. La relación de drivers.external_id con sub de Keycloak es lógica y externa; no es una FK entre las dos bases.')
    for name,title in [('00_erd_negocio_global','Modelo completo del negocio'),('01_recursos_planificacion','Recursos y planificación'),('02_pedidos_evidencias','Pedidos y evidencias'),('03_incidencias_mejora','Incidencias y mejora'),('04_analitica_catalogos','Analítica y catálogos')]:
        heading(doc,title,2);figure(doc,DIAG/f'{name}.png',title,max_height=8)
    evidence_section(doc)
    heading(doc,'Índice completo de pantallas y componentes',1)
    table(doc,['Pantalla o componente','Tecnología','Ruta','Archivo'],[(s['name'],s['framework'],', '.join(r['path'] for r in s['routes']),s['id']) for s in INV['screens']],[1.3,.65,1.25,3.45])
    heading(doc,'Consulta de funciones y contratos',1)
    para(doc,'El catálogo navegable index.html contiene los métodos localizados, el código fuente por función con número de línea, los endpoints y sus asociaciones estáticas con tareas y tablas. La presencia de un método no acredita una prueba de usuario: los resultados ejecutados se consultan por separado en evidencias. El inventario se regenera después de integrar todas las modificaciones para evitar referencias obsoletas.')
    heading(doc,'Condiciones para el piloto y producción',1)
    para(doc,'El piloto requiere validar roles funcionales con MICOTRANS, reglas de cobertura y ventanas, privacidad y retención de imágenes, manejo de dispositivos y aceptación real de los reportes. Debe medirse carga y disponibilidad con población representativa. Las pruebas actuales no acreditan SLA, seguridad exhaustiva, firma digital certificada ni reducción causal de incidencias.')
    para(doc,'El perfil prod separa credenciales AWS de las credenciales de emulación, deshabilita Swagger, requiere Redis y Keycloak configurados y no habilita el respaldo local de imágenes. No se ha desplegado producción ni configurado TLS o Power BI real. El almacenamiento H2 de Keycloak y LocalStack corresponden al entorno de desarrollo, no a una recomendación operativa pública.')
    path=OUT/'EcoRoute_Manual_Integral_Implementacion_y_Pruebas.docx';doc.save(path);return path

def build_speech():
    doc=Document();setup(doc)
    cover(doc,'Guion técnico de exposición de EcoRoute','Este guion ofrece una narración para explicar al profesor las decisiones de negocio, las pantallas y la ejecución del código. Debe acompañarse de los diagramas y de una demostración con datos de ensayo, sin presentar pruebas técnicas como aceptación empresarial.')
    source=OUT/'Guion_Tecnico_Paso_a_Paso.md'
    if source.exists():markdown(doc,source)
    else:
        heading(doc,'Apertura de la exposición',1)
        para(doc,'“Profesor, voy a explicar EcoRoute comenzando por el trabajo de MICOTRANS. Primero mostraré qué decide cada responsable, después las pantallas que apoyan esa decisión y finalmente las funciones que validan y guardan los datos. Separaré las pruebas técnicas de la aceptación real de la empresa, que todavía debe documentarse con sus usuarios.”')
        heading(doc,'Explicación de arquitectura',1)
        para(doc,'“La web está desarrollada con React y TypeScript y la aplicación del conductor con Flutter y BLoC. Ambas consumen el mismo backend Spring Boot WebFlux. Es un monolito modular: sus módulos se organizan por capas y puertos, pero se despliegan como una sola aplicación de negocio. PostgreSQL conserva la operación. Keycloak administra identidades de forma independiente; Redis distribuye GPS y aplica controles de acceso y frecuencia. OSRM calcula un recorrido sobre la red vial, sin afirmar tráfico en vivo.”')
        para(doc,'“Una petición sale de un control de la pantalla, pasa por el servicio HTTP y llega al controlador REST. El servidor valida la sesión y el propietario de la operación. El caso de uso aplica las reglas de negocio y llama a un puerto de persistencia. Su adaptador R2DBC escribe en PostgreSQL. Un error de regla devuelve 409 y la interfaz debe mostrarlo; ocultarlo o inventar un resultado exitoso rompería el flujo que estamos demostrando.”')
        for process in INV['processes']:
            heading(doc,process['id']+' '+process['name'],1)
            for text in NARRATIVE[process['id']][2:]:para(doc,'“'+text+'”')
        manual=OUT/'Manual_Pantallas_y_Codigo.md'
        if manual.exists():
            source=manual.read_text(encoding='utf8')
            sections=re.split(r'(?m)^## ',source)
            for section in sections:
                first,_,body=section.partition('\n')
                if not re.search(r'\b[WM]\d\d\b',first):continue
                display=re.sub(r'^\d+\.\s*','',first)
                display=re.sub(r'`[^`]+`','',display)
                heading(doc,'Demostración '+display,1)
                for paragraph in re.split(r'\n\s*\n',body):
                    paragraph=paragraph.strip()
                    if not paragraph or paragraph.startswith('```'):continue
                    if paragraph.startswith('#'):heading(doc,paragraph.lstrip('# '),2)
                    else:para(doc,'“'+paragraph.replace('\n',' ').replace('`','').replace('**','')+'”')
        heading(doc,'Explicación detallada del ERD',1)
        para(doc,'“Comienzo por drivers y vehicles. Cada ruta reserva un conductor y un vehículo para una fecha. En routes se conserva su estado y sus tiempos reales. Las restricciones de reserva evitan que el mismo recurso tenga dos jornadas activas el mismo día. No dibujo una relación solo porque dos campos se llamen parecido: las líneas continuas del ERD corresponden a claves foráneas que existen en PostgreSQL.”')
        para(doc,'“Orders es el centro de los compromisos de entrega. Antes de planificar puede existir sin route_id. Su revisión comercial se separa en order_business_data, porque estar pendiente no significa que el cliente haya aceptado. Al asignarlo se conserva el ID. Order_status_history registra el resultado, el motivo, el actor y la fecha. Route_dispatch_reviews conserva aprobación y aceptación de la jornada, y se invalida si cambia el plan.”')
        para(doc,'“Delivery_proofs se vincula con el pedido y conserva referencias privadas de foto y firma, receptor, DNI, coordenadas, fecha de captura y fecha de recepción. Los hashes permiten comparar los bytes. No afirmamos una firma digital certificada. Las imágenes no se incrustan como datos simulados en la pantalla: el controlador las recupera del almacenamiento y exige autorización.”')
        para(doc,'“Incidents conserva el caso y su resolución. Order_attempts une el pedido original con otro intento. Así, una reprogramación no transforma un fracaso histórico en una entrega exitosa ni destruye su evidencia. Improvement_actions conserva las decisiones de mejora y analytics_goals y su historial permiten contrastar metas. Las vistas analytics no duplican la operación: organizan sus datos para consultas de lectura.”')
        para(doc,'“Keycloak no es otra tabla del reparto. Su subject UUID se relaciona lógicamente con drivers.external_id, pero no existe una clave foránea física entre ambas bases. Las tablas CORP y las de autorización local heredada se excluyen de la vista central del negocio. Local_evidence_backups es soporte de recuperación del emulador en desarrollo, no una nueva entidad de la gestión logística.”')
        heading(doc,'Cierre y pruebas',1)
        para(doc,f'“Para respaldar lo mostrado, la regresión de API conservada contiene {API.get("summary",{}).get("total",len(API["tests"]))} casos. Revisamos por separado las pruebas del cliente, la compilación y la consulta de registros y hashes. Un PASS técnico no es una aprobación de MICOTRANS. El plan de aceptación debe ejecutarse con sus responsables y conservar sus observaciones y decisiones reales.”')
    heading(doc,'Preguntas de contraste para la demostración',1)
    for question,answer in [
        ('Por qué no se crea un pedido con ruta','La confirmación comercial precede a la planificación. El mismo ID PENDING se asigna después; así no se duplican pedidos ni se fuerza una ruta antes de acordar el servicio.'),
        ('Qué ocurre si no hay internet','La aplicación conserva la evidencia por cuenta y comunica pendiente de sincronización. Solo al confirmar el servidor se muestra entrega confirmada. La ejecución automática depende de que el proceso pueda ejecutarse o vuelva a abrirse.'),
        ('Qué diferencia hay entre reintento y reprogramación','El reintento ocurre en la jornada todavía activa. La reprogramación exige una ruta planificada y nuevas condiciones con el cliente. Ambos crean otro intento vinculado y preservan el resultado del original.'),
        ('Por qué no todos los nodos tienen una pantalla','Un inicio, una compuerta o un final expresa control del negocio. El contacto con el cliente y las decisiones físicas son actividades humanas. Las pantallas apoyan ese trabajo sin sustituirlo.'),
        ('Es Power BI','La consola tiene vistas analíticas con filtros y gráficos. Las vistas PostgreSQL permiten preparar un consumo Power BI, pero una interfaz parecida no demuestra una conexión Power BI ejecutada.'),
        ('Es una arquitectura de microservicios','El backend es un monolito modular reactivo con separaciones de dominio y adaptadores. Keycloak, PostgreSQL, Redis, OSRM y emuladores AWS son servicios auxiliares, no microservicios del dominio de EcoRoute.'),
        ('Los hashes convierten el POD en evidencia legal inmutable','Los hashes permiten contrastar bytes. El comprobante no es una firma digital certificada ni una autoridad de sellado de tiempo. El control legal y la retención deben validarse antes del piloto.')]:
        heading(doc,question,2);para(doc,'“'+answer+'”')
    manifest=OUT/'_qa/ppt/deck_manifest.json'
    if manifest.exists():
        slides=json.loads(manifest.read_text(encoding='utf8')).get('slides',[])
        heading(doc,'Narración de la presentación integrada',1)
        para(doc,'Este recorrido sigue el orden de la presentación vigente. Las explicaciones previas permiten ampliar cada pantalla y función durante las preguntas del profesor.')
        for slide in slides:
            heading(doc,f'Lámina {slide.get("number",slide.get("slide"))}: {slide["title"]}',2)
            notes=slide.get('notes','')
            if isinstance(notes,list):notes='\n\n'.join(str(n) for n in notes)
            # Narration and technical provenance have different readers. Retain the
            # spoken exposition here; complete notes/provenance remain in the PPTX.
            if '\n\nDetalle técnico y fuentes\n' in notes:
                notes=notes.split('\n\nDetalle técnico y fuentes\n',1)[0]
            else:
                notes=notes.split('\n## ',1)[0]
                notes='\n'.join(line for line in notes.splitlines() if not (
                    line.startswith(('EcoRoute / MICOTRANS.', '#', 'NARRATIVE:', 'Fuentes:', 'Raíz del repositorio:', 'Inventario generado:'))
                    or line.strip()=='Detalle técnico y fuentes'))
            for paragraph in re.split(r'\n\s*\n',str(notes)):
                if paragraph.strip():para(doc,'“'+paragraph.strip()+'”')
    path=OUT/'EcoRoute_Guion_Tecnico_Exposicion.docx';doc.save(path);return path

def build_academic():
    source=ROOT/'docs/Integrador I_Grupo3.docx';doc=Document(source)
    body=doc._element.body
    paragraphs=doc.paragraphs
    start=next(p._p for p in paragraphs if 'CAPÍTULO 2' in p.text)
    end=next(p._p for p in paragraphs if 'CAPÍTULO 3' in p.text)
    start_index=list(body).index(start);end_index=list(body).index(end)
    for el in list(body)[start_index:end_index]:body.remove(el)
    before=set(body)
    heading(doc,'CAPÍTULO 2 DISEÑO',1)
    para(doc,'El diseño presenta los procesos de negocio de MICOTRANS, sus interfaces de apoyo y el modelo físico de datos. Los controles humanos se diferencian de las funciones automatizadas. La revisión incorpora la confirmación comercial, la asignación posterior del pedido, los intentos de entrega vinculados y los indicadores con filtros uniformes.')
    heading(doc,'Diagramas BPMN',2)
    para(doc,'Se agrupan nueve procesos en seis macroprocesos: Gestión de pedidos; Gestión de recursos logísticos; Gestión de operaciones de reparto; Gestión del desempeño logístico; Gestión de seguridad de acceso; Gestión de compras. Los procesos hoja mantienen un objetivo único y nombre en infinitivo. Los carriles representan responsabilidades humanas, no el sistema EcoRoute.')
    for process in INV['processes']:process_section(doc,process,True,3)
    heading(doc,'Prototipos Wireframes',2)
    para(doc,'Los wireframes conservan la distribución de la consola y la aplicación. El diseño de los estados incorpora formularios, validación, carga, error, vacío, sesión vencida y operación pendiente. El inventario completo identifica todas las páginas y componentes del cliente, además de su relación con las tareas BPMN. Los wireframes son decisiones de diseño; las capturas se identifican por separado como vistas de implementación o escenarios simulados.')
    wire_dir=ROOT/'docs/entregables_procesos_ecoroute'
    for file in sorted(wire_dir.rglob('*wireframes*alineado.png')):
        figure(doc,file,f'Wireframe de referencia {file.stem.replace("_"," ")}',max_height=5)
    heading(doc,'Prototipos Mockups',2)
    para(doc,'Las interfaces se agrupan por proceso completo y no por una selección reducida de ocho pantallas. Para cada proceso se presentan las vistas que apoyan sus actividades y los estados de excepción. El manual integral describe la conexión con frontend, API, reglas del servidor y tablas. La decisión o comunicación humana que no tiene pantalla se declara expresamente.')
    table(doc,['Proceso','Interfaces de apoyo'],[
        ('P01','Pedidos, alta sin ruta, revisión comercial, CSV, cambio de ventana'),('P02','Rutas, backlog, asignación, revisión de salida, aceptación móvil y mapa'),
        ('P03','Ruta móvil, detalle de pedido, cámara, firma, resultado, cola e historial; POD en web'),('P04','Resultado de no entrega, incidencias, resolución y nuevos intentos'),
        ('P05','Dashboard, reportes, cinco vistas BI, metas y acciones de mejora'),('P06','Conductores, vehículos e identidad vinculada'),
        ('P07','Mapa, historial GPS, estado del canal y notificaciones operativas'),('P08','Login web, Keycloak, login móvil, sesión, roles y habilitación'),('P09','Compras y registro administrativo; presupuesto y recepción manuales')],[.65,6])
    capture_gallery(doc,3)
    heading(doc,'Modelo de Datos',2)
    para(doc,f'El modelo parte del catálogo físico PostgreSQL con migraciones {MIGRATION_RANGE}. Se separan las tablas de negocio, las vistas analíticas, el soporte técnico y el contexto corporativo heredado. La base de Keycloak es independiente. Las claves foráneas dibujadas existen como restricciones; no se infieren por compartir nombres de columnas.')
    for name,title in [('00_erd_negocio_global','Vista completa del negocio'),('01_recursos_planificacion','Recursos y planificación'),('02_pedidos_evidencias','Pedidos y evidencias'),('03_incidencias_mejora','Incidencias y mejora'),('04_analitica_catalogos','Analítica y catálogos')]:figure(doc,DIAG/f'{name}.png',title,max_height=7.8)
    heading(doc,'Diccionario de Datos',2)
    # Subsections use source styles without altering the chapter-one typography.
    allowed=set(ERD['objects'])
    for item in SCHEMA['tables']:
        if item['name'] not in allowed:continue
        heading(doc,item['name'],3)
        para(doc,'Vista analítica de lectura.' if item['kind']=='v' else (item.get('comment') or 'Entidad del modelo de negocio.'))
        table(doc,['Campo','Tipo PostgreSQL','Permite nulo','Valor por defecto'],[(c['name'],c['type'],'Sí' if c['nullable'] else 'No',c.get('default') or '') for c in item['columns']],[1.8,1.65,.65,2.55])
        constraints=item.get('constraints') or []
        if constraints:para(doc,'Restricciones: '+'; '.join(c['definition'] for c in constraints))
    appended=[el for el in body if el not in before and el.tag!=qn('w:sectPr')]
    for el in appended:body.remove(el)
    position=list(body).index(end)
    for offset,el in enumerate(appended):body.insert(position+offset,el)
    # Preserve the original chapter text while correcting inherited pagination:
    # a heading/description stays with its table and headers repeat on continuation.
    for p in doc.paragraphs:
        following=p._p.getnext()
        if p.style.name.startswith('Heading') or (p.text.strip() and following is not None and following.tag==qn('w:tbl')):
            p.paragraph_format.keep_with_next=True
    for t in doc.tables:
        if not t.rows:continue
        header=t.rows[0]
        if header._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:
            header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell in header.cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    path=ROOT/'docs/Integrador I_Grupo3_Correccion_Integral.docx';doc.save(path);return path

def build_html():
    def esc(text):return html.escape(str(text))
    content=['<h1>EcoRoute documentación integral</h1><p>Revisión técnica del 20 de octubre de 2026. Validación humana de MICOTRANS pendiente. El catálogo distingue evidencia estática del código y pruebas realmente ejecutadas.</p>']
    if (ROOT/'docs/control_integral_20261022/index.html').exists():
        content.append('<p><a href="../control_integral_20261022/index.html">Control ampliado del 22 de octubre: requisitos, nodos, controles, medición y aceptación</a></p>')
    content.append('<p>Entregables vigentes: <a href="../Integrador%20I_Grupo3_Correccion_Integral.docx">Informe académico</a> · <a href="EcoRoute_Manual_Integral_Implementacion_y_Pruebas.docx">Manual de implementación y pruebas</a> · <a href="EcoRoute_Guion_Tecnico_Exposicion.docx">Guion de exposición</a> · <a href="EcoRoute_Arquitectura_y_Flujos_Verificados.pptx">Presentación</a> · <a href="../base_datos/EcoRoute_BD_Completa_V1-V15.sql">SQL completo V1 a V15</a></p>')
    content.append('<nav><a href="#arquitectura">Arquitectura</a> <a href="#procesos">Procesos</a> <a href="#pantallas">Pantallas</a> <a href="#contratos">Contratos</a> <a href="#codigo">Código</a> <a href="#pruebas">Pruebas</a></nav>')
    content.append('<h2 id="arquitectura">Arquitectura</h2>')
    for file in sorted(DIAG.glob('*.svg')):content.append(f'<details><summary>{esc(file.stem)}</summary><a target="_blank" href="diagramas/{file.name}">Ampliar fuente editable</a><img loading="lazy" src="diagramas/{file.name}" alt="{esc(file.stem)}"></details>')
    content.append('<h2 id="procesos">Procesos y condiciones completas</h2>')
    for process in INV['processes']:
        pid=process['id'];content.append(f'<section><h3>{pid} {esc(process["name"])}</h3>')
        for text in NARRATIVE[pid]:content.append(f'<p>{esc(text)}</p>')
        image=f'../entregables_procesos_ecoroute/negocio_revisado/imagenes/{pid}_negocio.png'
        bpmn=f'../entregables_procesos_ecoroute/negocio_revisado/bpmn/{pid}_negocio.bpmn'
        content.append(f'<a href="{bpmn}">BPMN 2 editable</a> <a target="_blank" href="{image}">Abrir diagrama completo</a><img loading="lazy" src="{image}" alt="{pid}">')
        for node in process['nodes']:
            branches=[]
            for output in node['salidas']:
                condition=next((c for c in node['conditions'] if c['flowId']==output['flowId']),{})
                branches.append(f'{condition.get("label") or "Continuar"}: {output.get("to")}')
            trace=next((t for t in INV['traces'] if t['nodeId']==node['id']),{})
            content.append(f'<details><summary>{esc(node["id"])} {esc(node["name"])}</summary><p>Responsable {esc(node["actor"])}. Tipo {esc(node["type"])}. {esc("; ".join(branches))}</p><pre>{esc(json.dumps(trace,ensure_ascii=False,indent=2))}</pre></details>')
        content.append('</section>')
    content.append('<h2 id="pantallas">Todas las páginas y componentes localizados</h2>')
    for file in current_captures():
        relative=file.relative_to(OUT).as_posix()
        content.append(f'<details><summary>{capture_process(file)} · Captura {esc(file.stem)}</summary><a target="_blank" href="{relative}">Abrir tamaño completo</a><img loading="lazy" src="{relative}" alt="{esc(file.stem)}"></details>')
    for screen in INV['screens']:
        content.append(f'<details><summary>{esc(screen["name"])} {esc(screen["framework"])}</summary><p>{esc(screen["id"])}</p><p>Rutas {esc(", ".join(r["path"] for r in screen["routes"]))}</p><p>Los estados y llamadas se localizaron estáticamente; su ejecución se acredita por el manifiesto de cada prueba.</p><pre>{esc(json.dumps(screen,ensure_ascii=False,indent=2))}</pre></details>')
    content.append('<h2 id="contratos">Contratos y dependencias</h2>')
    for endpoint in INV['endpoints']:content.append(f'<details><summary>{esc(endpoint["id"])}</summary><pre>{esc(json.dumps(endpoint,ensure_ascii=False,indent=2))}</pre></details>')
    content.append('<h2 id="codigo">Código fuente por función</h2><p>El código y las líneas corresponden al inventario de esta revisión. Los métodos heredados o resueltos dinámicamente pueden requerir navegación adicional en el IDE.</p>')
    for method in INV['methods']:
        file=ROOT/method['file']
        if not file.exists():continue
        lines=file.read_text(encoding='utf8').splitlines();excerpt='\n'.join(f'{i+1}: {lines[i]}' for i in range(method['line']-1,min(method.get('endLine',method['line']),len(lines))))
        content.append(f'<details><summary>{esc(method["name"])} en {esc(method["file"])} línea {method["line"]}</summary><pre>{esc(excerpt)}</pre></details>')
    content.append('<h2 id="pruebas">Evidencia de ejecución y aceptación</h2><p>La UAT humana permanece pendiente. No se atribuyen pruebas a MICOTRANS ni se inventan firmas.</p>')
    content.append(f'<pre>{esc(json.dumps(API,ensure_ascii=False,indent=2))}</pre>')
    for file in sorted(E.glob('*.json')):content.append(f'<p><a href="evidencias/{file.name}">{esc(file.name)}</a></p>')
    for file in ['CIERRE_TECNICO.md','Plan_Pruebas_Usuario.md','Acta_UAT_Plantilla.md','Manual_Pantallas_y_Codigo.md','Guion_Tecnico_Paso_a_Paso.md','inventario_trazabilidad.json']:
        if (OUT/file).exists():content.append(f'<p><a href="{file}">{esc(file)}</a></p>')
    page='<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>EcoRoute documentación integral</title><style>body{font:17px/1.6 Arial;color:#172c3b;max-width:1300px;margin:35px auto;padding:0 25px}nav{position:sticky;top:0;background:#eff4f8;padding:15px;display:flex;gap:20px;flex-wrap:wrap}h1,h2,h3{color:#172c3b}img{width:100%;height:auto;background:white}details{border:1px solid #ccd7df;padding:12px;margin:10px 0}summary{cursor:pointer;overflow-wrap:anywhere}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:13px/1.5 Consolas,monospace}a{color:#185a86}section{margin-top:30px}</style><body>'+''.join(content)+'</body></html>'
    (OUT/'index.html').write_text(page,encoding='utf8')

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    # Regenerate inventory immediately before this authoring step once agents have stopped edits.
    outputs=[build_manual(),build_speech(),build_academic()]
    build_html()
    print(json.dumps({'documents':[str(p) for p in outputs],'html':str(OUT/'index.html')},ensure_ascii=False))
