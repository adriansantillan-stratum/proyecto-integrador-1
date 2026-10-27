"""Build the auditable EcoRoute technical manual and editable diagram sources."""
from pathlib import Path
import json, html, textwrap, shutil, re, sys
from collections import Counter
from PIL import ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/auditoria_end_to_end_20261019'
DIAG=OUT/'diagramas'; DIAG.mkdir(exist_ok=True)
E=OUT/'evidencias'
REQ=json.loads((E/'requisitos_y_bpmn.json').read_text(encoding='utf-8'))
API=json.loads((E/'pruebas_api.json').read_text(encoding='utf-8'))['tests']
EXTRA=json.loads((E/'pruebas_adicionales.json').read_text(encoding='utf-8'))
BLUE='#204d78'; INK='#17324d'; RED='#b23637'; GRAY='#65758b'; PALE='#edf3f8'
BLOCKS=[]; FIGURES={}

def wrap_pixels(text,max_width,size):
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',size)
    result=[]
    for paragraph in text.split('\n'):
        line=''
        for word in paragraph.split():
            candidate=(line+' '+word).strip()
            if line and font.getlength(candidate)>max_width:result.append(line);line=word
            else:line=candidate
        result.append(line)
    return result

def add(kind,**data): BLOCKS.append({'kind':kind,**data})
def heading(text,level=1): add('h',text=text,level=level)
def p(text): add('p',text=text)
def table(headers,rows,widths=None): add('table',headers=headers,rows=rows,widths=widths)
def steps(items):
    for i,item in enumerate(items,1): p(f'{i}. {item}')
def page():add('break')
def fig(name,caption,width=6.9,reference_only=False):add('figure',name=name,caption=caption,width=width,reference_only=reference_only)
def shot(name,caption,mobile=False):add('screen',name=name,caption=caption,mobile=mobile)
def esc(s):return html.escape(str(s))

class SVG:
    def __init__(self,title,w=1400,h=900):
        self.w=w;self.h=h
        self.parts=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><defs><marker id="a" markerWidth="10" markerHeight="10" refX="8" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="none" stroke="{BLUE}" stroke-width="1.5"/></marker><marker id="r" markerWidth="10" markerHeight="10" refX="8" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="none" stroke="{RED}" stroke-width="1.5"/></marker></defs><rect width="100%" height="100%" fill="white"/>']
        self.text(35,48,title,32,bold=True)
    def text(self,x,y,t,size=25,color=INK,bold=False,anchor='start'):
        self.parts.append(f'<text x="{x}" y="{y}" font-family="Arial,sans-serif" font-size="{size}" fill="{color}" font-weight="{700 if bold else 400}" text-anchor="{anchor}">{esc(t)}</text>')
    def lines(self,x,y,t,width=42,size=24,color=INK):
        ls=[]
        for paragraph in str(t).split('\n'):ls.extend(textwrap.wrap(paragraph,width=width) or [''])
        for i,line in enumerate(ls):self.text(x,y+i*(size+7),line,size,color)
        return len(ls)*(size+7)
    def box(self,x,y,w,h,title,body='',warn=False):
        color=RED if warn else BLUE
        self.parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="12" fill="{("#fff4f3" if warn else PALE)}" stroke="{color}" stroke-width="2"/>')
        self.text(x+18,y+35,title,26,color,True)
        if body:
            size=23
            while size>18 and len(wrap_pixels(body,w-36,size))*(size+6)>h-64:size-=1
            lines=wrap_pixels(body,w-36,size)
            assert len(lines)*(size+6)<=h-58,(title,lines)
            for i,line in enumerate(lines):self.text(x+18,y+69+i*(size+6),line,size)
    def line(self,points,label=None,warn=False,dash=False):
        points=list(points);color=RED if warn else BLUE
        ds=' stroke-dasharray="9 6"' if dash else ''
        self.parts.append(f'<polyline points="{" ".join(f"{x},{y}" for x,y in points)}" fill="none" stroke="{color}" stroke-width="2.5" marker-end="url(#{"r" if warn else "a"})"{ds}/>')
        if label:
            x,y=points[len(points)//2];self.text(x+10,y-12,label,22,color)
    def save(self,name):
        file=DIAG/(name+'.svg');file.write_text('\n'.join(self.parts)+"</svg>",encoding='utf-8');FIGURES[name]=file

def sequence(name,title,actors,events):
    w=1400; top=145; row=110; h=top+row*len(events)+55
    s=SVG(title,w,h); centers=[165+i*(w-330)/(len(actors)-1) for i in range(len(actors))]
    for i,a in enumerate(actors):
        x=centers[i];bw=min(265,(w-330)/(len(actors)-1)-15)
        s.parts.append(f'<rect x="{x-bw/2}" y="74" width="{bw}" height="65" rx="8" fill="{PALE}" stroke="{BLUE}"/>')
        for j,line in enumerate(textwrap.wrap(a,width=18)):
            s.text(x,100+j*26,line,23,bold=True,anchor='middle')
        s.parts.append(f'<line x1="{x}" y1="142" x2="{x}" y2="{h-25}" stroke="#b9c7d6" stroke-width="2" stroke-dasharray="7 7"/>')
    for n,(a,b,label,*flag) in enumerate(events):
        y=top+n*row;warn=bool(flag and flag[0]);color=RED if warn else INK
        # Each explanation has its own full-width row above its message arrow.
        s.parts.append(f'<rect x="28" y="{y+4}" width="1344" height="62" fill="white" opacity=".96"/>')
        for j,line in enumerate(wrap_pixels(f'{n+1:02d}  {label}',1305,24)):
            s.text(40,y+29+j*29,line,24,color)
        if a==b:s.line([(centers[a],y+80),(centers[a]+45,y+80),(centers[a]+45,y+97),(centers[a],y+97)],warn=warn)
        else:s.line([(centers[a],y+85),(centers[b],y+85)],warn=warn,dash=warn)
    s.save(name)

def build_diagrams():
    s=SVG('EcoRoute  Arquitectura del sistema en ejecución',1400,1000)
    s.box(40,95,395,180,'Panel web','Dispatcher y jefatura\nReact 19 y TypeScript\nNginx  localhost 3000')
    s.box(965,95,395,180,'Aplicación de campo','Conductor\nFlutter y BLoC\nEmulador Android')
    s.box(505,95,390,180,'Keycloak 23','OIDC y JWT con roles\nWeb PKCE  móvil password grant\nH2 propia  localhost 8080')
    s.line([(435,150),(505,150)]);s.line([(965,150),(895,150)])
    s.box(280,345,840,180,'Backend único Java 17','Spring Boot 3.5.8  WebFlux y Reactor\nControladores REST  seguridad JWT  WebSocket GPS\nCasos de uso  puertos  adaptadores  R2DBC\nAPI 8081  no son microservicios de negocio')
    s.line([(235,275),(235,300),(420,300),(420,345)])
    s.line([(1155,275),(1155,300),(980,300),(980,345)])
    s.line([(700,275),(700,345)],label='Validación JWT')
    s.box(35,650,305,220,'PostgreSQL 14','33 tablas  8 migraciones\nPedidos y rutas\nPOD y GPS  compras\nPuerto host 5433')
    s.box(375,650,305,220,'OSRM Perú','MLD  trip y route\nOrden y polilínea\nFallback haversine\nPuerto 5000')
    s.box(715,650,305,220,'LocalStack 3.8','S3  foto y firma\nSQS  eventos de entrega\nSin consumidor localizado\nPuerto 4566',True)
    s.box(1055,650,305,220,'Redis 7','Cliente y health check\nSin caché de negocio\nverificada\nPuerto interno 6379')
    for x in [188,528,868,1208]:s.line([(700,525),(700,585),(x,585),(x,650)])
    s.lines(40,900,'Fuera del backend: teselas abiertas y geocodificación Nominatim en navegador.\nWebSocket GPS directo a 8081; el proxy Nginx actual no publica /ws.',95,25)
    s.save('01_arquitectura_completa')
    s=SVG('Backend  Capas y dependencias reales',1400,900)
    s.box(40,110,1320,130,'Adaptadores de entrada','REST Controllers  DTO y validación  SecurityConfig  GpsWebSocketHandler')
    s.box(40,320,630,180,'Aplicación','OrderUseCasesImpl  RouteUseCasesImpl\nCreateDeliveryProofUseCaseImpl\nRoutePlannerService  NotificationService')
    s.box(730,320,630,180,'Dominio','Order  Route  DeliveryProof  GpsLocation\nPuertos de entrada de casos de uso\nPuertos de salida OrderRepository y RoutingPort')
    s.line([(360,240),(360,320)]);s.line([(1035,240),(1035,320)]);s.line([(670,410),(730,410)])
    s.box(40,610,1320,150,'Adaptadores de salida','R2DBC repositories y mappers  PostgreSQL\nOsrmRoutingAdapter y WebClient  S3Service  SDK AWS SQS')
    s.line([(360,500),(360,610)]);s.line([(1040,610),(1040,500)],label='Implementa puertos')
    s.lines(45,822,'Matiz: DashboardService depende de Spring Data y DTO de infraestructura.\nLa separación hexagonal organiza el proyecto, pero todavía no es estricta en analítica.',98,25,RED)
    s.save('02_capas_backend')
    sequence('03_acceso_web','Acceso web  Autorización con PKCE',['Usuario y React','Keycloak','Axios y API','PostgreSQL'],[
        (0,0,'Login genera code_verifier y challenge; guarda el verificador en sessionStorage'),
        (0,1,'Redirección OIDC con code_challenge S256; el formulario pertenece a Keycloak'),
        (1,1,'Keycloak valida contraseña en su almacenamiento H2 de desarrollo'),
        (1,0,'Retorno con code; exchangeCodeForTokens envía el code_verifier'),
        (0,0,'AuthContext conserva tokens en localStorage y extrae roles permitidos'),
        (0,2,'GET /orders con Authorization Bearer; Spring valida firma y roles JWT'),
        (2,3,'Consulta R2DBC y devolución de datos autorizados por ruta HTTP'),
        (2,0,'Ante 401 Axios intenta una renovación y repite la petición una sola vez'),
        (0,0,'Brecha: renovar al recibir 401 no equivale a renovar antes de expirar',True)])
    sequence('04_pedido_despacho','Pedido y despacho  Flujo real y punto faltante',['Dispatcher web','API y casos de uso','PostgreSQL','OSRM'],[
        (0,1,'POST /orders sin routeId permite confirmar un pedido en PENDING'),
        (1,2,'createOrder guarda orders con route_id nulo; no registra aprobación comercial'),
        (0,1,'POST /routes selecciona conductor, vehículo y fecha; estado inicial PLANNED'),
        (0,1,'Vincular un pedido ya creado a la ruta no tiene operación REST implementada',True),
        (0,1,'Alternativa actual de prueba: crear otro pedido con routeId; nace ASSIGNED'),
        (0,1,'PATCH /routes/{id}/status?status=IN_PROGRESS inicia el despacho'),
        (1,3,'RoutePlannerService calcula secuencia y geometría por la red vial'),
        (1,2,'Transacción guarda ruta, stop_sequence y todos los pedidos IN_TRANSIT'),
        (1,0,'Devuelve ruta activa; no debe confundirse la alternativa con el flujo To Be')])
    sequence('05_ruteo','Ruteo  Orden de visita y camino por calles',['Flutter','RouteController','RoutePlannerService','OSRM'],[
        (0,1,'GET /routes/{id}/path incluye fromLat y fromLng si hay posición disponible'),
        (1,1,'Selecciona solo pedidos PENDING o IN_TRANSIT con coordenadas'),
        (1,2,'Construye Waypoint por pedido y el punto de salida'),
        (2,3,'OsrmRoutingAdapter llama /trip con roundtrip=true para proponer el orden'),
        (2,3,'Luego llama /route con el orden obtenido y retorno solicitado'),
        (3,2,'Retorna polyline5, distancia en km y duración estimada en minutos'),
        (2,0,'Flutter decodifica, numera marcadores y resalta el camino'),
        (2,0,'Si OSRM falla: haversine sin geometría y duración cero; source explícito',True),
        (0,0,'No resuelve capacidad, ventanas horarias ni tráfico vivo como restricciones')])
    sequence('06_pod','Prueba de entrega  Captura y persistencia',['Conductor y Flutter','API y caso de uso','S3 y SQS','PostgreSQL'],[
        (0,0,'Cámara toma foto; SignaturePad convierte el trazo confirmado en PNG base64'),
        (0,1,'OrderBloc llama repositorio y POST /delivery-proofs con foto y receptor'),
        (0,1,'Defecto: el repositorio no pasa GPS; la petición envía latitude=0 longitude=0',True),
        (1,2,'S3Service sube foto y firma con UUID; devuelve URI s3://bucket/key'),
        (1,3,'CreateDeliveryProofUseCaseImpl guarda POD y cambia pedido a DELIVERED'),
        (1,2,'NotificationService publica el pedido en SQS; no hay consumidor encontrado'),
        (1,3,'Si no quedan pedidos activos, cambia la ruta a COMPLETED'),
        (1,0,'HTTP correcto muestra éxito y la app vuelve a cargar la ruta'),
        (0,1,'Sin red falla el envío; recuperar red no reintenta automáticamente',True),
        (1,2,'La transacción de BD no revierte objetos S3 ni eventos externos ya enviados',True)])
    sequence('07_gps','GPS  Captura persistencia y canal en vivo',['Sensor y Flutter','REST GPS','PostgreSQL','WebSocket y React'],[
        (0,0,'GpsBloc solicita permiso y escucha el stream de Geolocator'),
        (0,1,'GpsRepository envía POST /gps/ping con vehículo, conductor y coordenadas'),
        (1,2,'GpsUseCasesImpl guarda vehicle_gps_history mediante R2DBC'),
        (1,3,'GpsController publica la posición en un Sinks.Many por vehicleId en memoria'),
        (3,3,'El cliente debe abrir /ws/gps y enviar el vehicleId como primer mensaje'),
        (3,3,'Defecto web: Routes.tsx reemplaza onopen e impide enviar la suscripción',True),
        (0,0,'Defecto móvil: vehicleId se asigna igual a driverId, no al vehículo de la ruta',True),
        (3,3,'Riesgo: /ws/** permite acceso sin token; se comprobó recepción anónima',True),
        (3,2,'GET /gps/history/{vehicleId} permite consultar puntos históricos')])
    sequence('08_analitica','Analítica  Desde el dato hasta el gráfico',['Jefatura y React','Controladores','DashboardService','PostgreSQL'],[
        (0,1,'Dashboard solicita conteos, puntualidad, conductores y distritos'),
        (1,2,'DashboardService delega consultas agregadas a SpringDataOrderRepository'),
        (2,3,'COUNT y GROUP BY sobre orders; puntualidad usa delivery_proofs y ventana final'),
        (3,0,'Chart.js dibuja tarjetas, barras y comparación de puntualidad'),
        (0,1,'Reports ejecuta cuatro peticiones juntas usando Promise.all'),
        (1,1,'route-efficiency usa Map.of y falla si el distrito agrupado es nulo',True),
        (1,0,'Un HTTP 500 cancela la actualización conjunta; la pantalla queda en ceros',True),
        (0,0,'CSV solo exporta distrito y cantidad si hay filas; no es un modelo de Power BI'),
        (2,3,'No se encontraron vistas analíticas ni módulo completo de tendencias y metas',True)])
    s=SVG('Despliegue local  Contenedores y persistencia',1400,925)
    s.box(35,105,415,180,'Máquina Windows','Web localhost 3000\nKeycloak localhost 8080\nBackend localhost 8081\nAndroid usa 10.0.2.2')
    s.box(520,105,845,180,'Red de Docker Compose','ecoroute-web-admin Nginx 80  /api hacia backend\necoroute-backend 8081  aplicaciones Spring\npostgres 5432  keycloak 8080  osrm 5000\nlocalstack 4566  redis 6379')
    s.line([(450,190),(520,190)])
    s.box(35,420,415,185,'Volúmenes persistentes','postgres_data  negocio\nkeycloak_data  H2 identidad\nredis_data\nOSRM  extracto de Perú')
    s.box(520,420,400,185,'Arranque','PostgreSQL health check\nFlyway valida V1 a V8\nBackend crea conexiones\nWeb sirve el bundle')
    s.box(960,420,405,185,'Brecha operativa','Init de LocalStack falló\nS3 y SQS creados a mano\nSin volumen de objetos\nRecrear puede perder POD',True)
    s.line([(940,285),(940,350),(720,350),(720,420)])
    s.line([(725,605),(725,670),(250,670),(250,605)])
    s.lines(40,755,'Entorno de desarrollo: HTTP local, Keycloak start-dev y H2.\nNo acredita TLS, recuperación, carga concurrente, alta disponibilidad ni publicación móvil.\nDocker build usa -x test: una imagen compilada no certifica las pruebas automatizadas.',92,25)
    s.save('09_despliegue')
    sequence('10_historial','Historial móvil  Lectura y alcance actual',['Conductor','HistoryPage y BLoC','HistoryRemoteDataSource','API y PostgreSQL'],[
        (0,1,'Abre Historial desde la barra inferior de navegación'),
        (1,2,'El evento de carga ejecuta el repositorio de historial'),
        (2,3,'GET /orders obtiene pedidos sin filtro por conductor ni fecha'),
        (2,2,'Filtra DELIVERED y FAILED en el dispositivo y crea DeliveryHistoryModel'),
        (2,1,'Entrega tarjetas de resultado con fecha y estado'),
        (1,0,'Aparecen registros antiguos y de otras rutas en la prueba'),
        (2,1,'Si ocurre un error devuelve una lista vacía en vez de diferenciar el fallo',True),
        (3,3,'La propiedad de los datos debe verificarse en backend y no solo en pantalla',True)])
    sequence('11_sesion_movil','Aplicación móvil  Capas y sesión',['LoginPage','AuthBloc y caso de uso','Repositorio y Dio','Keycloak y API'],[
        (0,1,'Envía LoginRequested con usuario y contraseña'),
        (1,2,'LoginUseCase llama AuthRepositoryImpl y AuthRemoteDataSource'),
        (2,3,'POST token usa password grant del cliente mobile-app'),
        (3,2,'Retorna JWT; se intenta resolver driverId con claim y consulta de conductores'),
        (2,2,'AuthLocalDataSource guarda usuario y token en flutter_secure_storage'),
        (2,1,'AuthBloc emite autenticado y GoRouter permite Mis rutas'),
        (2,3,'AuthInterceptor agrega Bearer a las llamadas de Dio'),
        (2,2,'No hay renovación móvil ni validación de expiración al recuperar el usuario',True),
        (2,2,'Logout limpia almacenamiento local; no revoca la sesión de Keycloak',True)])

# Explicit audit decisions. V=execution, P=partial, C=code, N=absent, T=not tested.
RF={
1:('V','Web real con PKCE y token legítimo; S01 y prueba web.'),2:('V','Login real en el emulador; S04 y initial-DRIVER.'),3:('P','Axios renueva al recibir 401, no antes del vencimiento.'),4:('C','flutter_secure_storage guarda la sesión; no se ensayó restauración tras reinicio.'),5:('P','Web redirige al logout OIDC; móvil solo elimina la sesión local.'),6:('P','T03 y T04 bloquean rutas; T26 demuestra falta de aislamiento por conductor.'),7:('P','UserController escribe auth_user_role local, sin sincronizar roles JWT de Keycloak.'),8:('C','RateLimitFilter limita 120 por IP en memoria; no se forzó 429.'),
9:('V','Se registró conductor sintético en T07; restricciones revisadas en DTO.'),10:('P','Listado visible; paginación en navegador, no en consulta de servidor.'),11:('C','PUT de conductor y formulario localizados; edición no ejecutada.'),12:('C','DELETE existe; no se eliminaron recursos del usuario.'),13:('V','Se registró vehículo sintético en T07.'),14:('P','Listado web visible; pagina datos ya descargados.'),15:('C','PUT de vehículo localizado; no se modificó flota existente.'),16:('C','DELETE existe; sin prueba destructiva.'),
17:('P','T06 y alta web crean PENDING; T05 rechaza teléfono inválido; control comercial es manual.'),18:('P','Geocodificación por botón Nominatim; coordenadas predeterminadas no bloquean registro.'),19:('P','Alta con routeId funciona T11; falta asignar después un PENDING existente, T10.'),20:('P','Vista previa web; T19 importó una fila válida y omitió otra sin informe de errores.'),21:('V','Búsqueda AUDIT y estados visibles en S02.'),22:('V','Detalle real de pedido 8 abierto en emulador, S05.'),23:('P','T25 cambia a FAILED; backend no valida transiciones ni conserva reason.'),24:('C','DELETE de pedido implementado; no se ejecutó para preservar datos.'),25:('P','Existe URI S3; no se encontró endpoint firmado ni vista web completa de POD recuperado.'),
26:('V','T08 creó ruta 3 con conductor, vehículo y fecha.'),27:('P','T09 detecta conflicto pero devuelve 500; sin garantía concurrente en BD.'),28:('P','API filtra fecha; web usa día UTC y puede ocultar rutas del día en Lima.'),29:('P','T13 despacha; T27 también permite ruta vacía.'),30:('V','T13 y SQL confirman cascada a IN_TRANSIT.'),31:('P','T31 cerró ruta 3; criterios terminales difieren entre POD y actualización manual.'),32:('C','DELETE de ruta implementado; eliminación no ensayada.'),33:('P','T14 OSRM ordena y usa calles; no incorpora capacidad, tráfico ni ventanas como restricciones.'),34:('V','SQL confirma geometría, 12.62 km y 21.5 min para ruta 3.'),35:('V','SQL confirma stop_sequence 1, 2 y 3 en pedidos de ruta 3.'),36:('P','API admite posición; clave móvil de recálculo no cambia con cada coordenada.'),37:('V','X06 con OSRM detenido devuelve source=haversine, sin geometría y duración cero.'),
38:('P','S04 muestra hoy y filtra driverId en cliente; API permite consultar recursos ajenos.'),39:('V','S04 muestra polilínea por calles y marcadores numerados.'),40:('V','S04 muestra 6.4 km y 9 min estimados para pedidos activos.'),41:('P','Lista visible; contador de pendientes usa todos los pedidos y puede incluir terminales.'),42:('V','S05 presenta datos y dirección del pedido móvil 8.'),43:('P','Entrega móvil confirmada; API acepta transiciones inválidas.'),44:('P','Detalle marca terminal como solo lectura; T24 permite revertir DELIVERED por API.'),45:('P','S06 muestra historial; consulta global y sin filtro de fecha o conductor.'),46:('C','Perfil y licencia están en código; no se validó vigencia documental.'),47:('C','Tema oscuro implementado; contraste completo no ensayado.'),
48:('V','Cámara abierta y foto real del escenario virtual capturada.'),49:('P','UI exige foto al entregar; T23 acepta POD vacío por API.'),50:('P','Firma capturada y guardada en S3; su obligatoriedad no se valida en backend.'),51:('P','Receptor y DNI enviados; servidor permite datos vacíos.'),52:('P','API acepta GPS; móvil persistió 0,0 y fecha enviada por cliente.'),53:('V','S3 contiene foto y firma para pedidos 7 y 8 después de crear el bucket.'),54:('V','T20 y entrega móvil cambiaron pedido a DELIVERED y guardaron POD.'),55:('P','T21 devuelve PDF válido, pero textual, sin fotos incrustadas ni firma criptográfica.'),
56:('P','GpsBloc emite por Geolocator; vehicleId=driverId; no hubo muestras móviles verificadas.'),57:('V','T15 guarda y consulta dos muestras GPS por vehículo 2.'),58:('P','T16 recibe WS; web no envía suscripción y canal carece de autenticación.'),59:('P','Mapa web y progreso visibles; suscripción rota y trayecto planeado rectilíneo.'),60:('P','Se publicó evento SQS; no hay consumidor ni notificación final persistida.'),61:('P','API y panel consultan no leídas; sin eventos de negocio que alimenten la tabla.'),62:('C','Endpoint de lectura implementado; sin notificación sintética creada para probarlo.'),
63:('V','Conteos 200 y tarjetas visibles en S07.'),64:('P','Calcula POD contra ventana final, sin aplicar el intervalo de fecha de otros gráficos.'),65:('P','Gráfico por conductor responde; no comparte todos los filtros de fecha.'),66:('P','Agrupación por distrito responde; hay valores nulos y no aplica todos los filtros.'),67:('P','Resumen de flota calculado en web; fecha UTC y datos completos en cliente.'),68:('P','T29 detecta filtro de fecha inconsistente entre conteos y puntualidad.'),69:('P','Dos servicios responden; route-efficiency falla 500 por distrito nulo.'),70:('P','Exportador existe; prueba web no produce descarga debido al reporte fallido.'),
71:('V','T28 registró compra sintética de importe 1; no es una transacción financiera.'),72:('P','Listado web visible; paginación de cliente.'),73:('C','Actualización de compra implementada; no se alteraron compras existentes.'),74:('C','Eliminación implementada; prueba no ejecutada.'),75:('N','No se encontró módulo integrado que calcule IID, CHR y TDE por periodo.'),76:('N','No se localizaron series temporales analíticas de estos indicadores.'),77:('N','No hay gestión persistida de metas de indicadores.'),78:('P','La UI solicita motivo, pero no hay filas de historial en la prueba T25.'),79:('N','No hay agregación operativa de causas de no entrega.'),80:('N','route-efficiency no calcula km por entrega ni tiempo por parada.'),81:('N','No se localizaron vistas analíticas de solo lectura para Power BI.')}

RNF={
1:('P','X03 p95 20 ms y X04 p95 17 ms; 15 consultas secuenciales locales por recurso, no carga concurrente ni reportes sanos.'),2:('T','Hay latencia observada del recorrido, pero no serie p95 suficiente para certificar 200 ms.'),3:('P','Token conductor 56 ms; una autenticación admin 829 ms. No se certifica umbral general de 500 ms.'),4:('C','Límite de memoria configurado; no se envió un archivo de 10 MB y base64 añade sobrecarga.'),5:('C','120 por IP en memoria; no se realizó prueba de saturación.'),6:('P','WebClient acota llamadas OSRM; dos llamadas secuenciales pueden superar 5 s totales.'),7:('C','Dio declara timeout de conexión y recepción; cada fase tiene su límite.'),8:('C','Índices definidos en migraciones; no se ejecutó EXPLAIN con volumen representativo.'),9:('P','JWT sin sesión de negocio, pero Sinks WS y rate limit están en memoria por instancia.'),10:('P','X08 no obtiene OpenAPI; logs muestran NoSuchMethodError de ControllerAdviceBean.'),11:('P','T16 recibe datos, pero la integración web está rota y no exige token.'),12:('V','S3 local guardó archivos de prueba después de inicialización manual.'),13:('V','SQS local recibió mensajes; no equivale a notificación consumida.'),14:('P','T19 valida caso básico, no CSV con comillas ni reporte de filas inválidas.'),15:('P','PDF válido; CSV no descargó y PDF no incluye imagen ni firma.'),16:('C','OSM y OSRM sin clave; los servidores públicos de teselas y geocodificación siguen siendo externos.'),17:('P','Interfaz principal en español; estados y mensajes de servicios externos contienen inglés.'),18:('P','Validación de teléfono concreta; varios fallos móviles se transforman en mensaje genérico de red.'),19:('P','Paginación en componentes web; descarga previa de todo el listado.'),20:('P','Controles visibles, pero fecha UTC y filtros analíticos no uniformes.'),21:('V','Capturas muestran etiquetas y colores de estado.'),22:('C','ThemeBloc y control de perfil; no se comprobó contraste de todas las vistas.'),23:('V','Mapa móvil con secuencia visible; no equivale al mapa de despacho web.'),24:('P','Solo lectura en UI; T24 demuestra modificación por API.'),25:('V','X06 confirma respuesta degradada con origen explícito al apagar OSRM.'),26:('V','T01 UP con R2DBC y Redis; una lectura no demuestra disponibilidad mensual.'),27:('C','Compose usa healthcheck de PostgreSQL y condición service_healthy.'),28:('C','Políticas restart declaradas; no se provocó caída del backend.'),29:('P','SQL verifica 27 FK; faltan garantías sobre estados, asignación y validez del POD.'),30:('V','BD confirma V1 a V8 con success=true.'),31:('P','@Transactional y cascada verificadas en éxito; no se inyectó fallo parcial de BD.'),32:('P','Autenticación delegada; entorno local conserva credenciales demo y Keycloak H2.'),33:('V','Login web real muestra code_challenge_method=S256.'),34:('P','DRIVER no ve reportes, pero sí pedidos ajenos y WS anónimo.'),35:('V','T03 401 sin token; X02 401 token ficticio; WS es excepción insegura.'),36:('C','Keystore mediante secure_storage; no se hizo análisis forense del dispositivo.'),37:('P','DTO de pedido valida; POD y transiciones de estado permiten entradas inválidas.'),38:('P','Métricas autenticadas verificadas con ADMIN; no se acreditó exclusión de otros roles.'),39:('P','T23 acepta evidencia vacía; móvil GPS 0,0; sin inmutabilidad demostrada.'),40:('P','Prod exige variables; Compose y scripts locales incluyen credenciales de demostración.'),41:('P','X01 token 28800 s; sin renovación móvil ni validación al restaurar sesión.'),42:('P','Puertos y adaptadores presentes; DashboardService acopla DTO y repositorios de infraestructura.'),43:('C','Seis features móviles con capas data, domain y presentation.'),44:('C','Spring y GetIt inyectan dependencias; no se ensayaron dobles de todos los adaptadores.'),45:('V','V8 aplicada y geometría persistida; no se ensayó recuperación de backup.'),46:('P','X07 expone 62 nombres en /actuator/metrics; no confirma la cifra histórica de 69 en Prometheus.'),47:('C','Perfiles local/docker DEBUG y prod INFO revisados.'),48:('P','gradlew test no compila: tres errores en OrderUseCasesImplTest.'),49:('P','ESLint 34 errores y 4 advertencias; Flutter analyze 52 observaciones informativas.'),50:('P','Compose levanta contenedores; init de LocalStack falló y requirió crear bucket y cola.'),51:('C','application.yml, application-docker.yml y application-prod.yml presentes.'),52:('P','Backend parametrizado; Flutter fija 10.0.2.2 para API y Keycloak.'),53:('P','Un código Flutter; Android compilado e instalado, iOS no verificado.'),54:('V','OSRM Perú local responde; fallback X06 y posterior restauración comprobados.')}
STATUS={'V':'Verificado en ejecución','P':'Parcial o no conforme','C':'Revisado en código','N':'No implementado localizado','T':'No certificado'}

FINDINGS=[
('H01','Alta','P01 P02','RF-19 26 29','No existe asignación posterior del pedido','T10 devuelve 404; se requirió crear pedidos nuevos con routeId.','Implementar asignar y reasignar pedidos existentes con validación y auditoría; un PENDING debe llegar a ASSIGNED sin duplicarse.'),
('H02','Alta','P03 P08','RF-06 38 45','Un conductor consulta pedidos de otro','T26 HTTP 200; historial descarga GET /orders global.','Aplicar propiedad por usuario y ruta en backend; probar conductor A contra recursos B con 403 o 404.'),
('H03','Alta','P03','RF-23 44 49 50 51','Estados y POD carecen de controles obligatorios','T22 entrega sin POD; T23 POD vacío; T24 DELIVERED vuelve a PENDING.','Validar máquina de estados y evidencia en servidor; devolver 4xx y conservar estado previo ante datos inválidos.'),
('H04','Alta','P03','RF-52','El POD móvil guarda GPS 0,0','SQL del pedido 8 y cadena OrderRepositoryImpl sin lat/lng.','Transmitir posición con precisión y fecha de captura; rechazar coordenadas de reserva y verificar en BD ubicación real.'),
('H05','Alta','P03 P04','RF-54 y excepción de P03','No existe cola offline conectada al flujo','Sin red aparece fallo; recuperar conexión no actualiza; reintento manual sí funciona.','Persistir borrador y outbox idempotente; probar reinicio sin red y sincronización única al recuperar conexión.'),
('H06','Alta','P07 P08','RF-58 59','WebSocket GPS anónimo y sin suscripción web','T16 recibe sin JWT; prueba web registra cero mensajes enviados.','Autenticar handshake y autorizar vehículo; conservar onopen de suscripción y verificar recepción en pantalla.'),
('H07','Alta','P07','RF-56','Vehículo incorrecto en seguimiento móvil','routes_page.dart asigna vehicleId igual a driverId; ruta 3 usa vehículo 2.','Tomar vehicleId de la ruta activa; comprobar ping con conductor 1 y vehículo 2 desde el emulador.'),
('H08','Alta','P05','RF-69 70','Reportes falla y oculta todos los resultados','X09 HTTP 500 por Map.of con distrito nulo; exportación web sin descarga.','Aceptar distrito sin dato, aislar fallos por gráfico y probar CSV con filas reales y fechas.'),
('H09','Alta','P05','RF-64 65 66 68 75','Los filtros no definen una misma población','T29 periodo 2099 sin pedidos pero con entregas puntuales históricas.','Fijar zona Lima y fecha de negocio; compartir filtros en SQL de todos los indicadores.'),
('H10','Media','P04','RF-78 79','Se pierde el motivo de no entrega','T25 envía reason; order_status_history sin filas.','Persistir transición, actor, fecha y catálogo de causas; verificar historial inalterable de correcciones.'),
('H11','Alta','P03 y transversal','RNF-12 13 50','El arranque no prepara el almacenamiento','POD inicial 500 por bucket ausente; init-aws.sh no ejecutado.','Corregir inicializador y persistencia de LocalStack; un entorno limpio debe crear S3 y SQS sin intervención.'),
('H12','Media','P03','RF-60 61','El evento SQS no llega al panel de notificaciones','Mensajes en cola; notifications sin registros de prueba; sin consumidor localizado.','Consumidor idempotente que escriba notificación y maneje reintentos y cola de errores.'),
('H13','Media','P02','RF-28','Hoy en web es la fecha UTC','Captura hallazgo_fecha_web y toISOString; móvil usa fecha local.','Usar fecha de operación America/Lima; verificar entre 19 y 24 h y al cambiar de día.'),
('H14','Media','P02 P07','RF-33 36 39 59','El recorrido web no usa la geometría OSRM','Routes.tsx conecta puntos con heurística euclidiana; móvil sí dibuja calles.','Consumir la misma secuencia y polilínea de backend en ambos clientes y señalar el fallback.'),
('H15','Media','P02','RF-27 29 31','Despacho permite estados y recursos incoherentes','T27 inicia ruta vacía; conflicto de recurso da 500; terminales sin reglas fuertes.','Validar ruta con pedidos y recurso habilitado; 409 en conflicto; transiciones seguras e idempotentes.'),
('H16','Media','P08','RF-03 05 07','Sesión y roles no son uniformes','Web renueva tras 401; móvil no renueva; rol local no modifica Keycloak.','Definir fuente de autoridad Keycloak, refresco móvil, revocación y pruebas de cambio efectivo de rol.'),
('H17','Media','Transversal','RNF-10 48 49','La calidad automatizada no está aprobada','Tres errores Java; ESLint 34 errores; OpenAPI NoSuchMethodError.','Actualizar pruebas al modelo y springdoc compatible; CI con test, build, lint y contrato OpenAPI.'),
('H18','Media','P05','RF-75 a RF-81','La inteligencia de negocios todavía no está completa','No se localizó módulo de metas, tendencias, causas ni vistas de Power BI integrado.','Definir fichas de IID, CHR y TDE con MICOTRANS, implementar vistas, series, metas y validación con datos de empresa.'),
('H19','Media','P03','RF-25 55 RNF-39','La evidencia no constituye certificado inmutable','PDF solo texto; URI s3 no es URL pública; timestamp cliente; no hay sello criptográfico.','Recuperar objetos con autorización y URL temporal, integrar evidencia y política de integridad y retención.'),
('H20','Media','Transversal','RNF-09 40 52','El entorno local no equivale a producción','H2 start-dev, HTTP, WS y rate limit en memoria; direcciones móviles de emulador.','Configurar TLS y secretos, identidad con BD adecuada, mensajería distribuida y configuración por entorno.'),
]

SCREENS=[
{'id':'S01','title':'Acceso al panel web','actor':'Administrador o dispatcher','bpmn':'P08','rf':'RF-01 RF-03 RF-05 RF-06','image':'S01_login_web','support':'S01b_keycloak','mobile':False,'diagram':'03_acceso_web',
'purpose':'Identificar al trabajador y habilitar únicamente el acceso acorde con sus funciones. El proveedor de identidad autentica; EcoRoute consume la identidad y decide permisos sobre su API.',
'design':'La portada separa identidad del producto y acción de acceso. El botón lleva al formulario externo de Keycloak; no envía la contraseña al backend de logística. El estado autenticado habilita el menú lateral y las rutas protegidas. Ocultar un menú no reemplaza una autorización de servidor.',
'steps':['Abrir localhost:3000/login y elegir el acceso. redirectToKeycloakLogin crea el verificador PKCE y la URL OIDC del realm ecoroute.','Completar el formulario de Keycloak. El navegador retorna al callback con un código, no con la contraseña.','exchangeCodeForTokens intercambia el código usando code_verifier. AuthContext conserva tokens y roles y cambia el estado de React.','Abrir Pedidos. api.ts agrega Authorization Bearer; Spring verifica el JWT y consulta los datos por R2DBC.','Si una petición devuelve 401, el interceptor intenta refresh una sola vez. Si falla, debe finalizar la sesión.'],
'front':'web-admin/src/context/AuthContext.tsx; services/keycloak.ts y api.ts; pages/Login.tsx; components/ProtectedRoute.tsx. React usa Context para sesión y react-router para navegación. El control PKCE usa sessionStorage; los tokens del panel quedan en localStorage.',
'back':'SecurityConfig configura resource server y convierte realm_access.roles. Keycloak conserva usuarios, credenciales y sesiones en H2 de desarrollo. Las tablas auth_* del backend no son la base de Keycloak y editar sus roles no cambia automáticamente los roles del JWT.',
'contract':'GET /realms/ecoroute/protocol/openid-connect/auth con response_type=code y challenge S256; POST /protocol/openid-connect/token con code y verifier. Las operaciones logísticas usan Bearer JWT, nunca la contraseña.',
'result':'Login real aprobado; /orders sin token respondió 401 y /reports con DRIVER respondió 403. La renovación previa al vencimiento y la revocación móvil no quedaron cumplidas. No se imprimen tokens ni credenciales en el documento.',
'speech':'El acceso tiene dos responsabilidades distintas. Keycloak comprueba quién es el usuario y entrega un token firmado. Nuestro backend valida ese token antes de ejecutar el caso de uso. React usa el rol para organizar la interfaz, pero la restricción efectiva debe estar también en el servidor.'},
{'id':'S02','title':'Registrar y consultar pedidos','actor':'Dispatcher','bpmn':'P01 y entrada a P02','rf':'RF-17 a RF-25','image':'S02_pedidos','support':'S02b_alta_pedido','mobile':False,'diagram':'04_pedido_despacho',
'purpose':'Convertir una solicitud validada del cliente en un pedido de reparto identificable. Confirmar el servicio no debe depender de tener una ruta previamente creada; la planificación corresponde a la siguiente fase del negocio.',
'design':'La tabla expone tracking, cliente, destino y estado. La búsqueda y el filtro reducen la carga visual; la paginación es de cliente. El modal diferencia datos obligatorios, dirección, contacto, ventanas y ruta opcional. El mapa y la geocodificación ayudan a revisar la dirección, pero no prueban su exactitud.',
'steps':['En P01 el dispatcher comprueba cobertura, condiciones y aceptación del cliente; estos acuerdos son tareas humanas, no aprobaciones registradas automáticamente.','Abrir Nuevo pedido e ingresar tracking único, destinatario, dirección y teléfono válido. La ruta debe permanecer vacía mientras el pedido no esté planificado.','Consultar la dirección mediante Nominatim y revisar sus coordenadas. No aceptar automáticamente el punto predeterminado como ubicación confirmada.','Guardar. Orders.tsx ejecuta createOrder; Axios envía JSON a POST /orders. CreateOrderRequest valida y OrderRestMapper construye el dominio.','OrderUseCasesImpl asigna PENDING si routeId es nulo y guarda con OrderRepository. React vuelve a obtener el listado y muestra el nuevo registro.','Continuar a P02 para agrupar y asignar. El paso de vincular el pedido ya existente no está implementado; la auditoría no lo presenta como si funcionara.'],
'front':'web-admin/src/pages/Orders.tsx, services/orderService.ts y api.ts. useState mantiene formulario, filtros y página. fetchData combina pedidos y rutas con Promise.all. El CSV tiene vista previa y un parser simple de columnas, no un importador robusto de cualquier variante de CSV.',
'back':'OrderController.createOrder usa @Valid; OrderUseCasesImpl.createOrder decide estado y RepositoryAdapter persiste la entidad orders. orders.route_id es nullable. No existe un evento de confirmación comercial ni un endpoint de asignación posterior al plan. La versión con routeId crea ASSIGNED inmediatamente.',
'contract':'POST /orders {trackingNumber, recipientName, deliveryAddress, recipientPhone, latitude, longitude, routeId:null, ...}. La respuesta de alta web fue id=14 y status=PENDING. PATCH /orders/{id}/route devuelve 404 en esta versión.',
'result':'Alta web y API aprobadas y teléfono inválido rechazado. CSV omitió una fila inválida sin informar. La actualización de estado no conserva reason ni impide revertir terminales; la baja no se ejecutó para preservar datos.',
'speech':'Esta pantalla pertenece al inicio del flujo de negocio. Primero confirmamos un pedido y lo dejamos en cartera. La ruta es una decisión operativa posterior. El desarrollo permite el registro sin ruta, pero todavía falta la operación que vincula ese mismo pedido al plan sin volver a crearlo.'},
{'id':'S03','title':'Planificar y despachar la jornada','actor':'Dispatcher con autorización de jefatura','bpmn':'P02 y P07','rf':'RF-26 a RF-37 RF-58 RF-59','image':'S03_despacho','support':None,'mobile':False,'diagram':'04_pedido_despacho',
'purpose':'Disponer conductor, vehículo y pedidos para ejecutar una jornada viable. La jefatura autoriza el plan y el conductor revisa carga y recursos; esos controles del BPMN no son equivalentes a pulsar Iniciar.',
'design':'El panel combina lista de rutas, filtros de fecha, selección de recurso, estados, avance y mapa. La acción de despacho cambia el estado operativo. La línea planeada del mapa web une puntos; la línea histórica representa GPS. Deben distinguirse para no prometer una geometría vial que la web todavía no consume.',
'steps':['Validar que conductor y vehículo estén habilitados en P06 y que el operador tenga acceso en P08.','Crear una ruta con fecha de negocio, conductor y vehículo. POST /routes devuelve PLANNED.','Vincular pedidos confirmados y revisar carga, ventana y secuencia; la asignación posterior es la brecha H01. La prueba usó pedidos nuevos con routeId como alternativa técnica, no como proceso aprobado.','Autorizar la salida y revisar hoja y carga. La aplicación no registra una aprobación separada de jefatura ni una lista de chequeo del conductor.','Elegir Iniciar. PATCH /routes/{id}/status?status=IN_PROGRESS llama RouteUseCasesImpl. Calcula recorrido, guarda geometría y stop_sequence, y cambia pedidos a IN_TRANSIT en una transacción.','Supervisar avance y GPS. El canal requiere corrección de suscripción y autenticación antes de declararse seguimiento en vivo confiable.'],
'front':'Routes.tsx utiliza estado React, servicios de rutas, conductores, vehículos y GPS, y un componente Leaflet. connectGpsWebSocket configura la suscripción, pero Routes.tsx sustituye ws.onopen y la elimina. El filtro Hoy se calcula con toISOString, que usa UTC.',
'back':'RouteUseCasesImpl.createRoute detecta recurso ocupado por fecha y estado. updateRouteStatus invoca RoutePlannerService, persiste ruta, duración, distancia, geometría y secuencia de pedidos. El conflicto genera RuntimeException y 500. No se rechaza una ruta vacía y no hay una máquina de estados completa.',
'contract':'POST /routes {driverId:1,vehicleId:2,routeDate:"2026-10-19",estimatedStartTime:...}. PATCH /routes/3/status?status=IN_PROGRESS. GET /routes?date=2026-10-19. GET /routes/3/path para consultar recorrido activo.',
'result':'Ruta 3 creada, despachada y cerrada; SQL conservó 12.62 km, 21.5 min y geometría. Se comprobó conflicto de recurso, pero con 500. La captura se tomó en Histórico para evitar el filtro de fecha UTC, documentado en H13.',
'speech':'El despacho es el cambio de una jornada planificada a una jornada en ejecución. En backend se calcula la secuencia y se actualizan los pedidos dentro de una transacción. La planificación de capacidad, la autorización y la asignación posterior todavía necesitan controles para corresponder por completo al BPMN.'},
{'id':'S04','title':'Consultar la ruta del conductor','actor':'Conductor','bpmn':'P03 y P07','rf':'RF-33 RF-36 a RF-43 RF-56','image':'S04_rutas_movil','support':None,'mobile':True,'diagram':'05_ruteo',
'purpose':'Indicar al conductor las paradas de la jornada y el camino de referencia para atenderlas. La distancia y duración corresponden al plan estimado de paradas activas, no a un compromiso de llegada basado en tráfico en vivo.',
'design':'El mapa ocupa el área principal, con camino resaltado y marcadores numerados. La lista inferior relaciona el orden con destinatario y estado. Las tarjetas de distancia y ETA resumen la carga pendiente. La barra inferior comunica Rutas, Historial y Perfil.',
'steps':['Al entrar, RoutesPage solicita permisos de ubicación y obtiene el driverId autenticado. RouteBloc ejecuta GetDriverRoutesUseCase.','RouteRemoteDataSource consulta /routes?date con la fecha local del dispositivo; el repositorio filtra por conductor en el cliente.','Obtiene pedidos de la primera ruta y solicita /routes/{id}/path con la última posición conocida cuando existe.','El servidor selecciona pedidos PENDING e IN_TRANSIT, genera Waypoints y llama a OSRM. /trip propone el orden y /route devuelve el camino.','RoutePathModel interpreta orderedOrderIds, polyline, distanceKm, durationMin y source. polyline_decoder convierte polyline5 a coordenadas Leaflet.','El conductor toca una parada para ver el detalle. Si OSRM falla, la aplicación debe distinguir la aproximación; no debe mostrar un tiempo cero como ETA confiable.'],
'front':'Flutter StatefulWidget con RouteBloc y GpsBloc; GetIt resuelve repositorios y Dio. flutter_map pinta teselas, PolylineLayer y MarkerLayer. _fetchStreetPath consulta el backend; _optimizeSequence actúa como apoyo local. La clave de recálculo no incluye cada nueva coordenada GPS.',
'back':'RouteController.getRoutePath filtra PENDING e IN_TRANSIT, por lo que excluye ASSIGNED antes de despacho. RoutePlannerService usa un hub fijo de Lima si falta origen y OsrmRoutingAdapter consulta el extracto Perú. El fallback usa vecino cercano y haversine, sin geometría vial.',
'contract':'GET /routes/3/path?fromLat=-12.046374&fromLng=-77.042793. Respuesta: {polyline,orderedOrderIds,distanceKm,durationMin,source:"osrm"}. El endpoint de camino en vivo no guarda la nueva geometría; la persistencia ocurre al despachar.',
'result':'Captura real muestra 6.4 km y 9 min en la parte activa de la ruta. X06, con OSRM apagado, devolvió source=haversine, polyline nula y durationMin=0. No se certificó optimización global de capacidad, tráfico o ventanas.',
'speech':'Aquí se observa el camino por calles para las paradas activas. Flutter no calcula por sí solo la red vial; pide una propuesta al backend y este consulta OSRM. El orden, la polilínea y el resumen se devuelven juntos para que mapa y lista sean coherentes.'},
{'id':'S05','title':'Confirmar la entrega con evidencia','actor':'Conductor y receptor externo','bpmn':'P03 y excepción hacia P04','rf':'RF-42 a RF-44 RF-48 a RF-55','image':'S05a_formulario_entrega','support':'S05e_firma_confirmada','mobile':True,'diagram':'06_pod',
'purpose':'Registrar que el destinatario recibió el pedido y conservar los elementos que respaldan el resultado. Una prueba incompleta o una falla de comunicación debe ir a control de excepción, no interpretarse como entrega certificada.',
'design':'El detalle identifica pedido, dirección y estado. Al elegir Entregado aparecen foto, receptor, DNI y firma. La cámara ocupa una pantalla nativa; la firma usa un lienzo táctil y confirmación explícita. Guardar debe mostrar progreso, éxito o un error que permita continuar sin perder trabajo.',
'steps':['Abrir el pedido 8 desde Mis rutas y revisar destinatario y dirección. Seleccionar Entregado.','Capturar una foto con la cámara del emulador. La imagen mostrada corresponde al escenario virtual de Android, no a un cliente real.','Ingresar nombre de prueba y DNI ficticio; dibujar y confirmar la firma. SignaturePad genera PNG base64 a partir del lienzo.','Guardar activa UpdateOrderStatusEvent en OrderBloc. El caso de uso delega en OrderRepositoryImpl: si hay foto y estado DELIVERED, ejecuta createDeliveryProof.','Dio envía el JSON. El servidor sube foto y firma a S3, guarda la prueba y el estado del pedido y publica un evento en SQS.','Al cerrar la última parada terminal, el servidor marca la ruta COMPLETED. Flutter presenta éxito y recarga pedidos.','En el ensayo sin red el envío falló y no se sincronizó automáticamente al volver. El mismo formulario pudo completarse con reintento manual.'],
'front':'order_detail_page.dart conserva estado, foto y datos del receptor. image_picker utiliza la cámara y reduce calidad a 70. SignaturePad usa GestureDetector, CustomPainter y RepaintBoundary. BLoC separa interacción del caso de uso; Dio transporta el JSON base64. LocalCache existe, pero no participa de esta cadena.',
'back':'DeliveryProofController recibe el DTO sin @Valid. CreateDeliveryProofUseCaseImpl combina subidas S3, escritura de pedido y prueba, publicación SQS y chequeo de cierre. @Transactional cubre PostgreSQL, no revierte los servicios externos. S3Service devuelve URI s3. El PDF se genera con texto.',
'contract':'POST /delivery-proofs {orderId:8,imageUrl:"<base64>",signatureDataUrl:"<base64>",receiverName:"Prueba MOVIL",receiverDni:"00000000",latitude:0,longitude:0,verifiedAt:"...Z"}. Las coordenadas cero son el defecto observado, no el formato de negocio esperado.',
'result':'Entrega real desde móvil aprobada: pedido 8 DELIVERED, prueba 5 con dos objetos S3 y ruta 3 COMPLETED. Sin embargo, la ubicación fue 0,0. T23 aceptó prueba vacía por API y T22 permitió entregar sin POD. No se considera aceptación integral de P03.',
'speech':'La evidencia atraviesa interfaz, BLoC, repositorio, API, almacenamiento de objetos y base de datos. La operación sí se completó durante la prueba, pero el registro de GPS está defectuoso y el servidor todavía no obliga todos los elementos. Por ello mostramos tanto el recorrido exitoso como los controles que deben corregirse.'},
{'id':'S06','title':'Consultar el historial de entregas','actor':'Conductor','bpmn':'Consulta de P03 y soporte de P05','rf':'RF-45 RF-44 RF-06','image':'S06_historial_movil','support':None,'mobile':True,'diagram':'10_historial',
'purpose':'Permitir al conductor comprobar los resultados de su trabajo y apoyar la conciliación de la jornada. El alcance correcto es su propia operación y el periodo solicitado, con una referencia recuperable de evidencia.',
'design':'Las tarjetas separan entregados y fallidos mediante texto y color e incluyen fecha. La pantalla mantiene la misma navegación inferior. Una lista vacía debería distinguir ausencia de entregas de un fallo de carga; el código actual no lo hace.',
'steps':['Elegir Historial desde la barra inferior. HistoryBloc dispara la consulta a su caso de uso y repositorio.','HistoryRemoteDataSource llama GET /orders sin filtros de identidad ni fecha.','El dispositivo retiene estados DELIVERED y FAILED y los convierte a modelos de historial.','La pantalla pinta tarjetas con tracking, fecha y resultado. Se observaron registros antiguos y no exclusivamente del conductor.','Para que la consulta sea apta para el negocio, el servidor debe imponer propietario y periodo, y la UI debe permitir consultarlos con mensajes de error diferenciados.'],
'front':'features/history separa presentation, domain y data. HistoryPage representa el estado de carga y tarjetas. La fuente remota atrapa errores y devuelve [], lo que puede aparentar que no hubo entregas. La protección de acceso debe resolverse fuera del filtro visual.',
'back':'No hay un endpoint específico de historial propio. Se reutiliza OrderController.getAllOrders y R2DBC. Las relaciones orders.route_id y routes.driver_id permiten filtrar correctamente, pero el flujo actual no usa ese control en servidor.',
'contract':'GET /orders con Bearer conductor. El resultado es una lista de pedidos; la selección DELIVERED o FAILED ocurre en el cliente. La auditoría comprobó lectura de un pedido de otro conductor en T26.',
'result':'La pantalla abre y muestra entregas, incluida la operación de prueba. No cumple la restricción de historial propio ni la de jornada. La captura es una observación del sistema, no un reporte de productividad validado.',
'speech':'El historial reutiliza los pedidos del sistema y filtra estados en el teléfono. Esa reutilización simplifica el desarrollo, pero no basta para seguridad ni para gestión diaria: el backend debe limitar qué datos recibe cada conductor y qué periodo se está consultando.'},
{'id':'S07','title':'Evaluar el panel de control logístico','actor':'Jefatura logística','bpmn':'P05','rf':'RF-63 a RF-68','image':'S07_dashboard','support':None,'mobile':False,'diagram':'08_analitica',
'purpose':'Resumir la operación para identificar carga pendiente, entregas y diferencias entre conductores o zonas. El panel actual es un control operativo; no acredita por sí solo el módulo completo de inteligencia de negocios.',
'design':'La jerarquía visual empieza con filtros por conductor y fechas, continúa con cuatro tarjetas de pedidos y termina en gráficos de puntualidad, conductores, distritos y disponibilidad de flota. Chart.js representa agregados ya calculados; no sustituye reglas de negocio ni una población estadística consistente.',
'steps':['La jefatura abre Dashboard con rol autorizado y elige conductor y rango.','Dashboard.tsx dispara seis peticiones en paralelo: estados, flota, puntualidad, conductores, rendimiento y distritos.','DashboardController delega en DashboardService; Spring Data ejecuta agregaciones de orders y uniones con POD o rutas.','React suma estados y actualiza tarjetas; los componentes Charts dibujan barras y comparación de puntualidad.','La disponibilidad de flota se calcula en navegador con recursos y rutas. La fecha usada puede variar por UTC.','Antes de interpretar un KPI se debe comprobar que numerador, denominador y todos los gráficos respondan al mismo periodo y conductor. Esto no está garantizado en la versión actual.'],
'front':'Dashboard.tsx, dashboardService.ts y components/Charts.tsx. useEffect escucha filtros y Promise.all espera seis fuentes. loading representa el estado inicial. El manejo general de error deja información previa sin un indicador de datos incompletos por gráfico.',
'back':'DashboardService usa SpringDataOrderRepository directamente y DTO de infraestructura. getOrderStats filtra created_at por intervalo; consultas de puntualidad comparan verified_at con estimated_delivery_window_end y no aceptan el mismo intervalo. Distritos y flota tienen alcance diferente.',
'contract':'GET /dashboard/orders-by-status?driverId=&startDate=&endDate=; /dashboard/delivery-performance?driverId=; /dashboard/deliveries-by-driver?driverId=; /dashboard/orders-by-district. Flota combina /drivers, /vehicles y /routes.',
'result':'Gráficos y conteos cargan. T29 demuestra que un periodo sin pedidos puede convivir con entregas puntuales históricas. Los datos sintéticos de auditoría no son indicadores reales de MICOTRANS ni evidencia de mejora empresarial.',
'speech':'Las tarjetas se alimentan de consultas agrupadas y los gráficos de componentes Chart.js. Para tomar decisiones no basta con que un gráfico aparezca: todas las consultas deben usar la misma fecha de negocio y el mismo alcance. Esa consistencia es uno de los pendientes que identificamos.'},
{'id':'S08','title':'Consultar y exportar reportes','actor':'Jefatura logística','bpmn':'P05','rf':'RF-69 RF-70 RF-75 a RF-81','image':'S08_reportes','support':None,'mobile':False,'diagram':'08_analitica',
'purpose':'Obtener información consolidada para revisar cumplimiento y preparar análisis posteriores. La pantalla se conserva tal como respondió durante el ensayo, incluidos sus errores; los ceros visibles no significan que no existan operaciones.',
'design':'La página combina filtros, tarjetas, gráfico por conductor, tabla por distrito y botón CSV. La tabla dispone de paginación local. Debe diferenciar explícitamente error de servicio y valor cero, y permitir que un gráfico sano se mantenga aunque otro servicio falle.',
'steps':['Abrir Reportes con usuario administrador. Reports.tsx ejecuta Promise.all de resumen, conductores, eficiencia y catálogo de conductores.','/reports/orders-summary combina conteos y puntualidad; /driver-performance agrupa por ruta y conductor.','/route-efficiency solicita agrupación por distrito, no eficiencia kilométrica. Al aparecer un distrito nulo, Map.of lanza NullPointerException.','La petición devuelve HTTP 500 y se rechaza Promise.all. Ningún setState conjunto se aplica, por lo que las tarjetas quedan en su estado inicial.','Al pulsar CSV, handleExportCsv detecta tabla vacía y retorna sin crear descarga. Si hubiera filas, generaría Blob con Distrito y Cantidad de Pedidos.','Para Power BI se necesita un modelo analítico estable y de solo lectura, además de filtros y fichas de KPI. Es una mejora pendiente, no una integración existente.'],
'front':'Reports.tsx y reportService.ts. Filtros como estado React, gráfico react-chartjs-2 y paginación en cliente. handleExportCsv construye texto, Blob, URL temporal y enlace download. La falla de un servicio bloquea el conjunto de datos.',
'back':'ReportController usa DashboardService. getRouteEfficiency construye Map.of con dto.getName(), que puede ser null. Las fechas de ese método no se usan en la consulta. No hay cálculo real de km por entrega o tiempo por parada ni vistas analíticas implementadas en las migraciones revisadas.',
'contract':'GET /reports/orders-summary; GET /reports/driver-performance; GET /reports/route-efficiency?startDate=&endDate=. El último devolvió 500. La exportación CSV ocurre en navegador, no existe una descarga de reporte analítico completo desde servidor.',
'result':'Acceso a módulo y restricción de rol funcionan; reporte integral y exportación no aprobados. La captura muestra la falla real. RF-75 a RF-81 no se consideran completados por la existencia de tarjetas o un archivo HTML de prototipo.',
'speech':'El reporte actual es un conjunto de consultas operativas. En la prueba falló la consulta por distrito y eso dejó la página en ceros. La mejora debe corregir ese error, separar los estados de carga y diseñar el modelo analítico antes de afirmar que existe inteligencia de negocios completa.'},
]

def create_content():
    add('title',text='EcoRoute arquitectura verificación y flujos del sistema')
    p('Manual técnico y funcional para el proyecto aplicado a Grupo MICOTRANS S A C')
    p('Fecha de verificación 19 de octubre de 2026 · Zona de operación America Lima')
    p('El flujo básico sí pudo completar una entrega desde la aplicación móvil, almacenar foto y firma y cerrar automáticamente una ruta. La aceptación integral del sistema no está aprobada: persisten brechas de asignación, seguridad por conductor, evidencia, funcionamiento sin conexión y analítica. Este manual explica el desarrollo real y separa lo comprobado de los controles aún pendientes del BPMN.')
    p('Destinatarios: equipo de desarrollo, responsable del proyecto y docente. El documento permite recorrer los componentes, reproducir las pruebas y explicar ocho pantallas desde la acción del usuario hasta los datos persistidos.')
    table(['Resultado','Alcance'],[['Flujo de entrega básico','Aprobado después de crear manualmente el bucket S3 y la cola SQS locales.'],['Conformidad completa con requerimientos y BPMN','No aprobada. Los resultados negativos se incluyen con evidencia y criterio de corrección.'],['Datos usados','Pedidos AUDIT, receptor y DNI ficticios, GPS de Lima y foto del escenario virtual del emulador. No representan desempeño real de MICOTRANS.'],['Cambios realizados','Construcción y arranque local, instalación Android y creación de recursos de prueba. No se aplicaron correcciones al código de negocio.']], [1.65,5.2])
    heading('Contenido',1)
    for x in ['1 Verificación del flujo completo','2 Correspondencia con los procesos de negocio','3 Arquitectura del desarrollo','4 Funcionamiento de los componentes','5 Recorrido de ocho pantallas principales','6 Trazabilidad de los requerimientos','7 Reproducción y criterios de aceptación','8 Catálogo de evidencias y fuentes']:p(x)
    page();heading('1 Verificación del flujo completo')
    p('La auditoría reconstruyó las imágenes del backend y del panel desde el código presente en la carpeta. También compiló un APK Android x64, lo instaló en el emulador y ejecutó autenticación, consulta de ruta, cámara, firma y envío. Se contrastaron respuestas HTTP, interfaz, archivos S3, mensajes SQS y consultas SQL.')
    table(['Componente','Comprobación','Resultado'],[
        ['Backend Java','gradlew test --no-daemon','No compila el código de pruebas: tres errores por constructor antiguo de Order en OrderUseCasesImplTest.'],['Imagen backend','Docker build con -x test','Construye. Omitir tests no acredita calidad funcional.'],['Panel web','npm run build','Compila; bundle principal de 715.36 kB genera advertencia de tamaño.'],['Calidad web','npm run lint','34 errores y 4 advertencias.'],['Aplicación móvil','flutter analyze --no-pub','52 observaciones informativas; no fueron 52 errores de compilación.'],['Android x64','flutter build apk y adb install -r','Compilado, instalado y abierto. APK universal inicial no cabía; x64 sí se instaló.'],['BPMN','validate_business_bpmn.py','Nueve archivos pasan comprobación estructural local. No es certificación XSD ni ejecución de un motor BPMN.']], [1.2,1.9,3.75])
    heading('1 1 Casos ejecutados y límites de la comprobación',2)
    c=Counter(t['pass'] for t in API);cx=Counter(t['pass'] for t in EXTRA)
    p(f'La batería API registra {len(API)} ejecuciones: {c[True]} satisfactorias y {c[False]} hallazgos o resultados no conformes. Incluye seis autenticaciones repetidas entre fases. Las {len(EXTRA)} verificaciones adicionales registran {cx[True]} satisfactorias y {cx[False]} no satisfactorias. Las pruebas web y móviles complementan estas cifras; no se suman para inventar un porcentaje de cumplimiento.')
    p('Los tiempos p95 observados fueron 20 ms para /orders y 17 ms para /routes, con quince peticiones secuenciales locales por recurso. No se evaluó concurrencia, carga de producción, autonomía de batería, recuperación ante desastre, conectividad en campo ni distribución iOS. La validez de una respuesta 200 no demuestra por sí sola cumplimiento de una regla de negocio.')
    table(['Paso','Caso y evidencia','Conclusión'],[
        ['Identidad','Login web PKCE y login móvil real; T03, T04, X02','Tokens válidos y bloqueo de rutas básicos funcionan.'],['Confirmación','T06 y alta por formulario web id 14','Pedido PENDING sin ruta funciona; negociación y aceptación son humanas.'],['Asignación','T10 PATCH de asignación da 404','Falta conectar cartera existente al plan.'],['Despacho','T08 T11 T13 SQL ruta 3','Funciona con pedidos creados ya asociados a la ruta.'],['Camino','T14 y S04; X06 con OSRM detenido','OSRM y degradación funcionan; web y móvil no usan el mismo camino.'],['POD','T20 y envío manual móvil pedido 8','Foto, firma, actualización y cierre funcionan; GPS móvil incorrecto.'],['Sin red','Capturas hallazgo_sin_conexion y reconexión','Falla el envío; no hay sincronización automática observada.'],['Cierre','T31 y SQL ruta 3 COMPLETED','Se cerró al quedar pedidos DELIVERED o FAILED.'],['Control','S07, T29 y X09','Panel carga, pero filtros no son homogéneos y Reportes falla.']], [1.05,3.0,2.8])
    heading('1 2 Registro de hallazgos',2)
    for hid,priority,process,req,title,evidence,action in FINDINGS:
        heading(hid+' '+title,3)
        p(f'Prioridad {priority.lower()} · Procesos {process} · Requerimientos {req}.')
        p('Evidencia: '+evidence)
        p('Criterio de corrección: '+action)
    page();heading('2 Correspondencia con los procesos de negocio')
    p('El modelo BPMN representa el trabajo de MICOTRANS. El cliente es participante externo; dispatcher, conductor, administrador y jefatura son responsabilidades internas. EcoRoute es una herramienta de apoyo y no una persona o actor del negocio. Las lanes deben mostrar quién responde por la tarea, aunque utilice web o móvil.')
    p('Los nueve procesos conservan nombres de acción con un resultado identificable. La agrupación Gestión de la distribución puede usarse como macroproceso del proyecto; no obliga a renombrar cada proceso BPMN como Gestión. Tampoco existe una prohibición sintáctica de la palabra y: lo importante es la unidad del objetivo y un alcance medible.')
    coverage=[
        ['P01','Confirmar pedidos','Solicitud → pedido PENDING sin ruta','Cliente aclara o acepta; si faltan datos o no hay viabilidad, no admitir o desistir. Registro final sí funciona; acuerdos son manuales.'],
        ['P02','Preparar la jornada de reparto','Cartera y recursos → ruta IN_PROGRESS','Si no hay recurso o asignación válida, conservar cartera o reprogramar. Falta vincular PENDING; aprobación y capacidad no están automatizadas.'],
        ['P03','Ejecutar entregas','Ruta activa → paradas cerradas y ruta cerrada','Ausencia, rechazo, acceso bloqueado o POD incompleto pasan a P04. Entrega funciona, pero GPS y validación de prueba no cumplen.'],
        ['P04','Resolver incidencias','Excepción → reintento, nueva fecha o fallo cerrado','Reintento vuelve a P03; nueva fecha a P02; fallo final a P05. No hay expediente operativo ni motivo persistido completo.'],
        ['P05','Evaluar resultados','Jornada cerrada → evaluación y acción','Si faltan datos, corregir; si hay brecha, analizar causa, responsable y plazo. Dashboard parcial, reporte con error y BI avanzada pendiente.'],
        ['P06','Habilitar recursos de reparto','Necesidad → conductor o vehículo disponible','Documento inválido bloquea; indisponibilidad pasa a espera. CRUD existe; habilitación y verificación documental siguen siendo humanas.'],
        ['P07','Supervisar la jornada','Ruta activa → seguimiento o alerta','Falta de señal lleva a contacto y P04 si persiste; desvío confirmado también. GPS API funciona; canal web y vehículo móvil tienen defectos.'],
        ['P08','Administrar accesos','Solicitud de acceso → habilitado o revocado','Solicitud no autorizada se rechaza; rol incorrecto se corrige antes de habilitar. Keycloak autentica, pero rol local y JWT no se sincronizan.'],
        ['P09','Registrar compras operativas','Necesidad → recurso adquirido registrado','Si hay existencias, no comprar; si no se autoriza, cancelar o ajustar; bien observado se reclama. CRUD probado, presupuesto y aprobación son humanos.']]
    table(['Proceso','Nombre','Inicio y resultado','Condición de salida y correspondencia'],coverage,[.5,1.35,1.65,3.35])
    heading('2 1 Encadenamiento y finalización',2)
    steps(['P08 habilita el acceso y P06 aporta recursos; P09 puede abastecer a P06. Son procesos de soporte, no pasos que deban repetirse por cada pedido.','P01 comienza con la solicitud y termina al confirmar el pedido sin ruta o con una salida comercial no admitida, no viable o desistida. Solo un pedido validado debe entrar a P02.','P02 termina cuando el plan queda despachado. Una salida suspendida o sin asignación no puede pasar a P03. La aplicación actual necesita la operación de vinculación para cumplir esta transferencia.','P03 atiende las paradas; una excepción abre P04. P07 se desarrolla mientras la jornada esté activa y puede activar P04, pero no implica una entrega física adicional.','P04 debe decidir reintento en la misma jornada, nueva planificación o cierre fallido. No toda excepción significa finalizar la ruta ni todo no significa error técnico.','P05 empieza cuando existe un resultado consolidable. Si falta evidencia, el cierre analítico es provisional; no deben calcularse indicadores como si la información estuviera completa.'])
    heading('2 2 Revisión formal y límites de BPMN',2)
    p('Los XML existentes distinguen pool de MICOTRANS y lanes de trabajadores, tareas de envío o recepción y ramas Sí o No. El validador local comprueba referencias, conexiones, asignación de nodos a lanes y convergencias. No prueba que el código implemente esas actividades ni certifica toda la especificación BPMN 2.0.')
    p('Una terminación comercial normal puede usar un evento de fin simple. Error End Event sirve para lanzar un error BPMN capturable; Terminate End Event elimina todos los tokens del alcance. No se deben usar solo porque una rama sea No o se rechace un pedido. Una tarea puede recibir un mensaje entre pools; Receive Task hace explícita la espera. Las condiciones de salida deben ser exhaustivas y la convergencia debe evitar entradas ambiguas.')
    p('Los nueve BPMN originales se incluyen en el paquete, con PNG y XML editables. El índice HTML permite abrir cada diagrama a tamaño completo; no se reducen en Word a una miniatura ilegible. La tabla anterior y las pruebas muestran qué tarea se ejecuta en software y cuál sigue siendo un control humano o una mejora.')
    for proc in REQ['processes']:
        heading(proc['id']+' '+proc['title'],3)
        p('Responsables: '+', '.join(proc['lanes'])+'.')
        p('Secuencia principal de actividades: '+'; '.join(proc['tasks'])+'.')
        p('Resultados de terminación: '+'; '.join(proc['ends'])+'.')
    page();heading('3 Arquitectura del desarrollo')
    p('EcoRoute está formado por dos clientes y un backend de negocio único. El panel web y Flutter son aplicaciones distintas que consumen la misma API. PostgreSQL conserva los datos operativos, Keycloak gestiona identidad, OSRM calcula recorridos y LocalStack simula servicios AWS. Ejecutar varios contenedores no convierte el backend en una arquitectura de microservicios.')
    fig('01_arquitectura_completa','Figura 1  Contenedores y conexiones de la versión comprobada')
    table(['Parte','Tecnología','Responsabilidad'],[['Web','React 19 TypeScript Vite Axios Leaflet Chart.js','Interacción administrativa, planificación y control.'],['Móvil','Flutter Dart BLoC GetIt Dio flutter_map','Trabajo del conductor, ubicación y captura de POD.'],['Backend','Java 17 Spring Boot 3.5.8 WebFlux Reactor','API, casos de uso, seguridad y coordinación de adaptadores.'],['Persistencia','PostgreSQL 14 R2DBC Flyway','Integridad referencial y datos de negocio.'],['Identidad','Keycloak 23 OIDC JWT H2 dev','Usuarios, credenciales, roles emitidos y sesiones.'],['Ruteo','OSRM MLD extracto Perú','Orden de visita y camino estimado por calles.'],['Archivos y eventos','LocalStack S3 SQS','Objetos de evidencia y mensajes de entrega.'],['Redis','Redis 7','Conexión configurada y salud; sin uso de caché de negocio verificado.']], [1.0,2.4,3.45])
    heading('3 1 Capas internas del backend',2)
    fig('02_capas_backend','Figura 2  Puertos adaptadores y dependencias observadas')
    p('La entrada REST recibe JSON, convierte DTO y valida campos donde existe @Valid. Los casos de uso aplican decisiones y usan interfaces del dominio. Los adaptadores implementan acceso a PostgreSQL, OSRM y almacenamiento. Mono expresa un resultado asincrónico y Flux una secuencia. WebFlux permite composición no bloqueante, pero el rendimiento depende también de cada integración y no garantiza por sí solo escalabilidad.')
    p('La organización es hexagonal con dependencias todavía imperfectas: DashboardService conoce Spring Data y DTO de infraestructura; los servicios S3 y notificación se invocan desde aplicación sin aislamiento transaccional externo. La mejora no consiste en dibujar capas ideales, sino en reflejar estas dependencias y luego reducirlas.')
    heading('3 2 Despliegue y persistencia',2)
    fig('09_despliegue','Figura 3  Entorno Docker y configuración Android')
    table(['Acceso desde Windows','Destino','Observación'],[['localhost:3000','Nginx web puerto interno 80','/api se reenvía al backend; no hay proxy upgrade para /ws.'],['localhost:8081','Backend','REST, health y WebSocket GPS.'],['localhost:8080','Keycloak','Realm ecoroute; H2 en keycloak_data.'],['localhost:5433','PostgreSQL interno 5432','BD de negocio, no tablas internas de Keycloak.'],['localhost:5000','OSRM','API vial del extracto Perú.'],['localhost:4566','LocalStack','S3 y SQS locales; falta persistencia de objetos configurada.'],['10.0.2.2:8081 y :8080','Host visto desde emulador','Direcciones fijadas en Flutter; no sirven tal cual en teléfono físico.']], [1.85,1.55,3.45])
    heading('3 3 Datos de negocio e identidad',2)
    p('La base operativa tiene 33 tablas de aplicación y 27 claves foráneas verificadas; además existe flyway_schema_history para el control de migraciones. Las ocho migraciones están aplicadas. El modelo de negocio previamente diagramado selecciona 18 tablas de esas 33; no debe confundirse con un inventario completo del esquema ni con la base separada de Keycloak.')
    p('El núcleo de este flujo usa drivers y vehicles hacia routes; routes hacia orders mediante route_id nullable; orders hacia delivery_proofs mediante un vínculo único; y vehículo o conductor hacia vehicle_gps_history. La geometría y duración viven en routes y la secuencia de parada en orders. Los archivos no están como binarios dentro de PostgreSQL: delivery_proofs almacena sus referencias S3.')
    p('order_status_history y notifications existen, pero los escenarios auditados no generaron sus registros esperados. La existencia de una tabla no demuestra que su funcionalidad esté conectada. Las tablas corporativas, usuarios locales y compras no participan necesariamente de cada entrega. Keycloak no usa esas tablas auth_* como repositorio de sus credenciales.')
    page();heading('4 Funcionamiento de los componentes')
    for title,diagram,paragraphs in [
        ('4 1 Identidad y sesión web','03_acceso_web',[
            'PKCE protege el intercambio del código de autorización para un cliente público. El verificador permanece en el navegador durante el login y el challenge S256 viaja a Keycloak. Después, la API valida el token firmado y su emisor. El navegador no debe entregar contraseñas al servicio de pedidos.',
            'La UI decodifica roles para elegir navegación, pero no debe confiar en un driverId enviado libremente. La seguridad actual usa reglas amplias por URL; falta comprobar la propiedad del pedido y del vehículo. La emisión de roles de Keycloak y la edición de auth_user_role local son mecanismos distintos.']),
        ('4 2 Arquitectura y sesión móvil','11_sesion_movil',[
            'Cada feature separa presentación, dominio y datos. Una página emite eventos al BLoC; el BLoC ejecuta un caso de uso; el caso usa un repositorio abstracto; la implementación llama a Dio o al almacenamiento local. GetIt conecta implementaciones al iniciar. Esta separación facilita pruebas si las dependencias se sustituyen de forma consistente.',
            'La autenticación móvil usa password grant en esta configuración. El token se guarda mediante secure_storage, pero la recuperación de usuario no verifica vencimiento y no existe refresh móvil integrado. Hay una resolución de driverId por claim o consulta y un fallback de demostración; no debe usarse ese fallback como vínculo de identidad en producción.']),
        ('4 3 Pedido y transición a despacho','04_pedido_despacho',[
            'El pedido representa la obligación de atender una entrega. La ruta agrupa un trabajo operativo con recurso y fecha. Son entidades distintas y deben poder relacionarse después de la confirmación comercial. Actualmente createOrder decide entre PENDING y ASSIGNED según routeId al crear; no existe asignación posterior expuesta.',
            'Al iniciar una ruta, el backend calcula recorrido y actualiza pedidos a IN_TRANSIT. El comportamiento se comprobó en éxito y la transacción está declarada, pero no se inyectó un fallo de escritura para acreditar rollback. Una transición genérica sin validar origen puede reabrir terminales o despachar estados indebidos.']),
        ('4 4 Cálculo y representación del recorrido','05_ruteo',[
            'OSRM recibe coordenadas y devuelve una propuesta sobre la red vial disponible en el extracto Perú. La implementación usa /trip para ordenar y /route para dibujar el camino en el orden elegido. La geometría es polyline5 y las medidas se convierten a kilómetros y minutos. El dato es una estimación de red, no una garantía de llegada.',
            'El plan guardado al despachar puede incluir retorno y pedidos distintos del camino vivo solicitado después. Por eso 12.62 km de plan persistido y 6.4 km de paradas activas no son necesariamente contradictorios. La web dibuja otra heurística y debe alinearse. Capacidad, cobertura, ventanas y tráfico vivo no son restricciones resueltas por esta implementación.']),
        ('4 5 Captura y prueba de entrega','06_pod',[
            'La foto y firma se codifican en base64 y viajan en JSON. S3Service decodifica y crea claves UUID; el servidor guarda sus URI junto a receptor, coordenadas y fecha. Cambia el pedido y evalúa si la ruta puede cerrarse. El código etiqueta los objetos como PNG aunque la cámara puede entregar JPEG, un detalle a corregir al validar contenido y MIME.',
            'La transacción de PostgreSQL no abarca S3 ni SQS. Si una fase posterior falla, pueden quedar objetos huérfanos o eventos que deban reconciliarse. Una outbox transaccional y un consumidor idempotente son alternativas propuestas, no funciones presentes. El PDF actual solo muestra texto y no certifica inmutabilidad legal.']),
        ('4 6 Seguimiento de ubicación','07_gps',[
            'La ruta de escritura GPS es sensor, GpsBloc, repositorio, POST /gps/ping, R2DBC y emisión al canal en memoria. La lectura histórica usa el vehículo como clave. La lectura viva necesita abrir el socket y enviar ese vehicleId para seleccionar el stream.',
            'La API guardó y transmitió puntos de prueba. No se confirmó emisión real del móvil durante esta jornada. El vehicleId incorrecto y el onopen sobrescrito explican riesgos distintos: ubicación asociada al vehículo equivocado y cliente web que no se suscribe. Además, un socket anónimo pudo recibir datos; debe corregirse antes de exposición pública.']),
        ('4 7 Reportería y decisión de negocio','08_analitica',[
            'Las consultas actuales cuentan estados, agrupan por conductor o distrito y comparan fecha de prueba con ventana final. Para medir cumplimiento se necesita escoger qué fecha define la jornada, zona horaria, pedidos excluidos, tratamiento de reintentos y población de cada indicador. No se puede mezclar conteos del periodo con puntualidad histórica.',
            'El CSV actual contiene distrito y cantidad; es una salida operativa limitada. Una integración Power BI profesional requiere vistas o dataset con granularidad de pedido, parada y ruta, fechas de negocio, dimensiones, acceso de solo lectura y datos conciliados. IID, CHR y TDE deben conservar sus fichas y fórmulas aprobadas con la empresa; no se inventan cálculos equivalentes para llenar tarjetas.'])]:
        heading(title,2);fig(diagram,title)
        for txt in paragraphs:p(txt)
    heading('4 8 Módulos de soporte y funciones no conectadas',2)
    p('Conductores, vehículos y compras cuentan con controladores, repositorios y pantallas CRUD. La auditoría creó un conductor, un vehículo y una compra de prueba sin eliminar recursos existentes. No se verificaron aprobaciones presupuestales, mantenimiento preventivo, validez documental, reserva de capacidad ni vigencia de licencias automatizadas. Los procesos P06 y P09 incluyen decisiones humanas que no se convierten en funcionalidades por tener un formulario.')
    p('Notificaciones consulta y marca filas de su tabla, pero NotificationService solo publica el pedido en SQS. No se encontró consumidor que inserte esas filas. Redis participa del health check, no de un bus GPS compartido. Isar y LocalCache aparecen en dependencias o clases móviles, pero no están conectados al envío de POD. El proyecto debe describir uso efectivo y no solo librerías declaradas.')
    page();heading('5 Recorrido de ocho pantallas principales')
    p('Se seleccionaron ocho pantallas lógicas que cubren acceso, pedido, despacho, ruta de campo, prueba, historial y control. Un modal, un formulario externo o la captura de firma se incluyen como estados de la misma pantalla, no como pantallas adicionales para inflar la selección. Las capturas proceden de la versión ejecutada; no son wireframes ni promesas de una corrección pendiente.')
    table(['Pantalla','Proceso de negocio','Interacción principal'],[[s['id']+' '+s['title'],s['bpmn'],s['actor']] for s in SCREENS],[2.8,1.5,2.55])
    for s in SCREENS:
        page();heading('5 '+str(int(s['id'][1:]))+' '+s['title'])
        p('Responsable: '+s['actor']+'. Procesos: '+s['bpmn']+'. Requerimientos: '+s['rf']+'.')
        p(s['purpose']);shot(s['image'],s['id']+'  Captura real de la interfaz',s['mobile'])
        heading('Diseño e interacción',2);p(s['design'])
        if s['support']:shot(s['support'],s['id']+'  Estado complementario del mismo flujo',s['mobile'])
        heading('Paso a paso del proceso',2);steps(s['steps'])
        heading('Frontend y patrón de arquitectura',2);p(s['front'])
        heading('Backend y persistencia',2);p(s['back'])
        heading('Contrato y datos intercambiados',2);p(s['contract'])
        heading('Diagrama de la interacción',2)
        if s['id']=='S06':
            fig(s['diagram'],'Secuencia técnica de '+s['title'].lower())
        else:
            fig(s['diagram'],'Secuencia técnica de '+s['title'].lower(),reference_only=True)
        heading('Resultado y controles pendientes',2);p(s['result'])
        heading('Explicación sugerida para la exposición',2);p(s['speech'])
    page();heading('6 Trazabilidad de los requerimientos')
    p('Cada fila usa el texto del documento académico vigente y una evaluación explícita. Verificado en ejecución significa que se comprobó el caso indicado, no todas las variantes posibles. Revisado en código no equivale a aprobado. Parcial señala una diferencia observada o una parte pendiente. No implementado localizado se refiere a esta versión de trabajo, sin asumir que otra rama contenga una integración terminada.')
    heading('6 1 Requerimientos funcionales',2)
    table(['Código','Requerimiento','Estado y sustento'],[[r['id'],r['description'],STATUS[RF[int(r['id'].split('-')[1])][0]]+'. '+RF[int(r['id'].split('-')[1])][1]] for r in REQ['requirements'] if r['id'].startswith('RF-')],[.55,2.55,3.75])
    heading('6 2 Requerimientos no funcionales',2)
    p('Las mediciones históricas escritas en el capítulo 1 no se reutilizan como si se hubieran obtenido hoy. Esta revisión registra el valor observado o limita su alcance. No se certifica producción ni cumplimiento de ISO 25010 completo por haber clasificado los requisitos con esa norma.')
    table(['Código','Requerimiento','Estado y sustento'],[[r['id'],r['description'],STATUS[RNF[int(r['id'].split('-')[1])][0]]+'. '+RNF[int(r['id'].split('-')[1])][1]] for r in REQ['requirements'] if r['id'].startswith('RNF-')],[.55,2.55,3.75])
    page();heading('7 Reproducción y criterios de aceptación')
    heading('7 1 Secuencia de verificación local',2)
    steps(['Revisar docker compose ps y conservar los volúmenes existentes. Construir ecoroute-backend y ecoroute-web-admin desde el código que se pretende auditar; no asumir que una imagen antigua representa la carpeta actual.','Iniciar postgres, redis, keycloak, localstack, osrm y los dos servicios de aplicación. Comprobar salud y aplicación de V1 a V8.','Verificar que existan ecoroute-proofs y ecoroute-notifications. En este ensayo hubo que crear ambos manualmente porque el inicializador falló.','Compilar e instalar el APK Android x64. Usar America/Lima en el emulador y ubicación de prueba. Las direcciones 10.0.2.2 solo son apropiadas para el emulador.','Autenticar con cuentas locales de prueba y conservar los tokens solo en memoria durante scripts. No publicar volcados de headers ni secretos.','Ejecutar los scripts audit_api.mjs por sus fases y audit_web.mjs únicamente en datos de prueba; crean registros y la fase inicial no es idempotente. No repetirlos contra producción.','Capturar estados móviles, simular falta de red, recuperar conexión y comprobar SQL, objetos y ruta. El prefijo AUDIT permite distinguir los registros creados.','Para repetir el caso de resiliencia se detiene solo OSRM y se restaura con finally. No detener bases ni servicios ajenos.','Revisar matrices y contrastar el mismo pedido desde UI, HTTP y BD. Registrar los controles no ensayados por separado.'])
    heading('7 2 Correcciones necesarias antes de aceptación integral',2)
    steps(['Cerrar H01 para ejecutar P01 → P02 sin recrear el pedido. Validar propietario, asignación, recurso y transición con pruebas positivas y negativas.','Cerrar H02, H03, H04 y H06 antes de abrir acceso público: propiedad del dato, POD obligatorio válido, GPS real y canal autenticado.','Implementar y probar la excepción offline del BPMN con persistencia y reintento idempotente, incluido reinicio del teléfono.','Corregir suscripción GPS y vehicleId. Comparar una muestra móvil con lo mostrado en la web y el vehículo persistido.','Reparar reportes, filtros, fecha Lima y notificaciones. Verificar un periodo vacío, un distrito sin dato y una exportación con filas.','Aprobar test Java, lint web, contrato OpenAPI y pruebas de recorrido. La compilación Docker con -x test no puede ser criterio único.','Acordar fórmulas y fuentes de indicadores con MICOTRANS e implementar RF-75 a RF-81 según alcance aprobado. Después repetir la auditoría y actualizar la documentación académica con resultados reales.'])
    heading('7 3 Configuración y archivos para reproducir',2)
    p('Los comandos se ejecutan desde la raíz del proyecto. No incluyen eliminación de datos ni credenciales. Para probar endpoints protegidos se obtiene un JWT legítimo de Keycloak; el secreto no debe agregarse al manual.')
    for cmd in ['docker compose ps','docker compose build ecoroute-backend ecoroute-web-admin','docker compose up -d postgres redis keycloak localstack osrm ecoroute-backend ecoroute-web-admin','docker exec ecoroute-localstack awslocal s3 ls','docker exec ecoroute-localstack awslocal sqs list-queues','.\\gradlew.bat test --no-daemon','cd web-admin  →  npm run build  →  npm run lint','cd mobile-app  →  flutter analyze --no-pub','flutter build apk --debug --no-pub --target-platform android-x64','python tools/validate_business_bpmn.py']:
        add('code',text=cmd)
    page();heading('8 Catálogo de evidencias y fuentes')
    heading('8 1 Pruebas HTTP ejecutadas',2)
    table(['Caso','Comprobación','Resultado observado'],[[t['id'],t['name'],('Conforme al caso' if t['pass'] else 'Hallazgo')+' · '+str(t.get('http',''))+' '+str(t.get('detail','')).replace('\n',' ')[:350]] for t in API],[.85,2.0,4.0])
    heading('8 2 Verificaciones adicionales',2)
    table(['Caso','Verificación','Resultado'],[[t['id'],t['name'],('Correcto' if t['pass'] else 'No aprobado')+' · '+str(t.get('detail',t.get('http',''))) ] for t in EXTRA],[.65,2.1,4.1])
    heading('8 3 Archivos que respaldan el manual',2)
    table(['Ubicación dentro del paquete','Contenido'],[['evidencias/pruebas_api.json','38 ejecuciones HTTP con resultado y referencia a requerimiento.'],['evidencias/pruebas_adicionales.json','Token, latencias, ruteo degradado, métricas y reporte.'],['evidencias/pruebas_web.json','Navegación real, alta, exportación y mensajes WebSocket.'],['evidencias/verificacion_base_datos.json','Estados, pruebas S3, coordenadas, migraciones, tablas y FK.'],['evidencias/casos_prueba.json','Identificadores sintéticos del ensayo.'],['evidencias/requisitos_y_bpmn.json','81 RF, 54 RNF y nueve procesos extraídos del Word y XML.'],['evidencias/resultado_movil.json','Secuencia móvil observada y límites de ejecución.'],['pantallas','Ocho pantallas lógicas y capturas de sus estados y fallos.'],['diagramas','Once SVG editables y sus PNG.'],['bpmn','Nueve XML BPMN y sus PNG vigentes.'],['index.html','Manual navegable con figuras ampliables y acceso a evidencias.'],['trazabilidad_requerimientos.json','Evaluación de los 135 requerimientos con sustento.']], [3.1,3.75])
    heading('8 4 Fuentes de implementación',2)
    p('Todas las rutas siguientes son relativas a la carpeta raíz del repositorio. La versión exacta de los archivos utilizados queda registrada con SHA256 en evidencias/manifest_fuentes.json. El repositorio ya contenía modificaciones locales; este manual no atribuye esos cambios a la auditoría.')
    for folder,files in [
        ('src/main/java/com/ecoroute/backend/', ['infrastructure/config/SecurityConfig.java','infrastructure/input/rest/OrderController.java','infrastructure/input/rest/RouteController.java','infrastructure/input/rest/DeliveryProofController.java','infrastructure/input/rest/ReportController.java','application/usecases/OrderUseCasesImpl.java','application/usecases/RouteUseCasesImpl.java','application/usecases/CreateDeliveryProofUseCaseImpl.java','application/services/RoutePlannerService.java','application/services/S3Service.java','application/services/NotificationService.java','application/services/PdfService.java','application/services/DashboardService.java','infrastructure/output/routing/OsrmRoutingAdapter.java']),
        ('web-admin/src/', ['context/AuthContext.tsx','services/keycloak.ts','services/api.ts','services/gpsService.ts','pages/Orders.tsx','pages/Routes.tsx','pages/Dashboard.tsx','pages/Reports.tsx']),
        ('mobile-app/lib/', ['core/di/injection_container.dart','features/routes/presentation/pages/routes_page.dart','features/orders/presentation/pages/order_detail_page.dart','features/orders/data/repositories/order_repository_impl.dart','features/orders/data/datasources/order_remote_data_source.dart','features/history/data/datasources/history_remote_data_source.dart','features/gps/presentation/bloc/gps_bloc.dart','features/auth/data/repositories/auth_repository_impl.dart'])]:
        p(folder)
        for f in files:add('code',text=f)
    p('Fuentes de negocio: documento Integrador I_Grupo3_Capitulo2_BPMN_Formal.docx y nueve archivos P01 a P09 de negocio_revisado/bpmn. Este manual complementa el capítulo 2 y no sobrescribe el informe, la presentación ni el guion académico existentes.')

def copy_bpmn():
    dest=OUT/'bpmn';dest.mkdir(exist_ok=True)
    base=ROOT/'docs/entregables_procesos_ecoroute/negocio_revisado'
    for i in range(1,10):
        stem=f'P{i:02}_negocio'
        shutil.copy2(base/'bpmn'/(stem+'.bpmn'),dest/(stem+'.bpmn'))
        shutil.copy2(base/'imagenes'/(stem+'.png'),dest/(stem+'.png'))

def export_audit_data():
    rows=[]
    for r in REQ['requirements']:
        status,note=(RNF if r['id'].startswith('RNF') else RF)[int(r['id'].split('-')[1])]
        rows.append({**r,'status':STATUS[status],'evidence':note})
    (OUT/'trazabilidad_requerimientos.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
    (E/'resultado_movil.json').write_text(json.dumps({'date':'2026-10-19','device':'emulator-5554 Medium_Phone_API_36.1','orderId':8,'routeId':3,'tests':[
        {'case':'Login real','pass':True}, {'case':'Ruta por calles y detalle','pass':True}, {'case':'Cámara y firma','pass':True}, {'case':'Entrega manual y cierre','pass':True}, {'case':'Sincronización automática después de falta de red','pass':False,'evidence':'No cambió hasta reintento manual'}, {'case':'Coordenadas reales POD','pass':False,'evidence':'BD latitude=0 longitude=0'}, {'case':'Historial propio y del día','pass':False,'evidence':'GET /orders global; muestra fechas antiguas'}, {'case':'Emisión GPS móvil','pass':None,'evidence':'Sin muestras móviles verificadas; no se acredita'}]},ensure_ascii=False,indent=2),encoding='utf-8')
    import hashlib
    files=[]
    patterns=['src/main/java/**/*.java','src/main/resources/db/migration/*.sql','web-admin/src/**/*.ts','web-admin/src/**/*.tsx','mobile-app/lib/**/*.dart']
    selected=[ROOT/'docker-compose.yml',ROOT/'web-admin/package.json',ROOT/'mobile-app/pubspec.yaml',ROOT/'build.gradle']
    for pattern in patterns:selected.extend(ROOT.glob(pattern))
    for f in sorted(set(selected)):
        if f.exists():files.append({'path':f.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
    (E/'manifest_fuentes.json').write_text(json.dumps(files,indent=2),encoding='utf-8')

def export_html():
    nav=[]; body=[]
    for i,b in enumerate(BLOCKS):
        k=b['kind'];t=esc(b.get('text',''))
        if k=='title':body.append(f'<h1>{t}</h1>')
        elif k=='h':
            level=b['level']+1;body.append(f'<h{level} id="s{i}">{t}</h{level}>')
            if b['level']==1:nav.append(f'<a href="#s{i}">{t}</a>')
        elif k in ('p','code'):body.append(f'<{("pre" if k=="code" else "p")}>{t}</{("pre" if k=="code" else "p")}>')
        elif k=='table':body.append('<div class="table"><table><thead><tr>'+''.join('<th>'+esc(x)+'</th>' for x in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(x)+'</td>' for x in row)+'</tr>' for row in b['rows'])+'</tbody></table></div>')
        elif k in ('figure','screen'):
            name=('diagramas/'+b['name']+'.svg') if k=='figure' else ('pantallas/'+b['name']+'.png')
            body.append(f'<figure class="{("mobile" if b.get("mobile") else "")}"><a href="{name}" target="_blank"><img loading="lazy" src="{name}" alt="{esc(b["caption"])}"></a><figcaption>{esc(b["caption"])}</figcaption></figure>')
    body.append('<h2 id="bpmn">BPMN completos del negocio</h2>')
    for proc in REQ['processes']:
        stem=proc['id']+'_negocio';body.append(f'<h3>{esc(proc["id"]+" "+proc["title"])}</h3><a href="bpmn/{stem}.bpmn">XML BPMN editable</a><figure><a target="_blank" href="bpmn/{stem}.png"><img loading="lazy" src="bpmn/{stem}.png" alt="{esc(proc["title"])}"></a></figure>')
    body.append('<h2 id="evidence">Evidencias</h2><ul>')
    for f in sorted(E.glob('*.json')):body.append(f'<li><a href="evidencias/{f.name}">{esc(f.name)}</a></li>')
    body.append('</ul>')
    css='''body{margin:0;color:#172c42;background:#f0f3f7;font:16px/1.65 Segoe UI,Arial,sans-serif}nav{position:fixed;inset:0 auto 0 0;width:255px;background:#132d48;color:white;padding:28px 20px;overflow:auto}nav strong{font-size:23px}nav a{display:block;color:#e3edfa;margin:14px 0;text-decoration:none}main{margin-left:295px;max-width:1150px;background:white;padding:48px 58px}h1{font-size:36px;line-height:1.2}h2{border-top:1px solid #cbd5e1;padding-top:32px;margin-top:58px}h3{margin-top:30px}figure{margin:28px 0;text-align:center}figure img{max-width:100%;height:auto;border:1px solid #e1e7ee}figure.mobile img{max-height:680px;width:auto}figcaption{font-size:14px;color:#5a6b7f}table{border-collapse:collapse;width:100%;font-size:14px}th{background:#204d78;color:white;text-align:left}th,td{padding:12px;border:1px solid #d9d9d9;vertical-align:top}tr:nth-child(even){background:#f4f7fa}.table{overflow:auto}pre{background:#f2f5f8;padding:12px;white-space:pre-wrap;font-size:13px}a{color:#175ca0}@media(max-width:850px){nav{position:static;width:auto}main{margin:0;padding:25px}}@media print{nav{display:none}main{margin:0;padding:0}h2{break-before:page}img{max-height:800px}}'''
    content='<!DOCTYPE html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>EcoRoute arquitectura y auditoría</title><style>'+css+'</style><nav><strong>EcoRoute</strong><p>Arquitectura y auditoría</p>'+''.join(nav)+'<a href="#bpmn">BPMN completos</a><a href="#evidence">Evidencias</a></nav><main>'+''.join(body)+'</main></html>'
    (OUT/'index.html').write_text(content,encoding='utf-8')

def export_docx():
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    for name,size in [('Normal',10.5),('Title',27),('Heading 1',18),('Heading 2',13),('Heading 3',11.5),('Caption',9)]:
        st=doc.styles[name];st.font.name='Aptos';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        pf=st.paragraph_format;pf.space_after=Pt(7);pf.line_spacing=1.08
        if name.startswith('Heading'):pf.space_before=Pt(13);pf.keep_with_next=True
        for border in list(st.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
        for underline in list(st.element.findall('.//'+qn('w:u'))):underline.getparent().remove(underline)
    doc.styles['Normal'].paragraph_format.widow_control=True
    # Footer provides useful page location in a long technical manual.
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('EcoRoute  ·  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    doc.core_properties.title='EcoRoute arquitectura verificación y flujos del sistema'
    doc.core_properties.subject='Auditoría local y recorrido técnico de ocho pantallas'
    doc.core_properties.author='Proyecto EcoRoute'
    figcount=0
    for b in BLOCKS:
        k=b['kind']
        if k=='title':doc.add_paragraph(b['text'],'Title')
        elif k=='h':doc.add_heading(b['text'],b['level'])
        elif k=='p':doc.add_paragraph(b['text'])
        elif k=='code':
            pr=doc.add_paragraph();pr.paragraph_format.space_after=Pt(3)
            r=pr.add_run(b['text']);r.font.name='Consolas';r.font.size=Pt(8.5)
        elif k=='break':doc.add_page_break()
        elif k=='table':
            tb=doc.add_table(rows=1,cols=len(b['headers']));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
            widths=b.get('widths') or [6.9/len(b['headers'])]*len(b['headers'])
            for j,w in enumerate(widths):tb.columns[j].width=Inches(w)
            for j,h in enumerate(b['headers']):tb.rows[0].cells[j].text=h
            for row in b['rows']:
                for j,val in enumerate(row):tb.add_row() if False else None
                cells=tb.add_row().cells
                for j,val in enumerate(row):cells[j].text=str(val)
            for i,row in enumerate(tb.rows):
                trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
                if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
                for j,cell in enumerate(row.cells):
                    cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),BLUE[1:] if i==0 else ('F3F6F9' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                    tcpr.append(borders);marg=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');marg.append(el)
                    tcpr.append(marg)
                    for pr in cell.paragraphs:
                        pr.paragraph_format.space_after=Pt(3);pr.paragraph_format.line_spacing=1.02
                        if j==0 and len(b['headers'])>=3:pr.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        for run in pr.runs:run.font.size=Pt(8.7);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '172C42');run.bold=i==0
            doc.add_paragraph().paragraph_format.space_after=Pt(3)
        elif k in ('figure','screen'):
            if b.get('reference_only'):
                label={'03_acceso_web':'4 1 Identidad y sesión web','04_pedido_despacho':'4 3 Pedido y transición a despacho','05_ruteo':'4 4 Cálculo y representación del recorrido','06_pod':'4 5 Captura y prueba de entrega','08_analitica':'4 7 Reportería y decisión de negocio'}[b['name']]
                doc.add_paragraph('La secuencia completa aparece en '+label+'. El manual HTML la muestra también junto a esta pantalla y permite ampliarla. El archivo editable es diagramas/'+b['name']+'.svg.')
                continue
            file=DIAG/(b['name']+'.png') if k=='figure' else OUT/'pantallas'/(b['name']+'.png')
            if not file.exists():raise FileNotFoundError(file)
            pr=doc.add_paragraph();pr.alignment=WD_ALIGN_PARAGRAPH.CENTER;pr.paragraph_format.keep_with_next=True
            if b.get('mobile'):pr.add_run().add_picture(str(file),height=Inches(4.25))
            else:pr.add_run().add_picture(str(file),width=Inches(b.get('width',6.85)))
            shape=doc.inline_shapes[-1];shape._inline.docPr.set('descr',b['caption'])
            cap=doc.add_paragraph(b['caption'],'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    file=OUT/'EcoRoute_Arquitectura_Verificacion_y_Flujos.docx';doc.save(file)
    print('DOCX',file)

if __name__=='__main__':
    build_diagrams();copy_bpmn();export_audit_data();create_content();export_html()
    if '--docx' in sys.argv:export_docx()
    (OUT/'contenido.json').write_text(json.dumps(BLOCKS,ensure_ascii=False,indent=2),encoding='utf-8')
    print(len(BLOCKS),'bloques',len(FIGURES),'diagramas',len(RF),'RF',len(RNF),'RNF')
