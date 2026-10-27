"""Actualiza el capítulo y el guion de exposición con los nueve BPMN revisados."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
IMAGES = DOCS / 'entregables_procesos_ecoroute/negocio_revisado/imagenes'
REPORT_SOURCE = DOCS / 'Integrador I_Grupo3_Capitulo2_BPMN_Negocio.docx'
REPORT_DEST = DOCS / 'Integrador I_Grupo3_Capitulo2_BPMN_Formal.docx'
SPEECH_SOURCE = DOCS / 'Guion_Detallado_Exposicion_Capitulo_2_EcoRoute_BPMN_Negocio.docx'
SPEECH_DEST = DOCS / 'Guion_Detallado_Exposicion_Capitulo_2_EcoRoute_BPMN_Formal.docx'


def set_text(paragraph, value):
    paragraph.text = value
    paragraph.paragraph_format.keep_together = True


report = Document(REPORT_SOURCE)
p = report.paragraphs
assert 'Nivel de modelado' in p[219].text
set_text(p[219],
         'Nivel de modelado: los nueve BPMN describen procesos de negocio objetivo (To-Be). '
         'Los carriles identifican responsabilidades internas de MICOTRANS. En P01, MICOTRANS es '
         'el pool con el carril Dispatcher y el cliente solicitante es un pool externo de caja negra. '
         'La propuesta sale de una tarea de envío (Send Task) y la respuesta llega a una tarea de '
         'recepción (Receive Task) mediante flujos de mensaje; ningún flujo de secuencia cruza pools. '
         'En los nueve diagramas, las compuertas XOR divergentes separan Sí y No; cuando caminos '
         'alternativos vuelven a una actividad, una XOR convergente los reúne antes de la tarea. '
         'Las preguntas en los rombos facilitan la lectura y los nombres de las ramas explicitan la '
         'condición. Los rechazos, desistimientos y casos no viables terminan con eventos de fin '
         'ordinarios porque son resultados comerciales, no errores BPMN ni terminaciones de flujos '
         'paralelos. EcoRoute y Keycloak son herramientas de soporte, no actores del negocio. '
         'Referencia de notación: OMG, BPMN 2.0.2.')
set_text(p[229],
         'Figura 2.1. BPMN de P01 Confirmar pedidos. MICOTRANS contiene el carril Dispatcher; '
         'el cliente solicitante es un pool externo. La propuesta y la respuesta se modelan con '
         'tareas de envío y recepción, flujos de mensaje y convergencias XOR. Fuente: elaboración '
         'propia y análisis de la aplicación.')
for index in range(9):
    picture = p[228 + 8 * index]
    for run in list(picture.runs):
        picture._p.remove(run._element)
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(IMAGES / f'P{index+1:02d}_negocio.png'), width=Inches(6.1))
# La entrada P01 del índice heredado conserva un PAGEREF a un marcador ausente.
# Restaurarlo en su título evita «Error! Marcador no definido» al actualizar campos.
assert p[225].text == 'P01 Confirmar pedidos'
bookmark = OxmlElement('w:bookmarkStart')
bookmark.set(qn('w:id'), '90001')
bookmark.set(qn('w:name'), '_Toc241417675')
first_run = p[225].runs[0]._element
p[225]._p.insert(p[225]._p.index(first_run), bookmark)
bookmark_end = OxmlElement('w:bookmarkEnd')
bookmark_end.set(qn('w:id'), '90001')
p[225]._p.append(bookmark_end)
for link in report.element.body.xpath('.//w:hyperlink'):
    if link.get(qn('w:anchor')) == '_Toc241417675':
        cached = link.xpath('.//w:t')
        if cached and cached[-1].text == '27':
            cached[-1].text = '28'
report.save(REPORT_DEST)

speech = Document(SPEECH_SOURCE)
g = speech.paragraphs
set_text(g[4],
         'Criterio de rigor  El pedido PENDING, la ruta PLANNED y el despacho IN_PROGRESS son '
         'estados distintos. Los nueve BPMN representan la operación objetivo y muestran los '
         'responsables internos de MICOTRANS. P01 incorpora al cliente como pool externo: el '
         'Dispatcher envía condiciones y espera una respuesta en una tarea de recepción. Las '
         'decisiones usan XOR divergentes y las rutas alternativas se reúnen con XOR convergentes '
         'antes de la siguiente actividad. Los finales comerciales negativos son finales normales. '
         'EcoRoute y Keycloak apoyan el trabajo, sin actuar como trabajadores del negocio.')
set_text(g[33],
         'Aquí distinguimos dos niveles. PMBOK dirige el proyecto de desarrollo mediante el acta, '
         'la EDT y el cronograma. BPMN describe la operación de MICOTRANS. Un pool delimita cada '
         'participante; un carril asigna tareas dentro de la empresa. En P01 se ve el pool '
         'MICOTRANS con el carril Dispatcher y, aparte, el pool de caja negra del cliente. La '
         'comunicación cruza la frontera mediante un flujo de mensaje, no mediante un flujo de '
         'secuencia. El rombo naranja divide alternativas y el azul las reúne. EcoRoute y Keycloak '
         'son soporte técnico, por eso no ocupan carriles de negocio.')
set_text(g[39],
         'En pantalla  Empiece por los límites de los dos pools de P01. Dentro de MICOTRANS '
         'señale el carril Dispatcher. Ubique la tarea de envío de condiciones y siga la línea '
         'discontinua hasta el cliente. Luego siga la línea de respuesta hasta la tarea de '
         'recepción. Muestre una vuelta por datos incompletos: el flujo No solicita corrección; '
         'si el cliente corrige, la XOR azul reúne el retorno con el avance ordinario antes de '
         'revisar los datos. Señale después los finales normales de no admisión, inviabilidad, '
         'desistimiento y pedido validado.')
set_text(g[43],
         'La última decisión pregunta si el cliente confirma las condiciones. El Dispatcher '
         'comunica la propuesta mediante una tarea de envío. La respuesta del cliente se espera '
         'en una tarea de recepción y recién entonces se evalúa Sí o No. Si responde No, el '
         'caso termina como desistido con un evento de fin ordinario: no se trata de un error '
         'técnico. Si responde Sí, se registra el pedido sin ruta y queda PENDING. Los rombos '
         'azules del diagrama reúnen retornos alternativos antes de repetir una tarea. '
         'La confirmación comercial aún no tiene estado CONFIRMED ni evidencia automática.')
for index in (47, 55, 63, 71, 79, 87, 95, 103):
    g[index].add_run(' Las convergencias azules unen solo caminos alternativos antes de continuar; '
                     'no agregan una tarea ni crean otra ejecución.')
set_text(g[460], 'docs/Integrador I_Grupo3_Capitulo2_BPMN_Formal.docx, capítulo 2, apartados 2.1 a 2.5.')
set_text(g[461], 'docs/presentacion_capitulo2/final/EcoRoute_Capitulo_2_Diseno_BPMN_Formal_v4.pptx, 56 láminas y notas del expositor.')
speech.save(SPEECH_DEST)

print(REPORT_DEST)
print(SPEECH_DEST)
