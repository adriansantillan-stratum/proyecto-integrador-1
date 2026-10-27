"""Build the eight-screen defence guide from inspected source-backed content."""
from pathlib import Path
import json, hashlib, textwrap, re
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUT = ROOT / 'docs/exposicion_8_pantallas'
QA = OUT / '_qa'
ASSETS = OUT / 'diagramas'
for p in (OUT, QA, ASSETS): p.mkdir(parents=True, exist_ok=True)
MEMBERS = {1:'Coronel Márquez Enzo Alejandro',2:'Espinoza Cristóbal Luis Enrique',3:'Namuche Huarcaya Renzo Aldair',4:'Santillan Velasquez Adrian Nicolas'}
SCREENS = []
for n in range(1,5):
    pair=json.loads((HERE / f'alumno{n}.json').read_text(encoding='utf-8-sig'))
    for screen in pair:
        screen['student']=n
        for key in ['processes','actor','platform','route']:
            if isinstance(screen.get(key),list):screen[key]='; '.join(str(x) for x in screen[key])
        for excerpt in screen['code']:
            excerpt['explanation']=[x for x in excerpt['explanation'] if not x.lstrip().startswith('Código literal actual')]
    SCREENS += pair
assert len(SCREENS) == 8
assert [int(s['id']) for s in SCREENS] == list(range(1,9))
SCREENS[0]['code'].insert(1,{
 'file':'src/main/java/com/ecoroute/backend/application/usecases/OrderUseCasesImpl.java','start':50,'end':71,
 'explanation':[
  'El caso de uso rechaza routeId en el alta y crea otro Order con ID nulo, PENDING y fechas de servidor. La expresión ASSIGNED queda inaccesible para una entrada válida por la guarda previa.',
  'El puerto orderRepository.save devuelve el registro guardado. flatMap encadena audit.record y thenReturn devuelve el mismo pedido después del historial. El adaptador R2DBC, no este método, implementa la tecnología de persistencia.',
  '@Transactional contiene las escrituras participantes de este caso de uso. El PATCH comercial posterior es otra petición y otra frontera transaccional.'
 ]})
SCREENS[7]['code'].insert(0,{
 'file':'web-admin/src/pages/BusinessIntelligence.tsx','start':76,'end':93,
 'explanation':[
  'El componente conserva los filtros y la pestaña en estado React y construye una URL que identifica el informe. useQuery realiza la consulta y separa datos, carga, error y recarga.',
  'Las tarjetas derivan del contrato recibido y las pestañas representan distintos bloques del mismo informe. Cambiar de vista no ejecuta una entrega ni modifica la tabla orders.',
  'Seguir /reports/bi hasta ReportController, DashboardService y AnalyticsQueryPort. La consulta SQL prepara las fracciones; el componente aplica su representación visual.'
 ]})
SCREENS[1]['code'][0]['end']=109
SCREENS[1]['code'][1]['start']=88
SCREENS[3]['handoff']=SCREENS[3]['handoff'].replace('Los nombres del resto de pantallas quedan a la guía general, sin inventarlos aquí.','').strip()

FONTDIR = Path('C:/Windows/Fonts')
def font(size=26,bold=False):
    return ImageFont.truetype(str(FONTDIR / ('segoeuib.ttf' if bold else 'segoeui.ttf')),size)

def wrap_pixels(draw,text,f,max_width):
    lines=[]
    for para in str(text).split('\n'):
        words=para.split(); line=''
        for word in words:
            trial=(line+' '+word).strip()
            if draw.textlength(trial,font=f)>max_width and line:
                lines.append(line); line=word
            else: line=trial
        lines.append(line)
    return lines

FLOW_LABELS = {
 '01': [('Dispatcher','Registra la solicitud sin ruta'),('React','CreateOrder.submit valida y construye JSON'),('HTTP y seguridad','POST /api/orders con Bearer; Nginx → /orders'),('Backend','OrderController → CreateOrderUseCase → OrderUseCasesImpl'),('Persistencia','OrderRepository → adaptador R2DBC → orders e historial'),('Acuerdo comercial','PATCH business-data → servicio → revisión y auditoría'),('Traspaso P01 a P02','Mismo pedido PENDING con revisión completa')],
 '02': [('Dispatcher','Crea ruta PLANNED y asigna pedidos confirmados'),('Jefatura autorizada','POST review verifica plan y carga con ADMIN'),('Conductor propietario','POST acceptance desde móvil tras aprobación'),('Dispatcher','PATCH status=IN_PROGRESS solicita salida'),('Backend','RouteUseCasesImpl exige revisión y valida paradas'),('Persistencia y ruteo','Plan OSRM o fallback; pedidos IN_TRANSIT y ruta IN_PROGRESS'),('Traspaso P02 a P03','Hoja despachada lista para ejecutar entregas')],
 '03': [('Conductor','Abre Mis rutas con su sesión'),('Flutter','Página → RouteBloc → caso de uso y repositorio'),('Servidor','GET rutas y revisión con identidad y fecha'),('Decisión de negocio','Acepta hoja aprobada que le pertenece'),('Hoja activa','Carga pedidos y consulta /path para paradas activas'),('Mapa y posición','Dibuja plan; transmite GPS durante ruta activa'),('Traspaso a la parada','Abre detalle del pedido y vuelve recargando estado')],
 '04': [('Conductor','Abre pedido en tránsito de su propia ruta'),('Captura','Foto, firma, receptor, DNI, ubicación y fecha'),('Persistencia local','Guarda POD en cola privada de la cuenta'),('Sincronización','Intenta envío; recupera sesión y reintenta según error'),('Backend','Autoriza, valida y procesa POD idempotente'),('Persistencia','Evidencia, integridad y estado DELIVERED con auditoría'),('Respuesta y retorno','Confirma servidor, retira cola y actualiza ruta')],
 '05': [('Dispatcher','Consulta incidencias y contexto del pedido'),('Decisión humana','Elige resolución y documenta motivo'),('React y API','Formulario → servicio → petición autenticada'),('Backend','BusinessOperationsService valida resolución permitida'),('Persistencia','Cierra incidencia; crea nuevo intento si corresponde'),('Conservación de historia','Pedido original terminal no se reabre'),('Traspaso','Nuevo intento vuelve a ejecución o planificación')],
 '06': [('Responsable de mejora','Identifica causa e indicador que debe mejorar'),('React','Registra acción, responsable, plazo y medida'),('HTTP y seguridad','Servicio web → /improvement-actions'),('Backend','Valida campos y transición de la acción'),('Persistencia','BusinessOperationsPort → improvement_actions'),('Seguimiento','OPEN → IN_PROGRESS → VERIFIED con medición'),('Evaluación P05','Compara resultado observado sin inferir causalidad')],
 '07': [('Jefatura con ADMIN','Selecciona periodo y conductor'),('React','Filtros compartidos → peticiones de dashboard'),('Backend','DashboardController → servicio de análisis'),('Consulta','Adaptador analítico calcula población consistente'),('PostgreSQL','Pedidos, rutas, evidencias y fuentes analíticas'),('Respuesta','Resumen y distribuciones con valores nulos explícitos'),('Decisión P07 y P05','Interpreta avance y excepciones; deriva a detalle o BI')],
 '08': [('Jefatura con ADMIN','Define periodo y conductor de análisis'),('React','BusinessIntelligence carga fuentes independientes'),('Backend','ReportController → reglas de indicadores'),('PostgreSQL','Vistas y funciones analíticas con fecha de Lima'),('Calidad del dato','Denominadores, metas y calidad GPS condicionan cálculo'),('Visualización','Indicadores, tendencias, causas y eficiencia'),('Mejora continua','Acción medible y seguimiento; exportación conserva vacíos')]
}

def flow_image(screen):
    rows=FLOW_LABELS[str(screen['id']).zfill(2)]
    im=Image.new('RGB',(1440,1120),'white'); d=ImageDraw.Draw(im)
    d.text((50,25),f"Flujo de la pantalla {screen['id']}  {screen['title']}",font=font(34,True),fill='#172b43')
    d.text((50,79),'Secuencia de responsabilidades y datos',font=font(25),fill='#45576a')
    for i,(role,action) in enumerate(rows):
        y=135+i*132
        d.rounded_rectangle((60,y,1380,y+104),radius=13,fill='#f0f5fa',outline='#bac9d9',width=2)
        d.ellipse((77,y+31,121,y+75),fill='#254e77')
        d.text((89,y+35),str(i+1),font=font(24,True),fill='white')
        d.text((147,y+12),role,font=font(25,True),fill='#172b43')
        lines=wrap_pixels(d,action,font(26),1180)
        for j,line in enumerate(lines[:2]):d.text((147,y+48+j*29),line,font=font(26),fill='#23384e')
        if i<len(rows)-1:
            d.line((720,y+106,720,y+129),fill='#59758c',width=4)
            d.polygon([(712,y+120),(728,y+120),(720,y+130)],fill='#59758c')
    path=ASSETS/f"flujo_{screen['id']}.png"; im.save(path);return path

def architecture_image():
    im=Image.new('RGB',(1600,1120),'white');d=ImageDraw.Draw(im)
    d.text((45,25),'Arquitectura del recorrido explicado',font=font(38,True),fill='#172b43')
    boxes=[(50,110,740,275,'Web React y TypeScript',['Componentes, hooks y servicios Axios','Nginx publica la web y /api → backend']),
           (860,110,1550,275,'App Flutter y Dart',['BLoC, casos de uso, GetIt y Dio','Caché local y cola de evidencia por cuenta']),
           (300,360,1300,625,'Backend Spring Boot WebFlux',['Entrada REST y autorización JWT','Aplicación y dominio con puertos de entrada y salida','Adaptadores R2DBC, OSRM, evidencias y eventos']),
           (50,735,740,935,'PostgreSQL de logística',['Pedidos, rutas, POD, GPS, incidencias y acciones','Migraciones Flyway y fuentes analíticas']),
           (860,735,1550,935,'Servicios de apoyo',['Keycloak emite tokens y mantiene su identidad separada','OSRM calcula el recorrido propuesto','Almacenamiento y mensajería según entorno local'])]
    for x,y,x2,y2,title,lines in boxes:
        d.rounded_rectangle((x,y,x2,y2),radius=16,fill='#f0f5fa',outline='#afc2d5',width=3)
        d.text((x+22,y+20),title,font=font(29,True),fill='#172b43')
        for j,line in enumerate(lines):
            for k,ln in enumerate(wrap_pixels(d,line,font(23),x2-x-44)):
                d.text((x+22,y+68+j*43+k*26),ln,font=font(23),fill='#23384e')
    for a,b in [((395,278),(650,355)),((1205,278),(950,355)),((650,630),(395,730)),((950,630),(1205,730))]:
        d.line([a,b],fill='#466b8e',width=5)
        x,y=b;d.polygon([(x-9,y-17),(x+9,y-17),(x,y)],fill='#466b8e')
    d.text((50,996),'Un backend modular con arquitectura hexagonal y dos clientes',font=font(30,True),fill='#172b43')
    d.text((50,1044),'Los procesos de negocio organizan el recorrido; EcoRoute es la herramienta que los apoya.',font=font(25),fill='#45576a')
    p=ASSETS/'arquitectura_exposicion.png';im.save(p);return p

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.footer_distance=Inches(.25)
styles=doc.styles
for style in styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Normal'].paragraph_format.line_spacing=1.07
for name,size in [('Title',27),('Heading 1',19),('Heading 2',14),('Heading 3',11.5)]:
    styles[name].font.size=Pt(size);styles[name].font.bold=True
    styles[name].paragraph_format.keep_with_next=True
    styles[name].paragraph_format.space_before=Pt(13)
    styles[name].paragraph_format.space_after=Pt(7)
styles['Subtitle'].font.size=Pt(15)
styles['Subtitle'].font.italic=False
for name in ['Caption']:
    styles[name].font.size=Pt(9);styles[name].font.color.rgb=RGBColor.from_string('45576A')
code_style=styles.add_style('CodeExcerpt',1)
code_style.font.name='Consolas';code_style.font.size=Pt(8.2)
code_style.paragraph_format.line_spacing=1.0
code_style.paragraph_format.space_after=Pt(1)
code_style.paragraph_format.left_indent=Inches(.08)
code_style.paragraph_format.right_indent=Inches(.08)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('EcoRoute  |  Guía de exposición  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='EcoRoute Guía de ocho pantallas para cuatro integrantes'
doc.core_properties.subject='Explicación de negocio y código del recorrido end to end'
doc.core_properties.author='Grupo 3'
MD=[]; manifest={'date':'2026-10-26','screens':[],'source_files':{},'warnings':[]}
PENDING_PAGE_BREAK=False

def para(text,style=None):
    p=doc.add_paragraph(str(text),style);MD.append(str(text)+'\n');return p
def heading(text,level=1):
    global PENDING_PAGE_BREAK
    text=re.sub(r'[^\w\s]',' ',text).replace('_',' ')
    text=re.sub(r'\s+',' ',text).strip()
    p=doc.add_heading(text,level)
    if PENDING_PAGE_BREAK:
        p.paragraph_format.page_break_before=True
        PENDING_PAGE_BREAK=False
    MD.append('#'*(level+1)+' '+text+'\n')
def bullets(items,numbered=False):
    for i,t in enumerate(items,1):
        if numbered:t=re.sub(r'^\s*\d+[.)]\s+','',str(t))
        p=doc.add_paragraph(style=None)
        p.paragraph_format.left_indent=Inches(.16)
        p.paragraph_format.first_line_indent=Inches(-.16)
        p.add_run((f'{i}. ' if numbered else '• ')+str(t))
        MD.append((f'{i}. ' if numbered else '- ')+str(t)+'\n')
def track(path):
    p=ROOT/path
    if not p.exists():raise FileNotFoundError(path)
    manifest['source_files'][str(path)]=hashlib.sha256(p.read_bytes()).hexdigest()
    return p
def image_block(path,caption,max_h=3.2):
    p=Path(path)
    if not p.is_absolute():p=ROOT/p
    with Image.open(p) as im:w,h=im.size
    width=min(7.05,max_h*w/h)
    pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.keep_with_next=True
    pp.add_run().add_picture(str(p),width=Inches(width))
    para(caption,'Caption')
    MD.append(f'![{caption}]({p.as_posix()})\n')
def break_page():
    global PENDING_PAGE_BREAK
    PENDING_PAGE_BREAK=True
    MD.append('\n---\n')
def table(headers,rows,widths=None,keep_header=False):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
    hp=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');hp.append(repeat)
    for vals in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,vals):c.text=str(v).replace('/','/\u200b')
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[ci])
            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');tcpr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'CED7DF');borders.append(b)
            tcpr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                if ri==0 and (headers[0]=='Capa' or keep_header):p.paragraph_format.keep_with_next=True
                for r in p.runs:r.font.size=Pt(9.5);r.bold=ri==0
    MD.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n')
    for vals in rows:MD.append('| '+' | '.join(str(x).replace('\n','<br>') for x in vals)+' |\n')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def code_excerpt(c):
    path=track(c['file']);lines=path.read_text(encoding='utf-8-sig').splitlines()
    assert 1<=c['start']<=c['end']<=len(lines),(c['file'],c['start'],c['end'])
    label=f"{c['file']}  líneas {c['start']} a {c['end']}"
    p=para(label);p.paragraph_format.keep_with_next=True
    for r in p.runs:r.font.size=Pt(9);r.bold=True
    MD.append('```'+path.suffix.lstrip('.')+'\n')
    for n in range(c['start'],c['end']+1):
        line=lines[n-1].replace('\t','    ')
        chunks=textwrap.wrap(line,width=94,replace_whitespace=False,drop_whitespace=False,break_long_words=True,break_on_hyphens=False) or ['']
        for j,chunk in enumerate(chunks):
            p=doc.add_paragraph((f'{n:>4}  ' if j==0 else '      ')+chunk,'CodeExcerpt')
            p.paragraph_format.keep_with_next=n<c['end'] or j<len(chunks)-1
        MD.append(f'{n:>4}  {line}\n')
    MD.append('```\n')
    bullets(c['explanation'])

para('EcoRoute','Title')
para('Guía de ocho pantallas para cuatro integrantes','Subtitle')
para('Explicación del negocio y del código de extremo a extremo')
para('Grupo MICOTRANS S A C   |   Grupo 3   |   Integrador I')
para('26 de octubre de 2026')
para('Cada integrante presenta dos pantallas reales conectadas al mismo recorrido de reparto. La guía indica qué mostrar, qué decir y qué código abrir para explicar cómo una acción de la interfaz llega al backend, se valida, persiste y vuelve a la pantalla.')
para('La distribución siguiente es una propuesta académica. Los roles de negocio pertenecen a las tareas del proceso y no cambian porque otro alumno explique su pantalla. Las capturas se conservan con la fecha de su ensayo; los pasos de demostración describen cómo repetir el recorrido con datos de prueba.')
table(['Integrante propuesto','Sus dos pantallas','Entrega al siguiente'],[
    [MEMBERS[1],'01 Pedidos\n02 Gestión de despacho','Pedido validado y ruta preparada'],
    [MEMBERS[2],'03 Mis rutas móvil\n04 Detalle de pedido y POD móvil','Resultado de entrega y evidencia'],
    [MEMBERS[3],'05 Incidencias\n06 Acciones de mejora','Excepción resuelta y acción medible'],
    [MEMBERS[4],'07 Dashboard logístico\n08 Inteligencia de negocio','Lectura de resultados y seguimiento']
],[2.25,2.45,2.4])
para('Orden sugerido de defensa: contexto compartido, pantallas 01 a 08 y preguntas. El guion completo sirve para estudiar; durante la revisión, cada integrante demuestra sus dos pantallas y abre los fragmentos que solicita el profesor. Ensayar también las transiciones y mantener un reparto equilibrado del tiempo disponible.')

break_page();heading('Preparación de la demostración')
bullets([
 'Abrir el proyecto en VS Code y ubicar las carpetas web-admin, mobile-app y src/main/java/com/ecoroute/backend. Las referencias de esta guía parten de la raíz del repositorio; Ctrl+P abre un archivo y Ctrl+G lleva a la línea indicada.',
 'Preparar una sesión web con permisos adecuados y una sesión móvil del conductor asignado. No proyectar contraseñas, tokens, archivos .env ni exportaciones de usuarios.',
 'Para una demostración nueva, usar un tracking único, fecha operativa vigente en Lima, una ruta PLANNED con recursos disponibles y dos pedidos de prueba: uno para entrega con POD y otro para un resultado fallido con motivo. El ensayo histórico conserva sus fechas e IDs originales.',
 'La aceptación móvil precede al despacho web. Coordinar ese cruce entre integrantes 1 y 2; no narrar que el conductor ya repartía antes de autorizar la salida.',
 'Si el profesor permite solo lectura, usar las capturas y el video del ensayo para mostrar cambios de estado y abrir el código para explicar su ejecución. Consultar una evidencia histórica no equivale a ejecutar una nueva transacción.',
 'Cada alumno debe explicar al menos una validación exitosa y una negativa de sus pantallas, el dato que entra, la tabla que cambia o se consulta y la forma en que la respuesta actualiza la interfaz.',
 'Los nombres de métodos y líneas corresponden al código inspeccionado para esta edición. Si se modifica el desarrollo, volver a generar la guía y revisar los fragmentos antes de exponer.'
],True)
heading('Hilo de negocio compartido',2)
para('P01 confirma la solicitud y entrega a P02 el mismo pedido PENDING con acuerdo vigente. P02 asigna, revisa y autoriza la jornada; P03 ejecuta entregas y registra evidencia o un resultado fallido. P04 resuelve incidencias conservando historia. P07 supervisa el avance y P05 evalúa resultados para proponer y comprobar acciones de mejora. P06 recursos, P08 accesos y P09 compras apoyan la operación; no se presentan como pantallas adicionales de esta selección.')
table(['Objeto','Secuencia que debe reconocer el grupo'],[
 ['Pedido','PENDING → ASSIGNED → IN_TRANSIT → DELIVERED, FAILED o RETURNED. PICKED_UP existe en el modelo y las transiciones manuales permitidas; no se fuerza como paso intermedio del despacho en cascada.'],
 ['Confirmación comercial','Aceptación, dirección y cobertura se guardan aparte; no existe un estado técnico CONFIRMED.'],
 ['Ruta','PLANNED → IN_PROGRESS → COMPLETED. CANCELLED es un cierre permitido bajo sus condiciones; una ruta cerrada no se reabre.'],
 ['Incidencia','La resolución documenta el resultado. Reintentar puede crear otro pedido vinculado al original; no sobrescribe un pedido terminal.'],
 ['Acción de mejora','OPEN → IN_PROGRESS → VERIFIED según validaciones y medición registradas.']
],[1.35,5.75])

break_page();heading('Arquitectura que todos deben explicar')
image_block(architecture_image(),'Arquitectura lógica del recorrido. El backend es una aplicación modular con puertos y adaptadores; los servicios de apoyo se despliegan por separado.',max_h=5.2)
para('La web usa React y TypeScript. Un evento de formulario modifica estado local, llama un servicio Axios y recibe JSON. La app móvil usa Flutter y Dart; sus páginas se coordinan con BLoC, casos de uso y repositorios inyectados mediante GetIt. Dio realiza la comunicación y los componentes de caché y sincronización conservan información local cuando corresponde.')
para('El backend combina Spring Boot, WebFlux y R2DBC con arquitectura hexagonal. Los controladores son adaptadores de entrada. Los puertos de entrada expresan operaciones del negocio y los casos de uso las implementan. Los puertos de salida definen necesidades como persistencia o ruteo. Los adaptadores R2DBC y OSRM conectan esas necesidades con tecnología externa. La independencia se refiere a responsabilidades y dependencias de código; no significa que cada clase sea un microservicio.')
para('PostgreSQL conserva datos logísticos y fuentes analíticas con migraciones Flyway. Keycloak administra identidad y emite tokens JWT; la base de identidad es distinta de las tablas de negocio. OSRM propone un recorrido sobre calles. La configuración local también contiene almacenamiento y mensajería de apoyo a evidencias y notificaciones; no toda lectura de pantalla pasa por una cola.')

heading('Cómo viaja una petición',2)
table(['Capa','Responsabilidad y ejemplo real'],[
 ['Interfaz','Orders.tsx recibe submit o una página Flutter despacha un evento a su BLoC. Muestra carga, error y resultado.'],
 ['Cliente HTTP','orderService.ts → api.ts añade Bearer. En móvil, Dio usa interceptores y configuración de entorno.'],
 ['Transporte local','Web: /api/orders pasa por Nginx a /orders del backend. En Android Emulator, 10.0.2.2 representa el host según ApiConfig.'],
 ['Seguridad y entrada','Spring valida JWT y rol. El controlador mapea HTTP y DTO; OperationalAccess controla pertenencia cuando corresponde.'],
 ['Aplicación y dominio','El caso de uso comprueba condiciones del negocio y coordina puertos. No debe depender de un botón ni de un select.'],
 ['Salida y base de datos','El adaptador convierte objetos y ejecuta consultas parametrizadas o repositorios Spring Data R2DBC. PostgreSQL aplica FK, unicidad y demás restricciones.'],
 ['Retorno','Mono o Flux entregan resultados a WebFlux; se serializan a JSON. React recarga consultas o BLoC emite estado y la página se reconstruye.']
],[1.3,5.8])
heading('Seguridad y programación reactiva',2)
para('Un token identifica la sesión y sus roles; no reemplaza comprobar que un conductor sea propietario de una ruta o pedido. En web, el acceso se organiza en App.tsx y ProtectedRoute; el backend vuelve a aplicar SecurityConfig, anotaciones y OperationalAccess. ADMIN, DISPATCHER y DRIVER son permisos técnicos. El negocio conserva responsabilidades de Jefatura, Dispatcher y Conductor.')
para('Mono representa una operación con cero o un resultado y Flux una secuencia de resultados. map transforma un valor; flatMap encadena una operación reactiva que depende del valor; then espera completar una operación y continúa; concatMap conserva ejecución secuencial dentro de la cadena. @Transactional delimita una operación de base de datos reactiva cuando usa la configuración transaccional correspondiente. Un encadenamiento de dos peticiones HTTP no se convierte por ello en una sola transacción.')
para('RestExceptionHandler diferencia argumentos inválidos 400, conflictos de negocio o integridad 409 y recursos ausentes 404. La seguridad responde 401 o 403 según autenticación o autorización. El frontend conserva mensajes y estados de carga por fuente; una respuesta fallida no debe dibujarse como un cero de negocio ni una entrega exitosa.')

heading('Ejemplo de una petición y su resultado',2)
para('El siguiente JSON es ilustrativo y usa datos de prueba. Representa el contrato de alta del pedido; no registra una ejecución adicional ni prescribe reutilizar el tracking. El navegador enviaría POST /api/orders y el controlador recibiría POST /orders. La cabecera de autorización se omite para proteger la sesión.')
sample=json.dumps({'trackingNumber':'EXPO-PEDIDO-001','routeId':None,'recipientName':'Receptor de prueba','recipientPhone':'999123456','recipientEmail':'prueba@example.invalid','deliveryAddress':'Dirección de demostración','deliveryCity':'Lima','deliveryDistrict':'Lima','latitude':-12.0464,'longitude':-77.0428,'priority':1},ensure_ascii=False,indent=2)
for line in sample.splitlines():
    p=doc.add_paragraph(line,'CodeExcerpt');p.paragraph_format.keep_with_next=True
MD.append('```json\n'+sample+'\n```\n')
para('La respuesta contiene un id entero asignado por la persistencia, status PENDING y routeId null. El componente conserva ese id y envía los campos recipientRuc, customerAccepted, addressVerified, coverageVerified y agreementNote a PATCH /orders/{id}/business-data. El actor y la fecha de confirmación se determinan en servidor. Después, la asignación envía {"routeId": idDeRuta} a PATCH /orders/{id}/route. Así se puede seguir un solo pedido entre peticiones sin confundir los DTO ni crear registros duplicados para cada proceso.')
heading('Cómo mostrar la trazabilidad en la defensa',2)
bullets(['En la web, abrir las herramientas de desarrollo del navegador y filtrar Network por orders, routes o reports. Mostrar método, ruta, cuerpo sanitizado y respuesta, evitando cabeceras con tokens.', 'En VS Code, abrir el componente del botón y seguir la llamada al servicio; buscar la misma ruta en @RequestMapping, @PostMapping o @PatchMapping. Identificar el DTO recibido y el método del caso de uso.', 'Seguir el puerto de salida hasta su implementación y localizar la entidad, mapper o consulta SQL. Explicar qué FK enlaza los objetos y qué evidencia registra la operación.', 'Volver a la interfaz y mostrar cómo llega el resultado: recarga de useQuery en React o cambio de estado/relectura en Flutter. Si falla, explicar qué parte quedó persistida y qué permanece pendiente.'],True)

SOURCE_MAP={
 '01': [('web-admin/src/pages/Orders.tsx','CreateOrder.submit y BusinessEditor','Formulario, revisión y detalle'),('web-admin/src/services/orderService.ts','createOrder y saveBusinessData','Transporte JSON'),('web-admin/src/hooks/useQuery.ts','useQuery','Lecturas y recarga'),('src/main/java/com/ecoroute/backend/infrastructure/input/rest/OrderController.java','createOrder y assign','Entrada REST'),('src/main/java/com/ecoroute/backend/infrastructure/input/rest/CreateOrderRequest.java','CreateOrderRequest','Validación DTO'),('src/main/java/com/ecoroute/backend/application/usecases/OrderUseCasesImpl.java','createOrder y assignOrder','Reglas y transacción'),('src/main/java/com/ecoroute/backend/application/services/BusinessWorkflowService.java','confirm','Acuerdo comercial'),('src/main/java/com/ecoroute/backend/domain/ports/out/OrderRepository.java','OrderRepository','Puerto de persistencia'),('src/main/java/com/ecoroute/backend/infrastructure/output/persistence/R2dbcOrderRepository.java','save','Adaptación dominio a entidad'),('src/main/java/com/ecoroute/backend/infrastructure/output/persistence/R2dbcBusinessWorkflowAdapter.java','saveConfirmation','UPSERT de revisión y auditoría')],
 '02': [('web-admin/src/pages/Routes.tsx','CreateRouteForm y ReviewPanel','Preparación y despacho'),('web-admin/src/services/routeService.ts','createRoute, approveRoute, updateRouteStatus','Peticiones'),('src/main/java/com/ecoroute/backend/infrastructure/input/rest/RouteController.java','createRoute, updateStatus, getRoutePath','Entrada REST de ruta'),('src/main/java/com/ecoroute/backend/infrastructure/input/rest/BusinessWorkflowController.java','approve y accept','Roles y DTO de revisión'),('src/main/java/com/ecoroute/backend/application/usecases/RouteUseCasesImpl.java','updateRouteStatus','Controles y cascada'),('src/main/java/com/ecoroute/backend/application/services/BusinessWorkflowService.java','approve y accept','Aprobación y aceptación'),('src/main/java/com/ecoroute/backend/application/services/RoutePlannerService.java','plan','Ruteo y degradación'),('src/main/java/com/ecoroute/backend/infrastructure/output/routing/OsrmRoutingAdapter.java','optimizePath','Integración de ruteo'),('src/main/java/com/ecoroute/backend/infrastructure/output/persistence/R2dbcRouteRepository.java','save y updatePathInfo','Persistencia de ruta'),('src/main/java/com/ecoroute/backend/infrastructure/output/persistence/R2dbcBusinessWorkflowAdapter.java','approve, accept, invalidate','Revisión vigente')]
}

for s in SCREENS:
    sid=str(s['id']).zfill(2); s['id']=sid
    break_page();heading(f"Pantalla {sid} {s['title']}")
    para(f"Integrante {s['student']}  {MEMBERS[int(s['student'])]}")
    para(s['purpose'])
    table(['Aspecto','Aplicación al negocio'],[['Plataforma y acceso',s['platform']+' · '+s['route']],['Actor responsable',s['actor']],['Proceso y conexión',s['processes']]],[1.35,5.75])
    shot=s.get('screenshot')
    if shot and shot.get('file'):
        capture=track(shot['file'])
        caption=shot['caption']
        if sid=='06':
            # Show the usable area of the historically redacted screenshot.
            # The source stays untouched and no hidden values are reconstructed.
            fragment=ASSETS/'captura_06_contexto.png'
            with Image.open(capture) as original:original.crop((0,0,original.width,316)).save(fragment)
            capture=fragment
            caption='Fragmento superior de la pantalla real de Acciones de mejora. La captura histórica conserva ocultos sus registros; se muestran el acceso, el botón Nueva acción y las columnas del listado. Los formularios y transiciones se explican con el código actual.'
        image_block(capture,caption,max_h=5.3 if s['platform'].lower().find('móvil')>=0 or sid in ['03','04'] else 4.1)
    heading('Condiciones de entrada',2);bullets(s['preconditions'])
    heading('Demostración paso a paso',2);bullets(s['demo'],True)
    heading('Flujo de extremo a extremo',2)
    image_block(flow_image(s),'Secuencia explicativa del recorrido de esta pantalla. No sustituye el diagrama BPMN de negocio; detalla las capas del desarrollo.',max_h=5.0)
    bullets(s['flow'],True)
    heading('Peticiones que deben reconocer',2)
    table(['Método','Ruta del backend','Finalidad'],[[x['method'],x['path'][4:] if x['path'].startswith('/api/') else x['path'],x['purpose']] for x in s['endpoints']],[.85,2.5,3.75],keep_header=sid in ['05','07'])
    para('En la web, el navegador antepone /api a estas rutas y Nginx lo retira al enviar al backend. En móvil se utiliza la URL base configurada. Los valores entre llaves identifican recursos y no autorizan por sí mismos su acceso.','Caption')
    for title,key in [('Interfaz y estado del frontend','frontend'),('Reglas y ejecución del backend','backend'),('Persistencia y datos relacionados','persistence')]:
        heading(title,2)
        for text in s[key]:para(text)
        if key=='frontend' and sid in ['01','02']:
            para('En el diseño web, global.css utiliza .operation-form como CSS Grid con columnas auto-fit y un mínimo de 210 px. .full ocupa el ancho disponible, .operation-toolbar usa Flexbox y .table-scroll permite desplazamiento horizontal. El modal tiene ancho máximo de 960 px, altura máxima de 92vh y scroll interno. Los mensajes .source-error y los botones deshabilitados comunican el estado de la operación. Estas clases organizan la presentación; los permisos y las reglas del negocio se vuelven a comprobar en backend.')
    heading('Archivos para abrir en VS Code',2)
    source_map=s.get('source_map') or [dict(zip(['file','symbol','purpose'],x)) for x in SOURCE_MAP.get(sid,[])]
    if not source_map:
        source_map=[{'file':c['file'],'symbol':'Fragmento explicado a continuación','purpose':'Seguir el código real'} for c in s['code']]
    for position,item in enumerate(source_map,1):
        track(item['file'])
        p=para(item['file']);p.paragraph_format.keep_with_next=True
        for run in p.runs:run.font.size=Pt(9.5);run.bold=True
        purpose=item.get('purpose','')
        # Long secondary navigation notes remain complete in the companion source guide.
        if position>12:
            short=f"Abrir {item.get('symbol','el archivo')}."
            para(short)
            MD.append(purpose+'\n')
        else:
            para(f"Abrir {item.get('symbol','el archivo')}. {purpose}.")
    heading('Fragmentos de código para defender',2)
    for i,c in enumerate(s['code'],1):
        heading(f'Fragmento {i}',3);code_excerpt(c)
    heading('Validaciones y recorridos negativos',2);bullets(s['exceptions'])
    heading('Guion para decir al profesor',2)
    for text in s['speech']:para(text)
    heading('Preguntas probables del profesor',2)
    for item in s['questions']:
        p=para(item['q']);p.paragraph_format.keep_with_next=True
        for r in p.runs:r.bold=True
        para(item['a'])
    heading('Pruebas y evidencia para localizar',2)
    for t in s.get('tests',[]):
        if t.get('file'):
            f=track(t['file']);content=f.read_text(encoding='utf-8-sig') if f.suffix not in ['.png','.jpg'] else ''
            name=t.get('name','')
            names=[part.strip() for part in name.split(';') if part.strip()]
            if names and not all(part in content for part in names) and '/test' in t['file']:manifest['warnings'].append({'screen':sid,'test_symbol_not_literal':name,'file':t['file']})
        p=para(t.get('name','Prueba relacionada'));p.paragraph_format.keep_with_next=True
        for r in p.runs:r.bold=True
        description=para(t.get('scope',''))
        if t.get('file'):description.paragraph_format.keep_with_next=True
        if t.get('file'):para(t['file'],'Caption')
    para('Al abrir una prueba, identificar los datos preparados, la acción que se ejecuta y la aserción que comprueba el resultado. Distinguir los mocks o respuestas simuladas de una consulta al servidor real.','Caption')
    heading('Entrega al siguiente expositor',2);para(s['handoff'])
    manifest['screens'].append({'id':sid,'title':s['title'],'student':s['student'],'snippets':len(s['code']),'endpoints':len(s['endpoints'])})

break_page();heading('Ensayo integrado del grupo')
table(['Momento','Responsable','Qué debe poder demostrar'],[
 ['Ingreso del pedido',MEMBERS[1],'Alta sin ruta, revisión comercial, mismo ID y bloqueo de asignación cuando faltan condiciones.'],
 ['Preparación',MEMBERS[1],'Ruta PLANNED, asignación ASSIGNED y revisión administrativa de la composición vigente.'],
 ['Aceptación cruzada',MEMBERS[2],'Ruta propia en móvil, aceptación después de aprobación; retorno al integrante 1 para despachar.'],
 ['Ejecución',MEMBERS[2],'Ruta IN_PROGRESS, pedido IN_TRANSIT, captura POD y diferencia entre cola local y confirmación de servidor.'],
 ['Excepción',MEMBERS[3],'Incidencia con motivo, resolución permitida y conservación del pedido original.'],
 ['Mejora',MEMBERS[3],'Acción con responsable y seguimiento; medición registrada sin atribuir causalidad automática.'],
 ['Control y análisis',MEMBERS[4],'Filtros consistentes, indicador con denominador y calidad GPS visible; explicar el dato fuente.']
],[1.05,1.85,4.2])
heading('Lista de comprobación antes de exponer',2)
bullets([
 'Cada integrante ubica sin ayuda sus dos componentes o páginas y el método que ejecuta la acción principal.',
 'Cada integrante sigue al menos una petición desde el servicio del cliente hasta el controlador y la persistencia, y explica el camino de vuelta.',
 'El grupo conoce los IDs de ensayo, distingue una demostración nueva de una captura histórica y no renombra estados para ajustarlos al relato.',
 'La interfaz, el BPMN y el código se explican en niveles diferentes: responsabilidad humana del negocio, control automatizado y mecanismo técnico.',
 'No se confunde distancia planificada con observada, un dato nulo con cero, una evidencia local con confirmación remota ni un cierre de ruta con éxito total.',
 'Las preguntas sobre implementación deben responderse con el fragmento y sus precondiciones. Si una operación pertenece a otro proceso, señalar el traspaso y el registro que lo conecta.'
])
heading('Alcance de la evidencia',2)
para('La exposición se apoya en código del workspace y en capturas del ensayo local del 22 de octubre de 2026. Ese ensayo documentó una ruta con una entrega con POD y otra fallida; sus verificaciones de servidor registraron 26 comprobaciones. Esta guía no ejecuta nuevamente ese recorrido. La aceptación identificada por MICOTRANS y las pruebas de campo físico conservan su estado pendiente; no se deducen de una prueba en emulador. La conexión nativa de Power BI tampoco se da por comprobada al explicar el módulo web de inteligencia de negocio.')
para('Las fuentes de consulta incluyen docs/control_integral_20261022/control_procesos.md, docs/videos/ecoroute_e2e_20261022/README.md y docs/control_actualizacion/EXCEPCIONES_Y_TRAZABILIDAD.md. Los diagramas de esta guía explican el recorrido técnico de las ocho pantallas y sus traspasos; no alteran los nueve procesos BPMN ni sustituyen sus condiciones de cierre.')
heading('Términos que conviene dominar',2)
table(['Término','Explicación para la defensa'],[
 ['Componente y hook','El componente construye interfaz y maneja eventos; el hook encapsula estado y efectos reutilizables.'],
 ['BLoC','Recibe eventos, coordina operaciones y emite estados para reconstruir la interfaz Flutter.'],
 ['DTO y entidad','El DTO describe datos de transporte; la entidad describe persistencia. No son sinónimos de la regla del negocio.'],
 ['Puerto y adaptador','El puerto declara una necesidad u operación. El adaptador la conecta con HTTP, SQL u otro mecanismo concreto.'],
 ['Transacción','Agrupa escrituras participantes para confirmar o revertir conjuntamente dentro de su frontera.'],
 ['Idempotencia','Repetir una operación bajo una identidad y reglas definidas conserva su efecto lógico; no significa ignorar cualquier error.'],
 ['POD','Prueba de entrega con datos del receptor, evidencia y ubicación, validada por el servidor según el flujo implementado.'],
 ['KPI y denominador','El KPI resume una población definida. Sin denominador o calidad suficiente no corresponde mostrar un valor ficticio.']
],[1.35,5.75])

dest=OUT/'EcoRoute_Guia_8_Pantallas_4_Integrantes.docx'
for tree in (doc.element,doc.styles.element):
    for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(dest)
(OUT/'EcoRoute_Guia_8_Pantallas_4_Integrantes.md').write_text('\n'.join(MD),encoding='utf-8')
manifest['outputs']=[str(dest.relative_to(ROOT)),'docs/exposicion_8_pantallas/EcoRoute_Guia_8_Pantallas_4_Integrantes.md']
manifest['paragraphs']=len(doc.paragraphs);manifest['tables']=len(doc.tables)
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'outputs':manifest['outputs'],'screens':len(SCREENS),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'warnings':manifest['warnings']},ensure_ascii=False))
