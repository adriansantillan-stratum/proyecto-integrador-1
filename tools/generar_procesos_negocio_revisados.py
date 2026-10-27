"""Procesos de negocio EcoRoute/MICOTRANS: BPMN, memoria y guion alineados."""
from __future__ import annotations

from collections import defaultdict, deque
from html import escape
from pathlib import Path
from textwrap import wrap
import math
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs' / 'entregables_procesos_ecoroute'
OUT = BASE / 'negocio_revisado'
IMG = OUT / 'imagenes'
BPMN = OUT / 'bpmn'
for directory in (OUT, IMG, BPMN):
    directory.mkdir(parents=True, exist_ok=True)

NAVY = '#12395F'; BLUE = '#2B6CB0'; TEAL = '#0F766E'; ORANGE = '#C4551A'
RED = '#B42318'; DARK = '#1F2937'; MUTED = '#5C6B7B'; PALE = '#F4F8FB'

def N(id, kind, label, lane, col, offset=0):
    return {'id': id, 'kind': kind, 'label': label, 'lane': lane, 'col': col, 'offset': offset}

def F(source, target, label=''):
    return (source, target, label)

PROCESSES = [
  {
    'id':'P01', 'title':'Confirmar pedidos',
    'owner':'Dispatcher',
    'start':'Solicitud de servicio recibida',
    'success':'Pedido validado, registrado sin ruta y apto para planificar',
    'lanes':['Dispatcher'],
    'nodes':[
      N('S','start','Solicitud recibida',0,0), N('A1','task','Recibir solicitud y condiciones',0,1),
      N('A2','task','Revisar datos mínimos',0,2), N('G1','gateway','¿Datos completos?',0,3),
      N('A3','task','Solicitar corrección',0,4,145), N('G2','gateway','¿Cliente corrige a tiempo?',0,5,145),
      N('E1','end','Solicitud no admitida',0,6,270), N('A4','task','Revisar dirección y duplicidad',0,4,-140),
      N('G3','gateway','¿Dirección utilizable?',0,5,-140), N('A5','task','Aclarar ubicación con cliente',0,6,145),
      N('G4','gateway','¿Ubicación aclarada?',0,7,145), N('E2','end','Solicitud no viable',0,8,270),
      N('A6','task','Validar cobertura y ventana',0,6), N('G5','gateway','¿Servicio factible?',0,7),
      N('A7','task','Proponer fecha o condición alternativa',0,8,145),
      N('G6','gateway','¿Cliente acepta alternativa?',0,9,145), N('E3','end','Solicitud desistida',0,11,145),
      N('A8','send','Comunicar condiciones al cliente',0,8),
      N('A8b','receive','Recibir respuesta del cliente',0,9),
      N('G7','gateway','¿Cliente confirma?',0,10),
      N('A9','task','Registrar pedido sin ruta (PENDING)',0,11), N('E4','end','Pedido validado → P02',0,12),
    ],
    'flows':[
      F('S','A1'),F('A1','A2'),F('A2','G1'),F('G1','A4','Sí'),F('G1','A3','No'),
      F('A3','G2'),F('G2','A2','Sí'),F('G2','E1','No'),
      F('A4','G3'),F('G3','A6','Sí'),F('G3','A5','No'),F('A5','G4'),
      F('G4','A4','Sí'),F('G4','E2','No'),F('A6','G5'),F('G5','A8','Sí'),
      F('G5','A7','No'),F('A7','G6'),F('G6','A6','Sí'),F('G6','E3','No'),
      F('A8','A8b'),F('A8b','G7'),F('G7','A9','Sí'),F('G7','E3','No'),F('A9','E4'),
    ],
    'rules':[
      ('Datos incompletos','Se solicita corrección con plazo; si no se subsana, se cierra como no admitida.'),
      ('Dirección no utilizable','Se pide aclaración; si no se aclara, se cierra como no viable.'),
      ('Cobertura o ventana inviable','Se propone alternativa; la aceptación retorna a validación y el rechazo cierra la solicitud.'),
      ('Cliente no confirma condiciones finales','Se cierra la solicitud como desistida; no se registra un pedido por una aceptación supuesta.'),
    ],
    'handoff':'Solo E4 entrega a P02 un pedido registrado PENDING y validado por el Dispatcher. E1, E2 y E3 cierran la solicitud; el motivo y la aceptación comercial requieren evidencia fuera del estado actual de EcoRoute.',
    'current':'El formulario y CSV admiten routeId vacío: se crea PENDING. Elegir una ruta existente crea ASSIGNED. No hay estado CONFIRMED ni confirmación comercial automatizada; se requiere definir quién acepta condiciones y dónde se conserva la constancia.',
    'wireframes':['P01_wireframes_flujo_alineado.png'],
    'speech':'El Dispatcher recibe la solicitud, comprueba los datos y acuerda con el cliente una fecha y dirección viables. Luego registra el pedido sin ruta: EcoRoute lo guarda PENDING. La confirmación de negocio no es un estado separado de la aplicación. La ruta se crea en P02; la asignación posterior de un PENDING es una brecha del desarrollo actual.',
    'external':{'name':'Cliente solicitante','outgoing':['A8'],'incoming':['A8b']},
  },
  {
    'id':'P02', 'title':'Preparar la jornada de reparto',
    'owner':'Dispatcher con aprobación de Jefatura logística',
    'start':'Pedidos PENDING validados para una jornada',
    'success':'Ruta entregada al conductor e iniciada',
    'lanes':['Dispatcher','Jefatura logística','Conductor','EcoRoute (sistema)'],
    'nodes':[
      N('S','start','Pedidos validados PENDING',0,0),N('A1','task','Consolidar demanda del día',0,1),
      N('A2','task','Seleccionar conductor y vehículo',0,2),N('G1','gateway','¿Recursos disponibles?',0,3),
      N('A3','task','Ajustar carga o recursos',0,4,95),N('G2','gateway','¿Ajuste viable?',0,5,95),
      N('E1','end','Pedidos en cartera; nueva fecha',0,6,95),
      N('A4','task','Crear ruta PLANNED para la fecha',0,4,-95),
      N('A5','task','Vincular pedidos PENDING a ruta [objetivo]',0,5,-95),
      N('G3','gateway','¿Pedidos vinculados?',0,6,-95),
      N('A6','task','Corregir datos o asignación',0,7,75),
      N('E2','end','Asignación pendiente; no despachar',0,8,75),
      N('A7','service','Calcular recorrido sobre red vial',3,7,-95),
      N('A8','task','Revisar capacidad y ventanas',0,8,-95),
      N('G4','gateway','¿Plan viable?',0,9,-95),
      N('A9','task','Corregir secuencia o reparto',0,10,75),
      N('A10','task','Autorizar plan y salida',1,10,-95),
      N('G5','gateway','¿Despacho aprobado?',1,11,-95),
      N('A11','task','Emitir hoja de ruta',0,12,-95),
      N('A12','task','Revisar hoja, carga y vehículo',2,13,-95),
      N('G6','gateway','¿Conductor puede salir?',2,14,-95),
      N('A13','task','Gestionar sustituto o nueva fecha',0,15,75),
      N('G7','gateway','¿Sustituto disponible?',0,16,75),
      N('E3','end','Despacho suspendido; reprogramar',0,17,75),
      N('A14','task','Iniciar despacho en web',0,15,-95),
      N('E4','end','Ruta IN_PROGRESS → P03',0,16,-95),
    ],
    'flows':[
      F('S','A1'),F('A1','A2'),F('A2','G1'),F('G1','A4','Sí'),F('G1','A3','No'),
      F('A3','G2'),F('G2','A2','Sí'),F('G2','E1','No'),F('A4','A5'),F('A5','G3'),
      F('G3','A7','Sí'),F('G3','A6','No'),F('A6','E2'),F('A7','A8'),
      F('A8','G4'),F('G4','A10','Sí'),F('G4','A9','No'),F('A9','A7'),
      F('A10','G5'),F('G5','A11','Sí'),F('G5','A9','No'),F('A11','A12'),
      F('A12','G6'),F('G6','A14','Sí'),F('G6','A13','No'),
      F('A13','G7'),F('G7','A2','Sí'),F('G7','E3','No'),F('A14','E4'),
    ],
    'rules':[
      ('Recursos insuficientes','El Dispatcher ajusta carga o recursos; si no es viable, deja pedidos en cartera y acuerda nueva fecha.'),
      ('Pedidos no vinculados','Se corrige la asignación; si persiste, no se despacha. La vinculación posterior de PENDING requiere desarrollo.'),
      ('Plan no viable o no aprobado','Se corrige la secuencia o el reparto y se recalcula antes de autorizar la salida.'),
      ('Conductor no puede salir','Se busca sustituto y se vuelve a planificar; sin sustituto, se suspende el despacho y se acuerda nueva fecha.'),
    ],
    'handoff':'Solo E4 inicia P03. E1 conserva pedidos en cartera; E2 bloquea el despacho por falta de asignación; E3 suspende la jornada y exige nueva fecha. Ningún final negativo equivale a entrega fallida.',
    'current':'La web crea rutas PLANNED y permite registrar pedidos directamente con routeId; al iniciar despacho calcula OSRM y cambia los pedidos vinculados a IN_TRANSIT. No existe endpoint ni interfaz para asociar luego un PENDING a la ruta. La aprobación formal y el acuse del conductor son controles de negocio propuestos.',
    'wireframes':['P02_wireframes_flujo_alineado.png','P03_wireframes.png'],
    'speech':'Después de validar el pedido, el Dispatcher reúne los PENDING, elige recursos y crea una ruta PLANNED. El flujo objetivo vincula entonces esos pedidos a la ruta, revisa el recorrido y autoriza la salida. En la versión actual, la vinculación solo ocurre al registrar un pedido con una ruta ya existente; no se puede asignar posteriormente el PENDING. Por eso el BPMN marca este paso como objetivo y el despacho no debe darse por completo sin esa asociación.',
  },
  {
    'id':'P03', 'title':'Ejecutar entregas',
    'owner':'Conductor con supervisión del Dispatcher',
    'start':'Parada pendiente de una ruta iniciada',
    'success':'Paradas cerradas con resultado trazable',
    'lanes':['Conductor','EcoRoute (sistema)'],
    'nodes':[
      N('S','start','Ruta iniciada',0,0),N('A1','task','Tomar siguiente parada',0,1),
      N('A2','task','Desplazarse y verificar acceso',0,2),N('G1','gateway','¿Destino accesible?',0,3,-85),
      N('A3','task','Registrar impedimento',0,4,85),N('E1','end','Incidencia → P04',0,5,85),
      N('A4','task','Contactar al receptor',0,4,-85),N('G2','gateway','¿Receptor disponible?',0,5,-85),
      N('A5','task','Registrar ausencia',0,6,85),N('E2','end','No entrega → P04',0,7,85),
      N('A6','task','Entregar bultos y validar recepción',0,6,-85),N('G3','gateway','¿Entrega aceptada?',0,7,-85),
      N('A7','task','Registrar rechazo y motivo',0,8,85),N('E3','end','Rechazo → P04',0,9,85),
      N('A8','task','Capturar foto, firma y receptor',0,8,-85),N('G4','gateway','¿POD suficiente?',0,9,-85),
      N('A9','task','Completar evidencia',0,10,85),N('G5','gateway','¿Puede completarse?',0,11,85),
      N('A10','task','Registrar prueba incompleta',0,12,85),N('E4','end','POD observado → P04',0,13,85),
      N('A11','service','Asociar GPS/hora y sincronizar',1,10),N('G6','gateway','¿Sincronización confirmada?',1,11),
      N('A12','service','Conservar POD pendiente',1,12,75),N('E5','end','POD pendiente → P04',1,13,75),
      N('A13','task','Cerrar parada',0,12,-85),N('G7','gateway','¿Quedan paradas?',0,13,-85),
      N('E6','end','Ruta cerrada → P05',0,14),
    ],
    'flows':[
      F('S','A1'),F('A1','A2'),F('A2','G1'),F('G1','A4','Sí'),F('G1','A3','No'),F('A3','E1'),
      F('A4','G2'),F('G2','A6','Sí'),F('G2','A5','No'),F('A5','E2'),F('A6','G3'),
      F('G3','A8','Sí'),F('G3','A7','No'),F('A7','E3'),F('A8','G4'),F('G4','A11','Sí'),
      F('G4','A9','No'),F('A9','G5'),F('G5','A8','Sí'),F('G5','A10','No'),F('A10','E4'),
      F('A11','G6'),F('G6','A13','Sí'),F('G6','A12','No'),F('A12','E5'),
      F('A13','G7'),F('G7','A1','Sí'),F('G7','E6','No'),
    ],
    'rules':[
      ('Acceso, receptor o aceptación fallidos','Se documenta la causal y se abre P04 para decidir reintento, reprogramación o cierre.'),
      ('POD incompleto','El conductor intenta completar la evidencia; si no puede, la prueba queda observada y pasa a P04.'),
      ('Sin sincronización','Se conserva pendiente y P04 controla la recuperación; no se considera POD validado.'),
      ('Más paradas','La ruta regresa a la siguiente parada dentro de P03. Si no quedan, se cierra para análisis P05.'),
    ],
    'handoff':'El modelo se lee por parada. Una excepción abre P04 para esa parada; el Dispatcher puede permitir que el conductor siga atendiendo las demás. Cuando ya no quedan paradas, P05 evalúa el resultado del periodo.',
    'current':'La app móvil consulta rutas y permite cambio de estado con foto y firma; el backend registra POD. La cola offline garantizada y GPS obligatorio deben validarse antes de tratarlos como control operativo cumplido.',
    'wireframes':['P04_wireframes.png'],
    'speech':'La entrega se gestiona por parada. El conductor comprueba acceso, receptor y aceptación antes de registrar foto, firma y datos de entrega. Si hay rechazo, ausencia o evidencia insuficiente, el caso pasa a incidencias. La parada se cierra como entregada solo cuando la evidencia se confirma; luego se continúa a la siguiente parada o se cierra la ruta.',
  },
  {
    'id':'P04', 'title':'Resolver incidencias',
    'owner':'Dispatcher con escalamiento a Jefatura logística',
    'start':'Excepción de entrega, POD o despacho',
    'success':'Caso resuelto, reprogramado o cerrado con autorización',
    'lanes':['Conductor','Dispatcher','Jefatura logística','EcoRoute (sistema)'],
    'nodes':[
      N('S','start','Excepción registrada',0,0),N('A1','task','Describir causa y adjuntar sustento',0,1),
      N('A2','task','Revisar evidencia y urgencia',1,2),N('G1','gateway','¿Sustento suficiente?',1,3),
      N('A3','task','Solicitar complemento',0,4,65),N('G2','gateway','¿Puede completarse?',0,5,65),
      N('A4','service','Marcar evidencia incompleta',3,6,65),
      N('A5','task','Evaluar impacto y alternativas',1,6),N('G3','gateway','¿Recuperable hoy?',1,7),
      N('A6','task','Coordinar reintento o sustituto',1,8),N('G4','gateway','¿Cliente y recursos aceptan?',1,9),
      N('E1','end','Reintento autorizado → P03',1,10),
      N('A7','task','Acordar nueva fecha',2,10,65),N('G5','gateway','¿Reprogramación aceptada?',2,11,65),
      N('E2','end','Nueva fecha → P02',2,12,-95),
      N('A8','task','Autorizar cierre fallido con motivo',2,12),
      N('E3','end','Fallo cerrado → P05',2,13),
    ],
    'flows':[
      F('S','A1'),F('A1','A2'),F('A2','G1'),F('G1','A5','Sí'),F('G1','A3','No'),
      F('A3','G2'),F('G2','A1','Sí'),F('G2','A4','No'),F('A4','A5'),
      F('A5','G3'),F('G3','A6','Sí'),F('G3','A7','No'),F('A6','G4'),
      F('G4','E1','Sí'),F('G4','A7','No'),F('A7','G5'),F('G5','E2','Sí'),
      F('G5','A8','No'),F('A8','E3'),
    ],
    'rules':[
      ('Sustento insuficiente','Se pide complemento; si no puede obtenerse, se marca la brecha y se continúa la decisión bajo responsabilidad del Dispatcher.'),
      ('Reintento hoy posible','Debe existir aceptación del cliente y recursos disponibles; entonces el caso vuelve a P03.'),
      ('Reintento no viable','Se acuerda nueva fecha y vuelve a P02; si no se acepta, Jefatura autoriza el cierre fallido.'),
    ],
    'handoff':'E1 vuelve a P03; E2 entra a P02 en nueva fecha; E3 alimenta P05 para análisis de fallas. El caso no desaparece de la trazabilidad.',
    'current':'El sistema registra estados y razones y muestra alertas/reportes. La gestión integral de caso, aprobación de reprogramación y reasignación con historial requieren validación de flujo y control administrativo.',
    'wireframes':['P05_wireframes.png'],
    'speech':'Una excepción no es el final automático de un pedido. El Dispatcher revisa su sustento, intenta recuperar la entrega en la jornada y, si no puede, busca nueva fecha. El jefe de logística autoriza un cierre fallido cuando ya no hay alternativa aceptada. El resultado conduce de forma explícita a P03, P02 o P05.',
  },
  {
    'id':'P05', 'title':'Evaluar resultados',
    'owner':'Jefatura logística',
    'start':'Cierre de jornada o periodo de evaluación',
    'success':'Resultado revisado y acción verificada',
    'lanes':['Dispatcher','Jefatura logística','EcoRoute (sistema)'],
    'nodes':[
      N('S','start','Cierre de periodo',2,0),N('A1','service','Consolidar pedidos, rutas y POD',2,1),
      N('G1','gateway','¿Datos confiables?',2,2),N('A2','task','Solicitar correcciones',0,3,65),
      N('G2','gateway','¿Se corrigieron?',0,4,65),N('E1','end','Cierre provisional',0,5,65),
      N('A3','service','Calcular indicadores',2,3),N('A4','task','Comparar resultados y metas',1,4),
      N('G3','gateway','¿Hay desvío relevante?',1,5),N('A5','task','Registrar buena práctica',1,6,90),
      N('E2','end','Periodo conforme',1,7,90),N('A6','task','Analizar causa raíz',0,6,-90),
      N('A7','task','Proponer acción y responsable',0,7,-90),N('G4','gateway','¿Acción aprobada?',1,8),
      N('A8','task','Devolver para ajuste',0,9,65),N('A9','task','Asignar plazo y recursos',1,9),
      N('A10','task','Ejecutar y medir la acción',0,10),N('G5','gateway','¿Mejora efectiva?',1,11),
      N('A11','task','Revisar hipótesis y acción',0,12,65),N('A12','task','Cerrar acción y estandarizar',1,12),
      N('E3','end','Acción verificada',1,13),
    ],
    'flows':[
      F('S','A1'),F('A1','G1'),F('G1','A3','Sí'),F('G1','A2','No'),
      F('A2','G2'),F('G2','A1','Sí'),F('G2','E1','No'),F('A3','A4'),
      F('A4','G3'),F('G3','A6','Sí'),F('G3','A5','No'),F('A5','E2'),
      F('A6','A7'),F('A7','G4'),F('G4','A9','Sí'),F('G4','A8','No'),
      F('A8','A7'),F('A9','A10'),F('A10','G5'),F('G5','A12','Sí'),
      F('G5','A11','No'),F('A11','A6'),F('A12','E3'),
    ],
    'rules':[
      ('Datos no confiables','Se corrigen dentro del plazo; de lo contrario, el cierre queda provisional y la decisión se pospone.'),
      ('Sin desvío relevante','Se documenta la buena práctica y termina la revisión del periodo.'),
      ('Acción no aprobada o ineficaz','Se ajusta y se vuelve a evaluar; la acción solo se cierra al verificar su efecto.'),
    ],
    'handoff':'P05 se activa por corte diario/semanal/mensual y usa resultados de P01-P04. Las mejoras aprobadas pueden cambiar reglas de esos procesos en el siguiente ciclo.',
    'current':'Existen panel de control y reportes web. La línea base MICOTRANS y las metas requieren reconciliar definiciones, fuentes y periodos antes de gobernar decisiones con los cinco KPI.',
    'wireframes':['P06_wireframes.png'],
    'speech':'El proceso de mejora comienza al cerrar una jornada o periodo, no necesariamente después de un único pedido. Primero se comprueba la calidad del dato. Si hay un desvío, se analiza la causa, se aprueba una acción con dueño y plazo y se verifica si produjo efecto. Los dashboards son soporte para esta gestión.',
  },
  {
    'id':'P06', 'title':'Habilitar recursos de reparto',
    'owner':'Jefatura logística', 'start':'Solicitud de alta o cambio de recurso',
    'success':'Conductor o vehículo disponible para P02',
    'lanes':['Jefatura logística','Administrador','Dispatcher','EcoRoute (sistema)'],
    'nodes':[
      N('S','start','Necesidad de recurso',0,0), N('A1','task','Definir perfil o capacidad requerida',0,1),
      N('A2','task','Recibir documentos y datos',1,2), N('G1','gateway','¿Información completa?',1,3),
      N('A3','task','Solicitar subsanación',1,4,75), N('G2','gateway','¿Se subsana en plazo?',1,5,75),
      N('E1','end','Solicitud archivada',1,6,75),
      N('A4','task','Verificar licencia o estado del vehículo',0,4,-75),
      N('G3','gateway','¿Recurso apto?',0,5,-75),
      N('A5','task','Bloquear recurso e informar motivo',0,6,75), N('E2','end','Recurso no habilitado',0,7,75),
      N('A6','service','Registrar conductor o vehículo',3,6,-75),
      N('A7','task','Validar capacidad y disponibilidad',2,7,-75),
      N('G4','gateway','¿Disponible para jornada?',2,8,-75),
      N('A8','task','Programar disponibilidad futura',2,9,75), N('E3','end','Recurso en espera',2,10,75),
      N('A9','task','Liberar recurso para asignación',0,9,-75), N('E4','end','Recurso habilitado → P02',0,10,-75),
    ],
    'flows':[
      F('S','A1'),F('A1','A2'),F('A2','G1'),F('G1','A4','Sí'),F('G1','A3','No'),
      F('A3','G2'),F('G2','A2','Sí'),F('G2','E1','No'),F('A4','G3'),
      F('G3','A6','Sí'),F('G3','A5','No'),F('A5','E2'),F('A6','A7'),
      F('A7','G4'),F('G4','A9','Sí'),F('G4','A8','No'),F('A8','E3'),F('A9','E4'),
    ],
    'rules':[
      ('Datos incompletos','El Administrador solicita subsanación; sin respuesta en plazo, archiva sin habilitar.'),
      ('Recurso no apto','Jefatura bloquea la asignación y documenta el motivo de seguridad o capacidad.'),
      ('No disponible hoy','Dispatcher programa la disponibilidad futura; no lo ofrece para la jornada actual.'),
    ],
    'handoff':'Solo E4 habilita la asignación en P02; E1/E2 terminan sin habilitar y E3 conserva disponibilidad futura.',
    'current':'Existen altas, consultas y cambios de conductores y vehículos. La vigencia documental, aptitud y agenda de disponibilidad son controles de negocio por validar.',
    'wireframes':['P02_wireframes.png'],
    'speech':'Antes de planificar, la jefatura debe saber qué recursos realmente puede usar. Se revisa documentación, aptitud, capacidad y disponibilidad; un registro en la base no equivale por sí solo a autorización de salida.',
  },
  {
    'id':'P07', 'title':'Supervisar la jornada',
    'owner':'Dispatcher', 'start':'Ruta iniciada y posiciones recibidas',
    'success':'Jornada vigilada o alerta derivada a P04',
    'lanes':['Conductor','EcoRoute (sistema)','Dispatcher','Jefatura logística'],
    'nodes':[
      N('S','start','Ruta en ejecución',0,0), N('A1','service','Transmitir posición y estado',1,1),
      N('G1','gateway','¿Se reciben datos recientes?',1,2),
      N('A2','task','Contactar al conductor',2,3,75),N('G2','gateway','¿Se restablece seguimiento?',2,4,75),
      N('A3','task','Registrar pérdida de señal',2,5,75),N('E1','end','Alerta técnica → P04',2,6,75),
      N('A4','task','Comparar avance con ventana',2,3,-75),N('G3','gateway','¿Hay desvío o demora?',2,4,-75),
      N('A5','task','Confirmar causa con conductor',2,5,-75),N('G4','gateway','¿Se requiere intervención?',2,6,-75),
      N('A6','task','Escalar y registrar alerta',3,7,75),N('E2','end','Incidencia → P04',3,8,75),
      N('A7','task','Actualizar seguimiento',2,7,-75),N('G5','gateway','¿Ruta finalizada?',2,8,-75),
      N('E3','end','Resultado de jornada → P05',2,9,-75),
    ],
    'flows':[
      F('S','A1'),F('A1','G1'),F('G1','A4','Sí'),F('G1','A2','No'),
      F('A2','G2'),F('G2','A1','Sí'),F('G2','A3','No'),F('A3','E1'),
      F('A4','G3'),F('G3','A5','Sí'),F('G3','A7','No'),
      F('A5','G4'),F('G4','A6','Sí'),F('G4','A7','No'),F('A6','E2'),
      F('A7','G5'),F('G5','E3','Sí'),F('G5','A1','No'),
    ],
    'rules':[
      ('Sin posiciones recientes','Se contacta al conductor y se registra alerta si no se recupera el enlace; la ausencia de GPS no demuestra una parada no autorizada.'),
      ('Desvío o demora','Se contrasta con conductor y ventana; la intervención se deriva a P04 solo si requiere decisión.'),
      ('Ruta aún activa','Se repite el control; el cierre alimenta P05.'),
    ],
    'handoff':'E1/E2 abren P04; E3 entrega resultados a P05. P07 corre en paralelo a P03 y no sustituye la confirmación de POD.',
    'current':'Hay envío e historial GPS, mapa y notificaciones. Los umbrales de alerta, la frecuencia de reporte y la gestión de pérdida de señal deben acordarse.',
    'wireframes':['P03_wireframes.png','P06_wireframes.png'],
    'speech':'Mientras el conductor atiende paradas, el Dispatcher supervisa la jornada. El GPS y los estados son señales de control, no prueba de entrega. Si la señal falta o aparece una desviación, verifica la causa y deriva solo los casos que requieren intervención.',
  },
  {
    'id':'P08', 'title':'Administrar accesos',
    'owner':'Administrador del sistema', 'start':'Solicitud de alta, cambio o baja de usuario',
    'success':'Acceso concedido, modificado o revocado con trazabilidad',
    'lanes':['Jefatura logística','Administrador','Keycloak (sistema)','Usuario solicitante'],
    'nodes':[
      N('S','start','Solicitud de acceso',0,0),N('A1','task','Precisar cargo y necesidad',0,1),
      N('G1','gateway','¿Solicitud autorizada?',0,2),N('A2','task','Informar rechazo',1,3,75),
      N('E1','end','Acceso denegado',1,4,75),
      N('A3','task','Definir rol mínimo necesario',1,3,-75),N('G2','gateway','¿Es baja o revocación?',1,4,-75),
      N('A4','service','Desactivar cuenta y sesiones',2,5,75),N('E2','end','Acceso revocado',2,6,75),
      N('A5','service','Crear o ajustar identidad y rol',2,5,-75),
      N('A6','task','Comprobar acceso permitido',3,6,-75),N('G3','gateway','¿Permisos correctos?',3,7,-75),
      N('A7','task','Corregir rol y registrar cambio',1,8,75),
      N('A8','task','Confirmar activación al solicitante',1,8,-75),N('E3','end','Acceso habilitado',1,9,-75),
    ],
    'flows':[
      F('S','A1'),F('A1','G1'),F('G1','A3','Sí'),F('G1','A2','No'),F('A2','E1'),
      F('A3','G2'),F('G2','A4','Sí'),F('G2','A5','No'),F('A4','E2'),
      F('A5','A6'),F('A6','G3'),F('G3','A8','Sí'),F('G3','A7','No'),F('A7','A5'),F('A8','E3'),
    ],
    'rules':[
      ('Solicitud no autorizada','No se crea el usuario; se comunica el motivo.'),
      ('Baja solicitada','Se revoca acceso y se conserva la trazabilidad de la solicitud.'),
      ('Permisos incorrectos','Se corrige el rol antes de notificar la activación.'),
    ],
    'handoff':'E3 habilita uso de módulos según rol; E1/E2 cierran la solicitud sin acceso activo.',
    'current':'Keycloak y los roles ADMIN, DISPATCHER y DRIVER sustentan autenticación y autorización; la aprobación formal y auditoría de solicitudes deben organizarse.',
    'wireframes':[],
    'speech':'La seguridad no comienza en el login. Jefatura aprueba la necesidad, administración asigna el mínimo privilegio y comprueba el resultado. Una baja revoca la cuenta; un permiso erróneo se corrige antes de habilitar el trabajo.',
  },
  {
    'id':'P09', 'title':'Registrar compras operativas',
    'owner':'Administrador con aprobación de Jefatura logística',
    'start':'Necesidad de compra o reposición registrada',
    'success':'Compra documentada o solicitud cerrada con motivo',
    'lanes':['Dispatcher','Administrador','Jefatura logística','EcoRoute (sistema)'],
    'nodes':[
      N('S','start','Necesidad de compra',0,0),N('A1','task','Describir necesidad y prioridad',0,1),
      N('A2','task','Comprobar existencia y presupuesto',1,2),N('G1','gateway','¿Hay existencia utilizable?',1,3),
      N('A3','task','Entregar recurso existente',1,4,75),N('E1','end','Necesidad cubierta sin compra',1,5,75),
      N('A4','task','Evaluar costo y proveedor',1,4,-75),N('A5','task','Autorizar gasto',2,5,-75),
      N('G2','gateway','¿Gasto aprobado?',2,6,-75),
      N('A6','task','Ajustar o cancelar solicitud',0,7,75),N('G3','gateway','¿Se justifica nuevo intento?',0,8,75),
      N('E2','end','Solicitud cancelada',0,9,75),
      N('A7','task','Realizar compra y recibir bien',1,7,-75),N('G4','gateway','¿Recepción conforme?',1,8,-75),
      N('A8','task','Reclamar o devolver',1,9,75),N('E3','end','Compra observada',1,10,75),
      N('A9','service','Registrar compra y comprobante',3,9,-75),
      N('E4','end','Compra registrada → P06',3,10,-75),
    ],
    'flows':[
      F('S','A1'),F('A1','A2'),F('A2','G1'),F('G1','A3','Sí'),F('G1','A4','No'),
      F('A3','E1'),F('A4','A5'),F('A5','G2'),F('G2','A7','Sí'),F('G2','A6','No'),
      F('A6','G3'),F('G3','A4','Sí'),F('G3','E2','No'),F('A7','G4'),
      F('G4','A9','Sí'),F('G4','A8','No'),F('A8','E3'),F('A9','E4'),
    ],
    'rules':[
      ('Existe recurso utilizable','Se cubre la necesidad sin duplicar el gasto.'),
      ('Gasto no aprobado','Se ajusta el alcance o se cancela con motivo.'),
      ('Recepción no conforme','Se reclama o devuelve; no se registra como compra utilizable.'),
    ],
    'handoff':'Una compra conforme registrada puede alimentar P06 si aumenta capacidad; una compra observada queda pendiente de resolución administrativa.',
    'current':'El módulo de compras registra, consulta, modifica y elimina registros. Presupuesto, aprobación y recepción conforme son reglas operativas propuestas, no automatizaciones verificadas.',
    'wireframes':[],
    'speech':'El módulo de compras se modela como proceso de apoyo porque afecta recursos y costos. La necesidad se valida, el gasto se aprueba y la recepción se comprueba. El alta en el sistema ocurre solo después de una compra conforme.',
  },
]

# Vista de negocio: los carriles identifican responsabilidades de MICOTRANS.
# El software presta apoyo a estas tareas; no es trabajador ni actor del negocio.
BUSINESS_LANES = {
    'P02': (['Dispatcher','Jefatura logística','Conductor'], {0:0,1:1,2:2,3:0},
            {'A7':'Revisar recorrido propuesto'}),
    'P03': (['Conductor','Dispatcher'], {0:0,1:1},
            {'A11':'Revisar POD recibido','A12':'Controlar POD pendiente de sincronización'}),
    'P04': (['Conductor','Dispatcher','Jefatura logística'], {0:0,1:1,2:2,3:1},
            {'A4':'Registrar evidencia incompleta'}),
    'P05': (['Dispatcher','Jefatura logística'], {0:0,1:1,2:0},
            {'A1':'Consolidar datos de la jornada','A3':'Preparar indicadores de la jornada'}),
    'P06': (['Jefatura logística','Administrador','Dispatcher'], {0:0,1:1,2:2,3:1},
            {'A6':'Registrar conductor o vehículo'}),
    'P07': (['Conductor','Dispatcher','Jefatura logística'], {0:0,1:0,2:1,3:2},
            {'A1':'Compartir avance de la ruta'}),
    'P08': (['Jefatura logística','Administrador'], {0:0,1:1,2:1,3:1},
            {'A4':'Revocar acceso y sesiones','A5':'Crear o ajustar acceso y rol',
             'A6':'Comprobar acceso permitido'}),
    'P09': (['Dispatcher','Administrador','Jefatura logística'], {0:0,1:1,2:2,3:1},
            {'A9':'Registrar compra y comprobante'}),
}
for process in PROCESSES:
    if process['id'] not in BUSINESS_LANES:
        continue
    lanes, lane_map, labels = BUSINESS_LANES[process['id']]
    for node in process['nodes']:
        node['lane'] = lane_map[node['lane']]
        if node['id'] in labels:
            node['label'] = labels[node['id']]
        if node['kind'] == 'service':
            node['kind'] = 'task'
    process['lanes'] = lanes

OLD_TO_NEW = [('P01','P01'),('P02 + P03','P02'),('P04','P03'),('P05','P04'),('P06','P05')]
MACRO = [
    ('Gestión de pedidos', 'P01', 'Validar solicitudes y registrar pedidos PENDING antes de planificar.'),
    ('Gestión de recursos logísticos', 'P06', 'Habilitar conductores y vehículos aptos.'),
    ('Gestión de operaciones de reparto', 'P02, P03, P04, P07', 'Preparar, ejecutar, supervisar y recuperar la jornada.'),
    ('Gestión del desempeño logístico', 'P05', 'Evaluar indicadores y acciones de mejora.'),
    ('Gestión de seguridad de acceso', 'P08', 'Autorizar, asignar y revocar accesos.'),
    ('Gestión de compras', 'P09', 'Controlar la adquisición de recursos operativos.'),
]
RF_COVERAGE = [
    ('1. Seguridad y acceso', 'P08', 'Inicio de sesión, renovación, cierre y permisos: funciones transversales.'),
    ('2. Conductores', 'P06', 'Alta y cambios sustentan la habilitación.'),
    ('3. Vehículos', 'P06', 'Alta y cambios sustentan la aptitud y capacidad.'),
    ('4. Pedidos', 'P01, P02, P03, P04', 'Registro PENDING, vinculación a ruta, entrega y excepción; la vinculación posterior es brecha funcional.'),
    ('5. Planificación y despacho', 'P02', 'Jornada aprobada e iniciada.'),
    ('6. Optimización', 'P02', 'Cálculo y revisión del recorrido; el algoritmo es una tarea de apoyo.'),
    ('7. Operación móvil', 'P03, P07', 'Paradas, estados y seguimiento; perfil/tema son funciones de interfaz.'),
    ('8. Prueba de entrega', 'P03, P04', 'Captura, validación y tratamiento de POD observado.'),
    ('9. GPS', 'P07', 'Transmisión, historial y supervisión.'),
    ('10. Notificaciones', 'P03, P04, P07', 'Avisos asociados a resultados y alertas.'),
    ('11. Panel y reportes', 'P05', 'Consulta y exportación como insumos de decisión.'),
    ('12. Compras', 'P09', 'Solicitud, autorización, recepción y registro.'),
    ('13. Inteligencia de negocios', 'P05', 'Indicadores, tendencias, metas y mejora.'),
]

def validate(proc):
    nodes={n['id']:n for n in proc['nodes']}
    assert len(nodes)==len(proc['nodes'])
    assert sum(n['kind']=='start' for n in nodes.values())==1
    assert any(n['kind']=='end' for n in nodes.values())
    outgoing=defaultdict(list); incoming=defaultdict(list)
    for src,dst,label in proc['flows']:
        assert src in nodes and dst in nodes, (proc['id'],src,dst)
        outgoing[src].append((dst,label))
        incoming[dst].append((src,label))
    for id,node in nodes.items():
        if node['kind']=='end': assert not outgoing[id]
        elif node['kind']=='gateway': assert sorted(label for _,label in outgoing[id])==['No','Sí'], (proc['id'],id,outgoing[id])
        elif node['kind']=='merge':
            assert len(incoming[id]) >= 2 and len(outgoing[id]) == 1, (proc['id'],id)
            assert outgoing[id][0][1] == '', (proc['id'],id)
        else: assert len(outgoing[id])==1, (proc['id'],id,outgoing[id])
        if node['kind'] in ('task','send','receive','service'):
            assert len(incoming[id]) == 1, (proc['id'],id,'uncontrolled merge')
        if node['kind']=='gateway': assert len(incoming[id])==1, (proc['id'],id,'mixed gateway')
    seen=set(); queue=deque(['S'])
    while queue:
        id=queue.popleft()
        if id in seen: continue
        seen.add(id)
        queue.extend(dst for dst,_ in outgoing[id])
    assert seen==set(nodes), (proc['id'],set(nodes)-seen)
    for end in (n['id'] for n in nodes.values() if n['kind']=='end'):
        assert end in seen

def explicit_merges(proc):
    """Separa la convergencia XOR de la tarea que consume el siguiente token."""
    incoming = defaultdict(list)
    for index, (src, dst, label) in enumerate(proc['flows']):
        incoming[dst].append(index)
    additions = []
    for target_id, indices in incoming.items():
        target = next(node for node in proc['nodes'] if node['id'] == target_id)
        if len(indices) < 2 or target['kind'] in ('end', 'merge'):
            continue
        merge_id = f"M_{target_id}"
        additions.append(N(merge_id, 'merge', '', target['lane'], target['col'] - .5, target['offset']))
        for index in indices:
            src, _, label = proc['flows'][index]
            proc['flows'][index] = F(src, merge_id, label)
        proc['flows'].append(F(merge_id, target_id))
    proc['nodes'].extend(additions)


for process in PROCESSES:
    explicit_merges(process)
    validate(process)

def font(size, bold=False):
    names=['arialbd.ttf','arial.ttf'] if bold else ['arial.ttf','calibri.ttf']
    for name in names:
        try: return ImageFont.truetype(str(Path('C:/Windows/Fonts')/name),size)
        except OSError: pass
    return ImageFont.load_default()

def geometry(proc):
    lane_h=800 if proc['id']=='P01' else 300
    top=290 if proc.get('external') else 148
    left=400 if proc.get('external') else 260
    gap=340
    max_col=max(n['col'] for n in proc['nodes'])
    width=left+max_col*gap+340
    bottom=top+len(proc['lanes'])*lane_h
    loc={n['id']:(left+n['col']*gap,top+n['lane']*lane_h+lane_h//2+n['offset']) for n in proc['nodes']}
    loops=sum(loc[dst][0]<=loc[src][0] for src,dst,_ in proc['flows'])
    return width,bottom+max(170,75+loops*27),top,bottom,loc

def box(n,xy):
    x,y=xy
    if n['kind'] in ('start','end'): return (x-28,y-28,x+28,y+28)
    if n['kind'] in ('gateway','merge'): return (x-43,y-43,x+43,y+43)
    return (x-109,y-52,x+109,y+52)

def path_points(src,dst,loc,kind_src,kind_dst,loop_index,bottom):
    x1,y1=loc[src]; x2,y2=loc[dst]
    if x2<=x1:
        track=bottom+34+loop_index*27
        return [(x1,y1+52),(x1,track),(x2,track),(x2,y2+52)]
    right=43 if kind_src in ('gateway','merge') else 28 if kind_src in ('start','end') else 109
    left=43 if kind_dst in ('gateway','merge') else 28 if kind_dst in ('start','end') else 109
    start=(x1+right,y1); end=(x2-left,y2)
    if abs(y1-y2)<4: return [start,end]
    mid=(start[0]+end[0])/2
    return [start,(mid,y1),(mid,y2),end]

def draw_text_center(draw,text,cx,cy,size,maxwidth,fill=DARK,bold=False):
    f=font(size,bold); lines=[]
    for para in text.split('\n'):
        words=para.split(); line=''
        for word in words:
            trial=(line+' '+word).strip()
            if line and draw.textbbox((0,0),trial,font=f)[2]>maxwidth:
                lines.append(line); line=word
            else: line=trial
        lines.append(line)
    lh=size+6
    yy=cy-len(lines)*lh/2
    for line in lines:
        bb=draw.textbbox((0,0),line,font=f)
        draw.text((cx-(bb[2]-bb[0])/2,yy),line,font=f,fill=fill)
        yy+=lh

def render_png(proc,path):
    width,height,top,bottom,loc=geometry(proc)
    lane_h=(bottom-top)//len(proc['lanes'])
    im=Image.new('RGB',(width,height),'white'); d=ImageDraw.Draw(im)
    d.text((40,28),f"{proc['id']}  {proc['title']}",font=font(39,True),fill=NAVY)
    d.text((40,83),f"Responsable: {proc['owner']}  |  Inicio: {proc['start']}",font=font(22),fill=MUTED)
    if proc.get('external'):
        d.rectangle((24,150,width-22,245),fill='#F7FAFC',outline='#A7BCD0',width=3)
        d.rectangle((24,150,185,245),fill='#EAF2F8',outline='#B9CCDC',width=2)
        draw_text_center(d,proc['external']['name'],103,197,20,145,NAVY,True)
        draw_text_center(d,'Participante externo: recibe condiciones y responde',width/2,197,22,width-460,MUTED)
    d.rectangle((24,top,width-22,bottom),fill='white',outline='#A7BCD0',width=3)
    for i,lane in enumerate(proc['lanes']):
        y1=top+i*lane_h; y2=y1+lane_h
        x1=150 if proc.get('external') else 24
        x2=260 if proc.get('external') else 185
        d.rectangle((x1,y1,x2,y2),fill='#EAF2F8',outline='#B9CCDC',width=2)
        if i: d.line((x2,y1,width-22,y1),fill='#D3DFE9',width=2)
        draw_text_center(d,(lane+'\nCarril') if proc.get('external') else lane,(x1+x2)/2,(y1+y2)/2,
                         19 if proc.get('external') else 22,x2-x1-12,NAVY,True)
    if proc.get('external'):
        d.rectangle((24,top,150,bottom),fill='#DDEAF4',outline='#A7BCD0',width=2)
        draw_text_center(d,'MICOTRANS\nPool',87,(top+bottom)/2,17,116,NAVY,True)
    nodes={n['id']:n for n in proc['nodes']}
    loop_count=0
    for src,dst,label in proc['flows']:
        backwards=loc[dst][0]<=loc[src][0]
        pts=path_points(src,dst,loc,nodes[src]['kind'],nodes[dst]['kind'],loop_count,bottom)
        if backwards: loop_count+=1
        color=ORANGE if label=='No' or backwards else BLUE
        d.line(pts,fill=color,width=4,joint='curve')
        a,b=pts[-2],pts[-1]; angle=math.atan2(b[1]-a[1],b[0]-a[0])
        length=15
        sides=[(b[0]-length*math.cos(angle+0.50),b[1]-length*math.sin(angle+0.50)),
               (b[0]-length*math.cos(angle-0.50),b[1]-length*math.sin(angle-0.50))]
        d.polygon([b,*sides],fill=color)
        if label:
            px,py=pts[0]
            py-=30 if label=='Sí' else -13
            d.rounded_rectangle((px+6,py-4,px+50,py+27),6,fill='white',outline=color,width=2)
            d.text((px+13,py),label,font=font(21,True),fill=color)
    if proc.get('external'):
        for direction,ids in [('outgoing',proc['external']['outgoing']),
                              ('incoming',proc['external']['incoming'])]:
            for id in ids:
                x,y=loc[id]
                start_y=(y-52 if direction=='outgoing' else 245)
                end_y=(245 if direction=='outgoing' else y-52)
                a,b=sorted((start_y,end_y))
                for yy in range(int(a),int(b),20):
                    d.line((x,yy,x,min(yy+11,b)),fill=TEAL,width=3)
                d.ellipse((x-8,start_y-8,x+8,start_y+8),fill='white',outline=TEAL,width=3)
                tip=end_y
                if direction=='outgoing':
                    d.line((x-10,tip+15,x,tip),fill=TEAL,width=3)
                    d.line((x+10,tip+15,x,tip),fill=TEAL,width=3)
                else:
                    d.line((x-10,tip-15,x,tip),fill=TEAL,width=3)
                    d.line((x+10,tip-15,x,tip),fill=TEAL,width=3)
    for node in proc['nodes']:
        x,y=loc[node['id']]; kind=node['kind']; rect=box(node,(x,y))
        if kind=='start':
            d.ellipse(rect,fill='#D8F3E8',outline=TEAL,width=5)
            draw_text_center(d,node['label'],x,y-61,19,185,TEAL,True)
        elif kind=='end':
            d.ellipse(rect,fill='white',outline=NAVY,width=6)
            label_y=y+67 if proc['id']=='P01' and node['id'] in ('E1','E2','E3') else y-65
            draw_text_center(d,node['label'],x,label_y,18,230,NAVY,True)
        elif kind=='merge':
            d.polygon([(x,y-43),(x+43,y),(x,y+43),(x-43,y)],fill='white',outline=BLUE,width=3)
            d.line((x-13,y-13,x+13,y+13),fill=BLUE,width=5)
            d.line((x+13,y-13,x-13,y+13),fill=BLUE,width=5)
        elif kind=='gateway':
            d.polygon([(x,y-43),(x+43,y),(x,y+43),(x-43,y)],fill='#FFF3E7',outline=ORANGE,width=3)
            d.line((x-13,y-13,x+13,y+13),fill=ORANGE,width=5)
            d.line((x+13,y-13,x-13,y+13),fill=ORANGE,width=5)
            label_y=y+69 if proc['id']=='P01' and node['offset']>0 else (y-88 if y-88>top+node['lane']*lane_h+10 else y+69)
            draw_text_center(d,node['label'],x,label_y,20,245,ORANGE,True)
        else:
            fill='#EBF4FD' if kind=='service' else 'white'
            d.rounded_rectangle(rect,radius=13,fill=fill,outline=BLUE,width=3)
            draw_text_center(d,node['label'],x,y,23,199,DARK,True)
            if kind=='service': d.text((rect[0]+9,rect[1]+6),'S',font=font(16,True),fill=BLUE)
            if kind in ('send','receive'):
                ex,ey=rect[0]+13,rect[3]-33
                d.rectangle((ex,ey,ex+28,ey+19),fill=BLUE if kind=='send' else 'white',outline=BLUE,width=2)
                d.line((ex,ey,ex+14,ey+11,ex+28,ey),fill='white' if kind=='send' else BLUE,width=2)
    d.text((40,height-50),'XOR naranja = decisión  ·  XOR azul = convergencia  ·  Fin circular = resultado del caso',font=font(22),fill=MUTED)
    im.save(path)

BPMN_NS='http://www.omg.org/spec/BPMN/20100524/MODEL'
BPMNDI_NS='http://www.omg.org/spec/BPMN/20100524/DI'
DC_NS='http://www.omg.org/spec/DD/20100524/DC'
DI_NS='http://www.omg.org/spec/DD/20100524/DI'
ET.register_namespace('bpmn',BPMN_NS); ET.register_namespace('bpmndi',BPMNDI_NS)
ET.register_namespace('dc',DC_NS); ET.register_namespace('di',DI_NS)

def bpmn_xml(proc,path):
    nodes={n['id']:n for n in proc['nodes']}; width,height,top,bottom,loc=geometry(proc)
    lane_h=(bottom-top)//len(proc['lanes'])
    q=lambda ns,tag:f'{{{ns}}}{tag}'
    root=ET.Element(q(BPMN_NS,'definitions'),{'id':f"Definitions_{proc['id']}",'targetNamespace':'https://ecoroute.pe/procesos-negocio'})
    if proc.get('external'):
        ET.SubElement(root,q(BPMN_NS,'message'),{'id':f"Message_{proc['id']}_Propuesta",'name':'Propuesta de servicio'})
        ET.SubElement(root,q(BPMN_NS,'message'),{'id':f"Message_{proc['id']}_Respuesta",'name':'Respuesta del cliente'})
    process=ET.SubElement(root,q(BPMN_NS,'process'),{'id':f"Process_{proc['id']}",'name':proc['title'],'isExecutable':'false'})
    ET.SubElement(process,q(BPMN_NS,'documentation')).text='Proceso objetivo de negocio para MICOTRANS. '+proc['handoff']
    collaboration_id=None
    if proc.get('external'):
        collaboration_id=f"Collaboration_{proc['id']}"
        collaboration=ET.SubElement(root,q(BPMN_NS,'collaboration'),{'id':collaboration_id})
        ET.SubElement(collaboration,q(BPMN_NS,'participant'),{
            'id':f"Participant_{proc['id']}_MICOTRANS",'name':'MICOTRANS',
            'processRef':f"Process_{proc['id']}"})
        ET.SubElement(collaboration,q(BPMN_NS,'participant'),{
            'id':f"Participant_{proc['id']}_Cliente",'name':proc['external']['name']})
        for direction,ids in [('outgoing',proc['external']['outgoing']),
                              ('incoming',proc['external']['incoming'])]:
            for id in ids:
                target=f"Participant_{proc['id']}_Cliente"
                activity=f"{proc['id']}_{id}"
                ET.SubElement(collaboration,q(BPMN_NS,'messageFlow'),{
                    'id':f"Message_{proc['id']}_{id}",
                    'name':'Propuesta de servicio' if direction=='outgoing' else 'Respuesta del cliente',
                    'messageRef':f"Message_{proc['id']}_{'Propuesta' if direction=='outgoing' else 'Respuesta'}",
                    'sourceRef':activity if direction=='outgoing' else target,
                    'targetRef':target if direction=='outgoing' else activity})
    laneset=ET.SubElement(process,q(BPMN_NS,'laneSet'),{'id':f"LaneSet_{proc['id']}"})
    for i,lane in enumerate(proc['lanes']):
        lane_el=ET.SubElement(laneset,q(BPMN_NS,'lane'),{'id':f"Lane_{proc['id']}_{i}",'name':lane})
        for n in proc['nodes']:
            if n['lane']==i: ET.SubElement(lane_el,q(BPMN_NS,'flowNodeRef')).text=f"{proc['id']}_{n['id']}"
    kind_tags={'start':'startEvent','end':'endEvent','gateway':'exclusiveGateway','merge':'exclusiveGateway',
               'task':'userTask','service':'serviceTask','send':'sendTask','receive':'receiveTask'}
    incoming=defaultdict(list); outgoing=defaultdict(list)
    for i,(src,dst,_) in enumerate(proc['flows'],1):
        incoming[dst].append(f"Flow_{proc['id']}_{i}")
        outgoing[src].append(f"Flow_{proc['id']}_{i}")
    for n in proc['nodes']:
        attrs={'id':f"{proc['id']}_{n['id']}",'name':n['label']}
        if n['kind'] in ('gateway','merge'):
            attrs['gatewayDirection']='Diverging' if n['kind']=='gateway' else 'Converging'
        if n['kind']=='send': attrs['messageRef']=f"Message_{proc['id']}_Propuesta"
        if n['kind']=='receive': attrs['messageRef']=f"Message_{proc['id']}_Respuesta"
        element=ET.SubElement(process,q(BPMN_NS,kind_tags[n['kind']]),attrs)
        for flow_id in incoming[n['id']]: ET.SubElement(element,q(BPMN_NS,'incoming')).text=flow_id
        for flow_id in outgoing[n['id']]: ET.SubElement(element,q(BPMN_NS,'outgoing')).text=flow_id
    for i,(src,dst,label) in enumerate(proc['flows'],1):
        attrs={'id':f"Flow_{proc['id']}_{i}",'sourceRef':f"{proc['id']}_{src}",'targetRef':f"{proc['id']}_{dst}"}
        if label: attrs['name']=label
        flow=ET.SubElement(process,q(BPMN_NS,'sequenceFlow'),attrs)
        if label: ET.SubElement(flow,q(BPMN_NS,'documentation')).text=f"Rama de decisión: {label}"
    diagram=ET.SubElement(root,q(BPMNDI_NS,'BPMNDiagram'),{'id':f"Diagram_{proc['id']}"})
    plane=ET.SubElement(diagram,q(BPMNDI_NS,'BPMNPlane'),{
        'id':f"Plane_{proc['id']}",
        'bpmnElement':collaboration_id or f"Process_{proc['id']}"})
    if proc.get('external'):
        for suffix,x,y,h in [('MICOTRANS',24,top,bottom-top),('Cliente',24,150,95)]:
            shape=ET.SubElement(plane,q(BPMNDI_NS,'BPMNShape'),{
                'id':f"Shape_Participant_{proc['id']}_{suffix}",
                'bpmnElement':f"Participant_{proc['id']}_{suffix}",'isHorizontal':'true'})
            ET.SubElement(shape,q(DC_NS,'Bounds'),{
                'x':str(x),'y':str(y),'width':str(width-46),'height':str(h)})
    for i,lane in enumerate(proc['lanes']):
        shape=ET.SubElement(plane,q(BPMNDI_NS,'BPMNShape'),{'id':f"Shape_Lane_{proc['id']}_{i}",'bpmnElement':f"Lane_{proc['id']}_{i}",'isHorizontal':'true'})
        lane_x=150 if proc.get('external') else 24
        ET.SubElement(shape,q(DC_NS,'Bounds'),{'x':str(lane_x),'y':str(top+i*lane_h),
                                            'width':str(width-22-lane_x),'height':str(lane_h)})
    for n in proc['nodes']:
        shape=ET.SubElement(plane,q(BPMNDI_NS,'BPMNShape'),{'id':f"Shape_{proc['id']}_{n['id']}",'bpmnElement':f"{proc['id']}_{n['id']}"})
        x1,y1,x2,y2=box(n,loc[n['id']])
        ET.SubElement(shape,q(DC_NS,'Bounds'),{'x':str(x1),'y':str(y1),'width':str(x2-x1),'height':str(y2-y1)})
    loop_count=0
    for i,(src,dst,label) in enumerate(proc['flows'],1):
        edge=ET.SubElement(plane,q(BPMNDI_NS,'BPMNEdge'),{'id':f"Edge_{proc['id']}_{i}",'bpmnElement':f"Flow_{proc['id']}_{i}"})
        backwards=loc[dst][0]<=loc[src][0]
        points=path_points(src,dst,loc,nodes[src]['kind'],nodes[dst]['kind'],loop_count,bottom)
        if backwards: loop_count+=1
        for x,y in points: ET.SubElement(edge,q(DI_NS,'waypoint'),{'x':str(round(x)),'y':str(round(y))})
    if proc.get('external'):
        for direction,ids in [('outgoing',proc['external']['outgoing']),
                              ('incoming',proc['external']['incoming'])]:
            for id in ids:
                edge=ET.SubElement(plane,q(BPMNDI_NS,'BPMNEdge'),{
                    'id':f"Edge_Message_{proc['id']}_{id}",
                    'bpmnElement':f"Message_{proc['id']}_{id}"})
                x,y=loc[id]
                start_y=y-52 if direction=='outgoing' else 245
                end_y=245 if direction=='outgoing' else y-52
                for yy in (start_y,end_y):
                    ET.SubElement(edge,q(DI_NS,'waypoint'),{'x':str(round(x)),'y':str(round(yy))})
    ET.indent(root,space='  ')
    ET.ElementTree(root).write(path,encoding='utf-8',xml_declaration=True)

for proc in PROCESSES:
    render_png(proc,IMG/f"{proc['id']}_negocio.png")
    bpmn_xml(proc,BPMN/f"{proc['id']}_negocio.bpmn")

def cell_fill(cell,color):
    tc=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),color); tc.append(sh)

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    header_trpr=t.rows[0]._tr.get_or_add_trPr()
    header_trpr.append(OxmlElement('w:tblHeader'))
    for i,h in enumerate(headers):
        cell=t.rows[0].cells[i]; cell.text=h; cell_fill(cell,'12395F')
        for run in cell.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(9)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=str(value); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j%2: cell_fill(cells[i],'EFF5F9')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs: run.font.size=Pt(9)
    return t

def para(doc,text,lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    if lead: p.add_run(lead).bold=True
    p.add_run(text)
    return p

def setup(doc,landscape=False):
    s=doc.sections[0]
    if landscape:
        s.page_width=Inches(16.54); s.page_height=Inches(11.69)
        s.left_margin=s.right_margin=Inches(.65)
        s.top_margin=s.bottom_margin=Inches(.52)
    else:
        s.page_width=Inches(8.27); s.page_height=Inches(11.69)
        s.left_margin=s.right_margin=Inches(.76)
        s.top_margin=s.bottom_margin=Inches(.7)
    doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(10)
    for style in ('Title','Heading 1','Heading 2'):
        doc.styles[style].font.name='Aptos Display'
        doc.styles[style].font.color.rgb=RGBColor(0,0,0)
    # Word's built-in Title style can carry a colored bottom border.
    title_ppr=doc.styles['Title']._element.get_or_add_pPr()
    for border in title_ppr.findall(qn('w:pBdr')):
        title_ppr.remove(border)

def gate_rows(proc):
    nodes={node['id']:node for node in proc['nodes']}
    by_source=defaultdict(dict)
    for src,dst,label in proc['flows']:
        if label:
            next_id = dst
            if nodes[next_id]['kind'] == 'merge':
                next_id = next(target for origin,target,_ in proc['flows'] if origin == dst)
            by_source[src][label]=nodes[next_id]['label']
    gateways=sorted((node for node in proc['nodes'] if node['kind']=='gateway'),key=lambda node:node['col'])
    return [(node['label'],proc['lanes'][node['lane']],by_source[node['id']]['Sí'],by_source[node['id']]['No']) for node in gateways]

spec=Document(); setup(spec,True)
spec.add_paragraph('Procesos de negocio EcoRoute para MICOTRANS',style='Title')
para(spec,'Inventario ampliado de nueve procesos. P01-P05 forman el flujo principal del servicio; P06-P09 cubren recursos, supervisión, accesos y compras. Los macroprocesos usan «Gestión de…»; cada proceso ejecutable se nombra con verbo y objeto. La solución actual es soporte, no prueba de que todos los controles propuestos estén implementados.')
spec.add_heading('1. Mapa de procesos y límites',level=1)
table(spec,['Macroproceso','Procesos BPMN','Objetivo de negocio'],MACRO)
para(spec,'Los grupos de procesos de dirección de proyectos del PMI (inicio, planificación, ejecución, monitoreo y control, cierre) gobiernan el proyecto EcoRoute. No son los carriles ni los procesos de la operación de reparto de MICOTRANS; el acta, WBS y cronograma se documentan en el capítulo 1.')
table(spec,['Proceso anterior','Proceso revisado','Disparador','Salida que habilita otro proceso'],[
    (old,new,p['start'],p['handoff']) for (old,new),p in zip(OLD_TO_NEW,PROCESSES)
])
para(spec,'Las conexiones entre procesos son transferencias de caso, no flechas de secuencia BPMN entre diagramas independientes. Un retorno interno permanece en el mismo proceso. Los resultados pendientes o fallidos terminan con estado y responsable explícitos.')
para(spec,'Nivel de modelado: estos nueve BPMN representan procesos de negocio objetivo (To-Be), no la arquitectura informática. Los carriles internos expresan responsabilidades funcionales de MICOTRANS: Dispatcher, Jefatura logística, Conductor y Administrador. P01 muestra al cliente solicitante como pool externo de caja negra, con mensajes para la propuesta y la respuesta; no se inventa su proceso interno. EcoRoute y Keycloak apoyan las tareas, pero no aparecen como trabajadores ni actores del negocio. Esta elección metodológica no implica que BPMN prohíba carriles de sistemas en diagramas técnicos. Los puestos formales de MICOTRANS deben confirmarse antes del piloto.')
para(spec,'Convención BPMN 2.0: en P01 MICOTRANS es el pool y Dispatcher su carril; el cliente ocupa un pool externo de caja negra. El flujo de mensaje de salida se vincula a una Send Task y la respuesta a una Receive Task. Cada retorno alternativo que entra a una actividad pasa por una compuerta XOR convergente explícita; los rombos de decisión son XOR divergentes y las ramas se rotulan Sí/No. Las preguntas en los rombos son nombres descriptivos válidos, no una obligación ni una prohibición de la norma. Los cierres por rechazo, inviabilidad o desistimiento son eventos de fin ordinarios: describen resultados comerciales, no errores técnicos ni terminación de ramas paralelas. Referencia normativa: OMG, Business Process Model and Notation, versión 2.0.2.')
para(spec,'Secuencia verificada: el formulario permite registrar sin ruta y el backend genera PENDING. Si se elige una ruta existente, genera ASSIGNED. La ruta se crea PLANNED con fecha, conductor y vehículo. No existe en la interfaz ni en el API una operación de asignación posterior de un PENDING a una ruta; por ello, el orden ideal P01 → P02 contiene una brecha explícita. El recorrido end-to-end demostrable usa la ruta primero y registra después el pedido vinculado. La confirmación comercial no está representada por un estado CONFIRMED.')
spec.add_heading('Cobertura de los trece módulos funcionales',level=2)
table(spec,['Módulo RF','Proceso BPMN','Alcance'],RF_COVERAGE)
spec.add_heading('2. Criterios usados para evaluar el detalle',level=1)
table(spec,['Criterio','Resultado de la revisión'],[
    ('Enfoque de negocio','Los carriles internos indican quién decide, corrige, autoriza o comunica. El cliente aparece como participante externo en P01; EcoRoute y Keycloak son soporte informático descrito fuera de los carriles.'),
    ('Decisiones completas','Toda compuerta exclusiva tiene salida Sí y No; cada No corrige, reintenta, escala o cierra con motivo.'),
    ('Convergencia explícita','Ninguna actividad recibe dos flujos de secuencia directamente; los retornos alternativos pasan por una XOR convergente azul.'),
    ('Mensajes y participantes','Solo P01 necesita interacción con el cliente externo. El mensaje entra a una tarea de recepción y sale de una tarea de envío, sin flujos de secuencia entre pools.'),
    ('Fin tipificado','Los resultados de negocio usan eventos de fin ordinarios; Error y Terminate se reservan para excepciones BPMN que realmente requieren esas semánticas.'),
    ('Inicio y cierre','Cada proceso define un evento inicial, varios resultados posibles y el proceso siguiente para cada transferencia.'),
    ('Relación con desarrollo','Se indica soporte visible en código y control operativo que debe acordarse o verificarse durante el piloto.'),
])
for idx,p in enumerate(PROCESSES):
    heading=spec.add_heading(f"{p['id']} {p['title']}",level=1)
    if idx: heading.paragraph_format.page_break_before=True
    para(spec,f"{p['owner']}  |  Inicio: {p['start']}  |  Salida principal: {p['success']}",lead='Dueño del proceso: ')
    spec.add_picture(str(IMG/f"{p['id']}_negocio.png"),width=Inches(15.1))
    spec.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    spec.add_heading('Lectura de cada compuerta',level=2)
    table(spec,['Condición','Responsable','Si la respuesta es Sí','Si la respuesta es No'],gate_rows(p))
    spec.add_heading('Condiciones, tratamiento y cierre',level=2)
    table(spec,['Condición cuando es No','Tratamiento y decisión'],p['rules'])
    para(spec,p['handoff'],lead='Transferencia: ')
    para(spec,p['current'],lead='Soporte actual y control a validar: ')
    para(spec,'Wireframes relacionados: '+', '.join(p['wireframes'])+' (láminas previas basadas en pantallas web y móvil).')
section_three=spec.add_heading('3. Qué ya respalda el desarrollo y qué requiere acuerdo operativo',level=1)
section_three.paragraph_format.page_break_before=True
table(spec,['Proceso','Soporte identificable','Regla de gestión por validar'],[
    ('P01','Alta con ruta opcional, consulta e importación CSV; PENDING sin ruta.','Admisión, plazo de corrección y constancia de aceptación comercial.'),
    ('P02','Ruta PLANNED, conductor/vehículo, inicio de despacho y recorrido OSRM.','Asignar después los PENDING a la ruta; chequeo de carga y aprobación de salida.'),
    ('P03','Ruta móvil, estados, foto, firma y POD.','POD mínimo verificable, cola offline fiable y ubicación válida.'),
    ('P04','Motivo de cambio de estado y alertas.','Caso de incidencia, reintento, reasignación y cierre autorizado.'),
    ('P05','Panel y reportes web.','Definiciones de KPI, calidad de dato, responsables y seguimiento de acciones.'),
    ('P06','Mantenimiento de conductores y vehículos.','Vigencia, aptitud y disponibilidad para cada jornada.'),
    ('P07','GPS, historial y notificaciones.','Umbrales, pérdida de señal y confirmación de causas.'),
    ('P08','Autenticación y roles mediante Keycloak.','Solicitud aprobada, mínimo privilegio y revocación trazable.'),
    ('P09','Módulo de compras.','Presupuesto, aprobación y conformidad de recepción.'),
])
para(spec,'Referencias técnicas revisadas: OrderController, RouteController, DeliveryProofController, Orders.tsx, Routes.tsx, Reports.tsx, Dashboard.tsx, order_detail_page.dart y route_remote_data_source.dart. Estas referencias respaldan la existencia de funciones, pero no prueban por sí solas el cumplimiento de cada control de negocio.')
spec.add_heading('Acuerdos que deben cerrar los responsables antes del piloto',level=2)
table(spec,['Proceso','Acuerdo empresarial necesario'],[
    ('P01','Definir datos mínimos, plazo para subsanar, cobertura y evidencia de aceptación del cliente; PENDING no equivale a confirmado.'),
    ('P02','Priorizar la asignación posterior de pedidos PENDING a una ruta; definir límite de capacidad, aprobación y salida.'),
    ('P03','Definir POD mínimo, cuándo se considera entregado y cómo se controla el pendiente sin red.'),
    ('P04','Aprobar causales, criterios de reintento, límites para reprogramar y autorización de cierre fallido.'),
    ('P05','Validar fórmulas, línea base, metas, periodicidad y dueño de cada acción de mejora.'),
    ('P06','Definir requisitos de aptitud documental y disponibilidad de recursos.'),
    ('P07','Acordar frecuencia GPS, umbrales de demora y tratamiento de pérdida de señal.'),
    ('P08','Establecer aprobador de accesos y evidencia de bajas.'),
    ('P09','Definir límite de aprobación de compras y conformidad de recepción.'),
])
spec.add_heading('4. Entregables complementarios',level=1)
para(spec,'Se adjuntan nueve archivos BPMN 2.0 con carriles de responsabilidades, compuertas y flujos editables; P01 usa además una colaboración con cliente externo. Se incluyen imágenes a resolución completa y un visualizador HTML. Los cinco dashboards, cinco indicadores y dos reportes operativos conservan sus diagramas en el paquete anterior.')
spec_path=OUT/'EcoRoute_Procesos_Negocio_Revisados.docx'; spec.save(spec_path)

guide=Document(); setup(guide,False)
guide.add_paragraph('Guion de exposición de los procesos de negocio',style='Title')
para(guide,'EcoRoute para Grupo MICOTRANS S. A. C. Este guion acompaña nueve procesos organizados en seis macroprocesos de gestión. Cinco describen el flujo principal y cuatro completan las funciones de soporte y control.')
guide.add_heading('Apertura',level=1)
para(guide,'“Hemos revisado los procesos desde la operación logística. La planificación y el despacho forman ahora un solo proceso, porque una ruta no está lista para ejecutarse hasta que se apruebe, se entregue al conductor y este confirme la salida. Cada decisión del BPMN tiene tratamiento y el cierre indica si el pedido continúa, se corrige, se reprograma o termina.”')
guide.add_heading('Mapa y forma de leer los diagramas',level=1)
para(guide,'Explique que P01-P05 cubren el servicio principal y P06-P09 recursos, supervisión, acceso y compras. Los carriles muestran responsables de MICOTRANS, no aplicaciones. En P01, MICOTRANS es un pool con Dispatcher como carril y el cliente es otro pool de caja negra. Los mensajes punteados salen de una tarea de envío y llegan a una tarea de recepción. El rombo naranja divide por Sí/No y el azul reúne caminos alternativos antes de continuar. Un círculo de borde grueso señala el fin normal de ese caso, aun cuando el resultado sea rechazo o desistimiento. EcoRoute y Keycloak apoyan tareas sin convertirse en actores del negocio.')
for p in PROCESSES:
    if p['id']=='P02': guide.add_page_break()
    guide.add_heading(f"{p['id']} {p['title']}",level=1)
    para(guide,'“'+p['speech']+'”')
    for decision,owner,yes,no in gate_rows(p):
        para(guide,f"{owner}. Sí: {yes}. No: {no}.",lead=decision+' ')
    for condition,response in p['rules']:
        para(guide,response,lead=f"Si {condition.lower()}: ")
    para(guide,p['handoff'],lead='Dónde termina y qué sigue: ')
    para(guide,p['current'],lead='Si preguntan por la aplicación: ')
guide.add_page_break()
guide.add_heading('Cierre',level=1)
para(guide,'“La gestión completa distingue la aceptación del servicio por el cliente del estado PENDING guardado por EcoRoute. La ruta se crea PLANNED y solo pasa a ejecución cuando tiene pedidos asociados y se inicia el despacho. Hoy la aplicación permite asociarlos al registrarlos, pero no vincular después un PENDING existente. Esa brecha es la prioridad de diseño antes del piloto.”')
guide.add_heading('Ejemplo para explicar las transferencias',level=1)
para(guide,'“Supongamos que el cliente acepta una entrega con dirección y ventana válidas. El Dispatcher registra el pedido sin ruta y queda PENDING. En el flujo objetivo, P02 crea una ruta PLANNED y asocia ese pedido. Como el sistema todavía no permite esa asociación posterior, la demostración actual crea primero la ruta y registra luego el pedido indicando su routeId. Después se inicia el despacho, se atiende la parada y se captura la POD. El orden objetivo queda condicionado a resolver la brecha funcional, no se afirma que ya esté automatizado.”')
guide.add_heading('Preguntas frecuentes',level=1)
for question,answer in [
    ('¿Por qué P02 incluye despacho?','Porque el resultado de negocio es una ruta efectivamente entregada e iniciada, no solo un trazado calculado.'),
    ('¿P04 siempre termina en falla?','No. Puede devolver el caso a P03 para reintento o a P02 para nueva fecha. Solo Jefatura autoriza el cierre fallido.'),
    ('¿P05 empieza después de cada pedido?','No. Se activa por corte de jornada o periodo y usa resultados acumulados de P01 a P04.'),
    ('¿Todo el flujo ya está automatizado?','No. El end-to-end actual funciona si la ruta existe antes del alta del pedido vinculado. La asignación posterior de PENDING, la confirmación comercial y las autorizaciones formales son brechas o controles por validar.'),
]:
    para(guide,answer,lead=question+' ')
guide_path=OUT/'Guion_Procesos_Negocio_Revisados.docx'; guide.save(guide_path)

cards=[]
for p in PROCESSES:
    rules=''.join(f'<tr><td>{escape(a)}</td><td>{escape(b)}</td></tr>' for a,b in p['rules'])
    wire=''.join(f'<img src="../imagenes/{escape(path)}" alt="Wireframes asociados a {p["id"]}">' for path in p['wireframes'])
    cards.append(f'''<section id="{p['id']}"><h2>{p['id']} · {escape(p['title'])}</h2>
      <p><strong>Dueño:</strong> {escape(p['owner'])} · <strong>Inicio:</strong> {escape(p['start'])}</p>
      <a href="bpmn/{p['id']}_negocio.bpmn">Abrir BPMN 2.0 editable</a>
      <img class="diagram" src="imagenes/{p['id']}_negocio.png" alt="BPMN de {escape(p['title'])}">
      <h3>Decisiones y tratamiento</h3><table><tr><th>Condición</th><th>Tratamiento</th></tr>{rules}</table>
      <p><strong>Fin y transferencia:</strong> {escape(p['handoff'])}</p>
      <p><strong>Relación con la aplicación:</strong> {escape(p['current'])}</p>
      <h3>Wireframes relacionados</h3><div class="wire">{wire}</div></section>''')
links=''.join(f'<a href="#{p["id"]}">{p["id"]}</a>' for p in PROCESSES)
html=f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>EcoRoute procesos de negocio revisados</title>
<style>body{{margin:0;font:16px/1.5 Arial,sans-serif;color:#1f2937;background:#edf3f7}}header{{background:#12395f;color:white;padding:30px 5vw}}header h1{{margin:0}}nav{{position:sticky;top:0;background:white;padding:12px 5vw;box-shadow:0 2px 10px #0002;z-index:3}}nav a{{margin-right:24px;color:#12395f;font-weight:bold;text-decoration:none}}main{{max-width:1700px;margin:auto;padding:28px}}section{{background:white;border:1px solid #d8e3ec;border-radius:12px;padding:25px;margin:25px 0}}h2{{color:#12395f}}h3{{color:#0f766e}}.diagram{{display:block;width:100%;max-width:none;overflow:auto;border:1px solid #cddae5;margin:20px 0}}table{{border-collapse:collapse;width:100%}}th{{background:#12395f;color:white;text-align:left}}td,th{{border:1px solid #d7e1ea;padding:10px}}.wire img{{display:block;width:100%;margin:15px 0;border:1px solid #d7e1ea}}.note{{background:#eaf2f8;padding:14px;border-left:4px solid #0f766e}}</style></head>
<body><header><h1>EcoRoute · procesos de negocio MICOTRANS</h1><p>Nueve procesos en seis macroprocesos · BPMN por responsables · decisiones y finales explícitos</p></header><nav>{links}<a href="../Visualizador_Procesos_EcoRoute_MICOTRANS.html#D01">Dashboards y reportes</a></nav><main><section><h2>Mapa de transferencia</h2><p class="note">P01 valida la solicitud y registra PENDING sin ruta. P02 crea una ruta PLANNED y, en el flujo objetivo, vincula esos pedidos antes del despacho. Hoy esa vinculación posterior no está implementada: para demostrar el recorrido completo se crea la ruta primero y luego se registra el pedido con routeId. P03 entrega, P04 trata excepciones, P05 evalúa resultados. P06 habilita recursos, P07 supervisa, P08 gobierna accesos y P09 registra compras.</p></section>{''.join(cards)}</main></body></html>'''
(OUT/'Visualizador_Procesos_Negocio_Revisados.html').write_text(html,encoding='utf-8')
print(spec_path)
print(guide_path)
print(OUT/'Visualizador_Procesos_Negocio_Revisados.html')
