"""Create a reviewed Chapter 2 copy without replacing the user's working DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Integrador I_Grupo3.docx"
DEST = ROOT / "docs" / "Integrador I_Grupo3_Capitulo2_Alineado.docx"
ERD = ROOT / "docs" / "base_datos" / "modelo_erd_negocio" / "09_vista_completa_tipo_dbeaver.png"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    edge = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        edge.append(element)
    tc_pr.append(edge)


def inserted_paragraph(doc, anchor, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if style and style.startswith("Heading"):
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        p.paragraph_format.space_after = Pt(6)
    if text:
        p.add_run(text)
    anchor._p.addprevious(p._p)
    return p


def inserted_table(doc, anchor, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for row_index, values in enumerate([headers, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            borders(cell)
            if row_index == 0:
                shade(cell, "253E56")
            elif row_index % 2 == 0:
                shade(cell, "F3F6F9")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.0 if len(headers) >= 4 else 8.5)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    header = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)
    anchor._p.addprevious(table._tbl)
    return table


def main() -> None:
    doc = Document(SOURCE)
    # Existing paragraphs remain in place and retain the user's numbering and style.
    doc.paragraphs[357].text = (
        "PostgreSQL conserva el estado transaccional de EcoRoute. Las migraciones Flyway V1–V8 "
        "definen 33 tablas, 239 columnas y 27 claves foráneas físicas. El ERD del negocio "
        "selecciona las 18 tablas operativas (159 columnas y 17 FK) y deja fuera las nueve corp_* "
        "de referencia y las seis auth_* de autorización técnica. Keycloak administra su propia "
        "base de identidad, independiente de la base logística."
    )
    doc.paragraphs[359].text = (
        "Figura 2.29. Modelo lógico resumido del núcleo operativo. La vista física integral "
        "de las 18 tablas del negocio se presenta a continuación y se entrega también como SVG "
        "ampliable. Fuente: migraciones Flyway V1–V8 y contraste con PostgreSQL en ejecución."
    )
    doc.paragraphs[361].text = (
        "El ERD dibuja exclusivamente las claves foráneas declaradas. orders.driver_id es un campo "
        "BIGINT sin FK a drivers: la asignación verificable de un pedido se obtiene mediante "
        "orders.route_id → routes.driver_id. Hubs y purchases existen físicamente, pero no tienen "
        "FK al grafo de entregas. Las tablas corp_* conservan datos de demostración y no prueban "
        "una integración con MICOTRANS."
    )
    doc.paragraphs[362].text = (
        "La ruta guarda su geometría en routes.geometry, la duración calculada en "
        "routes.total_duration_min y el orden de cada parada en orders.stop_sequence. "
        "La evidencia se vincula con delivery_proofs.order_id, que es NOT NULL y UNIQUE: "
        "un pedido admite cero o una prueba de entrega. Los indicadores analíticos se derivarán "
        "de pedidos, rutas, historial, POD y costos con reglas de cálculo y cortes temporales "
        "aprobados por Jefatura logística antes de usarse para decidir."
    )
    doc.paragraphs[364].text = (
        "El diccionario documenta las 33 tablas y 239 columnas físicas de Flyway V1–V8. "
        "Las primeras 18 tablas corresponden al dominio operativo; auth_* y corp_* se registran "
        "aparte para no confundir autorización técnica o datos de referencia con el proceso de "
        "entrega. “Opcional” indica que el DDL admite NULL, no que el negocio deba omitir el dato."
    )

    anchor = doc.paragraphs[363]  # 2.5 Diccionario de Datos
    title = inserted_paragraph(doc, anchor, "Vista física integral del negocio", "Heading 3")
    title.paragraph_format.keep_with_next = True
    inserted_paragraph(doc, anchor,
        "La siguiente figura reproduce la vista completa tipo DBeaver: muestra todas las columnas, "
        "tipos PostgreSQL, obligatoriedad, PK, UQ y FK de las 18 tablas operativas. Las líneas "
        "unen la fila FK con la PK declarada. Las entidades aisladas se muestran sin enlaces "
        "inventados. Para leer cada campo a escala real se utiliza el archivo SVG vectorial "
        "docs/base_datos/modelo_erd_negocio/09_vista_completa_tipo_dbeaver.svg.")
    figure = inserted_paragraph(doc, anchor)
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(ERD), width=Inches(6.4))
    caption = inserted_paragraph(doc, anchor,
        "Figura 2.30. ERD físico completo del negocio EcoRoute, sin tablas corporativas ni RBAC técnico. "
        "Fuente: Flyway V1–V8; contrastado con PostgreSQL en ejecución (33 tablas, 239 columnas, 27 FK en el esquema total).")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = False

    inserted_paragraph(doc, anchor, "Trazabilidad entre procesos, pantallas y datos", "Heading 3")
    inserted_paragraph(doc, anchor,
        "La tabla vincula cada BPMN con su soporte visible y su base transaccional. “Actual” "
        "describe funciones que el código o las capturas respaldan; las reglas de aprobación, "
        "calidad de evidencia y gobierno del dato siguen siendo diseño objetivo hasta su "
        "validación en el piloto con MICOTRANS.")
    rows = [
        ("P01 Confirmar pedidos", "Web: pedidos, alta, importación y detalle", "orders; order_items; products; order_status_history", "RF-17–25"),
        ("P02 Preparar jornada", "Web: rutas, asignación y mapa; móvil: Mis rutas", "routes; drivers; vehicles; orders", "RF-26–37; RF-38"),
        ("P03 Ejecutar entregas", "Móvil: ruta, detalle, estado y POD", "orders; delivery_proofs; order_status_history", "RF-38–55"),
        ("P04 Resolver incidencias", "Web/móvil: estado y motivo; alertas", "incidents; orders; order_status_history; routes", "RF-44–47; RF-60–62"),
        ("P05 Evaluar resultados", "Web: panel, reportes y diseños BI", "orders; routes; delivery_proofs; incidents; route_expenses", "RF-63–70; RF-75–81"),
        ("P06 Habilitar recursos", "Web: conductores y vehículos", "drivers; vehicles; driver_shifts; driver_contracts; vehicle_maintenance_logs", "RF-09–16"),
        ("P07 Supervisar jornada", "Web: mapa y seguimiento; móvil: GPS", "vehicle_gps_history; routes; orders; notifications", "RF-56–62"),
        ("P08 Administrar accesos", "Web: usuarios/roles; móvil: login y perfil", "Keycloak (BD separada); auth_* (RBAC local)", "RF-01–08"),
        ("P09 Registrar compras", "Web: listado y formulario de compras", "purchases (sin FK al núcleo operativo)", "RF-71–74"),
    ]
    inserted_table(doc, anchor, ["Proceso BPMN", "Pantallas de referencia", "Datos implicados", "Requisitos"], rows, [1.22, 1.72, 2.45, 0.93])

    inserted_paragraph(doc, anchor, "Controles de negocio pendientes de validación", "Heading 3")
    controls = [
        ("Confirmación y despacho", "Pedido completo y aceptado; capacidad, ventana y recurso aptos; aprobación y acuse antes de iniciar la ruta.", "P01, P02; hoy predominan registro, asignación y cálculo de ruta."),
        ("Entrega y prueba", "Definir evidencia mínima por resultado, hora y ubicación confiables, responsable y política de corrección; evitar sobrescribir POD sin auditoría.", "P03, P04; la tabla POD permite NULL en varios campos y no garantiza estas reglas por sí sola."),
        ("Excepción y reprogramación", "Vincular incidente, decisión, autorización, nueva fecha y resultado con el mismo pedido; conservar historial.", "P04; los estados y motivos existen, pero el expediente completo es objetivo."),
        ("Identidad y autorización", "Alinear roles de Keycloak con permisos locales y auditar cambios; revisar el mapeo de auth_user_role frente al DDL.", "P08; Keycloak y auth_* son almacenes distintos."),
        ("Inteligencia de negocios", "Aprobar definiciones IID, CHR y TDE, fecha de corte, denominador, exclusiones y metas; publicar vistas analíticas de solo lectura.", "P05; panel operativo actual y diseños BI propuestos."),
    ]
    inserted_table(doc, anchor, ["Control", "Regla o decisión que debe aprobar MICOTRANS", "Alcance actual"], controls, [1.23, 3.05, 2.04])
    inserted_paragraph(doc, anchor,
        "El recorrido técnico pedido → ruta → consulta móvil → resultado/POD → panel permite demostrar "
        "un flujo de extremo a extremo. No equivale todavía a afirmar que cada salida “No” del BPMN, "
        "el trabajo sin red o la calidad legal de la evidencia estén cerrados. La aceptación del "
        "piloto debe probar expresamente esas condiciones y sus responsables. También debe "
        "resolver orders.driver_id frente a routes.driver_id antes de imponer una nueva FK, y "
        "definir si hubs y purchases necesitan vínculos físicos con el servicio de reparto.")

    doc.save(DEST)
    print(DEST)


if __name__ == "__main__":
    main()
