"""Alinea el capítulo 2 con el alta opcional de ruta y la brecha de asignación posterior."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from update_chapter2_alignment import inserted_paragraph, inserted_table

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SOURCE = DOCS / 'Integrador I_Grupo3_Capitulo2_Alineado.docx'
DEST = DOCS / 'Integrador I_Grupo3_Capitulo2_Flujo_Corregido.docx'
BUSINESS = DOCS / 'entregables_procesos_ecoroute/negocio_revisado'
WIRES = DOCS / 'entregables_procesos_ecoroute/imagenes'
SHOTS = DOCS / 'screenshots/procesos'


def set_cell(table, row: int, col: int, value: str) -> None:
    table.rows[row].cells[col].text = value


def replace_body_rows(table, rows: list[tuple[str, ...]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)
    for data in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, data):
            cell.text = value


def replace_picture(paragraph, path: Path, width: float) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def main() -> None:
    doc = Document(SOURCE)
    p = doc.paragraphs
    t = doc.tables

    p[215].text = (
        'Este capítulo modela la operación de última milla de Grupo MICOTRANS S. A. C. en nueve procesos BPMN. '
        'La secuencia de negocio propuesta es recibir y validar una solicitud, registrar un pedido PENDING sin ruta, '
        'crear la ruta PLANNED, vincular el pedido, despachar, entregar y evaluar. Se distingue este flujo objetivo '
        'del recorrido técnico actualmente demostrable, que requiere crear primero la ruta y registrar después el '
        'pedido con esa ruta. Las pantallas, BPMN, modelo de datos y diccionario se leen con esa misma distinción.'
    )
    p[218].text = (
        'La unidad principal es el pedido. P01-P05 cubren el ciclo del servicio y P06-P09 los controles de apoyo. '
        'Dispatcher, Jefatura logística, Conductor y Administrador son roles funcionales acordes al acceso y módulos '
        'existentes; el cliente remitente interviene externamente. EcoRoute y Keycloak son sistemas, no áreas de '
        'MICOTRANS. Los puestos formales y las autorizaciones deben validarse con la empresa antes del piloto.'
    )
    p[224].text = (
        'Dependencias: P08 habilita accesos; P06 habilita recursos para P02; P01 entrega pedidos validados PENDING a '
        'P02; P02 crea una ruta PLANNED, requiere pedidos asociados y solo entonces la inicia para P03; P07 supervisa '
        'P03; P03 o P07 abren P04 ante una excepción; P05 evalúa los resultados. P09 documenta compras que pueden '
        'afectar P06. La asignación posterior de un pedido PENDING a una ruta es la brecha funcional que impide '
        'ejecutar hoy esa secuencia ideal sin cambiar el desarrollo.'
    )

    set_cell(t[28],1,2,'Validar solicitudes y registrar pedidos PENDING antes de planificar.')
    set_cell(t[29],1,1,'Solicitud de servicio recibida')
    set_cell(t[29],1,2,'Pedido validado, registrado PENDING sin ruta → P02')
    set_cell(t[29],1,3,'Dispatcher')
    set_cell(t[29],2,1,'Pedidos PENDING validados para una jornada')
    set_cell(t[29],2,2,'Ruta PLANNED con pedidos vinculados, aprobada e iniciada → P03')
    set_cell(t[29],9,3,'Administrador con aprobación de Jefatura logística')
    set_cell(t[30],4,2,'Registro PENDING, vinculación a ruta, entrega y excepción; la vinculación posterior requiere desarrollo.')

    p[225].text = 'P01 Confirmar pedidos'
    p[226].text = (
        'Inicio: solicitud recibida. Resultado: condiciones aceptadas por el cliente y pedido registrado sin ruta '
        '(PENDING). Responsable: Dispatcher. La aceptación es una decisión de negocio; EcoRoute no guarda un estado '
        'CONFIRMED ni una constancia específica de esa confirmación.'
    )
    p[227].text = (
        'Cierre y transferencia. Solo E4 entrega a P02 un pedido PENDING validado. E1 no admite la solicitud; E2 '
        'cierra por dirección no viable; E3 registra el desistimiento. La conversación con el cliente y los motivos '
        'de estos tres cierres necesitan un mecanismo de evidencia acordado por MICOTRANS.'
    )
    replace_picture(p[228], BUSINESS/'imagenes/P01_negocio.png', 6.1)
    p[229].text = 'Figura 2.1. BPMN de P01 Confirmar pedidos, con el Dispatcher como responsable del flujo de negocio. Fuente: elaboración propia y código EcoRoute.'
    p[232].text = (
        'Soporte actual y brecha. El formulario “Registrar Nuevo Pedido” ofrece “-- Sin Ruta --”; el backend asigna '
        'PENDING cuando routeId es nulo y ASSIGNED cuando se elige una ruta existente. La importación CSV también '
        'admite routeId vacío. La confirmación comercial, el control de duplicados y la aceptación de alternativas '
        'no son estados ni validaciones automatizadas. El formulario conserva coordenadas predeterminadas si no se '
        'geocodifica la dirección; el Dispatcher debe verificarla antes de planificar.'
    )

    p[234].text = (
        'Inicio: pedidos PENDING validados para una fecha. Resultado: ruta PLANNED con pedidos asociados, revisada '
        'y luego iniciada. Responsable: Dispatcher; Jefatura logística decide la aprobación propuesta y el Conductor '
        'verifica la salida. La vinculación posterior de los PENDING figura como objetivo porque no existe aún en '
        'la interfaz ni en el API.'
    )
    p[235].text = (
        'Cierre y transferencia. Solo E4 inicia P03 con una ruta IN_PROGRESS. E1 mantiene pedidos en cartera para '
        'acordar nueva fecha; E2 detiene el despacho si los pedidos no están vinculados; E3 suspende la salida si no '
        'hay sustituto. Ninguno de esos cierres negativos representa una entrega fallida.'
    )
    replace_picture(p[236], BUSINESS/'imagenes/P02_negocio.png', 6.1)
    p[237].text = 'Figura 2.2. BPMN de P02 Preparar la jornada de reparto. La asignación posterior marcada como objetivo es una brecha funcional. Fuente: elaboración propia y código EcoRoute.'
    p[240].text = (
        'Soporte actual y brecha. La web crea rutas PLANNED con conductor, vehículo y fecha; permite iniciar el '
        'despacho y el backend cambia a IN_TRANSIT los pedidos ya vinculados. Al iniciar se calcula el recorrido '
        'OSRM. El único modo actual de vincular un pedido es elegir una ruta existente durante el alta: no hay '
        'operación para asociar después un PENDING. Por ello, el end-to-end demostrable sigue el orden técnico '
        'ruta → alta de pedido con routeId → despacho. Aprobar el plan, verificar capacidad y acusar salida son '
        'controles de negocio aún no evidenciados como automatización.'
    )

    replace_body_rows(t[31], [
        ('¿Datos completos?','Dispatcher','Revisar dirección y duplicidad','Solicitar corrección'),
        ('¿Cliente corrige a tiempo?','Dispatcher','Revisar datos mínimos','Solicitud no admitida'),
        ('¿Dirección utilizable?','Dispatcher','Validar cobertura y ventana','Aclarar ubicación con cliente'),
        ('¿Ubicación aclarada?','Dispatcher','Revisar dirección y duplicidad','Solicitud no viable'),
        ('¿Servicio factible?','Dispatcher','Confirmar condiciones con cliente','Proponer fecha o condición alternativa'),
        ('¿Cliente acepta alternativa?','Dispatcher','Validar cobertura y ventana','Solicitud desistida'),
        ('¿Cliente confirma?','Dispatcher','Registrar pedido sin ruta (PENDING)','Solicitud desistida'),
    ])
    replace_body_rows(t[32], [
        ('¿Recursos disponibles?','Dispatcher','Crear ruta PLANNED para la fecha','Ajustar carga o recursos'),
        ('¿Ajuste viable?','Dispatcher','Seleccionar conductor y vehículo','Pedidos en cartera; nueva fecha'),
        ('¿Pedidos vinculados?','Dispatcher','Calcular recorrido sobre red vial','Asignación pendiente; no despachar'),
        ('¿Plan viable?','Dispatcher','Autorizar plan y salida','Corregir secuencia o reparto'),
        ('¿Despacho aprobado?','Jefatura logística','Emitir hoja de ruta','Corregir secuencia o reparto'),
        ('¿Conductor puede salir?','Conductor','Iniciar despacho en web','Gestionar sustituto o nueva fecha'),
        ('¿Sustituto disponible?','Dispatcher','Seleccionar conductor y vehículo','Despacho suspendido; reprogramar'),
    ])

    p[300].text = (
        'Los wireframes se basan en la estructura de las pantallas web y móvil. Las nuevas láminas P01 y P02 '
        'corrigen estados y dependencias: el pedido puede nacer PENDING sin ruta y la ruta nace PLANNED. La '
        'tercera vista de P02 está rotulada como objetivo pendiente de implementar; ninguna otra vista de esa '
        'lámina se presenta como si la aprobación comercial estuviera automatizada. Las láminas preliminares '
        'P03-P06 conservan su numeración histórica y funcionan como referencia visual, no como prueba de flujo.'
    )
    replace_picture(p[301], WIRES/'P01_wireframes_flujo_alineado.png', 6.35)
    p[302].text = 'Figura 2.10. Wireframes de pedido: lista, alta con ruta opcional, CSV y consulta PENDING. Fuente: pantallas actuales de EcoRoute; diseño alineado al flujo.'
    replace_picture(p[303], WIRES/'P02_wireframes_flujo_alineado.png', 6.35)
    p[304].text = 'Figura 2.11. Wireframes de ruta PLANNED, asignación posterior propuesta y despacho con pedidos vinculados. Fuente: pantallas EcoRoute y brecha funcional identificada.'
    p[313].text = (
        'Las figuras 2.10 y 2.11 separan evidencia de interfaz actual y propuesta. La confirmación del pedido '
        'es una decisión del Dispatcher con el cliente, no una acción “Confirmar” disponible en el módulo Pedidos. '
        'Antes del piloto se debe habilitar la asignación posterior y definir la constancia de aceptación, la '
        'validación de ubicación, la aprobación de salida y el tratamiento de pedido sin ruta.'
    )
    p[317].text = 'P01: alta con vínculo a ruta opcional, estado PENDING sin ruta y consulta del pedido'
    p[321].text = 'P02: ruta PLANNED con conductor, vehículo y fecha; iniciar despacho solo con pedidos asociados'
    p[361].text = (
        'El ERD dibuja solamente claves foráneas declaradas. orders.route_id es nullable: permite guardar el '
        'pedido sin ruta y hace físicamente posible asignarlo después, aunque el API actual no ofrece esa acción. '
        'orders.driver_id es BIGINT sin FK a drivers; la asignación verificable del conductor se obtiene mediante '
        'orders.route_id → routes.driver_id. Hubs y purchases existen, pero no tienen FK al grafo de entregas. '
        'Las tablas corp_* no prueban integración con MICOTRANS.'
    )

    set_cell(t[41],1,1,'Web: lista, alta con ruta opcional, importación y detalle; no hay pantalla de confirmación comercial')
    set_cell(t[41],2,1,'Web: crear ruta PLANNED y despacho; asignación posterior de PENDING es diseño objetivo')
    set_cell(t[42],1,0,'Validación, asignación y despacho')
    set_cell(t[42],1,1,'Aceptar condiciones del cliente, registrar PENDING sin ruta, vincularlo a ruta PLANNED y aprobar salida con pedidos completos.')
    set_cell(t[42],1,2,'Alta sin ruta y alta con ruta funcionan; asignación posterior, confirmación comercial y aprobación formal no están implementadas.')

    for row in t[43].rows[1:]:
        name = row.cells[0].text.strip()
        if name == 'status':
            row.cells[3].text = 'PENDING si se crea sin ruta; ASSIGNED si se crea con routeId. No existe CONFIRMED.'
        elif name == 'route_id':
            row.cells[2].text = 'FK → routes.id; nullable'
            row.cells[3].text = 'Ruta opcional al registrar. La asignación posterior carece de API/UI actual.'

    # La captura del modal despeja la ambigüedad que originó esta revisión.
    anchor = p[320]
    new_title = inserted_paragraph(doc, anchor, 'Modal de alta del pedido', 'Heading 3')
    new_title.paragraph_format.keep_with_next = True
    inserted_paragraph(doc, anchor,
        'La captura original muestra el selector “Vincular a Ruta” con “-- Sin Ruta --” como opción inicial. '
        'Es un campo opcional; “Registrar” crea el pedido, pero no confirma comercialmente el servicio.')
    image_para = inserted_paragraph(doc, anchor)
    image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_para.add_run().add_picture(str(SHOTS/'P01-alta-pedido-focus.png'), width=Inches(5.9))
    inserted_paragraph(doc, anchor,
        'Figura 2.16a. Captura del modal Registrar Nuevo Pedido: ruta opcional. Fuente: aplicación web EcoRoute.')

    doc.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
