"""Prepare and build four new revisions without rewriting retained OOXML parts.

Run with the bundled Python. --prepare is read-only for DOCX sources and emits
the patch contract. --build needs reviewed current evidence configured below.
Artifact-operation marker is owned by the principal and must not be repeated.
"""
from __future__ import annotations
import argparse
from copy import copy, deepcopy
from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/control_actualizacion/documentos'
BASE = OUT / '_qa/baseline'
CONFIG = ROOT / 'tools/control_actualizacion/documents_evidence.json'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'
OUTPUT_NAMES = {
    'academico': 'EcoRoute_Academico_Revision_20261026.docx',
    'manual': 'EcoRoute_Manual_Revision_20261026.docx',
    'guion': 'EcoRoute_Guion_Revision_20261026.docx',
    'informe': 'EcoRoute_Informe_Tecnico_Actualizacion_20261026.docx',
}

def sha(value):
    return hashlib.sha256(value).hexdigest()

def load(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))

def text(p):
    return ''.join(p.xpath('.//w:t/text()', namespaces=NS))

def inventories():
    return {key: load(BASE / f'{key}_inventory.json') for key in ('academico', 'manual', 'guion')}

def known_patches(data):
    auth = ('CheckAuthStatusEvent restaura la sesión y SessionRefresh coordina su renovación. '
            'AuthInterceptor distingue el rechazo del refresh del error de la petición repetida: '
            'si la renovación tuvo éxito y el recurso responde 400, informa ese error sin cerrar la sesión; '
            'un segundo 401 sí activa la expiración de la misma cuenta. GoRouter mantiene la navegación '
            'según AuthBloc. Los tests session_navigation_auth_test.dart y session_navigation_router_test.dart '
            'cubren esta distinción. Son pruebas técnicas; UAT066–UAT068, UAT070 y UAT091 siguen '
            'requiriendo evidencia de aceptación humana.')
    queue = ('core/sync/proof_sync.dart mantiene una ejecución por cuenta (_flights), '
             'un timer de treinta segundos con el proceso abierto y reintentos al volver a primer plano '
             'o restaurar la cuenta. _drain elimina campos privados antes de POST /delivery-proofs, '
             'incluye Idempotency-Key y retira el pendiente solo después de una respuesta satisfactoria. '
             'Un fallo de red conserva el contenido. Los rechazos 4xx permanentes, salvo 401, 408 y 429, '
             'quedan visibles para revisión. La corrección de AuthInterceptor permite que un 400 del '
             'recurso después de renovar credenciales llegue a esta lógica sin convertirse en logout. '
             'No se garantiza sincronización con la aplicación terminada por el sistema operativo.')
    patches = {'academico': [], 'manual': [], 'guion': []}
    def add(key, index, replacement, reason, sources):
        p = data[key]['paragraphs'][index]
        patches[key].append({'paragraphIndex': index, 'expectedText': p['text'],
                             'expectedXmlSha256': p['xmlSha256'], 'replacement': replacement,
                             'reason': reason, 'sources': sources})
    for key in ('manual', 'guion'):
        add(key, 3, 'Revisión documental del 26 de octubre de 2026 con evidencia técnica local '
            'del 22 de octubre. La revisión técnica del 20 de octubre se conserva como antecedente. '
            'Los ensayos no sustituyen la aceptación de personas de MICOTRANS.',
            'Fecha documental y corte de pruebas separados', [])
    auth_sources = ['mobile-app/lib/core/network/auth_interceptor.dart',
                    'mobile-app/test/session_navigation_auth_test.dart',
                    'mobile-app/test/session_navigation_router_test.dart',
                    'docs/control_integral_20261022/e2e_cierre/calidad_mobile_final.json']
    add('manual', 1179, auth, 'Separar refresh rechazado de replay con error 400', auth_sources)
    add('manual', 1197, queue, 'Conectar la cola existente con el comportamiento de sesión corregido', auth_sources)
    add('guion', 130, '“' + auth + '”', 'Explicación de sesión en la demostración móvil', auth_sources)
    gps_source = ['docs/control_actualizacion/evidencias/api_v16.json']
    add('guion', 239, '“Flutter envía muestras GPS con contexto operativo y posición válida; el backend '
        'verifica identidad, vehículo y ruta. La web usa WebSocket autenticado y consulta HTTP de respaldo. '
        'OSRM aporta geometría vial. V16 controla coordenadas, tiempo, velocidad e intervalo antes de '
        'publicar distancia operativa. La suma observada queda como auditoría; una traza rechazada '
        'produce distancia y km por entrega nulos. El mapa no acredita una entrega, llamada o sanción.”',
        'Control de calidad GPS en la narración', gps_source)
    add('guion', 253, '“Eficiencia separa plan, GPS evaluable y duración operativa. V16 conserva '
        'Haversine y exige superar los controles de calidad. La ruta 20 mantiene IMPLAUSIBLE_SPEED, '
        'distancia operativa y km por entrega nulos; la ruta 21 conserva cero válido. Km por entrega '
        'divide por DELIVERED y minutos por parada por DELIVERED, FAILED y RETURNED. El CSV deja '
        'vacíos los GPS no evaluables y no publica la suma bruta. Los cinco KPI mantienen su '
        'población y fórmulas; no se rellenan ausencias con el plan.”',
        'Explicación V16 y CSV sin distancia espuria', gps_source)
    return patches

def prepare():
    OUT.mkdir(parents=True, exist_ok=True)
    data = inventories()
    plan = {'schemaVersion': 1, 'scope': 'LOCAL_DOCUMENTAL_NO_UAT',
            'operationMarker': 'Ejecutado por el principal edit expected 4 docx; no repetir',
            'sources': [{k: row[k] for k in ('key', 'source', 'sha256')} for row in data.values()],
            'outputs': OUTPUT_NAMES, 'localizedPatches': known_patches(data),
            'pendingSlots': [
                {'topic': 'mapa móvil', 'slots': {'manual': [1181, 1184, 1185], 'guion': [132, 135, 136]},
                 'gate': 'Reporte actual en docs/control_actualizacion y contraste con código; no afirmar corrección antes.'},
                {'topic': 'GPS servidor', 'slots': {'manual': [1184, 1239], 'guion': [135]},
                 'gate': 'Reporte del worker backend con IDs, SQL/HTTP, fechas y alcance.'},
                {'topic': 'cinco KPI', 'slots': {'manual': [1215, 1220, 1224, 1228, 1232, 1236, 1237, 1238, 1239, 1240]},
                 'gate': 'Reporte actual API/SQL del worker; no copiar cifras de otra población ni redefinir fórmulas.'}],
            'insertions': {
                'academico': 'Nueva actualización al final de capítulo 2, antes de CAPÍTULO 3; capítulos 1 y 2 retenidos sin supresiones.',
                'manual': 'Actualización de sesiones y cola in situ y sección final de evidencia local con fuentes concretas.',
                'guion': 'Corrección del texto móvil y sección final para narrar dos recorridos distintos y sus límites.',
                'informe': 'Informe técnico nuevo con problema, cambio, fuente, ejecución, límite y siguiente comprobación por tema.'},
            'preservation': ['Fuentes originales byte a byte', 'Todos los ZIP parts no editados',
                             'Párrafos no afectados en orden y XML idéntico', 'Secciones, estilos, medios, relaciones, tablas y BPMN'],
            'rendering': {'canonical': 'render_docx.py intentado: soffice.exe ausente en runtime Windows',
                          'fallback': 'Word COM serial solo este worker; sin LibreOffice del usuario',
                          'gate': 'Cada página PNG al 100 por ciento; no sustituir por extracción XML o miniaturas'},
            'status': 'PREPARADO_PENDIENTE_REPORTES_ACTUALES_Y_QA'}
    if CONFIG.exists():
        current = load(CONFIG)
        plan['currentGates'] = {key: value['status'] for key, value in current['currentGates'].items()}
        for slot, topic in zip(plan['pendingSlots'], ('mapa', 'gps', 'kpi')):
            slot['status'] = current['currentGates'][topic]['status']
        if all(v == 'REVISADO_CON_FUENTE_CONCRETA' for v in plan['currentGates'].values()):
            plan['status'] = 'FUENTES_REVISADAS_LISTO_PARA_GENERAR_Y_RENDERIZAR'
    (OUT / 'plan_parches.json').write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf8')
    return plan

def replace_text_preserving_runs(p, replacement):
    nodes = p.xpath('.//w:t', namespaces=NS)
    if not nodes or p.xpath('.//w:drawing|.//w:fldChar|.//w:instrText', namespaces=NS):
        raise ValueError('El parche requiere un párrafo de texto simple verificado')
    # Preserve paragraph/run properties and bookmarks. Redistribute text across
    # the original text nodes instead of recreating the paragraph or its runs.
    remaining = replacement
    for index, node in enumerate(nodes):
        length = len(node.text or '')
        node.text = remaining if index == len(nodes) - 1 else remaining[:length]
        remaining = '' if index == len(nodes) - 1 else remaining[length:]
        node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')

def paragraph_node(value, style=None, template=None, page_break=False):
    p = etree.Element(W + 'p')
    if template is not None:
        props = template.find(W + 'pPr')
        if props is not None:
            p.append(deepcopy(props))
    props = p.find(W + 'pPr')
    if props is None:
        props = etree.SubElement(p, W + 'pPr')
    # New update paragraphs inherit the source body rhythm. New headings are
    # kept black and are not added to the retained academic TOC hierarchy.
    if style:
        existing = props.find(W + 'pStyle')
        if existing is not None:
            props.remove(existing)
        etree.SubElement(props, W + 'pStyle').set(W + 'val', style)
    if page_break:
        etree.SubElement(props, W + 'pageBreakBefore')
    r = etree.SubElement(p, W + 'r')
    if template is not None:
        rpr = template.find('.//' + W + 'rPr')
        if rpr is not None:
            r.append(deepcopy(rpr))
    if style:
        rpr = r.find(W + 'rPr')
        if rpr is None:
            rpr = etree.SubElement(r, W + 'rPr')
        etree.SubElement(rpr, W + 'color').set(W + 'val', '000000')
    etree.SubElement(r, W + 't').text = value
    return p

def evidence_rows(config, final):
    rows = []
    for row in config['evidence']:
        path = ROOT / row['file']
        if not path.is_file():
            raise FileNotFoundError(path)
        if row.get('sha256') and sha(path.read_bytes()) != row['sha256']:
            raise ValueError('Evidencia cambió; revisar antes de incorporar: ' + row['file'])
        rows.append({**row, 'sha256': sha(path.read_bytes())})
    if final:
        for topic in ('mapa', 'gps', 'kpi'):
            gate = config['currentGates'].get(topic, {})
            if gate.get('status') != 'REVISADO_CON_FUENTE_CONCRETA' or not gate.get('file') or not gate.get('paragraphs'):
                raise ValueError('No finalizar sin reporte actual revisado y texto concreto de ' + topic)
            if not (ROOT / gate['file']).is_file():
                raise FileNotFoundError(gate['file'])
    return rows

def update_blocks(config, rows):
    blocks = [('Revisión documental del 26 de octubre de 2026', 'heading')]
    for paragraph in config['updateParagraphs']:
        blocks.append((paragraph, 'body'))
    for topic in ('mapa', 'gps', 'kpi'):
        gate = config['currentGates'].get(topic, {})
        if gate.get('status') == 'REVISADO_CON_FUENTE_CONCRETA':
            blocks.append((gate['title'], 'heading2'))
            blocks.extend((p, 'body') for p in gate['paragraphs'])
    blocks.append(('Relación con los nueve procesos y las excepciones', 'heading2'))
    blocks.extend((p, 'body') for p in config.get('processMapping', []))
    blocks.append(('Fuentes de la actualización', 'heading2'))
    for row in (r for r in rows if r.get('includeInDocument', True)):
        blocks.append((row['id'] + '. ' + row['description'] + '. Fuente: ' + row['file'] + '.', 'body'))
    return blocks

def patch_revision(key, baseline, patches, blocks, dest):
    source = ROOT / baseline['source']
    if sha(source.read_bytes()) != baseline['sha256']:
        raise ValueError('Fuente vigente cambió; no sobrescribir ni reutilizar locadores obsoletos: ' + str(source))
    with ZipFile(source) as archive:
        original = archive.read('word/document.xml')
        xml = etree.fromstring(original)
        paragraphs = xml.xpath('//w:p', namespaces=NS)
        for patch in patches:
            p = paragraphs[patch['paragraphIndex']]
            if text(p) != patch['expectedText'] or sha(etree.tostring(p)) != patch['expectedXmlSha256']:
                raise ValueError('Locator/hash de párrafo no coincide')
            replace_text_preserving_runs(p, patch['replacement'])
        unaffected = [(i, etree.tostring(p)) for i, p in enumerate(paragraphs)
                      if i not in {patch['paragraphIndex'] for patch in patches}]
        body = xml.find(W + 'body')
        body_template = paragraphs[1219 if key == 'academico' else 2]
        if key == 'academico':
            anchor = next(p for p in paragraphs if text(p).startswith('CAPÍTULO 3')
                          and p.getparent() is body)
        else:
            anchor = body.find(W + 'sectPr')
        for i, (value, role) in enumerate(blocks):
            style = 'Heading1' if role == 'heading' and key != 'academico' else 'Heading2' if role == 'heading2' and key != 'academico' else None
            node = paragraph_node(value, style=style, template=body_template, page_break=i == 0)
            if key == 'academico' and role != 'body':
                # Heading-like body text preserves the academic TOC and numbering.
                rpr = node.find('.//' + W + 'rPr')
                if rpr is None:
                    rpr = etree.SubElement(node.find(W + 'r'), W + 'rPr')
                etree.SubElement(rpr, W + 'b')
            anchor.addprevious(node)
        replacement = etree.tostring(xml, encoding='UTF-8', xml_declaration=True, standalone=True)
        with ZipFile(dest, 'w') as output:
            for info in archive.infolist():
                output.writestr(copy(info), replacement if info.filename == 'word/document.xml' else archive.read(info.filename))
        # Node identity survives insertion, so every unaffected original paragraph
        # must still be byte-identical in the edited XML tree and original order.
        assert all(etree.tostring(paragraphs[i]) == raw for i, raw in unaffected)
        with ZipFile(dest) as output:
            changed = [name for name in archive.namelist() if archive.read(name) != output.read(name)]
            assert changed == ['word/document.xml'], changed
    assert sha(source.read_bytes()) == baseline['sha256']
    return {'source': baseline['source'], 'sourceSha256': baseline['sha256'],
            'output': str(dest.relative_to(ROOT)), 'outputSha256': sha(dest.read_bytes()),
            'changedParts': changed, 'localizedPatches': patches,
            'unaffectedParagraphsIdentical': len(unaffected), 'addedParagraphs': len(blocks),
            'originalUnchanged': True}

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    widths = [2.1, 1.0, 2.8, 4.4, 6.3] if len(headers) == 5 else [16.6 / len(headers)] * len(headers)
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    for i, title in enumerate(headers):
        t.rows[0].cells[i].text = title
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                item = OxmlElement('w:' + edge)
                for attr, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                    item.set(qn('w:' + attr), value)
                borders.append(item)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), '100'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EEF4' if ri == 0 else 'FFFFFF'); tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(9); run.bold = ri == 0
        if ri == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph()

def report(config, rows, dest):
    doc = Document()
    sec = doc.sections[0]; sec.page_width = Cm(21.59); sec.page_height = Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(2); sec.left_margin = sec.right_margin = Cm(2.2)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]; style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    doc.add_paragraph('Informe técnico de actualización de EcoRoute', style='Title')
    doc.add_paragraph('Revisión documental del 26 de octubre de 2026', style='Subtitle')
    doc.add_paragraph('La revisión integra sesión, cartografía, calidad GPS, indicadores y trazabilidad con evidencia '
                      'técnica del 22 de octubre de 2026 en America/Lima. No se atribuye una nueva ejecución al 26 de octubre. '
                      'Permite localizar cada resultado en su reporte y distinguir el estado del producto, '
                      'la verificación técnica y la aceptación empresarial pendiente.')
    doc.add_heading('Estado de los recorridos locales', level=1)
    table(doc, ['Recorrido', 'Ruta', 'Entrega', 'Excepción', 'Cierre'], [
        ['E2E local', 20, '44 DELIVERED', '45 FAILED e incidencia 12', 'Reporte final identificado'],
        ['Video local', 21, '46 DELIVERED', '47 FAILED e incidencia 13', 'RESOLVED CLOSE_FAILED y 26 de 26 checks'],
    ])
    doc.add_paragraph('Los identificadores pertenecen a ensayos diferentes. No se mezclan sus contadores ni se presentan '
                      'como una aceptación de MICOTRANS. CLOSE_FAILED conserva el pedido fallido y no acredita un nuevo intento.')
    for item in config['reportSections']:
        doc.add_heading(item['title'], level=1)
        for paragraph in item['paragraphs']:
            doc.add_paragraph(paragraph)
    for topic in ('mapa', 'gps', 'kpi'):
        gate = config['currentGates'].get(topic, {})
        if gate.get('status') == 'REVISADO_CON_FUENTE_CONCRETA':
            doc.add_heading(gate['title'], level=1)
            for paragraph in gate['paragraphs']:
                doc.add_paragraph(paragraph)
    doc.add_heading('Relación con los nueve procesos y las excepciones', level=1)
    for paragraph in config.get('processMapping', []):
        doc.add_paragraph(paragraph)
    doc.add_heading('Registro de fuentes y alcance', level=1)
    for row in (r for r in rows if r.get('includeInDocument', True)):
        doc.add_paragraph(row['id'] + '. ' + row['description'], style='Heading 2')
        doc.add_paragraph('Archivo del repositorio: ' + row['file'])
        doc.add_paragraph('SHA256: ' + row['sha256'])
    doc.core_properties.created = datetime(2026, 10, 26, 0, 0)
    doc.core_properties.modified = datetime(2026, 10, 26, 0, 0)
    buffer = io.BytesIO(); doc.save(buffer)
    # Canonicalize member times for reproducible output bytes.
    with ZipFile(io.BytesIO(buffer.getvalue())) as archive, ZipFile(dest, 'w') as output:
        for info in archive.infolist():
            info.date_time = (2026, 10, 26, 0, 0, 0)
            output.writestr(copy(info), archive.read(info.filename))
    return {'output': str(dest.relative_to(ROOT)), 'outputSha256': sha(dest.read_bytes()), 'newReport': True}

def build(plan, revision, final=False):
    if not re.fullmatch(r'[a-z0-9_-]+', revision):
        raise ValueError('Use una revisión simple y nueva')
    config = load(CONFIG)
    rows = evidence_rows(config, final)
    folder = OUT / revision
    if (folder / 'build_manifest.json').exists():
        raise FileExistsError('No sobrescribir revisión completa: ' + str(folder))
    folder.mkdir(parents=True, exist_ok=True)
    baseline = inventories()
    patches = deepcopy(plan['localizedPatches'])
    for topic, gate in config['currentGates'].items():
        if gate.get('status') == 'REVISADO_CON_FUENTE_CONCRETA':
            for patch in gate.get('localizedPatches', []):
                key, index = patch['document'], patch['paragraphIndex']
                old = baseline[key]['paragraphs'][index]
                patches[key].append({'paragraphIndex': index, 'expectedText': old['text'],
                                     'expectedXmlSha256': old['xmlSha256'], 'replacement': patch['replacement'],
                                     'reason': topic, 'sources': [gate['file']]})
    records = []
    blocks = update_blocks(config, rows)
    for key in ('academico', 'manual', 'guion'):
        dest = folder / OUTPUT_NAMES[key]
        pending = folder / (OUTPUT_NAMES[key] + '.pending')
        record = patch_revision(key, baseline[key], patches[key], blocks, pending)
        if dest.exists():
            if dest.read_bytes() != pending.read_bytes():
                raise ValueError('La revisión parcial difiere; usar otra revisión: ' + str(dest))
            pending.unlink()
        else:
            pending.rename(dest)
        record['output'] = str(dest.relative_to(ROOT))
        records.append(record)
    records.append(report(config, rows, folder / OUTPUT_NAMES['informe']))
    manifest = {'scope': 'LOCAL_DOCUMENTAL_NO_UAT', 'revision': revision,
                'status': 'GENERADO_PENDIENTE_RENDER_Y_REVISION_PAGINA_A_PAGINA',
                'finalSourcesReviewed': final, 'visualReviewComplete': False,
                'sources': rows, 'documents': records}
    (folder / 'build_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf8')
    return manifest

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--prepare', action='store_true')
    parser.add_argument('--build', action='store_true')
    parser.add_argument('--revision', default='r01')
    parser.add_argument('--final-sources', action='store_true')
    args = parser.parse_args()
    plan = prepare()
    if args.build:
        manifest = build(plan, args.revision, args.final_sources)
        print(json.dumps({'status': manifest['status'], 'outputs': [d['output'] for d in manifest['documents']]}, ensure_ascii=False))
    else:
        print(json.dumps({'status': plan['status'], 'plan': str((OUT / 'plan_parches.json').relative_to(ROOT)),
                          'knownPatches': sum(len(v) for v in plan['localizedPatches'].values())}, ensure_ascii=False))

if __name__ == '__main__':
    main()
