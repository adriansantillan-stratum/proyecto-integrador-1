"""Guion oral íntegro para las 40 láminas del capítulo 2 de EcoRoute."""

from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Guion_Detallado_Exposicion_Capitulo_2_EcoRoute.docx"
DECK = "EcoRoute_Capitulo_2_Diseno_Exposicion.pptx"


def slide(n: int, title: str, cue: str, speech: str, bridge: str, seconds: int, section: str) -> dict:
    return dict(n=n, title=title, cue=cue, speech=dedent(speech).strip(),
                bridge=bridge, seconds=seconds, section=section)


SLIDES = [
    slide(1, "Portada del capítulo 2", "Abrir la exposición y mantener visible la pregunta central de la portada.", """
        Buenos días, profesor. Somos el Grupo 3 y presentamos el capítulo 2 de EcoRoute, aplicado al caso de Grupo MICOTRANS S. A. C. En este capítulo explicamos cómo se diseñaría una entrega que pueda planificarse, seguirse y verificarse. La pregunta de la portada nos obliga a conectar la operación de la empresa con las decisiones del sistema: primero definimos los procesos, luego las pantallas que los acompañan y finalmente los datos que conservan la trazabilidad.

        Mostraremos propuestas de proceso junto con interfaces que ya existen en el proyecto. Cuando hablemos de una regla de negocio, indicaremos si forma parte del diseño objetivo o si está respaldada por una pantalla actual. No presentaremos las cifras de las maquetas como resultados operativos ya obtenidos por MICOTRANS.
    """, "Pasemos al hilo que une las cinco partes del capítulo.", 50, "Apertura"),

    slide(2, "Relación entre pedido ruta entrega y datos", "Recorrer visualmente las tres palabras de izquierda a derecha: Pedido, Ruta y Entrega.", """
        El diseño parte de una secuencia simple. Primero se recibe una solicitud y se confirma que el pedido tiene datos suficientes y condiciones aceptadas. Después se incorpora a una jornada: se asignan conductor y vehículo, se ordenan las paradas y se decide si la ruta puede salir. Por último, el conductor atiende cada parada y registra un resultado verificable, sea entrega, ausencia, rechazo o incidencia.

        Alrededor de esta secuencia se necesitan controles. Un GPS muestra dónde está el vehículo, pero no acredita por sí mismo que el cliente recibió el pedido. Una foto aislada tampoco explica un rechazo o una reprogramación. Por eso el diseño combina estado, hora, ubicación, evidencia y responsable de cada decisión. Los nueve BPMN desarrollan estas reglas y muestran qué ocurre también cuando la respuesta es No.
    """, "Primero veremos cómo agrupamos esos nueve procesos.", 80, "Apertura"),

    slide(3, "Mapa de macroprocesos", "Señalar el macroproceso de operaciones de reparto y luego recorrer los otros cinco.", """
        Para ordenar todo el proyecto usamos seis macroprocesos de gestión. Gestión de pedidos contiene P01, que confirma la solicitud antes de planificar. Gestión de recursos logísticos contiene P06, porque un conductor o vehículo debe estar habilitado antes de asignarse. El núcleo es Gestión de operaciones de reparto: allí están P02, preparación de la jornada; P03, ejecución de entregas; P04, resolución de incidencias; y P07, supervisión en paralelo.

        Gestión del desempeño logístico contiene P05, que convierte los resultados de la operación en indicadores y acciones de mejora. Gestión de seguridad de acceso contiene P08; Gestión de compras contiene P09. Con este mapa podemos asociar los trece módulos y ochenta y un requerimientos funcionales del proyecto a procesos de negocio, sin inventar un BPMN para cada botón o método de programación.
    """, "Antes de leer las compuertas, aclaremos la convención metodológica.", 95, "Apertura"),

    slide(4, "PMBOK y BPMN", "Apuntar primero a Dirección del proyecto y luego a Operación de MICOTRANS.", """
        Aquí distinguimos dos niveles. PMBOK nos ayuda a dirigir el proyecto de desarrollo: definir alcance, planificar el trabajo, ejecutar, hacer seguimiento y cerrar entregables. En el capítulo 1 eso se refleja en el acta de constitución, la EDT o WBS y el cronograma. BPMN, en cambio, modela cómo funciona la operación de MICOTRANS: actores, tareas, decisiones, excepciones y resultados.

        Por eso llamamos Gestión de pedidos o Gestión de operaciones de reparto a los macroprocesos, mientras que cada BPMN concreto lleva verbo y objeto, como Confirmar pedidos o Resolver incidencias. No existe una prohibición literal de escribir la conjunción y en BPMN; elegimos estos nombres porque cada proceso debe producir un resultado claro y medible. Al revisar un diagrama, seguiremos siempre su evento inicial, el carril que decide, el camino Sí, el camino No y su evento final.
    """, "Con esa forma de lectura, empezamos por la admisión del pedido.", 105, "Apertura"),

    slide(5, "P01 Confirmar pedidos", "Seguir los carriles Atención comercial, Coordinación de pedidos y Soporte EcoRoute; detenerse en cada compuerta.", """
        P01 empieza cuando llega una solicitud de servicio. Atención comercial registra los datos y condiciones; Coordinación de pedidos revisa la información mínima. La primera pregunta es si los datos están completos. Si la respuesta es No, comercial solicita una corrección. Si el cliente responde dentro del plazo, se revisa otra vez; si no responde, el proceso termina como solicitud no admitida. Así evitamos que un pedido incompleto se convierta silenciosamente en una parada de ruta.

        Con datos completos, el soporte verifica ubicación y duplicidad. Si la dirección no es utilizable, se pide al cliente una aclaración. Si logra aclararse, se repite la verificación; si no, el caso termina como no viable. Luego coordinación valida cobertura y ventana horaria. Si el servicio no es factible, comercial propone una fecha u otra condición. Cuando el cliente acepta, se vuelve a comprobar factibilidad; si no acepta, la solicitud se cierra como desistida.

        Solo después de superar estas decisiones se confirman las condiciones y se comunica la aceptación. El evento final Pedido confirmado habilita P02. Los otros tres eventos finales conservan motivo y no envían el pedido a planificación. La consola web ya permite registrar, importar y consultar pedidos; plazos de subsanación y aceptación comercial son reglas objetivo que deben formalizarse con MICOTRANS.
    """, "Una vez confirmado el pedido, la pregunta es si la jornada puede salir.", 160, "Procesos BPMN"),

    slide(6, "P02 Preparar la jornada de reparto", "Recorrer Dispatcher, Jefatura logística, Conductor y Soporte EcoRoute; distinguir cálculo de ruta y autorización de salida.", """
        P02 comienza con pedidos confirmados para una jornada. El Dispatcher consolida la demanda y asigna conductor y vehículo. Si los recursos no alcanzan, ajusta carga o recursos. Si el ajuste es viable, vuelve a asignar; si no, los pedidos quedan pendientes para acordar otra fecha mediante P01. Esta es la primera salida No del proceso y no debe confundirse con una ruta despachada.

        Si hay recursos, el soporte calcula el recorrido sobre la red vial. El Dispatcher revisa capacidad y ventanas horarias. Si el plan no es viable, cambia secuencia o reparto y vuelve a calcular; no se continúa con una ruta observada. Cuando el plan es viable, Jefatura logística lo revisa y autoriza. Si no lo aprueba, la observación regresa al Dispatcher y se corrige dentro de P02.

        Con aprobación se emiten manifiesto y hoja de ruta. El conductor revisa la hoja, la carga y el estado del vehículo. Si no puede salir, se busca un sustituto o se ajusta la asignación. Si tampoco existe sustituto, el despacho se suspende y los pedidos retornan a cartera; si se consigue, vuelve la asignación y revisión. Solo el acuse de inicio del conductor cierra P02 como ruta en ejecución y abre P03. La web y OSRM ya apoyan el trazado; el control de capacidad, la autorización formal y el acuse son condiciones de negocio que deben validarse en el piloto.
    """, "Con una ruta iniciada, examinemos cada parada de entrega.", 180, "Procesos BPMN"),

    slide(7, "P03 Ejecutar entregas", "Seguir una sola parada; marcar las salidas por acceso, ausencia, rechazo, POD y sincronización.", """
        P03 se lee por parada de una ruta ya iniciada. El conductor toma la siguiente dirección, se desplaza y verifica si puede acceder. Si el destino es inaccesible, registra el impedimento y esa parada abre P04. Si llega, contacta al receptor. Un receptor ausente produce un resultado de no entrega y también abre P04. Si el receptor está presente, el conductor entrega los bultos y valida la recepción; un rechazo se documenta con motivo antes de pasar a incidencias.

        Cuando la recepción es aceptada, el conductor captura foto, firma y datos del receptor. La compuerta pregunta si la prueba de entrega o POD es suficiente. Si falta información, intenta completarla; si no puede, deja la prueba observada y envía el caso a P04. Cuando la evidencia es suficiente, el soporte asocia hora y ubicación e intenta sincronizar. Si la sincronización no se confirma, conserva el POD pendiente y abre seguimiento en P04. No debemos llamar entrega verificada a una evidencia que aún no llegó y no ha sido comprobada.

        Con sincronización confirmada se cierra la parada. Si quedan paradas, el conductor vuelve a tomar la siguiente; si no quedan, la ruta termina y sus resultados alimentan P05. Una excepción afecta a su parada: el Dispatcher puede permitir que el conductor continúe con otras entregas. La app ya muestra rutas y permite cambios de estado con foto y firma; offline garantizado y GPS obligatorio siguen siendo controles que necesitan prueba operativa.
    """, "Ahora veremos cómo se resuelve una parada que terminó con excepción.", 185, "Procesos BPMN"),

    slide(8, "P04 Resolver incidencias", "Señalar quién documenta, quién decide y cuándo interviene Jefatura.", """
        P04 comienza con una excepción de entrega, despacho o POD. El conductor describe la causa y adjunta el sustento. El Dispatcher revisa evidencia y urgencia. Si el sustento no alcanza, pide complementarlo. Cuando puede completarse, el caso vuelve a revisión; cuando no, se deja marcada la brecha de evidencia. El caso no desaparece ni se presenta como resuelto por el solo hecho de haber cambiado un estado.

        Después, el Dispatcher evalúa impacto y alternativas. Si es recuperable hoy, coordina reintento o sustitución. Solo cuando el cliente y los recursos aceptan se autoriza el reintento; el evento final lleva la parada a P03. Si no es recuperable hoy o alguien no acepta ese reintento, Jefatura logística participa en el acuerdo de una nueva fecha. Una fecha aceptada produce reprogramación y el pedido vuelve a P02 para planificarse en la nueva jornada.

        Si tampoco se acepta la reprogramación, Jefatura autoriza el cierre fallido con motivo. Ese cierre alimenta P05 para analizar la falla. Así quedan tres salidas distintas: reintentar, reprogramar o cerrar justificadamente. El sistema ya registra estados y razones, pero el expediente completo de incidencia, la aprobación de la reprogramación y la trazabilidad de una reasignación son reglas por validar con la empresa.
    """, "Las decisiones de cada jornada se revisan en el proceso de desempeño.", 165, "Procesos BPMN"),

    slide(9, "P05 Evaluar resultados", "Recorrer primero el control de calidad de datos y luego la cadena causa, acción y verificación.", """
        P05 no nace de un único pedido: se activa al cierre de una jornada o de un periodo de evaluación. El soporte analítico consolida pedidos, rutas y POD. La primera compuerta pregunta si los datos son confiables. Si no lo son, el Dispatcher recibe solicitudes de corrección. Cuando se corrigen, la consolidación se repite; si no, el cierre queda provisional y no se usa como base definitiva para una decisión.

        Con datos confiables se calculan los indicadores y el analista compara resultados con metas. Si no existe un desvío relevante, documenta la buena práctica y cierra el periodo conforme. Si aparece un desvío, analiza la causa raíz y propone una acción con responsable. Jefatura decide si aprueba la acción. Si no la aprueba, se ajusta la propuesta; si la aprueba, asigna plazo y recursos.

        El Dispatcher ejecuta la acción y mide su efecto. Si la mejora no funciona, el análisis vuelve a la causa y se replantea. Solo una mejora comprobada permite cerrar y estandarizar la acción. Aquí el dashboard no reemplaza el juicio de Jefatura: sirve para ver tendencias y sustentar la decisión. El panel y los reportes actuales son el punto de partida; fórmulas, metas y líneas base deben conciliarse con MICOTRANS antes de afirmar mejoras.
    """, "Ese ciclo necesita recursos aptos, no solo registros en pantalla.", 165, "Procesos BPMN"),

    slide(10, "P06 Habilitar recursos de reparto", "Seguir Jefatura, Administración y Dispatcher; destacar la diferencia entre alta y aptitud.", """
        P06 empieza cuando se solicita un conductor o un vehículo. Jefatura define el perfil o la capacidad requerida y Administración reúne datos y documentos. Si faltan documentos, solicita subsanación. Si llegan dentro del plazo, vuelve a revisarlos; si no, archiva la solicitud sin habilitar el recurso.

        Con información completa, Jefatura comprueba la licencia del conductor o el estado del vehículo. Si el recurso no es apto, bloquea su asignación e informa el motivo. Si es apto, el soporte lo registra y el Dispatcher verifica capacidad y disponibilidad para la jornada. Si no está disponible hoy, se agenda para una fecha futura; si sí está disponible, Jefatura lo libera para asignación.

        Solo el evento Recurso habilitado permite que P02 lo considere. Los otros cierres son solicitud archivada, recurso no habilitado y recurso en espera. La aplicación ya administra conductores y vehículos, pero el alta técnica no prueba que la licencia esté vigente, que la unidad haya pasado una revisión o que esté libre para esa fecha. Esos controles deben acordarse con la empresa y dejar evidencia.
    """, "Una vez que salen las rutas, el Dispatcher necesita supervisarlas.", 140, "Procesos BPMN"),

    slide(11, "P07 Supervisar la jornada", "Señalar que P07 corre en paralelo a P03 y no acredita una entrega.", """
        P07 empieza con una ruta en ejecución. El soporte recibe posición y estado transmitidos desde campo. Si no llegan datos recientes, el Dispatcher contacta al conductor. Cuando el seguimiento se recupera, vuelve al monitoreo; si no, registra la pérdida de señal y abre una alerta técnica en P04. La ausencia de GPS no demuestra por sí sola una parada no autorizada.

        Si los datos llegan, el Dispatcher compara el avance con la ventana de entrega. Un desvío o demora exige confirmar la causa con el conductor. Si esa causa requiere una decisión operativa, Jefatura recibe la alerta y el caso pasa a P04; si no, se actualiza el seguimiento. Cuando no hay desvío también se continúa el control normal. Al final se pregunta si la ruta terminó: mientras siga activa, el ciclo se repite; cuando termina, el resultado se remite a P05.

        P07 supervisa la jornada mientras P03 atiende paradas. Ni un punto GPS ni un estado en tránsito sustituyen foto, firma y resultado de una entrega. El proyecto dispone de envío e historial GPS, mapa y avisos; las frecuencias, umbrales de alerta y acciones ante pérdida de señal aún necesitan definición operativa.
    """, "También debe definirse quién puede ejecutar cada función.", 145, "Procesos BPMN"),

    slide(12, "P08 Administrar accesos", "Recorrer autorización, rol mínimo, baja y prueba de permisos.", """
        P08 inicia con una solicitud de alta, cambio o baja de usuario. Jefatura precisa cargo y necesidad, y decide si la solicitud está autorizada. Si no, el Administrador comunica el rechazo y el proceso termina sin acceso. Si sí, identifica el rol mínimo necesario para el trabajo.

        La siguiente compuerta separa baja de alta o cambio. En una baja, el soporte de identidad desactiva la cuenta y sus sesiones; el evento final es acceso revocado. En un alta o cambio, crea o ajusta identidad y rol, y el usuario comprueba el acceso permitido. Si los permisos son incorrectos, Administración corrige el rol y repite la prueba. Solo después de comprobarlos se confirma la activación.

        En este proyecto Keycloak administra la identidad y los roles del token. PostgreSQL contiene además tablas auth_* para permisos locales, pero esas asignaciones no actualizan por sí solas Keycloak. Por ello, al exponer este BPMN hablamos de la política objetivo de administración y de una integración por completar, no de un flujo totalmente automatizado. El resultado correcto exige aprobación, mínimo privilegio, revocación y una prueba del acceso efectivo.
    """, "El último proceso de apoyo controla compras que pueden habilitar capacidad.", 155, "Procesos BPMN"),

    slide(13, "P09 Registrar compras operativas", "Seguir necesidad, existencia, autorización, recepción y los cuatro finales.", """
        P09 comienza cuando un área solicita compra o reposición. Primero describe necesidad y prioridad. Administración revisa si ya existe un recurso utilizable. Si existe, lo entrega y el caso termina como necesidad cubierta sin compra; esto evita duplicar gasto. Si no existe, evalúa costo y proveedor y Jefatura logística decide sobre la autorización.

        Cuando el gasto no se aprueba, el área puede ajustar el pedido y justificar otro intento. Si lo hace, vuelve a la evaluación; si no, la solicitud se cancela con motivo. Cuando el gasto se aprueba, Administración compra y comprueba la recepción. Si el bien llega observado, se reclama o devuelve y el caso queda como compra observada. Solo la recepción conforme permite registrar compra y comprobante.

        Una compra registrada que incrementa capacidad puede activar P06 para habilitar el nuevo recurso. El módulo actual registra compras, pero el control de presupuesto, la aprobación y la recepción conforme son reglas de negocio propuestas. Por eso el evento final no afirma que cualquier registro de compra equivalga a un recurso inmediatamente disponible para repartir.
    """, "Con los nueve BPMN vistos, conectemos sus eventos finales.", 140, "Procesos BPMN"),

    slide(14, "Transferencias entre procesos", "Señalar el flujo principal y después el ciclo de recuperación.", """
        El flujo principal empieza con P01: solo un pedido confirmado puede pasar a P02. P02 termina cuando la ruta inicia y habilita P03. P03 atiende cada parada y registra entrega o excepción. P05 analiza los resultados al cierre de la jornada o periodo. P07 supervisa en paralelo a P03; no espera al cierre de cada parada para observar una demora.

        Cuando P03 o P07 detectan una excepción, P04 conserva el caso y decide su salida. Un reintento aceptado regresa a P03; una nueva fecha aceptada regresa a P02; un cierre fallido autorizado queda para análisis en P05. P06 habilita recursos que P02 puede utilizar. P08 permite el acceso según rol, y P09 puede proveer un recurso que todavía requiere habilitación en P06. En los BPMN separados, la transferencia se documenta como resultado final e inicio del siguiente proceso, no como una flecha de secuencia que cruza arbitrariamente diagramas independientes.
    """, "Veamos ahora cómo se traducen estos pasos a pantallas preliminares.", 115, "Procesos BPMN"),

    slide(15, "Wireframes de admisión y confirmación", "Recorrer los seis cuadros empezando por listado, nuevo pedido e importación CSV.", """
        En los wireframes mostramos estructura de pantalla y secuencia de trabajo, no una certificación de funcionamiento. La primera lámina parte de la gestión de pedidos que ya conocemos: listado con búsqueda, número de seguimiento, cliente, ruta y estado. Al lado aparece el alta de un pedido, donde el usuario registra cliente, dirección, ventana y contacto. La importación CSV responde al caso de muchos pedidos en una jornada.

        En la segunda fila se ve el detalle del pedido y las bandejas de observación y confirmación. Aquí el punto de diseño no es añadir otro botón Guardar, sino permitir al responsable distinguir una solicitud incompleta de un pedido comercialmente confirmado. El proceso P01 exige preguntar por datos, dirección y factibilidad antes de pasar a P02. Los seis recortes mantienen el aspecto de EcoRoute web para que el prototipo sea reconocible, pero las reglas de subsanación y confirmación aún deben acordarse.
    """, "Con pedidos confirmados, pasemos al plan de rutas.", 105, "Wireframes"),

    slide(16, "Wireframes de planificación de rutas", "Indicar el panel de rutas, el mapa y la lista de pedidos asignables.", """
        Este conjunto muestra la parte de planificación de P02. El Dispatcher ve rutas de despacho, un mapa con secuencia de paradas y un detalle de pedido para comprobar dirección y estado. El listado de pedidos ayuda a decidir qué demanda entrará a cada ruta. La comparación de alternativas permite revisar distancia y orden antes de publicarla.

        El camino óptimo que muestra el mapa es un apoyo del algoritmo, no una decisión automática de despachar. El responsable aún debe verificar ventanas horarias, capacidad y disponibilidad de conductor y vehículo. Por eso, en el diseño objetivo, el mapa debe convivir con un estado de revisión y una aprobación distinguible de la simple visualización del recorrido. Si la alternativa no es viable, P02 vuelve al cálculo y al ajuste de asignación.
    """, "La misma jornada continúa con el despacho efectivo.", 105, "Wireframes"),

    slide(17, "Wireframes de despacho y hoja de ruta", "Mostrar el panel web y después las pantallas móviles de rutas del día y detalle de pedido.", """
        Aquí vemos la transición entre el Dispatcher y el conductor. En web se prepara el despacho; en móvil se muestra Mis rutas de hoy, la línea del recorrido, las paradas numeradas y un resumen de distancia y tiempo. Al abrir un pedido, el conductor puede consultar dirección y estado antes de actuar. El perfil ayuda a reconocer con qué cuenta está trabajando.

        La lámina conserva el rótulo preliminar P03 Despacho, pero en el mapa de procesos revisado todo este tramo pertenece a P02, Preparar la jornada de reparto. El evento final de P02 no ocurre cuando el mapa se dibuja ni cuando se asigna un conductor. Requiere plan aprobado, hoja de ruta emitida y confirmación de que conductor, vehículo y carga pueden salir. Esa confirmación es la frontera que inicia P03 Ejecución de entregas.
    """, "Con la ruta iniciada, la interfaz móvil acompaña la entrega.", 110, "Wireframes"),

    slide(18, "Wireframes de entrega y POD", "Recorrer lista de paradas, detalle, cambio de estado y formulario de evidencia.", """
        Esta secuencia móvil sigue una parada de P03. El conductor elige un pedido de la lista ordenada, abre el detalle, actualiza su estado y llega al formulario de confirmar entrega. Allí se prevén datos del receptor, foto y firma. Después vuelve al resumen de la ruta para continuar con la siguiente parada.

        La lámina antigua se titula P04, pero hoy corresponde a P03 Ejecutar entregas. La prueba de entrega no debe entenderse como un campo ornamental: permite sustentar quién recibió, cuándo y dónde. Si el cliente está ausente, rechaza o falta evidencia, el conductor no debería cerrar la parada como entregada. El BPMN la envía a P04 Resolver incidencias. En un futuro piloto habrá que comprobar también qué hace la app cuando pierde conexión y cómo se recupera un POD que quedó pendiente.
    """, "Las paradas no resueltas se administran como casos de incidencia.", 105, "Wireframes"),

    slide(19, "Wireframes de tratamiento de incidencias", "Señalar formulario móvil, bandeja web y estados de reasignación o reprogramación.", """
        Esta lámina empieza en el campo: el conductor selecciona una causal, redacta observación y adjunta evidencia. El Dispatcher ve luego una bandeja de alertas con casos como cliente ausente, dirección no encontrada o demora por tráfico. La intención es que cada caso pueda revisarse, documentarse y cerrar con una decisión visible.

        Aunque el archivo preliminar lleva P05, en la versión de negocio este flujo es P04 Resolver incidencias. La diferencia importa: registrar la alerta no equivale a haber resuelto el caso. Debe quedar claro si se autorizó un reintento hoy, una nueva fecha o un cierre fallido, quién lo decidió y con qué sustento. Las vistas de reasignación e historial apuntan a esa trazabilidad, pero son diseños de trabajo para contrastar con MICOTRANS, no una afirmación de que toda la gestión ya esté automatizada.
    """, "El cierre de la operación nos conduce a la lectura de indicadores.", 100, "Wireframes"),

    slide(20, "Wireframes de análisis de resultados", "Señalar filtros, tarjetas, gráfico y acceso a reportes; aclarar la numeración antigua.", """
        La última lámina preliminar agrupa vistas de panel de control y reportes. Se reconocen filtros por conductor y fecha, contadores de pedidos entregados y pendientes, y un gráfico de rendimiento. Sirven para orientar al analista hacia una pregunta: qué cambió en la operación y por qué.

        Este archivo conserva el nombre P06 Analítica y mejora, pero corresponde al proceso actual P05 Evaluar resultados. No debemos leer una tarjeta de porcentaje como conclusión automática. Primero se define el universo de pedidos y el periodo; luego se compara con una meta y se examina la causa de los desvíos. P05 termina cuando el periodo queda conforme o cuando una acción correctiva demuestra efecto. Los wireframes sirven para acordar qué información necesita esa decisión, y los dashboards posteriores muestran un diseño analítico más completo.
    """, "Hasta aquí vimos estructura propuesta; ahora enseñamos la interfaz real que inspiró el diseño.", 105, "Wireframes"),

    slide(21, "Diferencia entre captura y regla de negocio", "Pausar en las columnas Referencia visual existente y Diseño objetivo a validar.", """
        Esta distinción es importante para interpretar las siguientes capturas. EcoRoute ya tiene pantallas web de pedidos, rutas, dashboard y reportes, y una app móvil para ruta y evidencia. Las mostramos como referencia visual de alta fidelidad: permiten comprobar jerarquía, navegación y continuidad entre actores.

        Sin embargo, ver una pantalla no prueba que una política del BPMN esté cerrada. Debemos validar con MICOTRANS la autorización del despacho, cuándo un POD es suficiente, cómo se reprograma una entrega y qué calidad necesitan los indicadores. También debemos distinguir datos de prueba del sistema de mediciones reales de la empresa. Con esa cautela, las capturas nos ayudan a diseñar el siguiente ciclo sin presentar el piloto como concluido.
    """, "Veamos primero la consola de pedidos.", 100, "Mockups y analítica"),

    slide(22, "Captura de la consola de pedidos", "Apuntar al menú Pedidos, búsqueda, importación y filas con estado.", """
        Esta es una captura de EcoRoute web. El menú sitúa al operador en Pedidos; arriba aparecen búsqueda, filtro de estado, importación CSV y alta de nuevo pedido. Cada fila expone número de seguimiento, cliente, ruta asignada, estado y acciones. Es la base visual para P01 y para el seguimiento posterior del pedido.

        Lo que queremos mejorar no es solo la presentación de la tabla. Para que P01 funcione como control de negocio, la consola debe distinguir solicitud recibida, solicitud observada y pedido confirmado. También debe conservar el motivo de una no admisión o una alternativa aceptada. Los nombres y estados que se ven son datos de prueba: no equivalen a una estadística de MICOTRANS.
    """, "Una vez admitidos los pedidos, el Dispatcher trabaja en rutas.", 95, "Mockups y analítica"),

    slide(23, "Captura de la consola de rutas", "Señalar tarjeta de ruta, conductor, estado y trazado sobre el mapa.", """
        La consola de rutas reúne la jornada, el conductor, el vehículo y el recorrido. La línea sobre el mapa muestra el orden de visita propuesto sobre calles reales; eso ayuda al Dispatcher a inspeccionar distancia y paradas. La captura conecta con P02, donde el algoritmo propone una secuencia que el operador todavía debe revisar.

        En la sustentación evitamos decir que el mapa autoriza por sí solo una salida. La aprobación del plan y el acuse del conductor son decisiones separadas. Si hay falta de capacidad, una ventana imposible o indisponibilidad de vehículo, el BPMN manda corregir o suspender, aunque la polilínea sea técnicamente correcta. Esta pantalla muestra una base útil; la regla de negocio completa sigue siendo el objeto de validación.
    """, "El conductor recibe la jornada en su app móvil.", 95, "Mockups y analítica"),

    slide(24, "Captura de la ruta del conductor", "Indicar ruta de hoy, paradas numeradas y recorrido resaltado.", """
        En la app del conductor se ve la ruta del día, la secuencia de paradas y el camino resaltado. El resumen de distancia y duración orienta la jornada, y las paradas numeradas facilitan saber qué pedido sigue. La pantalla se relaciona con el final de P02, cuando el conductor recibe una hoja de ruta, y con el inicio de P03, cuando empieza a atenderla.

        El mapa es una guía, no una prueba de entrega ni una garantía de cumplimiento horario. Para usarlo como parte del control operativo debemos confirmar fecha de la ruta, estado de cada parada, posible reordenamiento y respuesta ante falta de señal. La aprobación de salida se documenta fuera de esta visualización.
    """, "Al llegar al cliente, la app abre el registro de evidencia.", 95, "Mockups y analítica"),

    slide(25, "Captura del formulario de prueba de entrega", "Mostrar campos de receptor, fotografía y firma; no describir el formulario vacío como POD validado.", """
        El formulario móvil está orientado a confirmar el resultado de la parada. Permite identificar al receptor, tomar una foto y registrar firma. La captura enseña la estructura de esa interacción, pero aún no muestra una prueba completa y aceptada. Por eso no la usamos como evidencia de una entrega real.

        El diseño de P03 exige revisar la suficiencia de esos datos antes de cerrar el pedido. Una foto sin contexto, una firma ausente o un intento sin sincronización deben conservarse como pendientes u observados y pasar a P04 cuando corresponda. También necesitamos comprobar cómo se asociarán hora y ubicación. La regla que queremos validar es sencilla: el estado entregado requiere respaldo trazable, no únicamente pulsar Guardar.
    """, "Con los resultados de campo, el supervisor consulta el panel.", 100, "Mockups y analítica"),

    slide(26, "Captura del panel de control logístico", "Señalar filtros, tarjetas de pedidos y los gráficos visibles.", """
        El panel actual ofrece una lectura rápida de pedidos totales, entregados, en ruta y pendientes, con filtros y gráficos de rendimiento y distribución. Esto sirve para que el supervisor detecte qué ocurre hoy, pero todavía no responde todas las preguntas de gestión que planteó P05.

        Para decidir mejoras se necesita comparar periodos y metas, medir demoras por ruta, analizar causas de no entrega y saber si las pruebas de entrega son completas. Por eso proponemos una capa analítica adicional, de estilo Power BI, alimentada por datos definidos y reconciliados. No presentaremos un tablero ilustrativo como si fuera una medición operativa aprobada por MICOTRANS.
    """, "El panel se complementa con consultas y exportaciones detalladas.", 100, "Mockups y analítica"),

    slide(27, "Captura de reportes web", "Recorrer selector o filtros, resultados y opción de exportación.", """
        La pantalla de reportes permite consultar resultados y exportar información de la operación. Es útil para revisar un periodo y llevar un listado fuera del sistema, pero la exportación por sí sola no crea inteligencia de negocios. Un archivo CSV puede servir para intercambio puntual; para un análisis recurrente conviene una fuente PostgreSQL de solo lectura con vistas analíticas, reglas de cálculo y control de acceso.

        Para P05, el reporte debe responder qué pasó y permitir auditar los pedidos que componen cada indicador. Si una cifra de cumplimiento cambia, necesitamos bajar al estado, la fecha y la evidencia de cada servicio. Esa trazabilidad enlaza los reportes con el diccionario de datos que veremos al final.
    """, "Con esta referencia visual, presentamos cinco diseños analíticos.", 100, "Mockups y analítica"),

    slide(28, "Dashboard de resumen ejecutivo", "Señalar filtros superiores, indicadores y los cuatro gráficos; anunciar que las cifras son ilustrativas.", """
        Esta primera propuesta concentra la lectura de Jefatura. Los filtros permiten elegir periodo, sede, conductor o zona; así el mismo indicador puede compararse sin mezclar poblaciones distintas. Las tarjetas muestran IID, CHR, TDE y TNI. IID se plantea como proporción de servicios con información y evidencia completas; CHR compara entregas completadas con servicios programados; TDE mide la proporción de entregas completadas que dispone de POD digital. TNI refleja no entregas e incidencias sobre servicios programados. El quinto indicador definido para el proyecto es OTIF, entregas a tiempo y completas.

        Los cuatro gráficos propuestos contrastan indicadores con meta, muestran evolución semanal, composición de evidencia digital y estado de servicios. Aquí las cifras visibles pertenecen a una maqueta y no deben citarse como una línea base aprobada. Antes del piloto habrá que fijar numerador, denominador, corte, tratamiento de pendientes y fuente para cada KPI; por ejemplo, TDE debe dividir entre entregas completadas, no entre todos los servicios, si se mantiene esa definición. El valor del panel es convertir la revisión de P05 en una pregunta verificable sobre causa y acción.
    """, "De la vista ejecutiva bajamos al cumplimiento diario de entregas.", 160, "Mockups y analítica"),

    slide(29, "Dashboard de cumplimiento de entregas", "Recorrer filtros de fecha, conductor, zona y estado; después los cuatro gráficos.", """
        Esta vista ayuda al Dispatcher y a Jefatura a ubicar dónde se incumple la jornada. El primer gráfico distribuye entregas por estado; el segundo compara el cumplimiento de semana en semana; el tercero permite observar diferencias por zona; el cuarto mira las ventanas horarias. Los filtros evitan atribuir a un conductor una demora de otra ruta o comparar periodos de tamaños diferentes.

        Una caída del porcentaje de cumplimiento no explica la causa. Puede responder a ausencia de cliente, dirección imprecisa, falta de capacidad o a un atraso en la salida. El tablero permite seleccionar el grupo de casos y abrir su detalle para conectarlo con P02, P03 o P04. Los valores del diseño son ilustrativos: el resultado válido debe salir de estados y tiempos conciliados, con una definición aprobada de pedido programado y entrega efectiva.
    """, "La siguiente vista analiza si el recorrido utiliza bien tiempo y distancia.", 120, "Mockups y analítica"),

    slide(30, "Dashboard de eficiencia de rutas", "Señalar filtros por fecha, vehículo, conductor y zona; explicar cuatro comparaciones.", """
        La eficiencia se observa desde la ruta y no solo desde el número de entregas. Este diseño compara kilómetros por entrega, tiempo por parada, distancia planificada frente a distancia real y rutas con demora. El filtro por vehículo y conductor permite estudiar jornadas comparables; el filtro por zona evita sacar conclusiones erróneas cuando una ruta urbana y otra periférica tienen condiciones diferentes.

        La distancia propuesta por OSRM y la geometría almacenada ayudan a reconstruir el plan. Para comparar con recorrido real se requieren posiciones GPS suficientes y depuración de puntos anómalos. Una distancia mayor puede ser una incidencia vial legítima, no necesariamente un mal desempeño del conductor. P05 debe usar esta vista para investigar causas y probar una mejora, como cambiar una ventana, reasignar carga o ajustar la secuencia. Los números de esta maqueta son solo ejemplos visuales.
    """, "La eficiencia debe revisarse junto con la calidad de la entrega.", 125, "Mockups y analítica"),

    slide(31, "Dashboard de incidencias y calidad del POD", "Marcar las cuatro áreas: causales, completitud de POD, tendencia y zonas.", """
        Este tablero une el resultado fallido con el respaldo que permite explicarlo. Agrupa incidencias por causal, compara POD completo e incompleto, muestra la tendencia de fallas y distribuye casos por zona. Sus filtros por fecha, causal, conductor y cliente permiten investigar si un problema se concentra en una ruta, un tipo de dirección o un requisito de evidencia.

        Una entrega fallida y un POD incompleto son problemas distintos. La primera exige decisión de reintento, reprogramación o cierre según P04. La segunda exige completar o aclarar la prueba antes de darla por verificada. La vista debería permitir bajar al pedido y revisar foto, firma, GPS, hora y motivo. No implica que esos datos estén siempre completos hoy: la maqueta plantea la calidad que el piloto debe medir y mejorar.
    """, "La quinta vista convierte esos hallazgos en seguimiento de responsables y acciones.", 120, "Mockups y analítica"),

    slide(32, "Dashboard de rendimiento y mejora", "Recorrer filtros, ranking, productividad, acciones abiertas y meta contra resultado.", """
        Esta vista acompaña el tramo final de P05. El ranking de cumplimiento y la productividad por conductor ofrecen señales para investigar, pero no deben usarse para sancionar sin considerar carga, zona, ventanas y calidad del dato. Otra visualización muestra acciones correctivas abiertas, sus responsables y plazos; la cuarta compara la meta con el resultado observado después del cambio.

        El propósito es cerrar el ciclo de mejora. Si una acción prometía reducir retrasos y el indicador no mejora, Jefatura debe reexaminar la hipótesis en vez de declarar éxito por haber ejecutado la tarea. Los filtros por periodo, conductor, vehículo y cliente permiten mantener comparables los casos. También aquí los valores dibujados son ilustrativos, y las metas se deben acordar con MICOTRANS antes de utilizarlas como compromisos.
    """, "Además del análisis agregado, se necesitan dos listados para el control diario.", 125, "Mockups y analítica"),

    slide(33, "Reporte de manifiesto diario de despacho", "Señalar filtros, encabezados de la tabla y estados de salida.", """
        El primer reporte operativo es el manifiesto diario de despacho. El Dispatcher lo filtra por fecha, ruta, conductor, vehículo o estado y revisa una fila por pedido o parada. Las columnas propuestas muestran ruta, secuencia, pedido, cliente, dirección, ventana, paquetes, estado y hora de salida. Así puede comprobar qué servicio está en cada vehículo antes de iniciar la jornada.

        Este listado acompaña P02: debe reflejar el plan aprobado y el acuse de salida, no solo una ruta calculada. Si falta conductor, cambia la capacidad o una ventana no es viable, el reporte permite localizar los pedidos afectados. Exportar o imprimir es una salida de consulta; la decisión de despacho y su responsable deben permanecer en el registro del proceso. Las filas de la imagen son ejemplos, no operaciones reales de MICOTRANS.
    """, "El segundo listado cierra la jornada pedido por pedido.", 115, "Mockups y analítica"),

    slide(34, "Reporte de cierre de entregas y POD", "Recorrer resultado, foto, firma, GPS, fecha y hora, causal y observación.", """
        El reporte de cierre reconcilia cada pedido con su resultado y con la evidencia disponible. Jefatura puede filtrar por fecha, conductor, estado, cliente o causal. La tabla propone pedido, cliente, resultado, presencia de foto, firma, GPS, fecha y hora, motivo y observación. Esta vista permite localizar una entrega que aparece como completada sin todos los componentes del POD o un fallo que todavía no tiene causal clara.

        También permite alimentar P05 sin perder la capacidad de auditar el detalle. Un agregado de cien entregas solo es defendible si podemos explicar las cien filas que lo forman y sus reglas de exclusión. Los ejemplos visuales muestran el formato deseado, no son un certificado legal ni un cierre mensual real de MICOTRANS.
    """, "Para sostener esos procesos y reportes, pasamos al modelo de datos.", 115, "Mockups y analítica"),

    slide(35, "Modelo lógico operativo", "Señalar drivers y vehicles, routes, orders y delivery_proofs; luego historial, GPS y productos.", """
        Este diagrama resume el núcleo operativo. Un conductor y un vehículo pueden participar en rutas; cada ruta agrupa pedidos; un pedido puede tener una prueba de entrega y varios cambios de estado. La historia GPS pertenece a vehículos y conductores, y el detalle de pedido relaciona órdenes con productos. La imagen es deliberadamente un resumen para explicar el flujo de información, mientras que el ERD físico completo de las migraciones V1 a V8 se entrega aparte.

        La geometría del recorrido se conserva en routes.geometry; la duración calculada en routes.total_duration_min; la secuencia de visita en orders.stop_sequence. Estos campos permiten comparar el plan con la ejecución. Debemos distinguir las claves foráneas declaradas de asociaciones solo lógicas: por ejemplo, orders.driver_id está presente como número, pero en el SQL no tiene una FK a drivers. Esa diferencia define hasta dónde la base garantiza integridad por sí misma.
    """, "El siguiente resumen cuantifica el esquema y señala sus límites.", 135, "Modelo y diccionario"),

    slide(36, "Esquema físico y límites referenciales", "Leer 33 tablas y 239 campos; después señalar Límites declarados.", """
        El inventario de Flyway V1 a V8 contiene treinta y tres tablas, doscientos treinta y nueve campos y veintisiete claves foráneas físicas. Hay tablas de reparto, catálogos, flota, compras, notificaciones y permisos locales. El bloque corp_* de V4 pertenece a un contexto corporativo de demostración y contiene registros ficticios de otra empresa; no constituye el padrón de clientes de MICOTRANS ni se conecta físicamente con sus pedidos.

        Otras decisiones también deben explicitarse. orders.driver_id carece de FK; corp_maintenance_logs.vehicle_id tampoco la tiene. auth_user_role.user_external_id apunta lógicamente a un usuario de Keycloak, pero Keycloak maneja otra base, así que no existe FK entre ambos sistemas. El ERD completo separa estas relaciones lógicas de las veintisiete restricciones reales. Por eso decimos que la trazabilidad exige reglas de proceso y validación de aplicación además de tablas. El esquema es una base de diseño, no una prueba de que toda la operación ya usa cada tabla.
    """, "Por último, el diccionario aclara el significado de los campos principales.", 145, "Modelo y diccionario"),

    slide(37, "Diccionario de pedidos rutas y secuencia", "Señalar los cuatro rótulos; corregir oralmente que order_stop_sequence es un concepto, no una tabla.", """
        El diccionario describe campo por campo el modelo físico. En orders están seguimiento, estado, destinatario, dirección, coordenadas, ventana de entrega y ruta asociada. En routes se guardan fecha de jornada, conductor, vehículo, estado, distancia, geometría y duración. order_status_history conserva los cambios de estado con su momento, motivo y ubicación, lo cual ayuda a reconstruir qué pasó después de planificar.

        La cuarta línea de la lámina utiliza el rótulo conceptual order_stop_sequence para explicar el orden de visita. Aclaro que en el SQL no existe una tabla con ese nombre: el campo real es orders.stop_sequence, añadido en V8. Este orden planeado debe poder contrastarse con la ejecución y con los resultados de cada parada. En el informe, cada campo lleva tipo, nulabilidad, clave y valor por defecto cuando corresponde.
    """, "La prueba de entrega y los eventos de campo completan esa trazabilidad.", 125, "Modelo y diccionario"),

    slide(38, "Diccionario de POD GPS incidencias y avisos", "Distinguir evidencia de entrega, telemetría y excepción.", """
        delivery_proofs representa la prueba de entrega. Relaciona un pedido con imagen, firma, receptor, momento de verificación y coordenadas. vehicle_gps_history conserva posiciones sucesivas, velocidad, rumbo y hora de cada reporte. incidents registra tipo, severidad, descripción y tiempos de reporte y resolución para una ruta. notifications guarda avisos del sistema y si se han leído.

        Estas entidades cumplen funciones distintas. GPS ayuda a supervisar P07; POD sustenta el resultado de P03; una incidencia explica por qué P04 tuvo que intervenir. La presencia de un campo no garantiza que todas las filas lo tengan completo: el diccionario diferencia columnas obligatorias de opcionales y el BPMN define qué evidencia exige el negocio para aceptar una entrega. La calidad del dato que llegue a P05 depende de ambas capas.
    """, "El último grupo cubre permisos, recursos y módulos de apoyo.", 125, "Modelo y diccionario"),

    slide(39, "Diccionario de acceso y catálogos", "Señalar auth, drivers y vehicles, purchases y corp; explicar la separación con Keycloak.", """
        Las tablas auth_* organizan módulos, recursos, acciones, roles, permisos y asignaciones locales. Sin embargo, el backend autoriza las rutas web y API a partir de roles contenidos en el JWT emitido por Keycloak. Debemos integrar o sincronizar ambos modelos antes de afirmar que cambiar un rol en la tabla local modifica el permiso efectivo del usuario. Keycloak conserva la identidad y credenciales en su propia base, fuera del script PostgreSQL de EcoRoute.

        drivers y vehicles aportan datos de conductor, vehículo y capacidad para P06 y P02. purchases conserva descripción, proveedor, importe, categoría y estado del gasto. corp_* es un bloque corporativo de referencia, con datos de prueba que requieren depuración y autorización antes de asociarlos con MICOTRANS. La lectura del diccionario nos deja una conclusión concreta: una entidad puede existir en la base, pero la regla de negocio y la integración que la convierten en información confiable aún deben comprobarse.
    """, "Cerremos con las tres validaciones que abren el piloto.", 140, "Modelo y diccionario"),

    slide(40, "Cierre del capítulo 2", "Leer las tres decisiones una a una y mirar al profesor al formular el cierre.", """
        El capítulo 2 deja tres tareas de validación con MICOTRANS. La primera es aprobar reglas de admisión del pedido, despacho y suficiencia de la prueba de entrega. La segunda es ensayar las excepciones: datos incompletos, falta de recurso, ausencia, rechazo, POD pendiente, pérdida de señal y reprogramación, comprobando quién decide y a qué proceso vuelve cada caso. La tercera es conciliar datos e indicadores, de modo que un tablero explique decisiones a partir de registros auditables.

        En resumen, los BPMN definen resultados y responsabilidades, los prototipos muestran cómo trabajarían los actores, y el modelo de datos conserva el rastro de cada pedido. Tenemos una aplicación de referencia y un diseño objetivo; el piloto debe demostrar que ambos responden juntos a las necesidades reales de la empresa. Muchas gracias. Quedamos atentos a sus preguntas.
    """, "Fin de la exposición principal.", 105, "Cierre"),
]


QUESTIONS = [
    ("Por qué hay seis macroprocesos y nueve BPMN",
     "Los macroprocesos agrupan responsabilidades de gestión. Los nueve BPMN representan casos concretos con inicio, decisiones y resultado. Los trece módulos y 81 RF se trazan hacia ellos, pero no se convierte cada pantalla en proceso."),
    ("Por qué planificación y despacho están en P02",
     "La unidad de gestión es la preparación de una jornada lista para ejecutarse. Dentro de ella distinguimos calcular la ruta, revisar su factibilidad, aprobar el plan y obtener el acuse de salida. Esas decisiones son separables y medibles, aunque el caso termina solo cuando la ruta inicia o se suspende."),
    ("Cómo se modelan los caminos No y cuándo termina un proceso",
     "Cada compuerta exclusiva tiene salida Sí y No. El No puede volver a una tarea para corregir, abrir P04, o cerrar el caso con un motivo explícito. El evento final representa el resultado del caso; la transferencia a otro BPMN se explica como estado de salida y entrada, no como una flecha de secuencia entre diagramas independientes."),
    ("Por qué hay seis wireframes si se modelaron nueve procesos",
     "Los seis archivos se dibujaron antes de la revisión de procesos. Sus etiquetas P01 a P06 no son los códigos actuales: el antiguo P02 y P03 se integran en el nuevo P02, el antiguo P04 pasa al nuevo P03, el antiguo P05 al nuevo P04 y el antiguo P06 al nuevo P05. Los procesos P06 a P09 actuales todavía no tienen una lámina propia de interfaz."),
    ("Qué funciona hoy y qué sigue siendo diseño objetivo",
     "Las capturas documentan interfaces y datos de prueba del proyecto. Los BPMN añaden controles de operación como aprobación formal de salida, suficiencia de POD, recuperación de incidencias y metas de KPI. Esos controles requieren acuerdo y pruebas con MICOTRANS antes de presentarlos como operación validada."),
    ("Qué sucede cuando el POD no se sincroniza",
     "En el proceso objetivo, la evidencia queda pendiente u observada y se deriva a P04. No se declara una entrega verificada hasta confirmar que la prueba existe y es suficiente. El soporte offline y su sincronización automática deben ensayarse en el piloto; una pantalla de captura no demuestra esa garantía."),
    ("Qué diferencia hay entre panel operativo y módulo de inteligencia de negocios",
     "El panel actual ayuda a ver estados y actividad. El módulo analítico propuesto compara periodos y metas, analiza causas, permite filtrar y profundizar hasta pedidos concretos y sigue acciones correctivas. Cinco dashboards cumplen usos diferentes; dos listados permiten control y conciliación diaria."),
    ("Cómo se calcularían IID CHR TDE OTIF y TNI",
     "IID: servicios con datos y evidencia digital completos entre servicios totales. CHR: entregas completadas entre servicios programados. TDE: entregas con POD digital entre entregas completadas. OTIF: entregas dentro de ventana y completas entre programadas. TNI: servicios fallidos o con incidencia entre programados. Cada fórmula requiere definición exacta de estados, fecha de corte y exclusiones. Las cifras de los dashboards son ilustrativas; no citar TDE 51.3 % como línea base validada cuando el denominador definido son las entregas completadas."),
    ("Por qué no basta exportar CSV para Power BI",
     "El CSV sirve para un intercambio puntual, pero su repetición manual crea riesgo de versiones y filtros distintos. Para análisis recurrente se propone una conexión de solo lectura a vistas analíticas de PostgreSQL, con permisos, definición de KPI y actualización controlada. La arquitectura final se validará con TI y con la política de datos de MICOTRANS."),
    ("Qué representa el ERD completo y qué relaciones faltan",
     "Las migraciones V1 a V8 definen 33 tablas, 239 campos y 27 claves foráneas. orders.driver_id y corp_maintenance_logs.vehicle_id no tienen FK. auth_user_role.user_external_id apunta lógicamente a Keycloak, que usa otra base. El esquema corp_* contiene datos ficticios heredados y no está integrado físicamente al núcleo de pedidos de MICOTRANS."),
    ("Keycloak y las tablas auth son la misma base",
     "No. Keycloak mantiene su propia identidad, credenciales, clientes y roles. El backend valida el JWT y extrae roles de realm_access. PostgreSQL de EcoRoute incluye tablas auth_* locales; el formulario local de asignación no cambia por sí solo un rol efectivo en Keycloak. Esa integración se debe resolver antes del piloto."),
    ("Cuál es el siguiente paso después de este diseño",
     "Validar con MICOTRANS responsables y reglas de los BPMN, probar excepciones de extremo a extremo y aprobar un diccionario de KPI con fuente, fórmula, corte y meta. Solo después corresponde medir el piloto y decidir mejoras sobre resultados comprobados."),
]


def speaker_seconds(item: dict) -> int:
    # Ritmo hablado conservador: 125 palabras/minuto + pausa para señalar la lámina.
    return max(45, int(round((len(item["speech"].split()) / 125 * 60 + 12) / 5) * 5))


def minutes_label(seconds: int) -> str:
    return f"{seconds // 60}:{seconds % 60:02d} min"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    tc_pr.append(borders)


def set_cell_margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, amount in (("top", "100"), ("bottom", "100"), ("left", "125"), ("right", "125")):
        child = OxmlElement(f"w:{edge}")
        child.set(qn("w:w"), amount)
        child.set(qn("w:type"), "dxa")
        margins.append(child)
    tc_pr.append(margins)


def style_table(table) -> None:
    table.autofit = False
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_i == 0:
                set_cell_shading(cell, "173D5C")
            elif row_i % 2 == 0:
                set_cell_shading(cell, "F1F6FA")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_i == 0 else RGBColor(24, 42, 57)
                    if row_i == 0:
                        run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def add_page_number(footer) -> None:
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.style = "Footer"
    run = p.add_run("EcoRoute  ·  Capítulo 2  ·  Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(82, 101, 118)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def add_label_paragraph(doc: Document, label: str, value: str, *, italic=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    lead = p.add_run(label)
    lead.bold = True
    lead.font.size = Pt(10)
    rest = p.add_run(value)
    rest.font.size = Pt(10)
    rest.italic = italic


def make_doc() -> None:
    assert len(SLIDES) in (40, 48, 56) and [s["n"] for s in SLIDES] == list(range(1, len(SLIDES) + 1))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.85)
    section.top_margin, section.bottom_margin = Inches(0.82), Inches(0.72)
    section.footer_distance = Inches(0.38)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13
    for style_name, size, space_before, space_after in (("Title", 24, 0, 10), ("Heading 1", 16, 20, 9),
                                                       ("Heading 2", 12.5, 14, 5), ("Heading 3", 10.5, 9, 3)):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    title.add_run("Guion detallado de exposición del capítulo 2")
    # Word's built-in Title style may inherit a colored bottom rule. Explicitly
    # suppress it so the title relies on type and whitespace, not a border.
    title_ppr = title._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        rule = OxmlElement(f"w:{edge}")
        rule.set(qn("w:val"), "nil")
        border.append(rule)
    title_ppr.append(border)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("EcoRoute para Grupo MICOTRANS S. A. C.  |  Integrador I  |  Grupo 3")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(68, 82, 94)

    doc.add_paragraph(
        f"Este guion acompaña las {len(SLIDES)} láminas de la presentación del capítulo de Diseño. "
        "Está escrito para hablarle al profesor en nombre del Grupo 3 y seguir cada imagen de principio a fin: "
        "los nueve procesos BPMN, las pantallas de referencia, cinco dashboards, dos reportes, el modelo de datos "
        "y el diccionario. La tesis que debe quedar clara es que el proceso define la decisión, la interfaz la acompaña "
        "y los datos permiten verificarla."
    )
    doc.add_paragraph(
        "El texto de Qué decir se puede pronunciar casi literalmente. Las indicaciones de pantalla y las transiciones "
        "no se leen como contenido; sirven para mover el puntero y cambiar de lámina. La duración completa estimada "
        f"es de unos {round(sum(speaker_seconds(s) for s in SLIDES)/60)} minutos a ritmo moderado. Si el docente fija menos tiempo, se puede abreviar la descripción "
        "de pantallas, pero conviene conservar las ramas No y los eventos finales de cada BPMN."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run("Criterio de rigor  ")
    r.bold = True
    p.add_run("Diferenciar siempre el proceso objetivo de las funciones ya visibles en EcoRoute. "
              "Los datos de maquetas son ilustrativos y las políticas operativas deben validarse con MICOTRANS.")

    doc.add_heading("Ruta de exposición", 1)
    is_expanded = len(SLIDES) >= 48
    has_business_erd = len(SLIDES) == 56
    blocks = [
        ("Apertura", "1 a 4", "Alcance y convención PMBOK BPMN"),
        ("Procesos BPMN", "5 a 14", "Nueve procesos con decisiones y transferencia"),
        ("Wireframes", "15 a 20", "Estructura de pantallas y mapeo de códigos"),
        ("Mockups y analítica", "21 a 42" if is_expanded else "21 a 34", "Mockups P01–P09, cinco dashboards y dos reportes" if is_expanded else "Capturas, cinco dashboards y dos reportes"),
        ("Modelo y diccionario", "43 a 55" if has_business_erd else "43 a 47" if is_expanded else "35 a 39", "ERD del negocio, 17 FK físicas y diccionario" if has_business_erd else "ERD, límites físicos y campos críticos"),
        ("Cierre", "56" if has_business_erd else "48" if is_expanded else "40", "Validaciones necesarias para el piloto"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.columns[0].width, table.columns[1].width = Inches(1.28), Inches(0.78)
    table.columns[2].width, table.columns[3].width = Inches(3.71), Inches(0.95)
    for cell, text in zip(table.rows[0].cells, ["Bloque", "Láminas", "Enfoque", "Tiempo"]):
        cell.text = text
    for block, indices, focus in blocks:
        row = table.add_row().cells
        subtotal = sum(speaker_seconds(s) for s in SLIDES if s["section"] == block)
        for cell, text in zip(row, [block, indices, focus, f"{round(subtotal/60)} min"]):
            cell.text = text
    # The grid widths alone do not change each cell's preferred width. Set
    # both so Word does not expand this summary table beyond the page margin.
    widths = [Inches(1.28), Inches(0.78), Inches(3.71), Inches(0.95)]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    style_table(table)
    doc.add_paragraph()

    current_section = None
    for item in SLIDES:
        if item["section"] != current_section:
            current_section = item["section"]
            doc.add_heading(current_section, 1)
        heading = doc.add_heading(f'Lámina {item["n"]:02d}  {item["title"]}', 2)
        heading.paragraph_format.keep_with_next = True
        info = doc.add_paragraph()
        info.paragraph_format.space_after = Pt(6)
        run = info.add_run(f'Tiempo orientativo  {minutes_label(speaker_seconds(item))}')
        run.font.size = Pt(9)
        run.italic = True
        run.font.color.rgb = RGBColor(84, 103, 116)
        add_label_paragraph(doc, "En pantalla  ", item["cue"])
        label = doc.add_paragraph()
        label.paragraph_format.keep_with_next = True
        label.paragraph_format.space_after = Pt(3)
        label_run = label.add_run("Qué decir")
        label_run.bold = True
        label_run.font.size = Pt(10)
        for paragraph in item["speech"].split("\n\n"):
            p = doc.add_paragraph(paragraph.replace("\n", " "))
            p.paragraph_format.widow_control = True
        add_label_paragraph(doc, "Transición  ", item["bridge"], italic=True)

    doc.add_heading("Respuestas de apoyo para preguntas del profesor", 1)
    doc.add_paragraph("Estas respuestas no forman parte de la lectura principal. Sirven para sostener decisiones metodológicas, límites del prototipo y definiciones de datos si el profesor las pregunta.")
    for question, answer in QUESTIONS:
        doc.add_heading(question, 2)
        doc.add_paragraph(answer)

    doc.add_heading("Fuentes del guion", 1)
    for item in [
        "docs/Integrador I_Grupo3.docx, capítulo 2, apartados 2.1 a 2.5.",
        f"docs/presentacion_capitulo2/final/{DECK}, {len(SLIDES)} láminas y notas del expositor.",
        "docs/entregables_procesos_ecoroute/negocio_revisado/bpmn, diagramas P01 a P09.",
        "docs/entregables_procesos_ecoroute/imagenes, wireframes, dashboards D01 a D05 y reportes R01 a R02.",
        "docs/screenshots, capturas web y móvil de EcoRoute con datos de prueba.",
        "src/main/resources/db/migration, migraciones V1 a V8; docs/base_datos/modelo_erd_negocio, ERD del negocio y fuente Mermaid editable." if has_business_erd else "src/main/resources/db/migration, migraciones V1 a V8; docs/base_datos/modelo_erd, ERD completo.",
        "Project Management Institute, Process Groups A Practice Guide; Object Management Group, BPMN 2.0.2.",
    ]:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

    add_page_number(section.footer)
    doc.core_properties.title = "Guion detallado de exposición del capítulo 2 de EcoRoute"
    doc.core_properties.subject = "Sustentación de diseño para Grupo MICOTRANS"
    doc.core_properties.author = "Grupo 3"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"{OUT}\n{len(SLIDES)} láminas; {len(QUESTIONS)} preguntas; {sum(speaker_seconds(s) for s in SLIDES)/60:.1f} minutos estimados")


if __name__ == "__main__":
    make_doc()
