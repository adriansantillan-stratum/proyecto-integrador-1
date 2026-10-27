"""Align the Chapter 2 speech with the full business ERD and slide 52."""

from pathlib import Path
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Guion_Detallado_Exposicion_Capitulo_2_EcoRoute_ERD_Negocio.docx"
DEST = ROOT / "docs" / "Guion_Detallado_Exposicion_Capitulo_2_EcoRoute_Alineado.docx"


def main() -> None:
    doc = Document(SOURCE)
    p = doc.paragraphs
    p[315].text = "Lámina 43  Núcleo conectado del reparto: 16 tablas"
    p[319].text = (
        "Profesor, entramos al modelo de datos. Esta lámina resume el núcleo conectado del reparto: "
        "dieciséis tablas, ciento cuarenta y una columnas y diecisiete claves foráneas físicas. "
        "La vista integral tipo DBeaver que entregamos como anexo incluye además hubs y purchases, "
        "que existen en PostgreSQL pero no tienen FK con este núcleo. Por eso el dominio operativo "
        "completo suma dieciocho tablas y ciento cincuenta y nueve columnas. Una FK física significa "
        "que PostgreSQL puede comprobar la referencia al guardar la fila; no la deducimos por nombres parecidos."
    )
    p[321].text = (
        "Las tarjetas inferiores muestran qué información soporta cada fase de los BPMN: planificar, "
        "ejecutar, supervisar y analizar. Algunas tablas se repiten en las láminas siguientes para "
        "leer sus columnas a tamaño de proyección; no son copias físicas. El conteo de dieciséis "
        "pertenece al grafo conectado de esta lámina, mientras el anexo físico incluye dieciocho."
    )
    p[389].text = (
        "En pantalla  Leer las dos columnas y detenerse en orders.driver_id. Aclarar que la lámina 43 "
        "resume 16 tablas conectadas y que el ERD integral incluye 18 tablas operativas."
    )
    p[391].text = (
        "El alcance final del ERD integral de negocio es de dieciocho tablas, ciento cincuenta y "
        "nueve columnas y diecisiete FK declaradas. Dieciséis tablas forman el grafo de reparto; "
        "hubs y purchases son entidades aisladas. Una FK protege que el identificador referenciado "
        "exista, pero no garantiza que el pedido esté aprobado, que la foto sea suficiente o que "
        "el vehículo esté disponible. Esas decisiones pertenecen a los BPMN, a reglas de aplicación "
        "y a pruebas del piloto con MICOTRANS."
    )
    p[392].text = (
        "Orders.driver_id sigue siendo BIGINT sin FK a drivers; antes de agregarla debemos definir "
        "si el conductor se asigna únicamente mediante la ruta. Hubs y purchases aparecen en la "
        "vista integral sin enlaces inventados. Las tablas auth_* son RBAC técnico, Keycloak guarda "
        "identidad en otra base y corp_* es referencia heredada. Esos objetos se documentan en el "
        "diccionario, pero no se presentan como relaciones operativas ya integradas."
    )
    p[393].text = (
        "El SVG 09_vista_completa_tipo_dbeaver.svg deja ampliar las dieciocho cajas y seguir cada "
        "línea desde la fila FK hasta la PK. Los ocho diagramas de detalle permiten explicar "
        "relaciones concretas en pantalla. Así mostramos el SQL verdadero y distinguimos qué "
        "integridad garantiza la base de lo que MICOTRANS todavía debe aprobar."
    )
    p[427].text = (
        "La lámina 43 presenta las 16 tablas conectadas del reparto, con 141 columnas y 17 FK. "
        "La vista física integral adjunta incorpora hubs y purchases como tablas aisladas: 18 tablas "
        "operativas y 159 columnas, sin alterar el número de FK. El esquema PostgreSQL completo "
        "tiene 33 tablas, 239 columnas y 27 FK, excluyendo la tabla de control de Flyway. "
        "corp_* y auth_* no son entidades directas de entrega; Keycloak usa su propia base."
    )
    anchor = p[428]._p
    extra = doc.add_paragraph()
    extra.style = doc.styles['Normal']
    extra.add_run(
        "Para leer el ERD ampliado, empiezo en drivers y vehicles, sigo las dos FK de routes "
        "y luego la FK orders.route_id. Desde orders se derivan delivery_proofs, historial de estados, "
        "notificaciones y líneas de producto. Después identifico GPS, incidentes y gastos de ruta. "
        "Termino con hubs y purchases, sin conectores porque el DDL no los relaciona con reparto. "
        "Esto explica visualmente el paso del pedido confirmado a la jornada, la entrega y la evaluación."
    )
    anchor.addprevious(extra._p)
    doc.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
