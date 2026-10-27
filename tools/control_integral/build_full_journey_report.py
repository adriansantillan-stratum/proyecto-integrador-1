"""Publish the actually executed journey, not a retrospective test simulation."""
import json, html, hashlib, sys
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/control_integral_20261022/e2e_cierre'
def read(name):return json.loads((OUT/name).read_text(encoding='utf-8-sig'))
fixture=read('fixture.json'); nav=read('native_session_navigation.json'); quality=read('calidad_mobile_final.json'); diagnosis=read('session_diagnosis.json'); apk=read('apk_4004.json')
attempts=[(p,json.loads(p.read_text(encoding='utf-8'))) for p in OUT.glob('server_final*.json')]
final_path,final=max(((p,r) for p,r in attempts if r.get('result')=='PASS'),key=lambda pr:pr[1]['startedAt'])
assert final['summary']=={'passed':26,'failed':0} and nav['result']=='PASS' and quality['status']=='PASS'
assert final['fixture']['routeId']==fixture['routeId'] and apk['versionCode']==8004
assert final['incidents'][0]['status']=='RESOLVED' and final['incidents'][0]['outcome']=='CLOSE_FAILED'
web=[]
for stage in ['prepare','dispatch','resolve','final']:
    reports=[(p,read(p.name)) for p in OUT.glob('web_'+stage+'_*.json')]
    p,r=max(((p,r) for p,r in reports if r.get('result')=='PASS'),key=lambda pr:pr[1]['at'])
    web.append({'stage':stage,'file':p.name,'checks':len(r['checks']),'result':r['result']})
def native(name):return '\n'.join(n['text']+'\n'+n['description'] for n in read('mobile/'+name+'.json')['nodes'])
assert 'Pedido #44' in native('26_detalle_cola') and 'Pedido #44' in native('30_cola_actualizacion')
assert '1' in native('27_cola_tras_reinicio') and 'Evidencia pendiente de confirmación' in native('27_cola_tras_reinicio')
assert 'Sin evidencias pendientes' in native('31_sincronizado')
assert 'COMPLETED' in native('39_cierre_jornada') and '1 entregados · 1 fallidos' in native('39_cierre_jornada')
pod=next(s['evidence'] for s in final['steps'] if s['name']=='POD 44')
route=final['sqlAfter']['route']; proof=final['sqlAfter']['proofs'][0]
steps=[
('01','Jefatura','Consultar dashboard e indicadores con fecha de hoy y conductor 1','01_dashboard_antes.png'),
('02','Dispatcher','Crear pedidos 44 y 45 sin ruta; ambos quedan PENDING','04_pedido_pendiente_success.png'),
('03','Dispatcher','Declarar aceptación comercial, dirección, cobertura y ventana','03_formulario_success.png'),
('04','Dispatcher','Crear ruta 20 de hoy y asignar los pedidos conservando sus IDs','06_asignado_success.png'),
('05','Jefatura','Aprobar la salida; el despacho sigue bloqueado sin aceptación móvil','07_aprobacion_web_espera_conductor.png'),
('06','Conductor','Revisar y aceptar la hoja desde la app','mobile/07_hoja_aceptada.png'),
('07','Dispatcher','Iniciar despacho web; ruta IN_PROGRESS y pedidos IN_TRANSIT','10_despacho_y_recorrido.png'),
('08','Conductor','Tomar foto virtual, confirmar firma y receptor; capturar GPS simulado','mobile/22_firma_confirmada.png'),
('09','Conductor','Guardar offline; conservar el POD al reabrir y actualizar la APK','mobile/30_cola_actualizacion.png'),
('10','Conductor','Reconectar y reintentar; recibir confirmación y vaciar la cola','mobile/31_sincronizado.png'),
('11','Conductor','Registrar FAILED por cliente ausente en pedido 45; cerrar ruta','mobile/39_cierre_jornada.png'),
('12','Jefatura','Ver POD y PDF en web; resolver la incidencia como CLOSE_FAILED','26_excepcion_resuelta_ready.png'),
('13','Jefatura','Consultar dashboard, cinco indicadores y CSV; contrastar SQL','24_indicadores_final_ready.png'),
]
summary={'at':datetime.now(timezone.utc).isoformat(),'status':'LOCAL_SYNTHETIC_JOURNEY_VERIFIED','scope':'Dashboard → web operations → native Android QA → web evidence and BI',
 'fixture':fixture,'latestFinalServer':final_path.name,'serverFinalChecks':final['summary'],'webStages':web,'nativeNavigation':nav,'mobileQuality':quality['summary'],
 'kpis':final['bi']['kpis'],'incident':final['incidents'][0],'apk':{'buildNumber':4004,'versionCode':8004,'sha256':apk['sha256']},'businessAcceptance':'NOT_ATTESTED','physicalFieldTest':'NOT_EXECUTED',
 'limits':['One successful offline delivery and one failed delivery with administrative closure, not exhaustive product coverage.','Baseline and assigned states were observed through UI and HTTP responses; no retrospective SQL snapshots were generated.','A native exit during the first signature attempt was not conclusively attributed; retry at 3 GB completed.','The session defect was reproduced and fixed; historical login incident cannot be conclusively assigned to one cause.','Power BI Desktop and real field conditions remain outside this executed local journey.'],
 'visualEvidencePolicy':'Final gallery uses _ready captures after exact UI population and loaded state. Original transient loading/animation screenshots remain archived, not final-state evidence.',
 'steps':[{'step':n,'actor':a,'action':t,'evidence':e} for n,a,t,e in steps]}
(OUT/'cierre_recorrido.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')

doc=Document(); sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(23);doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 2'].font.size=Pt(12)
def p(text):return doc.add_paragraph(text)
def h(text):doc.add_heading(text,1)
def page(title):doc.add_page_break();h(title)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
            pr.append(b)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');margins.append(x)
            pr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3');pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
                if j==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(10);run.bold=i==0
    p('')
    return t

doc.add_paragraph('Validación del recorrido completo de EcoRoute',style='Title')
p('Dashboard web y aplicación móvil en entorno local')
p('22 de octubre de 2026 · Grupo 3 · Proyecto Integrador de Desarrollo de Software')
h('Resultado de la ejecución')
p('Verificamos el recorrido de una ruta nueva de hoy con dos pedidos sintéticos. El pedido 44 fue entregado mediante evidencia capturada en la app y sincronizada después de trabajar sin conexión. El pedido 45 terminó fallido por cliente ausente; la incidencia 12 se resolvió desde la web como CLOSE_FAILED. La ruta 20 quedó COMPLETED.')
p('El contraste final aprobó 26 controles API y SQL. La web mostró la foto, la firma y el comprobante PDF del pedido entregado, además del estado fallido del segundo pedido. Los cinco indicadores coincidieron con PostgreSQL para la misma fecha y conductor.')
table(['Elemento','Identificación del ensayo'],[['Fecha operativa','2026-10-22 · America/Lima'],['Ruta y recursos','Ruta 20 · conductor 1 · vehículo 1'],['Pedido entregado','44 · '+fixture['orders'][0]['trackingNumber']],['Pedido fallido','45 · '+fixture['orders'][1]['trackingNumber']],['Incidencia','12 · RESOLVED · CLOSE_FAILED'],['Aplicación instalada','Build Flutter 4004 · Android x86_64 versionCode 8004']],[1.5,5.5])
p('Esta validación técnica usa la web real, Keycloak local, la app instalada en el emulador y los servicios Docker existentes. No representa un servicio de reparto real, una prueba física en campo ni la aceptación empresarial de MICOTRANS.')

page('Secuencia de negocio ejecutada')
p('Las responsabilidades se mantienen en los actores del negocio. EcoRoute automatiza y registra sus decisiones; no sustituye al cliente, conductor, dispatcher ni jefatura como actor del BPMN.')
table(['Paso','Actor','Acción y resultado observado'],[(n,a,t) for n,a,t,e in steps],[.5,1.15,5.35])
p('Los pasos de pedido corresponden a P01; planificación, aceptación y despacho a P02; entrega y evidencia a P03; la no entrega se deriva a P04 para decidir su cierre. El panel consume los resultados persistidos. No se reabrió el pedido fallido ni se creó un intento adicional.')
p('La cuenta ADMIN local ejecutó la preparación, aprobación y despacho web; el conductor usó su cuenta DRIVER. Las responsabilidades de la tabla son roles del negocio. Este ensayo no acredita por sí solo separación entre personas ni los permisos exclusivos de una cuenta DISPATCHER.')

page('Entrega offline y cadena de evidencia')
p('En el pedido 44 seleccionamos DELIVERED desde la app. Se tomó una foto de la escena virtual del emulador, se trazó y confirmó una firma de ensayo y se completó el receptor ficticio. El dispositivo obtuvo una ubicación simulada reciente en Lima. La firma utilizada no pertenece a un cliente ni constituye una firma digital certificada.')
p('Con WiFi y datos móviles desactivados solo en el emulador, Guardar registró la evidencia en la cola de esa cuenta. La UI indicó pendiente de confirmación y mantuvo el pedido IN_TRANSIT. En ese momento el contraste API y SQL aprobó 26 controles y confirmó que no existía POD en el servidor.')
p('La cola conservó el pedido 44 después de cerrar y reabrir la app. También sobrevivió a la actualización de la APK sin borrar datos. Al restaurar la red y solicitar el reintento, el servidor confirmó la entrega y la cola quedó sin pendientes.')
table(['Dato','Resultado contrastado'],[['POD','ID 10 · un único registro para el pedido 44'],['Coordenadas',f"{pod['coordinates']['latitude']}, {pod['coordinates']['longitude']}"],['Captura',proof['captured_at']],['Recepción en servidor',proof['received_at']],['Verificación registrada',proof['verified_at']],['Foto',str(pod['media']['photo']['bytes'])+' bytes · JPEG'],['Firma',str(pod['media']['signature']['bytes'])+' bytes · PNG']],[2,5])
p('Los archivos privados recuperados por la API coinciden con sus hashes SHA-256 persistidos. La web descargó un PDF nuevo con foto y firma. Las fechas de captura, recepción y verificación se conservan como datos distintos; la demora del envío offline no se presenta como una entrega confirmada antes de recibirla.')
p('La cola reintenta con la app abierta, al reabrir y al reanudar. Este ensayo no acredita un servicio nativo que siga ejecutándose cuando Android mantiene la aplicación cerrada.')

page('Excepción e indicadores finales')
p('En el pedido 45 seleccionamos FAILED y registramos el motivo “SINTETICO cliente ausente E2E cierre tecnico”. El servidor conservó al conductor como actor del historial y creó la incidencia 12 vinculada al pedido y la ruta. No generó foto, firma ni un POD de entrega para el caso fallido.')
p('Al terminar ambas paradas, la ruta cerró automáticamente. Desde la web, la jefatura resolvió la incidencia con CLOSE_FAILED y una justificación identificada como ensayo. El pedido original permaneció FAILED y la incidencia quedó RESOLVED, sin duplicar pedidos ni reescribir el resultado de la entrega.')
table(['Indicador','Numerador','Denominador','Resultado'],[[k['code'],k['numerator'],k['denominator'],str(k['percentage'])+' %'] for k in final['bi']['kpis']],[2.5,1.3,1.4,1.8])
p('La población es conductor 1, del 22 al 22 de octubre de 2026, con fechas inclusivas y zona America/Lima. IID quedó 100 %; CHR, TDE y entregas exitosas quedaron 50 % por una entrega de dos pedidos. Puntualidad quedó 100 % sobre una sola entrega evaluable, no sobre los dos pedidos.')
p('No se configuraron metas durante el ensayo; los indicadores informaron NO_TARGET. Estos valores describen el ensayo sintético y no demuestran una mejora empresarial ni un efecto causal en MICOTRANS. Se descargó el CSV de indicadores y tendencias desde la interfaz.')

page('Corrección y regresión de sesión')
p('La investigación reprodujo un defecto concreto en AuthInterceptor: una solicitud recibía 401, la renovación de la sesión concluía correctamente y la solicitud reintentada devolvía 400. El código trataba ese error de validación del recurso como un rechazo de la renovación y ejecutaba onExpired, lo que redirigía al login.')
p('La corrección separa el error de la solicitud reintentada del error al renovar. Un 400, 403, 409, 422 o 500 del recurso se propaga sin cerrar por ello la sesión renovada. Un nuevo 401 mantiene el tratamiento de sesión no utilizable. No se relajaron la autorización por conductor ni las validaciones del servidor.')
table(['Ejecución','Resultado'],[['Casos originales antes del cambio','23 aprobados · 6 fallidos'],['Regresión ampliada después del cambio','34 aprobados · 0 fallidos'],['Suite Flutter completa final','81 aprobados · 8 opt-in omitidos en host'],['Contratos reales separados','8 aprobados · 0 omitidos'],['Análisis estático final','Sin problemas'],['Navegación nativa en la APK nueva','3 ciclos aprobados de Sync → volver → Historial → volver']],[3.4,3.6])
p('También se reprodujo, con transporte simulado, que una coordenada antigua de Historial podía tocar Cerrar sesión durante la animación de regreso. La automatización ahora vuelve a obtener la pantalla y rechaza una jerarquía inválida antes de tocar. Esto explica un mecanismo posible, pero no identifica con certeza la causa única del retorno histórico al login.')
p('La primera captura de firma terminó en el launcher del emulador, no en el login. La causa nativa no quedó demostrada. Con el AVD ampliado de 2 a 3 GB y la pantalla verificada, la captura posterior se completó. La nueva APK conserva el mismo paquete y datos; sus fuentes de lib y Android fueron comparadas con la copia usada para compilar fuera de OneDrive.')

page('Evidencias y alcance del cierre')
p('El catálogo index.html de la carpeta docs/control_integral_20261022/e2e_cierre reúne los pasos, capturas ampliables y resultados JSON. cierre_recorrido.json identifica el ensayo, el informe final de servidor, los indicadores y la APK. Cada archivo incluido en el manifiesto conserva su hash.')
table(['Evidencia','Qué acredita'],[['web_prepare y capturas 03 a 07','Alta PENDING, confirmación, asignación y aprobación desde la interfaz'],['web_dispatch y captura móvil 07','Aceptación del conductor antes del despacho'],['server_offline y capturas móviles 26 a 30','Sin POD en servidor mientras la evidencia permanece en la cola'],['web_final y comprobante_nuevo_ensayo.pdf','Retorno a la web, archivos recuperables y PDF'],['server_final más reciente','26 controles de estados, historial, incidencia, archivos e indicadores'],['web_resolve y captura 26','Cierre administrativo de la excepción sin reabrir el pedido'],['session_diagnosis y calidad_mobile_final','Defecto reproducido, corrección y regresión'],['native_session_navigation','Tres ciclos de navegación en el emulador con la APK instalada']],[3,4])
p('Los intentos fallidos de automatización y los errores de entorno permanecen guardados. La creación y asignación inicial se observaron en la UI y en respuestas HTTP; no se inventaron snapshots SQL anteriores. El intento de verificación bloqueado por permisos se conserva junto al reintento autorizado aprobado. Los resultados de salida se separaron de la definición inmutable del ensayo.')
p('Se cierra el recorrido técnico local solicitado: una entrega exitosa con POD offline y una no entrega con resolución administrativa, de la consola web a la app y de regreso a los reportes. No equivale a cobertura exhaustiva de los 135 requisitos, los nueve procesos ni los 96 casos UAT. La aceptación de MICOTRANS, las condiciones físicas de GPS y cámara y una sesión real de Power BI Desktop permanecen sin acreditar por este ensayo.')
for border in list(doc.element.xpath('.//w:pBdr'))+list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
# Explicitly suppress Word theme residue under the Title paragraph.
no_border=OxmlElement('w:pBdr')
for edge in ['top','left','bottom','right','between','bar']:
    node=OxmlElement('w:'+edge);node.set(qn('w:val'),'nil');no_border.append(node)
doc.paragraphs[0]._p.get_or_add_pPr().append(no_border)
target=ROOT/'docs/correccion_integral/EcoRoute_Validacion_End_to_End_Local.docx'
if '--catalogue-only' not in sys.argv:doc.save(target)

esc=html.escape
cards=''.join(f'<article><h3>{n} {esc(a)}</h3><p>{esc(t)}</p><a href="{esc(e)}"><img loading="lazy" src="{esc(e)}" alt="Paso {n}: {esc(t)}"></a></article>' for n,a,t,e in steps)
links=[('Informe Word','../../correccion_integral/EcoRoute_Validacion_End_to_End_Local.docx'),('Cierre verificable','cierre_recorrido.json'),('API SQL final',final_path.name),('Diagnóstico de sesión','session_diagnosis.md'),('Calidad móvil','calidad_mobile_final.json'),('PDF del nuevo POD','comprobante_nuevo_ensayo.pdf'),('CSV de indicadores','indicadores_nuevo_ensayo.csv'),('Manifiesto de archivos','manifest.json')]
markup=f'''<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>EcoRoute recorrido completo verificado</title><style>body{{font:16px Arial;background:#f5f7fa;color:#17212b;margin:30px auto;max-width:1400px;padding:0 24px}}h1{{font-size:30px}}nav{{display:flex;gap:18px;flex-wrap:wrap;margin:25px 0}}.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:24px}}article{{background:white;border:1px solid #d9d9d9;padding:20px;border-radius:10px}}img{{width:100%;height:370px;object-fit:contain;background:#edf1f4}}a{{color:#135e96}}table{{border-collapse:collapse}}td,th{{padding:12px;border:1px solid #d9d9d9}}small{{display:block;margin-top:18px}}</style><h1>EcoRoute recorrido completo verificado</h1><p>22 de octubre de 2026. Ruta 20: pedido 44 entregado mediante POD offline y pedido 45 fallido por ausencia. Incidencia 12 resuelta como CLOSE_FAILED. Entorno local y datos sintéticos.</p><p>26 controles finales API/SQL aprobados; 81 pruebas Flutter host y ocho contratos reales aprobados. La APK corregida superó tres ciclos de navegación nativa. No es aceptación empresarial ni cobertura exhaustiva del producto.</p><nav>{''.join(f'<a href="{esc(u)}">{esc(t)}</a>' for t,u in links)}</nav><h2>Indicadores de los dos pedidos del ensayo</h2><table><tr><th>Indicador</th><th>Fracción</th><th>Resultado</th></tr>{''.join(f"<tr><td>{k['code']}</td><td>{k['numerator']} / {k['denominator']}</td><td>{k['percentage']} %</td></tr>" for k in final['bi']['kpis'])}</table><h2>Pasos y capturas ampliables</h2><div class="grid">{cards}</div><h2>Condiciones de la validación</h2><p>GPS, cámara y firma corresponden al emulador. No se reabrió el pedido fallido ni se creó un nuevo intento. Se preservan los fallos de automatización y se distinguen de resultados aprobados. La caída nativa inicial de la captura de firma no tiene causa concluyente; la captura repetida se completó.</p><p>Baseline y asignación: evidencia UI y HTTP, sin SQL retrospectivo. La UAT de MICOTRANS y la prueba de campo siguen pendientes. La conexión Power BI Desktop no se ejecutó.</p></html>'''
(OUT/'index.html').write_text(markup,encoding='utf-8')
manifest=[]
for f in sorted(OUT.rglob('*')):
    if f.is_file() and f.name not in ['manifest.json','qa_catalogue.json'] and '_qa' not in f.parts:
        manifest.append({'file':str(f.relative_to(OUT)).replace('\\','/'),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
manifest.append({'file':'../../correccion_integral/'+target.name,'bytes':target.stat().st_size,'sha256':hashlib.sha256(target.read_bytes()).hexdigest()})
(OUT/'manifest.json').write_text(json.dumps({'at':datetime.now(timezone.utc).isoformat(),'files':manifest},ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'README.md').write_text('# Recorrido completo de EcoRoute\n\n[Catálogo de pasos y capturas](index.html) · [Informe Word](../../correccion_integral/EcoRoute_Validacion_End_to_End_Local.docx) · [Cierre verificable](cierre_recorrido.json).\n\nRuta 20 de hoy: pedido 44 entregado con POD offline, pedido 45 fallido, incidencia 12 resuelta como CLOSE_FAILED. Validación local sintética, no UAT empresarial. Los límites y los intentos no aprobados se conservan.\n',encoding='utf-8')
print(json.dumps({'docx':str(target),'html':str(OUT/'index.html'),'files':len(manifest),'server':final_path.name}))
