"""Amplía solo el apartado 2.1 del informe sin rehacer los otros capítulos."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from completar_capitulo2_diseno import add_table, figure, gateway_rows, para
from generar_procesos_negocio_revisados import PROCESSES, MACRO, RF_COVERAGE

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SOURCE = DOCS / 'Integrador I_Grupo3.docx'
OUTPUT = DOCS / 'Integrador I_Grupo3_Procesos_Ampliados.docx'
IMAGES = DOCS / 'entregables_procesos_ecoroute/negocio_revisado/imagenes'

doc = Document(SOURCE)
paragraphs = doc.paragraphs
ch2 = next(p for p in paragraphs if p.style.name == 'Heading 1' and p.text.startswith('CAPÍTULO 2'))
section = next(p for p in paragraphs if p.style.name == 'Heading 2' and p.text == 'Diagramas BPMN')
next_section = next(p for p in paragraphs[paragraphs.index(section)+1:] if p.style.name == 'Heading 2')
body = doc._body._body
start = list(body).index(section._element)
stop = list(body).index(next_section._element)
for element in list(body)[start+1:stop]:
    body.remove(element)
before_append = set(body)

intro = paragraphs[paragraphs.index(ch2)+1]
assert intro.style.name == 'Normal' and 'cinco procesos' in intro.text
intro.text = (
    'Este capítulo presenta la operación objetivo de última milla de Grupo MICOTRANS S. A. C. '
    'mediante nueve procesos BPMN agrupados en seis macroprocesos de gestión. '
    'El flujo principal cubre la admisión del pedido, la preparación de la jornada, la entrega, '
    'las incidencias y la evaluación; los procesos de apoyo cubren recursos, supervisión, accesos '
    'y compras. Los prototipos representan interfaces actuales y futuras; las migraciones V1–V8 '
    'sustentan el modelo de datos. Las políticas operativas propuestas requieren validación con la empresa.'
)

para(doc, 'Convención de nombres: “Gestión de…” identifica el macroproceso, mientras cada proceso BPMN se nombra con verbo en infinitivo y objeto. BPMN describe quién ejecuta y decide en la operación de MICOTRANS; los grupos del PMI dirigen el proyecto de desarrollo y se reflejan en el acta, la EDT/WBS y el cronograma del capítulo 1. No existe una regla de BPMN que prohíba literalmente la conjunción “y”; aquí se evita en los nombres de procesos para conservar un resultado medible. Fuente metodológica: PMI, Process Groups: A Practice Guide; OMG, BPMN 2.0.2.')
para(doc, 'La unidad principal de análisis es el pedido, pero para abarcar el proyecto también se modelan los controles transversales. P01-P05 describen el ciclo del servicio; P06-P09 son procesos de soporte y control. Una función de interfaz o un método del sistema no se convierte por sí solo en un proceso BPMN independiente.')
add_table(doc, ['Macroproceso', 'BPMN incluidos', 'Resultado esperado'], MACRO,
          [2.05, 1.05, 3.2], 8.2)

para(doc, 'Inventario de procesos. Cada evento final indica el estado que cierra el caso o habilita otro proceso; las transferencias entre diagramas no se representan como secuencia BPMN que cruza dos procesos independientes.')
add_table(doc, ['Proceso', 'Disparador', 'Resultado / transferencia', 'Responsable'], [
    (p['id']+' '+p['title'], p['start'], p['success'], p['owner']) for p in PROCESSES
], [1.5, 1.42, 1.92, 1.46], 7.8)

para(doc, 'Cobertura funcional. Los 13 módulos de los 81 requerimientos funcionales quedan relacionados con al menos un proceso. Las funciones técnicas transversales —por ejemplo, tema visual, perfil o rate limiting— se incluyen como soporte, no como un proceso de negocio artificial.')
add_table(doc, ['Módulo funcional', 'BPMN', 'Alcance'], RF_COVERAGE,
          [1.47, .82, 4.01], 7.7)

para(doc, 'Dependencias principales: P08 habilita acceso a los actores; P09 puede proveer recursos a P06; P06 habilita capacidad para P02; P01 confirma pedidos para P02; P02 inicia la ruta de P03; P07 supervisa P03 en paralelo; P03 o P07 abren P04 al detectar una excepción; P04 retorna a P03 para reintento, a P02 para nueva fecha, o cierra con motivo; P05 analiza los resultados al cierre de jornada o período.')

for idx, p in enumerate(PROCESSES, 1):
    heading = doc.add_heading(f"{p['id']} {p['title']}", level=3)
    heading.paragraph_format.keep_with_next = True
    lead = para(doc, f"Inicio: {p['start']}. Resultado principal: {p['success']}. Responsable: {p['owner']}.")
    lead.paragraph_format.keep_with_next = True
    para(doc, p['handoff'], bold_label='Cierre y transferencia. ')
    figure(doc, IMAGES / f"{p['id']}_negocio.png",
           f"Figura 2.{idx}. BPMN de {p['id']} {p['title']}. Fuente: elaboración propia sobre el proceso objetivo de MICOTRANS.",
           width=6.1, max_height=6.4)
    para(doc, 'Decisiones de la compuerta exclusiva. Las salidas “No” corrigen, derivan o cierran el caso de forma explícita.')
    add_table(doc, ['Condición', 'Actor', 'Sí', 'No'], gateway_rows(p['id']),
              [1.53, 1.12, 1.84, 1.81], 7.7, keep_all=True)
    para(doc, p['current'], bold_label='Soporte actual y regla pendiente. ')

para(doc, 'Los nueve diagramas BPMN 2.0 editables y las imágenes a resolución completa se conservan en docs/entregables_procesos_ecoroute/negocio_revisado. Las reglas de negocio no automatizadas se deben acordar con MICOTRANS antes del piloto.')
para(doc, 'Referencias metodológicas: Project Management Institute, Process Groups: A Practice Guide (https://www.pmi.org/); Object Management Group, Business Process Model and Notation 2.0.2 (https://www.omg.org/spec/BPMN/2.0.2/).')

# Mover únicamente el contenido nuevo antes del apartado 2.2; las relaciones de imágenes
# permanecen en el mismo paquete DOCX.
new_elements = [e for e in list(body) if e not in before_append and e.tag != qn('w:sectPr')]
for element in new_elements:
    next_section._element.addprevious(element)

current_paragraphs = doc.paragraphs
section_idx = next(i for i, q in enumerate(current_paragraphs) if q._element is section._element)
next_section_idx = next(i for i, q in enumerate(current_paragraphs) if q._element is next_section._element)
for i, p in enumerate(current_paragraphs):
    if p.text.startswith('Figura 2.'):
        match = re.match(r'Figura 2\.(\d+)\.', p.text)
        if match and int(match.group(1)) >= 6 and not (section_idx < i < next_section_idx):
            old = int(match.group(1))
            p.runs[0].text = p.runs[0].text.replace(f'Figura 2.{old}.', f'Figura 2.{old+4}.', 1)

wire_intro = current_paragraphs[next_section_idx+1]
if 'wireframes antiguos P02 y P03' in wire_intro.text:
    wire_intro.text = (
        'Los wireframes muestran composición, navegación y datos a partir de las pantallas web y móvil. '
        'Sus archivos conservan la numeración de seis láminas preliminares: las láminas antiguas P02 y P03 '
        'representan juntas el proceso actual P02; la lámina antigua P06 corresponde al proceso actual P05. '
        'P06-P09 de este inventario son procesos de apoyo y no tienen todavía una lámina de interfaz específica.'
    )

doc.save(OUTPUT)
print(OUTPUT)
print('BPMN:', len(PROCESSES), 'Macroprocesos:', len(MACRO), 'Módulos:', len(RF_COVERAGE))
