"""Publica el capítulo 2 con BPMN de negocio y participante cliente externo."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from generar_procesos_negocio_revisados import PROCESSES, gate_rows


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SOURCE = DOCS / 'Integrador I_Grupo3_Capitulo2_Flujo_Corregido.docx'
DEST = DOCS / 'Integrador I_Grupo3_Capitulo2_BPMN_Negocio.docx'
DIAGRAMS = DOCS / 'entregables_procesos_ecoroute/negocio_revisado/imagenes'


def replace_picture(paragraph, path: Path) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.1))


def replace_rows(table, rows) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value


doc = Document(SOURCE)
p = doc.paragraphs
t = doc.tables
p[218].text = (
    'La unidad principal es el pedido. P01-P05 cubren el ciclo del servicio y P06-P09 los controles de apoyo. '
    'Dispatcher, Jefatura logística, Conductor y Administrador son responsabilidades funcionales internas de '
    'MICOTRANS; el cliente solicitante interviene externamente. La estructura formal de puestos y las '
    'autorizaciones deben validarse con la empresa antes del piloto.'
)
p[219].text = (
    'Nivel de modelado: los nueve BPMN describen el proceso de negocio objetivo (To-Be). Sus carriles muestran '
    'responsables internos; EcoRoute y Keycloak son herramientas de soporte y no aparecen como actores del negocio. '
    'P01 muestra al cliente solicitante en un pool externo de caja negra y usa flujos de mensaje para la propuesta '
    'y su respuesta. No se modela un supuesto proceso interno del cliente. BPMN permite carriles de sistemas en '
    'diagramas técnicos, pero aquí se separan deliberadamente las responsabilidades empresariales de la automatización.'
)
p[219].paragraph_format.keep_together = True
p[229].text = (
    'Figura 2.1. BPMN de P01 Confirmar pedidos. MICOTRANS realiza el proceso y el cliente solicitante '
    'participa externamente mediante mensajes; la aceptación no equivale a un estado CONFIRMED en EcoRoute. '
    'Fuente: elaboración propia y análisis de la aplicación.'
)
p[237].text = (
    'Figura 2.2. BPMN de P02 Preparar la jornada de reparto. Los carriles corresponden a Dispatcher, '
    'Jefatura logística y Conductor; la vinculación posterior de PENDING se marca como diseño objetivo. '
    'Fuente: elaboración propia y análisis de la aplicación.'
)
responsible = t[29].rows[9].cells[3]
responsible.text = 'Administrador; Jefatura aprueba'
for paragraph in responsible.paragraphs:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.size = Pt(9)

for index, process in enumerate(PROCESSES):
    caption_index = 229 + 8 * index
    assert f"BPMN de {process['id']}" in p[caption_index].text
    replace_picture(p[caption_index - 1], DIAGRAMS / f"{process['id']}_negocio.png")
    replace_rows(t[31 + index], gate_rows(process))

doc.save(DEST)
print(DEST)
